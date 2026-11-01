"""Stage 4a and 4b files.

4a confirm delivers annotation.docx for the chat, download-list.docx as Word,
and a zip of real PDFs only. Cannot-open sources are also a links card in that
reply. 4b confirm delivers one Word file per paper in critical-reviews.zip.
Internal markdown under output/reviews/ stays for stage 5. The bytes are the
same whichever mid office writes them. Draft writes none of these files. The
word dig is not written into them.
"""

from __future__ import annotations

import base64
import os
import re
import zipfile
from pathlib import Path
from typing import Any

from essay_writer.confirm import refuse
from essay_writer.harvard_4a import harvard_4a_line, harvard_4a_parts
from essay_writer.known_paths import (
    allow_research_path,
    is_roam_root,
    known_recorded_paths,
    research_file,
)
from essay_writer.paths import relative, uploads_dir, work_dir
from essay_writer.pdf_year import apply_citation_year
from essay_writer.stage_llm import HOUSE_MODEL
from essay_writer.state import load_state
from essay_writer.student_voice import (
    annotation_prose,
    prose_for_surface,
    strip_word_labels,
)

DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP_MEDIA = "application/zip"
ANNOTATION_NAME = "annotation.docx"
DOWNLOAD_LIST_NAME = "download-list.docx"
PAPERS_ZIP_NAME = "papers.zip"
REVIEWS_ZIP_NAME = "critical-reviews.zip"
_MD_NAME = re.compile(r"[^A-Za-z0-9._-]+")


def chat_docx(path: Path) -> dict[str, Any]:
    """Word file the front office puts in the chat. The bytes match the file on disk."""
    data = path.read_bytes()
    return {
        "name": path.name,
        "path": relative(path),
        "deliver": "chat",
        "media_type": DOCX_MEDIA,
        "docx_base64": base64.b64encode(data).decode("ascii"),
    }


def chat_zip(path: Path) -> dict[str, Any]:
    """Zip the front office puts in the chat. The bytes match the file on disk."""
    data = path.read_bytes()
    return {
        "name": path.name,
        "path": relative(path),
        "deliver": "chat",
        "media_type": ZIP_MEDIA,
        "zip_base64": base64.b64encode(data).decode("ascii"),
    }


def _output(essay_id: str) -> Path:
    return work_dir(essay_id) / "output"


def _md_name(paper_id: str) -> str:
    cleaned = _MD_NAME.sub("-", paper_id).strip(".-") or "paper"
    return f"{cleaned}.md"


def _docx_name(paper_id: str) -> str:
    cleaned = _MD_NAME.sub("-", paper_id).strip(".-") or "paper"
    return f"{cleaned}.docx"


def _unique_zip_name(stem: str, suffix: str, taken: set[str]) -> str:
    name = f"{stem}{suffix}"
    if name not in taken:
        taken.add(name)
        return name
    counter = 2
    while f"{stem}-{counter}{suffix}" in taken:
        counter += 1
    name = f"{stem}-{counter}{suffix}"
    taken.add(name)
    return name


def _save_docx(path: Path, document) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".docx.tmp")
    document.save(str(temporary))
    os.replace(temporary, path)


def write_annotation_docx(
    path: Path,
    entries: list[dict[str, Any]],
    *,
    empty_note: str,
) -> None:
    """Annotated bibliography. Exact citation line, then the full sentences."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    from essay_writer.word import _font, _style_document

    document = Document()
    _style_document(document)
    document.core_properties.title = "Annotated bibliography"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    heading = document.add_heading(strip_word_labels("Annotated bibliography"), level=1)
    heading.paragraph_format.line_spacing = 1.5
    for run in heading.runs:
        _font(run, 14, bold=True)
    if not entries:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5
        note = paragraph.add_run(strip_word_labels(empty_note))
        _font(note, 12)
    for entry in entries:
        citation = document.add_paragraph()
        citation.alignment = WD_ALIGN_PARAGRAPH.LEFT
        citation.paragraph_format.line_spacing = 1.5
        citation.paragraph_format.left_indent = Inches(0.5)
        citation.paragraph_format.first_line_indent = Inches(-0.5)
        for text, italic in entry["parts"]:
            piece = citation.add_run(prose_for_surface(text, "annotation"))
            _font(piece, 12, italic=italic)
        prose = entry.get("paragraph")
        if isinstance(prose, str) and prose.strip():
            written_text = annotation_prose(prose.strip())
            citation_text = "".join(text for text, _italic in entry["parts"]).strip()
            if written_text.strip() and written_text.strip() != citation_text:
                body = document.add_paragraph()
                body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                body.paragraph_format.line_spacing = 1.5
                body.paragraph_format.space_after = Pt(12)
                written = body.add_run(written_text)
                _font(written, 12)
    _save_docx(path, document)


def _pdf_bytes(path: Path) -> bytes | None:
    from essay_writer.downloads import pdf_magic

    if path.suffix.lower() != ".pdf" or not path.is_file():
        return None
    data = path.read_bytes()
    if not pdf_magic(data):
        return None
    return data


def write_papers_zip(path: Path, candidates: list[Path]) -> list[str]:
    """Real PDFs only. Markdown and any other file stay out of the zip."""
    members: list[tuple[str, bytes]] = []
    taken: set[str] = set()
    for candidate in candidates:
        payload = _pdf_bytes(candidate)
        if payload is None:
            continue
        name = candidate.name
        stem = candidate.stem
        suffix = 2
        while name in taken:
            name = f"{stem}-{suffix}.pdf"
            suffix += 1
        taken.add(name)
        members.append((name, payload))
    _write_zip(path, members)
    return [name for name, _payload in members]


def _write_zip(path: Path, members: list[tuple[str, bytes]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members:
            archive.writestr(name, payload)
    os.replace(temporary, path)


def _slug(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"[^a-z0-9]+", "", value.casefold())


def _pdf_matches_row(pdf: Path, row: dict[str, Any]) -> bool:
    stem = _slug(pdf.stem)
    if len(stem) < 6:
        return False
    for key in ("doi", "title", "url", "id"):
        token = _slug(row.get(key))
        if len(token) >= 6 and (token in stem or stem in token):
            return True
    return False


def _source_has_pdf(source: dict[str, Any], saved_pdfs: list[Path]) -> bool:
    saved_names = {path.name.casefold() for path in saved_pdfs}
    locator = source.get("locator")
    if isinstance(locator, str) and locator.lower().endswith(".pdf"):
        name = Path(locator.replace("\\", "/")).name.casefold()
        if name in saved_names:
            return True
    return any(_pdf_matches_row(path, source) for path in saved_pdfs)


def _matches_failure(source: dict[str, Any], row: dict[str, Any]) -> bool:
    for key in ("doi", "title", "url"):
        left = _slug(source.get(key))
        right = _slug(row.get(key))
        if len(left) >= 6 and left == right:
            return True
    locator = source.get("locator")
    if isinstance(locator, str):
        stem = _slug(Path(locator.replace("\\", "/")).stem)
        url = _slug(row.get("url"))
        title = _slug(row.get("title"))
        if len(stem) >= 6 and (stem == url or stem == title or (url and stem in url)):
            return True
    return False


def _equal_token(left: object, right: object) -> bool:
    if not isinstance(left, str) or not isinstance(right, str):
        return False
    a = left.strip()
    b = right.strip()
    if len(a) < 6 or len(b) < 6:
        return False
    if a.casefold() == b.casefold():
        return True
    sa = _slug(a)
    sb = _slug(b)
    return len(sa) >= 6 and sa == sb


def same_source(left: dict[str, Any], right: dict[str, Any]) -> bool:
    """True when two rows name the same gathered source. Short tokens do not match."""
    for key in ("doi", "url", "title", "id"):
        if _equal_token(left.get(key), right.get(key)):
            return True
    for record, other in ((left, right), (right, left)):
        doi = record.get("doi")
        url = other.get("url")
        if not isinstance(doi, str) or not isinstance(url, str):
            continue
        token = _slug(doi)
        if len(token) >= 6 and token in _slug(url):
            return True
    return False


def _paper_paths(papers: list[dict[str, Any]]) -> list[Path]:
    paths: list[Path] = []
    for paper in papers:
        raw = paper.get("path")
        if isinstance(raw, str) and raw.lower().endswith(".pdf"):
            paths.append(Path(raw.replace("\\", "/")))
    return paths


def record_is_opened(
    record: dict[str, Any],
    saved_pdfs: list[Path],
    papers: list[dict[str, Any]] | None = None,
) -> bool:
    """True when this row already has a saved PDF. Filename or the same doi, title, or URL."""
    opened = [item for item in (papers or []) if isinstance(item, dict)]
    if _source_has_pdf(record, list(saved_pdfs) + _paper_paths(opened)):
        return True
    return any(same_source(record, paper) for paper in opened)


def abstract_only_ids(
    sources: list[dict[str, Any]],
    failures: list[dict[str, Any]],
    saved_pdfs: list[Path],
    papers: list[dict[str, Any]] | None = None,
) -> set[str]:
    """Classified abstracts stay off the bibliography until a PDF is opened."""
    banned: set[str] = set()
    abstracts = [item for item in failures if item.get("bucket") == "html-abstract"]
    for source in sources:
        if record_is_opened(source, saved_pdfs, papers):
            continue
        if any(_matches_failure(source, item) for item in abstracts):
            banned.add(str(source.get("id")))
    return banned


def bibliography_leftovers(
    sources: list[dict[str, Any]],
    saved_pdfs: list[Path],
    papers: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Gathered sources on the bibliography that do not yet have a saved PDF."""
    rows: list[dict[str, Any]] = []
    for source in sources:
        if record_is_opened(source, saved_pdfs, papers):
            continue
        row: dict[str, Any] = {
            "bucket": "leftover",
            "title": source.get("title") or source.get("id"),
        }
        if source.get("author"):
            row["author"] = source["author"]
        if source.get("year"):
            row["year"] = source["year"]
        if source.get("doi"):
            row["doi"] = source["doi"]
        if source.get("id"):
            row["id"] = source["id"]
        url = source.get("url")
        if isinstance(url, str) and url.strip().lower().startswith(("http://", "https://")):
            row["url"] = url.strip()
        rows.append(row)
    return rows


def gathered_pdfs(essay_id: str, sources: list[dict[str, Any]]) -> list[Path]:
    """PDF paths already saved for this essay. A markdown upload is not a paper."""
    from essay_writer.downloads import _load_state

    found: list[Path] = []
    seen: set[Path] = set()

    def add(raw: object) -> None:
        if not isinstance(raw, str) or not raw.strip():
            return

        decision = allow_research_path(essay_id, raw.strip())
        if decision.get("allowed") is not True:
            return
        path = Path(str(decision["path"]))
        if path in seen:
            return
        seen.add(path)
        if _pdf_bytes(path) is not None:
            found.append(path)

    for source in sources:
        locator = source.get("locator")
        if isinstance(locator, str) and locator.lower().endswith(".pdf"):
            add(locator)
    for item in _load_state(essay_id).get("papers") or []:
        if isinstance(item, dict):
            add(item.get("path"))
    return found


def _pdf_path(essay_id: str, raw: object) -> Path | None:
    if not isinstance(raw, str) or not raw.lower().endswith(".pdf"):
        return None

    path = research_file(essay_id, raw.strip())
    if path is None or not path.is_file():
        return None
    return path


def _with_pdf_year(essay_id: str, source: dict[str, Any]) -> dict[str, Any]:
    from essay_writer.cite_open import apply_opened_citation

    pdf = _pdf_path(essay_id, source.get("locator")) or _pdf_path(essay_id, source.get("path"))
    if pdf is None:
        return source
    return apply_opened_citation(apply_citation_year(source, pdf), essay_id)


def _kept_paragraph(source: dict[str, Any], notes: dict[str, str]) -> str | None:
    """The stored sentences. A missing note does not become a second citation line."""
    note = notes.get(source["id"]) if isinstance(notes, dict) else None
    if not isinstance(note, str) or not note.strip():
        stored = source.get("annotation")
        note = stored if isinstance(stored, str) else None
    if not isinstance(note, str) or not note.strip():
        return None
    return note.strip()


def _entries(essay_id: str, sources: list[dict[str, Any]], notes: dict[str, str]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for source in sources:
        cited = _with_pdf_year(essay_id, source)
        entries.append(
            {
                "parts": harvard_4a_parts(cited),
                "paragraph": _kept_paragraph(source, notes),
            }
        )
    return entries


def bibliography_sources(essay_id: str) -> list[dict[str, Any]]:
    """Gathered sources already stored for this sequence. Nothing is invented."""
    from essay_writer.downloads import _load_state

    found: list[dict[str, Any]] = []
    research = load_state(essay_id).get("research") or {}
    if isinstance(research, dict):
        for depth in ("summary", "full"):
            record = research.get(depth)
            if not isinstance(record, dict):
                continue
            sources = record.get("sources")
            if not isinstance(sources, list):
                continue
            for item in sources:
                if isinstance(item, dict) and item.get("id"):
                    found.append(item)
    if found:
        return found
    for paper in _load_state(essay_id).get("papers") or []:
        if not isinstance(paper, dict):
            continue
        ident = paper.get("title") or paper.get("path") or paper.get("url")
        if not isinstance(ident, str) or not ident.strip():
            continue
        row: dict[str, Any] = {"id": ident.strip()}
        for key in ("title", "doi", "url", "author", "year", "path"):
            value = paper.get(key)
            if isinstance(value, str) and value.strip():
                row[key] = value.strip()
        if isinstance(paper.get("path"), str):
            row["locator"] = paper["path"]
        found.append(row)
    if found:
        return found
    for pdf in gathered_pdfs(essay_id, []):
        found.append(
            {
                "id": pdf.stem,
                "title": pdf.stem,
                "locator": relative(pdf),
            }
        )
    return found


def real_4a_bibliography(essay_id: str) -> bool:
    """True when stage 4a has an annotation and at least one gathered source."""
    annotation = _output(essay_id) / ANNOTATION_NAME
    if not annotation.is_file() or annotation.stat().st_size == 0:
        return False
    return bool(bibliography_sources(essay_id))


def _gathered_texts(essay_id: str, sources: list[dict[str, Any]]) -> dict[str, str]:
    from essay_writer.invent import read_gathered_text

    texts: dict[str, str] = {}
    for source in sources:
        locator = source.get("locator")
        if not isinstance(locator, str) or not locator.strip():
            continue
        path = research_file(essay_id, locator.strip())
        if path is None:
            continue
        try:
            texts[locator] = read_gathered_text(path, essay_id)
        except (OSError, ValueError):
            continue
    return texts


def _supplied_note(source: dict[str, Any], notes: dict[str, str]) -> tuple[str, bool]:
    """Return the stored note and whether the caller actually supplied one.

    A missing note stays optional for older citation-line writers. An empty
    string, or a surname stub, was supplied and is not a real annotation.
    """
    ident = source.get("id")
    if isinstance(notes, dict) and isinstance(ident, str) and ident in notes:
        raw = notes.get(ident)
        return (raw if isinstance(raw, str) else ""), True
    if "annotation" in source and isinstance(source.get("annotation"), str):
        return source.get("annotation") or "", True
    return "", False


def ship_annotation_issues(
    essay_id: str,
    sources: list[dict[str, Any]],
    notes: dict[str, str],
) -> list[dict[str, str]]:
    """Same annotation bar as essay.research.summary. Refuse stubs before writing files."""
    from essay_writer.harvard_4a import annotation_note_issue, prose_issues

    texts = _gathered_texts(essay_id, sources)
    issues: list[dict[str, str]] = []
    for source in sources:
        if not isinstance(source, dict):
            continue
        ident = source.get("id")
        prose, explicit = _supplied_note(source, notes)
        if not explicit:
            continue
        locator = source.get("locator")
        gathered = texts.get(locator, "") if isinstance(locator, str) else ""
        label = f"{ident or 'source'} annotation"
        stub = annotation_note_issue(label, source, prose)
        if stub:
            issues.append(stub)
            continue
        issues.extend(
            prose_issues(label, source, gathered, prose.strip(), annotation=True)
        )
    return issues


def ship_4a(essay_id: str, sources: list[dict[str, Any]], notes: dict[str, str]) -> dict[str, Any]:
    """Write the 4a pack. Called only after confirm."""
    from essay_writer.downloads import _load_state, _write_download_list, download_links

    folder = _output(essay_id)
    state = _load_state(essay_id)
    pdfs = gathered_pdfs(essay_id, sources)
    papers = folder / PAPERS_ZIP_NAME
    opened = [item for item in (state.get("papers") or []) if isinstance(item, dict)]
    saved_urls = {item.get("url") for item in opened}
    failures = [
        item
        for item in (state.get("failures") or [])
        if isinstance(item, dict)
        and item.get("url") not in saved_urls
        and not record_is_opened(item, pdfs, opened)
    ]
    hidden = abstract_only_ids(sources, failures, pdfs, opened)
    bibliography = [source for source in sources if str(source.get("id")) not in hidden]
    issues = ship_annotation_issues(essay_id, bibliography, notes)
    if issues:
        first = issues[0]
        return refuse(
            str(first.get("code") or "annotation-prose"),
            str(first.get("reason") or "The annotation needs full sentences."),
            wrote=False,
        )
    names = write_papers_zip(papers, pdfs)
    annotation = folder / ANNOTATION_NAME
    write_annotation_docx(
        annotation,
        _entries(essay_id, bibliography, notes),
        empty_note="No gathered source is on this bibliography.",
    )
    leftovers = bibliography_leftovers(bibliography, pdfs, opened)
    listed = _write_download_list(
        essay_id,
        failures,
        [],
        leftovers=leftovers,
        opened=opened,
        pdfs=pdfs,
    )
    from essay_writer.assignment_handoff import papers_zip_pointer

    return {
        "stage": "4a",
        "annotation": chat_docx(annotation),
        "download_list": {
            "name": DOWNLOAD_LIST_NAME,
            "path": listed,
            "format": "docx",
        },
        "papers_zip": papers_zip_pointer(PAPERS_ZIP_NAME, relative(papers), entries=names),
        "links": download_links(failures, [], leftovers, opened=opened, pdfs=pdfs),
    }


def _review_body(essay_id: str, source: dict[str, Any], supplied: str | None) -> str:
    line = harvard_4a_line(_with_pdf_year(essay_id, source))
    note = f"Critical note: {HOUSE_MODEL}\n\n"
    if supplied:
        body = supplied.strip()
        if line not in body:
            body = f"{line}\n\n{body}"
        text = body if body.endswith("\n") else body + "\n"
    else:
        quote = source.get("quote")
        if isinstance(quote, str) and quote.strip():
            text = f"{line}\n\n\"{quote.strip()}\"\n"
        else:
            text = f"{line}\n"
    text = prose_for_surface(text, "review")
    if HOUSE_MODEL not in text:
        text = note + text
    return text


def open_critical_markdown(path: Path, essay_id: str | None = None) -> str:
    """Read the 4b markdown from disk. The note is not kept from memory alone."""
    if essay_id:
        from essay_writer.known_paths import research_file

        if research_file(essay_id, str(path)) is None:
            return ""
    return path.read_text(encoding="utf-8")


def ship_4b(
    essay_id: str,
    sources: list[dict[str, Any]],
    supplied: dict[str, str] | None,
) -> dict[str, Any]:
    """One Word file per paper in the zip. Markdown on disk is for stage 5 only.

    When supplied is None, an existing output/reviews markdown is kept and that
    text is zipped. A citation-line stub is written only when that file is
    absent. Stage 4b confirm uses this same call.
    """
    from essay_writer.word import write_critical_review_docx

    folder = _output(essay_id) / "reviews"
    folder.mkdir(parents=True, exist_ok=True)
    members: list[tuple[str, bytes]] = []
    files: list[dict[str, object]] = []
    md_taken: set[str] = set()
    zip_names: set[str] = set()
    for source in sources:
        paper_id = str(source["id"])
        md_stem = _MD_NAME.sub("-", paper_id).strip(".-") or "paper"
        md_name = _unique_zip_name(md_stem, ".md", md_taken)
        docx_name = _unique_zip_name(md_stem, ".docx", zip_names)
        md_path = folder / md_name
        from essay_writer.known_paths import research_file

        if research_file(essay_id, str(md_path), save=True) is None:
            continue
        if supplied is None and md_path.is_file():
            opened = open_critical_markdown(md_path, essay_id)
        else:
            prose = _review_body(
                essay_id,
                source,
                None if supplied is None else supplied.get(source["id"]),
            )
            md_path.write_text(prose, encoding="utf-8")
            opened = open_critical_markdown(md_path, essay_id)
            if HOUSE_MODEL not in opened:
                opened = f"Critical note: {HOUSE_MODEL}\n\n{opened}"
                md_path.write_text(opened, encoding="utf-8")
                opened = open_critical_markdown(md_path, essay_id)
        docx_path = folder / docx_name
        if research_file(essay_id, str(docx_path), save=True) is None:
            continue
        write_critical_review_docx(docx_path, opened)
        members.append((docx_name, docx_path.read_bytes()))
        files.append(
            {
                "paper_id": source["id"],
                "name": docx_name,
                "path": relative(docx_path),
                "markdown_path": relative(md_path),
                "opened": True,
                "model": HOUSE_MODEL,
            }
        )
    archive = _output(essay_id) / REVIEWS_ZIP_NAME
    _write_zip(archive, members)
    from essay_writer.first_run import copy_deliverable_to_mac

    mac_copy = copy_deliverable_to_mac(archive)
    return {
        "stage": "4b",
        "reviews": files,
        "reviews_zip": {
            "name": REVIEWS_ZIP_NAME,
            "path": relative(archive),
            "entries": [name for name, _payload in members],
            "chat": chat_zip(archive),
        },
        "mac_copy": mac_copy,
    }


def _pdfs_under(folder: Path | None) -> list[Path]:

    if folder is None or not folder.is_dir() or is_roam_root(folder):
        return []
    found: list[Path] = []
    for path in folder.rglob("*.pdf"):
        if not path.is_file() or path.name.startswith("."):
            continue
        if is_roam_root(path):
            continue
        if _pdf_bytes(path) is not None:
            found.append(path)
    return found


def _list_rows(essay_id: str, sources: list[dict[str, Any]]) -> list[dict[str, Any]]:
    from essay_writer.downloads import _load_state

    state = _load_state(essay_id)
    opened = [item for item in (state.get("papers") or []) if isinstance(item, dict)]
    pdfs = gathered_pdfs(essay_id, sources)
    rows: list[dict[str, Any]] = []
    for item in state.get("failures") or []:
        if isinstance(item, dict) and not record_is_opened(item, pdfs, opened):
            rows.append(item)
    hidden = abstract_only_ids(sources, rows, pdfs, opened)
    visible = [source for source in sources if str(source.get("id")) not in hidden]
    rows.extend(bibliography_leftovers(visible, pdfs, opened))
    return rows


def _append_citation(essay_id: str, annotation: Path, paper: dict[str, Any]) -> None:
    """Add one citation line. An existing paragraph on that line stays."""
    from docx import Document

    from essay_writer.word import _font

    cited = _with_pdf_year(essay_id, paper)
    line = harvard_4a_line(cited)
    if annotation.is_file():
        document = Document(str(annotation))
        if any(line in (paragraph.text or "") for paragraph in document.paragraphs):
            return
    else:
        write_annotation_docx(
            annotation,
            [{"parts": harvard_4a_parts(cited), "paragraph": None}],
            empty_note="No gathered source is on this bibliography.",
        )
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(prose_for_surface(line, "annotation"))
    _font(run, 12)
    _save_docx(annotation, document)


def _refresh_4a(essay_id: str, sources: list[dict[str, Any]], papers: list[dict[str, Any]]) -> dict[str, Any]:
    from essay_writer.downloads import _write_download_list

    pdfs = gathered_pdfs(essay_id, sources)
    archive = _output(essay_id) / PAPERS_ZIP_NAME
    names = write_papers_zip(archive, pdfs)
    from essay_writer.downloads import _load_state

    state = _load_state(essay_id)
    opened = [item for item in (state.get("papers") or []) if isinstance(item, dict)]
    saved_urls = {item.get("url") for item in opened}
    failures = [
        item
        for item in (state.get("failures") or [])
        if isinstance(item, dict)
        and item.get("url") not in saved_urls
        and not record_is_opened(item, pdfs, opened)
    ]
    hidden = abstract_only_ids(sources, failures, pdfs, opened)
    bibliography = [source for source in sources if str(source.get("id")) not in hidden]
    annotation = _output(essay_id) / ANNOTATION_NAME
    if not annotation.is_file():
        notes = {
            source["id"]: source["annotation"]
            for source in bibliography
            if isinstance(source.get("annotation"), str) and isinstance(source.get("id"), str)
        }
        write_annotation_docx(
            annotation,
            _entries(essay_id, bibliography, notes),
            empty_note="No gathered source is on this bibliography.",
        )
    for paper in papers:
        if isinstance(paper, dict):
            _append_citation(essay_id, annotation, paper)
    leftovers = bibliography_leftovers(bibliography, pdfs, opened)
    listed = _write_download_list(
        essay_id,
        failures,
        [],
        leftovers=leftovers,
        opened=opened,
        pdfs=pdfs,
    )
    return {
        "annotation": relative(annotation),
        "papers_zip": relative(archive),
        "download_list": listed,
        "zip_entries": names,
        "still_listed": [_cite_row(item) for item in leftovers] + [_cite_row(item) for item in failures],
    }


def _cite_row(item: dict[str, Any]) -> str:
    for key in ("title", "doi", "url", "id"):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return "Untitled download"


def reconcile_before_4b(
    essay_id: str,
    *,
    documents: Path | None = None,
) -> dict[str, Any]:
    """Compare the download list with PDFs on known paths.

    Product folders (uploads and the sequence workflow) are read. A PDF
    outside those folders is opened only when its path is already named on
    the download list, the annotation, or a known project file. Home, whole
    Documents, and whole Volumes are not searched. The papers zip is not the
    inventory. A listed source that now has a PDF is copied onto the essay
    disk when it was outside the sequence, then annotation.docx and the
    papers zip are rewritten.
    """
    import shutil

    from essay_writer.downloads import _load_state, _save_state, _state_path

    research = load_state(essay_id).get("research") or {}
    summary = research.get("summary") if isinstance(research, dict) else None
    sources = []
    if isinstance(summary, dict) and isinstance(summary.get("sources"), list):
        sources = [item for item in summary["sources"] if isinstance(item, dict)]
    rows = _list_rows(essay_id, sources)
    named_documents = None if documents is None or is_roam_root(documents) else documents
    seen: set[Path] = set()
    found: list[tuple[Path, str]] = []

    def _take(path: Path, where: str) -> None:
        decision = allow_research_path(essay_id, str(path))
        if decision.get("allowed") is not True:
            return
        resolved = path.resolve()
        if resolved in seen or _pdf_bytes(path) is None:
            return
        seen.add(resolved)
        found.append((path, where))

    for path in _pdfs_under(uploads_dir(essay_id)) + _pdfs_under(work_dir(essay_id)):
        _take(path, "disk")
    for path in known_recorded_paths(essay_id):
        if not path.is_file():
            continue
        where = "disk"
        under_product = False
        for folder in (uploads_dir(essay_id), work_dir(essay_id)):
            try:
                path.resolve().relative_to(folder.resolve())
                under_product = True
            except (OSError, ValueError):
                continue
        if not under_product:
            where = "Documents"
        elif named_documents is not None:
            try:
                path.resolve().relative_to(named_documents.resolve())
                where = "Documents"
            except (OSError, ValueError):
                where = "disk"
        _take(path, where)
    state = _load_state(essay_id)
    matched: list[dict[str, str]] = []
    papers = list(state.get("papers") or [])
    known_paths = {item.get("path") for item in papers if isinstance(item, dict)}
    for pdf, where in found:
        hit = next((row for row in rows if _pdf_matches_row(pdf, row)), None)
        if hit is None:
            continue
        dest = pdf
        if where == "Documents":
            folder = uploads_dir(essay_id)
            folder.mkdir(parents=True, exist_ok=True)
            dest = folder / pdf.name
            if dest.resolve() != pdf.resolve():
                shutil.copy2(pdf, dest)
        rel = relative(dest)
        if rel not in known_paths:
            row: dict[str, Any] = {"path": rel, "bucket": "paper", "url": hit.get("url") or rel}
            if hit.get("doi"):
                row["doi"] = hit["doi"]
            if hit.get("title"):
                row["title"] = hit["title"]
            papers.append(row)
            known_paths.add(rel)
        url = hit.get("url")
        if isinstance(url, str):
            state["failures"] = [
                item
                for item in state.get("failures") or []
                if not (isinstance(item, dict) and item.get("url") == url)
            ]
        matched.append({"name": pdf.name, "where": where})
    state["papers"] = papers
    existed = _state_path(essay_id).is_file()
    blank = (
        not state.get("papers")
        and not state.get("failures")
        and not state.get("attempts")
        and not state.get("seen_urls")
    )
    if existed or not blank:
        _save_state(essay_id, state)
    refreshed = _refresh_4a(essay_id, sources, papers) if matched else None
    if refreshed is None:
        still = [_cite_row(item) for item in rows]
    else:
        still = refreshed["still_listed"]
    option_c = "Download leftovers first"
    return {
        "checked": "disk-and-documents",
        "zip_only": False,
        "matched": matched,
        "still_listed": still,
        "refreshed": None
        if refreshed is None
        else {
            "annotation": refreshed["annotation"],
            "papers_zip": refreshed["papers_zip"],
            "download_list": refreshed["download_list"],
        },
        "option_c": option_c,
    }


def download_annotation(path: Path, papers: list[dict[str, Any]], essay_id: str | None = None) -> None:
    """Stop-path Word file. Same citation line. Failed downloads are not entries."""
    cited = [_with_pdf_year(essay_id or "", item) for item in papers]
    entries = [{"parts": harvard_4a_parts(item), "paragraph": None} for item in cited]
    write_annotation_docx(path, entries, empty_note="No PDF was saved.")
