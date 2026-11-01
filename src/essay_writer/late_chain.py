"""After stage 7 ask and edit, stages 8, 9, and 10 are separate runs.

One call does not run the later stages. A blocked independent review is
shown. Stage 10 writes the sequence Word file, refreshes that sequence's
-10 final from those bytes, and writes Documents/<project>-10/ when mounted.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from essay_writer import assignment_gate as essay_assignment_gate
from essay_writer import textutil as essay_textutil
from essay_writer.diary import (
    _cover_stored_title,
    _real_title,
    _stated_cover_values,
    assignment_on_cover,
    cover_residue_line,
)
from essay_writer.known_paths import refuse_repo_write
from essay_writer.paths import module_code, work_dir, workspace_root
from essay_writer.stage7 import comment_text, load_stage7
from essay_writer.state import load_state
from essay_writer.student_voice import say
from essay_writer.user_place import master_word

_LATE = frozenset({"7", "8", "9", "10"})


def documents_root(override: Path | None = None) -> Path | None:
    """Documents when this computer can see it. Tests pass the folder."""
    import os
    import sys

    if override is not None:
        return override
    env = os.environ.get("ESSAY_WRITER_MAC_DIR")
    if env is not None:
        text = env.strip()
        if not text or text.casefold() == "open":
            return None
        return Path(text).expanduser()
    if os.environ.get("PYTEST_CURRENT_TEST") or os.environ.get("ESSAY_WRITER_ROOT"):
        return None
    if sys.platform != "darwin":
        return None
    folder = Path.home() / "Documents"
    if folder.is_dir():
        return folder
    return None


def final_place(essay_id: str) -> str:
    from essay_writer.first_run import (
        sequence_mac_deliverables_dir,
        sequence_mac_handoff_place,
    )

    project = module_code(essay_id)
    if sequence_mac_deliverables_dir(essay_id) is not None:
        return sequence_mac_handoff_place(essay_id, f"{project}-final.docx")
    return f"Documents/{project}-10/{project}-final.docx"


def comments_place(essay_id: str) -> str:
    from essay_writer.first_run import (
        sequence_mac_deliverables_dir,
        sequence_mac_handoff_place,
    )

    project = module_code(essay_id)
    if sequence_mac_deliverables_dir(essay_id) is not None:
        return sequence_mac_handoff_place(essay_id, f"{project}-comments.docx")
    return f"Documents/{project}-10/{project}-comments.docx"


def findings_place(essay_id: str) -> str:
    """Independent-review findings. The project folder, not the stage 10 folder."""
    project = module_code(essay_id)
    return f"Documents/{project}/{project}-findings.docx"


def chain_armed(pointer: dict[str, Any]) -> bool:
    return pointer.get("chain_armed") is True


def next_run(essay_id: str, stage: str, *, armed: bool) -> dict[str, Any] | None:
    """The single following run. Never 8, 9, and 10 together."""
    if not armed or stage == "6":
        return None
    if stage == "7":
        return {"stage": "8", "run": "edit", "separate": True, "runs_in_this_step": 1}
    if stage == "8":
        return {"stage": "9", "run": "references", "separate": True, "runs_in_this_step": 1}
    if stage == "9":
        from essay_writer.spine import final_review

        text, verdict = final_review(essay_id)
        if text:
            return {
                "stage": "9",
                "run": "review",
                "separate": True,
                "runs_in_this_step": 1,
                "blocked": True,
                "eqr": verdict,
            }
        return {
            "stage": "10",
            "run": "final",
            "separate": True,
            "runs_in_this_step": 1,
            "eqr": verdict,
        }
    return None


def review_block(essay_id: str) -> str | None:
    from essay_writer.spine import final_review

    text, _verdict = final_review(essay_id)
    return text


STAGE10_UNMOUNTED = (
    "The Mac folder is not on this computer. "
    "The final Word stays in this chat and on this sequence. "
    "This tool does not submit."
)
FINDINGS_UNMOUNTED = (
    "The review opened the saved Word file. "
    "The Mac folder is not on this computer, so the findings stay in this chat and on this sequence."
)
QUIET_STEP = "Wait while the first draft is written."
FIRST_DRAFT_READY = "Read the first draft in this chat."
STAGE6_OPEN = "Your draft is saved. Start the first read in this chat."
_LEAKS = (
    "Essay-writer/",
    "uploads/",
    "tracking/",
    "sequences/",
    "~/",
    "/Users/",
    "/Volumes/",
    "ResearchDrive",
    "Apply every line",
)


def quiet_next(stage: str, *, armed: bool, blocked: bool, place: str) -> str | None:
    if not armed:
        return None
    if blocked:
        return "Read the independent review note in this chat before the final Word file."
    if stage == "10":
        return place
    return None


def chat_leaks(text: str) -> bool:
    """Checklist dump or a file path in the line a person would hear."""
    return any(token in text for token in _LEAKS)


def quiet_chat(text: str) -> str:

    spoken = say(text)
    if not spoken or chat_leaks(spoken):
        return QUIET_STEP
    return spoken


def completed_draft_chat(prose: str) -> str:
    """The finished first draft, or a short ready line when the prose would leak."""

    spoken = say(prose).strip() if prose else ""
    if not spoken or chat_leaks(spoken):
        return FIRST_DRAFT_READY
    return spoken


def draft_prose(pack: dict[str, Any] | None) -> str:
    if not isinstance(pack, dict):
        return ""
    sections = pack.get("sections")
    if not isinstance(sections, list):
        return ""
    parts: list[str] = []
    for section in sections:
        if isinstance(section, dict) and isinstance(section.get("body"), str):
            text = section["body"].strip()
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def mark_quiet(body: dict[str, Any], chat: str, *, step: str) -> dict[str, Any]:
    body["quiet"] = True
    body["step"] = step
    body["checklist_in_chat"] = False
    body["chat"] = chat
    body["next"] = chat
    body["confirm_line"] = chat
    body.pop("tracking_path", None)
    return body


def stage5_chat_line(text: str, fallback: str) -> str:
    """Stage 5 chat. A Basics checklist dump is not the line a person hears."""
    if not isinstance(text, str) or not text.strip():
        return fallback
    if chat_leaks(text) or "Basics checklist" in text or "Apply every line" in text:
        return fallback
    return text.strip()


def quiet_stage5(body: dict[str, Any], essay_id: str) -> dict[str, Any]:
    """Stage 5 speech. No save prompt. The finished draft is heard after Word exists."""

    word = work_dir(essay_id) / "output" / f"{essay_id}.docx"
    if body.get("stored") and word.is_file():
        line = stage5_chat_line(FIRST_DRAFT_READY, FIRST_DRAFT_READY)
    else:
        line = stage5_chat_line(STAGE6_OPEN, STAGE6_OPEN)
    quiet = mark_quiet(body, line, step="5")
    quiet["checklist_in_chat"] = False
    from essay_writer.cite_fit import speak_cite_ask

    return speak_cite_ask(quiet)


def quiet_word_step(body: dict[str, Any], pack: dict[str, Any] | None) -> dict[str, Any]:
    """Word preview is the quiet middle pass. Confirm shows the finished draft."""
    blocked = body.get("ok") is False or bool(body.get("blockers")) or bool(body.get("refused"))
    if blocked:
        body["quiet"] = False
        body["step"] = "5b"
        body["checklist_in_chat"] = False
        return body
    if body.get("wrote"):
        return mark_quiet(body, completed_draft_chat(draft_prose(pack)), step="5b")
    return mark_quiet(body, QUIET_STEP, step="5b")


def log_word_manifest(
    essay_id: str,
    pack: dict[str, Any],
    report: dict[str, Any],
    *,
    pack_path: str,
    product_log: bool = True,
) -> None:
    if not product_log:
        return
    from essay_writer.debug_log import (
        log_word_step,
        pack_log_assessment,
        pack_log_title,
    )

    cover = report.get("cover") if isinstance(report.get("cover"), list) else []
    log_word_step(
        "word_manifest",
        essay_id,
        pack_path=pack_path,
        title=pack_log_title(pack),
        assessment=pack_log_assessment(pack),
        conclusion=str(pack.get("conclusion") or ""),
        cover=cover,
        sources_count=len(report.get("sources") or []),
        refs_count=len(report.get("references") or []),
    )


def log_word_docx_written(
    essay_id: str,
    pack: dict[str, Any],
    report: dict[str, Any],
    cover: list[str],
    path: Path,
    *,
    product_log: bool = True,
) -> None:
    if not product_log:
        return
    from essay_writer.debug_log import (
        log_word_step,
        pack_log_assessment,
        pack_log_title,
    )

    log_word_step(
        "write_docx",
        essay_id,
        title=pack_log_title(pack),
        assessment=pack_log_assessment(pack),
        conclusion=str(pack.get("conclusion") or ""),
        cover=cover,
        sources_count=len(report.get("sources") or []),
        refs_count=len(report.get("sources") or []),
        output_path=path,
    )


def _saved_pack(essay_id: str) -> dict[str, Any] | None:

    pack = load_state(essay_id).get("saved_pack")
    if isinstance(pack, dict):
        return pack
    return None


def _infer_source_origin(locator: str) -> str:
    lowered = locator.strip().casefold()
    if lowered.startswith("uploads/"):
        return "upload"
    if lowered.startswith("essay-writer/") or lowered.startswith("kb1/"):
        return "kb1"
    return "brief"


def _pack_sources(essay_id: str, pack: dict[str, Any]) -> list[dict[str, Any]]:
    """Sources for Word: pack list, then the stored pack, then confirmed research."""

    raw = pack.get("sources")
    if isinstance(raw, list) and raw:
        return [item for item in raw if isinstance(item, dict)]
    if pack.get("recovered_from_docx") is True and isinstance(raw, list):
        return [item for item in raw if isinstance(item, dict)]
    saved = load_state(essay_id).get("saved_pack")
    if isinstance(saved, dict):
        stored = saved.get("sources")
        if isinstance(stored, list) and stored:
            return [item for item in stored if isinstance(item, dict)]
    from essay_writer.research_packs import bibliography_sources

    rows: list[dict[str, Any]] = []
    for item in bibliography_sources(essay_id):
        if not isinstance(item, dict):
            continue
        row = dict(item)
        locator = row.get("locator")
        if isinstance(locator, str) and locator.strip() and not row.get("origin"):
            row["origin"] = _infer_source_origin(locator)
        rows.append(row)
    return rows


def _word_pack(essay_id: str, pack: dict[str, Any]) -> dict[str, Any]:
    """Pack fields Word and inspect_pack need after a body-only recovery from .docx."""
    filled = dict(pack)
    sources = _pack_sources(essay_id, filled)
    if sources:
        filled["sources"] = sources

    assessment = filled.get("assessment")
    if not _cover_stored_title(assessment):
        title = _cover_stored_title(filled.get("title"))
        if title:
            filled["assessment"] = title
    if not filled.get("referencing"):
        filled["referencing"] = "harvard"
    return filled


_STACK_MARK = re.compile(
    r"(?i)^(?:registration number|counted words|stated word count)\s*:"
)


def _wrapped_draft(lines: list[str], cover: list[str]) -> bool:
    """True when a Draft section is an earlier Word file, not the essay."""
    known = {line.strip().casefold() for line in cover if line and line.strip()}
    for line in lines:
        folded = line.strip().casefold()
        if folded in {"draft", "references"} or folded in known:
            return True
        if _STACK_MARK.match(line.strip()):
            return True
    return False


def _split_named(lines: list[str], names: list[str]) -> list[dict[str, Any]] | None:
    """Split flattened paragraphs on the brief's part names, in that order."""
    lookup = {name.casefold(): name for name in names}
    present = {line.casefold() for line in lines}
    if any(name.casefold() not in present for name in names):
        return None
    buckets: dict[str, list[str]] = {name.casefold(): [] for name in names}
    current: str | None = None
    for line in lines:
        key = line.casefold()
        if key in lookup:
            current = key
            continue
        if current is None:
            continue
        buckets[current].append(line)
    return [{"heading": name, "lines": buckets[name.casefold()]} for name in names]


def _unwrap_draft(essay_id: str, lines: list[str], cover: list[str]) -> list[dict[str, Any]]:
    """Drop stacked covers and recover the essay sections inside a Draft wrap."""

    compared = list(cover)
    compared.extend(_stated_cover_values(essay_textutil.module_brief_text(essay_id)))
    kept = [line for line in lines if not cover_residue_line(line, compared)]
    record = essay_assignment_gate.read_assignment(essay_id)
    names: list[str] = []
    if record.get("lock") not in {None, "thesis", "purpose_and_scope"}:
        stored = record.get("named_entries")
        if isinstance(stored, list):
            names = [name.strip() for name in stored if isinstance(name, str) and name.strip()]
    if names:
        split = _split_named(kept, names)
        if split:
            return split
    if not kept:
        return []
    return [{"heading": kept[0], "lines": kept[1:]}]


_COUNTED_COVER = re.compile(r"(?i)^counted words:\s*(\d+)\s*$")
_REGISTRATION_COVER = re.compile(r"(?i)^registration number:\s*(.+)\s*$")
_REFERENCE_AUTHOR_YEAR = re.compile(r"^(.+?)\s*\((\d{4})")


def _programme_cover_lines() -> set[str]:
    from essay_writer.programme import PROGRAMME

    return {str(value).casefold() for value in PROGRAMME.values()}


def _hydrate_pack_from_docx_cover(
    essay_id: str, cover: list[str], pack: dict[str, Any]
) -> dict[str, Any]:
    """Carry cover values already printed on disk into the recovered pack."""

    programme = _programme_cover_lines()
    brief = essay_textutil.module_brief_text(essay_id)
    brief_assignment = assignment_on_cover({}, {}, brief)
    out = dict(pack)
    assignment_lines: list[str] = []
    for line in cover:
        text = line.strip()
        if not text:
            continue
        if text.casefold() in programme or text == essay_id:
            continue
        if text == "Conclusion: noted":
            out["conclusion"] = "noted"
            continue
        if text == "Conclusion: open — not a settled finding":
            out["conclusion"] = "open"
            continue
        match = _REGISTRATION_COVER.match(text)
        if match:
            out["registration_number"] = match.group(1).strip()
            continue
        if _COUNTED_COVER.match(text):
            continue
        if text.endswith(":") and not _real_title(text.rstrip(":")):
            continue
        if _real_title(text) and not cover_residue_line(text):
            assignment_lines.append(text)
    if assignment_lines and not _real_title(out.get("assessment")):
        chosen = None
        if isinstance(brief_assignment, str) and brief_assignment.strip():
            wanted = brief_assignment.strip()
            for line in assignment_lines:
                if line.strip() == wanted:
                    chosen = wanted
                    break
        if chosen is None:
            for line in assignment_lines:
                if not cover_residue_line(line):
                    chosen = line.strip()
                    break
        if chosen is None:
            chosen = assignment_lines[0].strip()
        out["assessment"] = chosen
        if not _real_title(out.get("title")):
            out["title"] = chosen
    return out


def _paragraph_text_style(paragraph) -> tuple[str, str]:
    text = paragraph.text.strip()
    style = ""
    if paragraph.style is not None and getattr(paragraph.style, "name", None):
        style = paragraph.style.name
    return text, style


def _reference_lines_from_docx(document) -> list[str]:
    """Reference lines under the References heading, including table cells."""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    current = False

    def add(text: str) -> None:
        if text and (not lines or lines[-1] != text):
            lines.append(text)

    def consume(paragraph) -> bool:
        nonlocal current
        text, style = _paragraph_text_style(paragraph)
        heading = style.startswith("Heading")
        if heading and text.casefold() == "references":
            current = True
            return False
        if not current:
            return False
        if heading and text:
            return True
        add(text)
        return False

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            if consume(Paragraph(child, document)):
                break
            continue
        if child.tag != qn("w:tbl") or not current:
            continue
        table = Table(child, document)
        stop = False
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                ident = id(cell._tc)
                if ident in seen:
                    continue
                seen.add(ident)
                for paragraph in cell.paragraphs:
                    if consume(paragraph):
                        stop = True
                        break
                if stop:
                    break
            if stop:
                break
        if stop:
            break
    return lines


def _source_for_reference_line(essay_id: str, line: str) -> dict[str, Any] | None:
    """Match a printed Harvard line to a gathered upload. Nothing is invented."""
    from essay_writer.invent import read_gathered_text

    match = _REFERENCE_AUTHOR_YEAR.match(line.strip())
    if match is None:
        return None
    author = match.group(1).strip().rstrip(",")
    year = match.group(2)
    surname = author.split(",", 1)[0].strip().split()[-1].casefold()
    if not surname:
        return None
    module = module_code(essay_id)
    upload_dir = workspace_root() / "uploads" / module
    if not upload_dir.is_dir():
        return None
    for path in sorted(upload_dir.iterdir()):
        if not path.is_file():
            continue
        try:
            text = read_gathered_text(path, essay_id)
        except (OSError, ValueError):
            continue
        folded = text.casefold()
        if surname not in folded or year not in text:
            continue
        quote = ""
        for sentence in re.split(r"(?<=[.!?])\s+", text):
            piece = sentence.strip()
            if surname in piece.casefold() and year in piece:
                quote = piece
                break
        if not quote:
            quote = text.strip()[:500]
        # The printed line is matched to the upload. Its title is not parsed
        # here, and the line itself is not a gathered title, so none is stored.
        return {
            "id": f"src-{path.stem}",
            "origin": "upload",
            "locator": f"uploads/{module}/{path.name}",
            "author": author,
            "year": year,
            "quote": quote[:500],
        }
    return None


def _finish_recovered_pack(
    essay_id: str,
    pack: dict[str, Any],
    reference_lines: list[str],
) -> dict[str, Any]:
    """Keep the reference lines printed in the Word file.

    When the file lists references, those lines are the reference list.
    Research that is not one of those lines is not added.
    """
    filled = dict(pack)
    filled["recovered_from_docx"] = True
    if reference_lines:
        filled["references_markdown"] = "\n".join(reference_lines)
    filled = _word_pack(essay_id, filled)
    if reference_lines:
        filled["references_markdown"] = "\n".join(reference_lines)
    if filled.get("sources") and not reference_lines:
        return filled
    matched: list[dict[str, Any]] = []
    seen: set[str] = set()
    for line in reference_lines:
        source = _source_for_reference_line(essay_id, line)
        if source is None:
            continue
        ident = str(source.get("id") or "")
        if not ident or ident in seen:
            continue
        seen.add(ident)
        matched.append(source)
    if reference_lines:
        filled["sources"] = matched
        return filled
    if matched:
        filled["sources"] = matched
    return filled


def _reference_lines_in_docx(path: Path) -> int:
    from docx import Document

    try:
        document = Document(str(path))
    except Exception:
        return 0
    return len(_reference_lines_from_docx(document))


def _pack_from_docx(essay_id: str, source: Path) -> dict[str, Any] | None:
    """Essay sections from an earlier Word file.

    The cover that write_docx prepends is not part of the essay. A previous
    rewrite that wrapped every paragraph under Draft is unwrapped, so the next
    write does not stack another cover or another Draft heading.
    """
    from docx import Document

    try:
        document = Document(str(source))
    except Exception:
        return None
    reference_lines = _reference_lines_from_docx(document)
    cover: list[str] = []
    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") and text:
            current = {"heading": text, "lines": []}
            sections.append(current)
            continue
        if current is None:
            if text:
                cover.append(text)
            continue
        if text:
            current["lines"].append(text)
    content = [item for item in sections if item["heading"].casefold() != "references"]
    if content and all(item["heading"].casefold() == "draft" for item in content):
        lines: list[str] = []
        for item in content:
            lines.extend(item["lines"])
        if len(content) > 1 or _wrapped_draft(lines, cover):
            content = _unwrap_draft(essay_id, lines, cover)
    elif not content:
        if _wrapped_draft(cover, []):
            content = _unwrap_draft(essay_id, cover, [])
        elif cover:
            content = [{"heading": cover[0], "lines": cover[1:]}]
    packed: list[dict[str, str]] = []
    for item in content:
        heading = str(item.get("heading") or "").strip()
        if not heading or heading.casefold() == "references":
            continue
        lines_raw = item.get("lines")
        body_lines: list[object] = lines_raw if isinstance(lines_raw, list) else []
        body = "\n".join(line.strip() for line in body_lines if isinstance(line, str) and line.strip())
        packed.append({"heading": heading, "body": body})
    if not packed:
        return None
    pack = {
        "essay_id": essay_id,
        "title": "Draft",
        "sections": packed,
        "conclusion": "noted",
        "referencing": "harvard",
    }
    pack = _hydrate_pack_from_docx_cover(essay_id, cover, pack)
    return _finish_recovered_pack(essay_id, pack, reference_lines)


def _refresh_sequence_final(source: Path, project: str) -> None:
    """Write output/<module>-10/<module>-final.docx from the sequence Word file just saved."""
    import os


    folder = source.parent / f"{project}-10"
    folder.mkdir(parents=True, exist_ok=True)
    final = folder / f"{project}-final.docx"
    refuse_repo_write(final)
    pending = final.with_name(f".{final.name}.writing")
    pending.write_bytes(source.read_bytes())
    os.replace(pending, final)


def _unmounted_documents(body: dict[str, Any]) -> None:
    body["reason"] = "mac-not-mounted"
    body["copied"] = False
    body["findings_place"] = ""
    body["findings_file"] = {
        "wrote": False,
        "place": "",
        "reason": "mac-not-mounted",
        "refused": "documents",
        "ok": False,
    }


def _footer_stop(body: dict[str, Any], *, refused: str, reason: str) -> dict[str, Any]:
    body["ok"] = False
    body["copied"] = False
    body["sequence_written"] = False
    body["refused"] = refused
    body["reason"] = reason
    return body


def write_stage10(essay_id: str, *, documents: Path | None = None) -> dict[str, Any]:
    """Write the sequence Word file, then refresh the sequence -10 final from it.

    The sequence file is Essay-writer/<module>/sequences/<n>/output/<id>.docx.
    Those same bytes are written to output/<module>-10/<module>-final.docx so a
    cloned parent final cannot stay stale. Documents/<module>-10/<module>-final.docx
    is written only when that folder is mounted. A missing registration number
    or an empty footer does not record the sequence file as the final.
    """
    import os

    project = module_code(essay_id)
    place = final_place(essay_id)
    notes = comments_place(essay_id)
    root = documents_root(documents)
    source = work_dir(essay_id) / "output" / f"{essay_id}.docx"
    body: dict[str, Any] = {
        "ok": True,
        "copied": False,
        "place": place,
        "comments_place": notes,
        "findings_place": "",
        "final_name": f"{project}-final.docx",
        "first_draft_name": source.name,
        "sequence_written": False,
    }
    saved_pack = _saved_pack(essay_id)
    pack_path = "saved_pack" if saved_pack is not None else "missing"
    pack = saved_pack
    existing_ref_lines = _reference_lines_in_docx(source) if source.is_file() else 0
    if pack is None and source.is_file():
        pack = _pack_from_docx(essay_id, source)
        if pack is not None:
            pack_path = "pack_from_docx"
    saved_pack_state = "present" if saved_pack is not None else "null"
    body["pack_path"] = pack_path
    body["saved_pack_state"] = saved_pack_state
    if pack is None:
        body["ok"] = False
        body["reason"] = "The saved draft Word file is not there yet."
        from essay_writer.debug_log import log_word_step

        log_word_step(
            "write_stage10",
            essay_id,
            stage="10",
            pack_path=pack_path,
            saved_pack=saved_pack_state,
            refused="missing-pack",
            reason=body["reason"],
        )
        if root is None:
            _unmounted_documents(body)
            body["ok"] = False
        return body
    from essay_writer.brief_contract import word_contract_refusal
    from essay_writer.debug_log import (
        log_word_step,
        pack_log_assessment,
        pack_log_title,
    )
    from essay_writer.word import (
        EmptyFooter,
        empty_footer_refusal,
        word_manifest,
        write_docx,
    )

    pack, report = word_manifest(essay_id, pack, pack_path=pack_path, product_log=False)
    refs_count = len(report.get("references") or [])
    sources_count = len(report.get("sources") or [])
    cover = report.get("cover") if isinstance(report.get("cover"), list) else []
    if (
        pack_path == "pack_from_docx"
        and existing_ref_lines > 0
        and refs_count == 0
    ):
        reason = (
            "The sequence Word file already lists references, but recovery from that file "
            "would write an empty References section. Restore saved_pack or confirm research first."
        )
        log_word_step(
            "write_stage10",
            essay_id,
            stage="10",
            pack_path=pack_path,
            saved_pack=saved_pack_state,
            title=pack_log_title(pack),
            assessment=pack_log_assessment(pack),
            conclusion=str(pack.get("conclusion") or ""),
            cover=cover,
            sources_count=sources_count,
            refs_count=refs_count,
            refused="word-recovery",
            reason=reason,
        )
        body["ok"] = False
        body["copied"] = False
        body["sequence_written"] = False
        body["refused"] = "word-recovery"
        body["reason"] = reason
        body["references"] = report.get("references") or []
        body["cover"] = cover
        body["sources"] = report.get("sources") or []
        return body
    blocked = word_contract_refusal(essay_id, pack, report)
    if blocked:
        body["ok"] = False
        body["copied"] = False
        body["refused"] = blocked.get("refused")
        body["reason"] = blocked.get("reason") or ""
        body["missing_fields"] = blocked.get("missing_fields")
        log_word_step(
            "write_stage10",
            essay_id,
            stage="10",
            pack_path=pack_path,
            saved_pack=saved_pack_state,
            title=pack_log_title(pack),
            assessment=pack_log_assessment(pack),
            conclusion=str(pack.get("conclusion") or ""),
            cover=cover,
            sources_count=sources_count,
            refs_count=refs_count,
            refused=str(blocked.get("refused") or "refused"),
            reason=str(blocked.get("reason") or ""),
        )
        return body
    if pack.get("word_count_acknowledged") is not True:
        blockers = report.get("blockers")
        if isinstance(blockers, list):
            for item in blockers:
                if not isinstance(item, dict) or item.get("code") != "word-count":
                    continue
                reason = str(item.get("reason") or "")
                log_word_step(
                    "write_stage10",
                    essay_id,
                    stage="10",
                    pack_path=pack_path,
                    saved_pack=saved_pack_state,
                    title=pack_log_title(pack),
                    assessment=pack_log_assessment(pack),
                    conclusion=str(pack.get("conclusion") or ""),
                    cover=cover,
                    sources_count=sources_count,
                    refs_count=refs_count,
                    refused="word-count",
                    reason=reason,
                )
                body["ok"] = False
                body["copied"] = False
                body["sequence_written"] = False
                body["refused"] = "word-count"
                body["reason"] = reason
                return body
    from essay_writer.word_references import stage10_reference_block

    blocked_refs = stage10_reference_block(
        report.get("sources") or [],
        pack,
        list(report.get("sections") or []),
    )
    if blocked_refs:
        body["ok"] = False
        body["copied"] = False
        body["sequence_written"] = False
        body["refused"] = blocked_refs.get("refused")
        body["reason"] = blocked_refs.get("reason") or ""
        log_word_step(
            "write_stage10",
            essay_id,
            stage="10",
            pack_path=pack_path,
            saved_pack=saved_pack_state,
            title=pack_log_title(pack),
            assessment=pack_log_assessment(pack),
            conclusion=str(pack.get("conclusion") or ""),
            cover=cover,
            sources_count=sources_count,
            refs_count=refs_count,
            refused=str(blocked_refs.get("refused") or "empty-references"),
            reason=str(blocked_refs.get("reason") or ""),
        )
        return body
    pending = source.with_name(f".{source.name}.writing")
    try:
        write_docx(essay_id, pack, report, pending, product_log=False)
    except EmptyFooter:
        pending.unlink(missing_ok=True)
        return _footer_stop(
            body,
            refused="footer",
            reason="The page number is missing. An empty footer is refused.",
        )
    empty = empty_footer_refusal(pending)
    if empty:
        pending.unlink(missing_ok=True)
        return _footer_stop(
            body,
            refused=str(empty.get("refused") or "footer"),
            reason=str(empty.get("reason") or ""),
        )
    if _reference_lines_in_docx(pending) == 0:
        pending.unlink(missing_ok=True)
        reason = (
            "The References section has no entries. Add reference lines before "
            "this step can close."
        )
        log_word_step(
            "write_stage10",
            essay_id,
            stage="10",
            pack_path=pack_path,
            saved_pack=saved_pack_state,
            title=pack_log_title(pack),
            assessment=pack_log_assessment(pack),
            conclusion=str(pack.get("conclusion") or ""),
            cover=cover,
            sources_count=sources_count,
            refs_count=0,
            refused="empty-references",
            reason=reason,
        )
        body["ok"] = False
        body["copied"] = False
        body["sequence_written"] = False
        body["refused"] = "empty-references"
        body["reason"] = reason
        return body
    from essay_writer.harvard_ref_check import writer_reference_block

    blocked_harvard = writer_reference_block(
        essay_id,
        pending,
        list(report.get("sources") or []),
    )
    if blocked_harvard:
        pending.unlink(missing_ok=True)
        reason = str(blocked_harvard.get("reason") or "")
        log_word_step(
            "write_stage10",
            essay_id,
            stage="10",
            pack_path=pack_path,
            saved_pack=saved_pack_state,
            title=pack_log_title(pack),
            assessment=pack_log_assessment(pack),
            conclusion=str(pack.get("conclusion") or ""),
            cover=cover,
            sources_count=sources_count,
            refs_count=refs_count,
            refused="harvard-references",
            reason=reason,
        )
        body["ok"] = False
        body["copied"] = False
        body["sequence_written"] = False
        body["refused"] = "harvard-references"
        body["reason"] = reason
        body["misses"] = blocked_harvard.get("misses") or []
        body["checked"] = True
        body["handbook"] = blocked_harvard.get("handbook") or ""
        return body
    body["harvard_references"] = {
        "checked": True,
        "ok": True,
        "handbook": "docs/referencing-harvard.md",
    }
    source.parent.mkdir(parents=True, exist_ok=True)
    os.replace(pending, source)
    _refresh_sequence_final(source, project)
    from essay_writer.assignment_handoff import attach_word_chat
    from essay_writer.word_references import reference_lines_origin

    attach_word_chat(body, source)
    body["sequence_written"] = True
    refs_from = reference_lines_origin(pack)
    log_word_step(
        "write_stage10",
        essay_id,
        stage="10",
        pack_path=pack_path,
        saved_pack=saved_pack_state,
        title=pack_log_title(pack),
        assessment=pack_log_assessment(pack),
        conclusion=str(pack.get("conclusion") or ""),
        cover=cover,
        sources_count=sources_count,
        refs_count=refs_count,
        refs_from=refs_from,
        output_path=source,
    )
    body["references"] = report.get("references") or []
    body["cover"] = cover
    body["sources"] = report.get("sources") or []
    if root is None:
        from essay_writer.ml_events import note_stage10_rewrite_paths

        note_stage10_rewrite_paths(essay_id)
        _unmounted_documents(body)
        body["ok"] = True
        body["sequence_written"] = True
        return body
    from essay_writer.first_run import (
        ensure_sequence_mac_folder,
        sequence_mac_deliverables_dir,
    )

    seq_mac = sequence_mac_deliverables_dir(essay_id)
    if seq_mac is not None:
        ensure_sequence_mac_folder(essay_id)
        folder = seq_mac
    elif not root.is_dir():
        from essay_writer.ml_events import note_stage10_rewrite_paths

        note_stage10_rewrite_paths(essay_id)
        _unmounted_documents(body)
        body["ok"] = True
        body["sequence_written"] = True
        return body
    else:
        folder = root / f"{project}-10"
    folder.mkdir(parents=True, exist_ok=True)
    final = folder / f"{project}-final.docx"
    comments = folder / f"{project}-comments.docx"
    try:
        write_docx(essay_id, pack, report, final, product_log=False)
    except EmptyFooter:
        if final.is_file():
            final.unlink()
        return _footer_stop(
            body,
            refused="footer",
            reason="The page number is missing. An empty footer is refused.",
        )
    empty = empty_footer_refusal(final)
    if empty:
        if final.is_file():
            final.unlink()
        return _footer_stop(
            body,
            refused=str(empty.get("refused") or "footer"),
            reason=str(empty.get("reason") or ""),
        )
    from essay_writer.review import place_findings_note

    note = place_findings_note(essay_id, documents=root)
    body["findings_file"] = note
    if note.get("wrote"):
        body["findings_place"] = note.get("place") or ""
    _write_comments(essay_id, comments)
    from essay_writer.quarantine import clear_mac_quarantine

    clear_mac_quarantine(final)
    clear_mac_quarantine(comments)
    if note.get("wrote"):
        clear_mac_quarantine(root / project / f"{project}-findings.docx")
    body["copied"] = True
    body["reason"] = ""
    body["sequence_written"] = True
    from essay_writer.ml_events import note_stage10_rewrite_paths

    note_stage10_rewrite_paths(essay_id)
    return body


def _write_comments(essay_id: str, path: Path) -> None:
    from docx import Document

    from essay_writer.specialists.eqr import eqr_record, eqr_verdict
    from essay_writer.word import _font, _style_document

    document = Document()
    _style_document(document)
    stored = eqr_record(essay_id, "9") or {}
    findings_raw = stored.get("findings")
    findings = findings_raw if isinstance(findings_raw, str) else ""
    if stored.get("word_opened") and findings.strip():
        review = findings.strip()
    else:
        verdict = eqr_verdict(essay_id, "9")
        if verdict == "PASS":
            review = "The independent review recorded a pass."
        elif verdict == "FAIL":
            review = "The independent review recorded a fail."
        else:
            review = "The independent review is not recorded."
    record = load_stage7(essay_id)
    notes = comment_text(record)
    for paragraph_text in (review, notes):
        if not paragraph_text.strip():
            continue
        paragraph = document.add_paragraph()
        run = paragraph.add_run(paragraph_text)
        _font(run, 12)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def stage6_fields(essay_id: str, *, master: str | None, named_folder: str | None) -> dict[str, Any]:
    report = master_word(essay_id, master=master, named_folder=named_folder)
    return {
        "place": report["place"],
        "master_ask": report["master_ask"],
        "choices": report["choices"],
        "expected_draft": report["expected"],
        "block_confirm": report["block_confirm"],
        "chain_started": False,
    }


def stage7_fields(essay_id: str) -> dict[str, Any]:
    record = load_stage7(essay_id)
    return {
        "questions": record.get("questions") or [],
        "draft_read": bool(record.get("asked")),
        "file_work": "chat",
        "delta": record.get("delta") or [],
        "chain_started": False,
    }


def finish_stage_reply(
    body: dict[str, Any],
    essay_id: str,
    stage: str,
    *,
    armed: bool,
) -> dict[str, Any]:
    """One stage only. Quiet mid steps once the chain is armed. Stage 10 names the final Word."""

    body["chain_started"] = False
    body["chain_armed"] = armed
    body["executed_runs"] = [stage] if body.get("stored") else []
    if stage == "9":
        from essay_writer.specialists.grok_bot import front_office_stage9

        stop = front_office_stage9(essay_id)
        body["remarked"] = False
        body["freeflow"] = False
        body["gap_fill"] = False
        body["fo_stop"] = stop["stopped"]
        body["submit_pass"] = stop.get("reason") or ""
    if stage == "7":
        for key, value in stage7_fields(essay_id).items():
            body.setdefault(key, value)
    if stage == "6":
        body["chain_started"] = False
        body.pop("next_run", None)
    if body.get("stored") and armed and stage in {"7", "8", "9"}:
        run = next_run(essay_id, stage, armed=True)
        if run:
            body["next_run"] = run
            blocked = bool(run.get("blocked"))
            line = quiet_next(stage, armed=True, blocked=blocked, place=final_place(essay_id))
            if line:
                body["next"] = say(line)
            if blocked:
                body["blocked"] = True
                body["eqr"] = run.get("eqr") or "missing"
    elif armed and stage in {"8", "9"} and not body.get("stored"):
        line = quiet_next(stage, armed=True, blocked=False, place=final_place(essay_id))
        if line:
            body["next"] = say(line)
    if stage == "10" and body.get("stored"):
        final = body.get("final_word")
        if not isinstance(final, dict) or "copied" not in final:
            final = write_stage10(essay_id)
        if isinstance(final, dict):
            final.pop("sequence_written", None)
        if final.get("reason") == "mac-not-mounted":
            final = dict(final)
            final["place"] = ""
            final["comments_place"] = ""
            final["delivery"] = "chat"
            body["place"] = ""
            body["next"] = say(STAGE10_UNMOUNTED)
        else:
            body["place"] = final["place"]
            findings = final.get("findings_place") or ""
            note = f" The review note is {findings}." if findings else ""
            body["next"] = say(
                f"Open {final['place']} and keep a copy.{note} This tool does not submit."
            )
        body["final_word"] = final
        delivery = final.get("delivery") if isinstance(final, dict) else None
        if isinstance(delivery, dict):
            body["delivery"] = delivery
    elif stage == "10" and armed and not body.get("stored"):
        from essay_writer.spine import final_review

        blocked_text, verdict = final_review(essay_id)
        if blocked_text:
            body["blocked"] = True
            body["next"] = say(blocked_text)
            body["eqr"] = verdict
    if stage == "9" and body.get("stored"):
        from essay_writer.specialists.grok_bot import STAGE9_OWN_WORDS

        next_raw = body.get("next")
        current = next_raw if isinstance(next_raw, str) else ""
        if STAGE9_OWN_WORDS not in current:
            body["next"] = say(f"{STAGE9_OWN_WORDS} {current}".strip())
    from essay_writer.cite_fit import in_line, speak_cite_ask

    speak_cite_ask(body)
    in_line(body, "student_outline_line")
    return body
