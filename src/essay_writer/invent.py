"""invent-forbid. A bibliographic claim may appear only when the gathered
text of this essay contains it. Sample reference-list lines are not sources.
"""

from __future__ import annotations

from typing import Any

from docx import Document

from essay_writer.canon import file_sha256
from essay_writer.cite_open import (
    author_kept,
    open_for_citation,
    page_span_printed,
    page_value_kind,
    pdf_from_source,
    printed_pages_for,
    quote_refusal,
)
from essay_writer.paths import TEXT_SUFFIXES, locator_target
from essay_writer.programme import PAID_RESOLVERS
from essay_writer.resolvers import RESOLVER_FILLABLE
from essay_writer.state import load_state
from essay_writer.vision import is_image_suffix, saved_photo_words

BIBLIO_FIELDS = (
    "author",
    "title",
    "year",
    "container",
    "volume",
    "issue",
    "pages",
    "doi",
    "url",
    "accessed",
    "publisher",
    "place",
    "edition",
)
CLAIM_FIELDS = BIBLIO_FIELDS + ("quote",)
FILE_BIBLIO_FIELDS = ("author", "title", "year", "edition")


def read_gathered_text(path: Any, essay_id: str | None = None) -> str:
    """Text of a gathered file. A PDF is opened, including one larger than 2MB."""
    if essay_id:
        from essay_writer.known_paths import research_file

        allowed = research_file(essay_id, str(path))
        if allowed is None:
            raise ValueError("Do not invent a student folder path.")
        path = allowed
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return open_for_citation(path, essay_id).text
    if suffix == ".docx":
        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    chunks.append(cell.text)
        return "\n".join(chunks)
    if is_image_suffix(suffix):
        return saved_photo_words(path)
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8")
    raise ValueError(
        f"cannot verify citations against {suffix or 'that file'}; bring a text or Word original"
    )


def _as_text(value: object) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, str) and value.strip():
        return value.strip()
    return None


def _paid(raw: dict[str, Any]) -> str | None:
    for key in ("resolver", "via", "provider", "api", "database"):
        text = _as_text(raw.get(key))
        if not text:
            continue
        lowered = text.casefold()
        for name in PAID_RESOLVERS:
            if name in lowered:
                return name
    return None


_LOOKUP_KEPT = (
    "doi",
    "volume",
    "issue",
    "pages",
    "url",
    "publisher",
    "place",
    "container",
    "accessed",
)


def _lookup_kept(essay_id: str, locator: str) -> dict[str, str]:
    """Fields already marked lookup on the stored source.

    The stored research record is the provenance. A pack that repeats those
    values is not the writer typing them. A different value still has to be
    in the gathered file.
    """
    from essay_writer.validate import research_index

    allowed, _digests = research_index(load_state(essay_id))
    prior = allowed.get(locator.strip().replace("\\", "/"))
    if not isinstance(prior, dict):
        return {}
    provenance = prior.get("field_sources")
    if not isinstance(provenance, dict):
        return {}
    kept: dict[str, str] = {}
    for field in _LOOKUP_KEPT:
        if provenance.get(field) != "lookup":
            continue
        value = prior.get(field)
        if isinstance(value, str) and value.strip():
            kept[field] = value.strip()
    return kept


def gather_source(essay_id: str, raw: object, index: int) -> tuple[dict[str, Any] | None, list[dict[str, str]]]:
    label = f"sources[{index}]"
    if not isinstance(raw, dict):
        return None, [{"code": "invent-forbid", "reason": f"{label} must be an object"}]
    origin = raw.get("origin")
    if origin in {"sample", "samples", "reference-list"}:
        return None, [
            {
                "code": "invent-forbid",
                "reason": f"{label} uses a sample reference list. Those lines were not gathered.",
            }
        ]
    paid = _paid(raw)
    if paid:
        return None, [
            {
                "code": "invent-forbid",
                "reason": f"{label} names the paid resolver {paid}. This door uses free reading of gathered files only.",
            }
        ]
    locator = raw.get("locator")
    if not isinstance(locator, str) or not locator.strip():
        claimed = [field for field in CLAIM_FIELDS if _as_text(raw.get(field))]
        if claimed:
            return None, [
                {
                    "code": "invent-forbid",
                    "reason": (
                        f"{label} has {', '.join(claimed)} but no gathered locator. "
                        "Do not invent authors, titles, years, quotations, pages, or DOIs."
                    ),
                }
            ]
        return None, [{"code": "invent-forbid", "reason": f"{label} has no gathered locator"}]
    try:
        kind, path = locator_target(essay_id, locator)
    except ValueError as exc:
        reason = str(exc)
        if reason == "cross-essay":
            return None, [
                {
                    "code": "cross-essay",
                    "reason": f"{label} points at a different essay id. Briefs and uploads do not cross.",
                }
            ]
        if reason == "sample-reference":
            return None, [
                {
                    "code": "invent-forbid",
                    "reason": f"{label} points at the samples study. A reference-list line there is not a gathered source.",
                }
            ]
        return None, [{"code": "invent-forbid", "reason": f"{label}: {reason}"}]
    if origin not in {"upload", "brief", "kb1"}:
        return None, [
            {
                "code": "invent-forbid",
                "reason": f"{label} origin must be upload, brief, or kb1.",
            }
        ]
    if origin == "kb1" and kind not in {"kb1", "kb1-file"}:
        return None, [{"code": "invent-forbid", "reason": f"{label} origin kb1 does not match the locator"}]
    if origin == "upload" and kind != "upload":
        return None, [{"code": "invent-forbid", "reason": f"{label} origin upload does not match the locator"}]
    if origin == "brief" and kind != "brief":
        return None, [{"code": "invent-forbid", "reason": f"{label} origin brief does not match the locator"}]

    if kind == "kb1-file":
        assert path is not None
        if not path.is_file():
            return None, [
                {
                    "code": "invent-forbid",
                    "reason": f"{label} KB1 note is not a file on disk: {locator}",
                }
            ]
        try:
            text = read_gathered_text(path, essay_id)
            sha = file_sha256(path)
        except ValueError as exc:
            return None, [{"code": "invent-forbid", "reason": f"{label}: {exc}"}]
    elif kind == "kb1":
        return None, [
            {
                "code": "invent-forbid",
                "reason": (
                    f"{label} is a KB1 note without a file. "
                    "An excerpt is not a gathered source. Leave it out rather than inventing it."
                ),
            }
        ]
    else:
        assert path is not None
        if not path.is_file():
            return None, [
                {
                    "code": "invent-forbid",
                    "reason": f"{label} locator is not a file on disk: {locator}",
                }
            ]
        try:
            text = read_gathered_text(path, essay_id)
            sha = file_sha256(path)
        except ValueError as exc:
            return None, [{"code": "invent-forbid", "reason": f"{label}: {exc}"}]

    opened = open_for_citation(path, essay_id) if path.suffix.lower() == ".pdf" else None
    if opened is not None:
        text = opened.text
    issues: list[dict[str, str]] = []
    kept: dict[str, Any] = {
        "id": _as_text(raw.get("id")) or f"src-{index + 1}",
        "origin": origin,
        "locator": locator.strip().replace("\\", "/"),
        "gathered": True,
        "sha256": sha,
    }
    folded = text.casefold()
    field_sources: dict[str, str] = {}
    lookup_kept = _lookup_kept(essay_id, locator)
    for field in BIBLIO_FIELDS:
        text_value = _as_text(raw.get(field))
        if text_value is None:
            continue
        if field == "author" and opened is not None:
            corrected = author_kept(text_value, opened)
            if corrected:
                kept["author"] = corrected
                field_sources[field] = "file"
                continue
        if field == "pages":
            kind = page_value_kind(text_value)
            from_lookup = lookup_kept.get("pages") == text_value
            if kind == "article" and from_lookup:
                kept["article_number"] = text_value
                field_sources["article_number"] = "lookup"
                continue
            if opened is not None:
                printed = page_span_printed(text_value, opened.printed_pages)
                if kind in {"range", "page"} and (from_lookup or printed):
                    # A matching look-up keeps its page, including a lone
                    # page. A page the file does not print and no look-up
                    # marks stays blank rather than being kept on a guess.
                    kept["pages"] = text_value
                    field_sources[field] = "lookup" if from_lookup else "file"
                    continue
                if printed:
                    kept["pages"] = text_value
                    field_sources[field] = "file"
                continue
        if field in lookup_kept and text_value == lookup_kept[field]:
            kept[field] = text_value
            field_sources[field] = "lookup"
            continue
        if text_value.casefold() not in folded:
            issues.append(
                {
                    "code": "invent-forbid",
                    "reason": (
                        f"{label} {field} {text_value!r} is not in the gathered text. "
                        "Leave it out rather than inventing it."
                    ),
                }
            )
            continue
        kept[field] = text_value
        if field in FILE_BIBLIO_FIELDS or field in RESOLVER_FILLABLE:
            field_sources[field] = "file"
    if opened is not None and opened.article_number:
        claimed_article = _as_text(raw.get("article_number"))
        if claimed_article is None or claimed_article == opened.article_number:
            kept["article_number"] = opened.article_number
        elif claimed_article.casefold() not in folded:
            issues.append(
                {
                    "code": "invent-forbid",
                    "reason": (
                        f"{label} article number {claimed_article!r} is not in the gathered text. "
                        "Leave it out rather than inventing it."
                    ),
                }
            )
    quote = _as_text(raw.get("quote"))
    if quote is not None:
        missing = quote_refusal(quote not in text, opened)
        if missing:
            issues.append({"code": "invent-forbid", "reason": f"{label} {missing}"})
        else:
            kept["quote"] = quote
    if field_sources:
        kept["field_sources"] = field_sources
    if issues:
        return None, issues
    return kept, []


def gather_sources(
    essay_id: str, sources: object
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    if sources is None:
        return [], []
    if not isinstance(sources, list):
        return [], [{"code": "invent-forbid", "reason": "sources must be a list"}]
    kept: list[dict[str, Any]] = []
    issues: list[dict[str, str]] = []
    seen: set[str] = set()
    for index, raw in enumerate(sources):
        source, source_issues = gather_source(essay_id, raw, index)
        issues.extend(source_issues)
        if source is None:
            continue
        if source["id"] in seen:
            issues.append(
                {"code": "invent-forbid", "reason": f"duplicate source id {source['id']}"}
            )
            continue
        seen.add(source["id"])
        kept.append(source)
    return kept, issues


def _ungathered_reason(source: dict[str, Any], essay_id: str | None = None) -> str | None:
    """Why this source cannot support a claim. None when the file was gathered."""
    origin = source.get("origin")
    if origin in {"sample", "samples", "reference-list"}:
        return "uses a sample reference list. Those lines were not gathered."
    paid = _paid(source)
    if paid:
        return f"names the paid resolver {paid}."
    locator = source.get("locator")
    if not isinstance(locator, str) or not locator.strip():
        return "has no gathered locator."
    parts = locator.strip().replace("\\", "/").split("/")
    if len(parts) < 2 or not parts[1]:
        return "has no gathered locator."
    essay = essay_id.strip() if isinstance(essay_id, str) and essay_id.strip() else parts[1]
    try:
        kind, path = locator_target(essay, locator)
    except ValueError as exc:
        return str(exc)
    if origin not in {"upload", "brief", "kb1"}:
        return "origin must be upload, brief, or kb1."
    if origin == "upload" and kind != "upload":
        return "origin upload does not match the locator."
    if origin == "brief" and kind != "brief":
        return "origin brief does not match the locator."
    if origin == "kb1" and kind not in {"kb1", "kb1-file"}:
        return "origin kb1 does not match the locator."
    if kind == "kb1" or path is None or not path.is_file():
        if origin == "kb1" or kind == "kb1-file":
            return "KB1 note is not a file on disk."
        return "locator is not a file on disk."
    return None


def check_claims(
    sources: list[dict[str, Any]],
    claims: object,
    texts: dict[str, str],
    essay_id: str | None = None,
) -> list[dict[str, str]]:
    if claims is None:
        return []
    if not isinstance(claims, list):
        return [{"code": "invent-forbid", "reason": "claims must be a list"}]
    by_id = {source["id"]: source for source in sources}
    issues: list[dict[str, str]] = []
    for index, claim in enumerate(claims):
        label = f"claims[{index}]"
        if not isinstance(claim, dict):
            issues.append({"code": "invent-forbid", "reason": f"{label} must be an object"})
            continue
        source = by_id.get(claim.get("source_id"))
        if source is None:
            issues.append(
                {
                    "code": "invent-forbid",
                    "reason": f"{label} cites a source that was not gathered.",
                }
            )
            continue
        refused = _ungathered_reason(source, essay_id)
        if refused:
            code = "cross-essay" if refused == "cross-essay" else "invent-forbid"
            issues.append({"code": code, "reason": f"{label} {refused}"})
            continue
        text = texts.get(source["locator"], "")
        path = pdf_from_source(source)
        opened = open_for_citation(path) if path is not None else None
        quote = claim.get("quote")
        if opened is not None:
            quote_missing = not isinstance(quote, str) or quote not in opened.text
        else:
            quote_missing = not isinstance(quote, str) or quote not in text
        missing = quote_refusal(quote_missing, opened)
        if missing:
            issues.append({"code": "invent-forbid", "reason": f"{label} {missing}"})
        page = claim.get("page")
        if page is not None:
            page_text = str(page).strip()
            printed = printed_pages_for(source)
            if printed is not None:
                allowed = bool(page_text) and page_span_printed(page_text, printed)
            else:
                allowed = bool(page_text) and page_text.casefold() in text.casefold()
            if not allowed:
                reason = (
                    f"{label} page is not printed in the opened file. Omit the page."
                    if printed is not None
                    else f"{label} page is not in the gathered text. Omit the page."
                )
                issues.append({"code": "invent-forbid", "reason": reason})
    return issues


def load_text_for(source: dict[str, Any]) -> str:
    module = source["locator"].split("/")[1]
    _kind, path = locator_target(module, source["locator"])
    if path is None:
        return str(source.get("excerpt") or "")
    return read_gathered_text(path, module)
