"""EQR specialist. One review door.

The door still records PASS or FAIL for stage 6 or 9. When the finished Word
file is already saved, this door opens that file, checks it against gathered
text, and keeps a findings note. When a mark scheme is present, a pass aims
at the highest band. That rule is HIGHEST_BAND, applied by assignment_gate
before a verdict is stored. Stage 9 also marks the top-band table cells for
argument, evidence, and critical thinking. Format and citations stay with
the writer. A pass is refused when the opened Word file misses a word band,
a cover field, or a diary heading already parsed from the brief, the marking
table, or the diary. An ask that was not parsed is not added. There is not
a second review function.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from essay_writer import assignment_gate as essay_assignment_gate
from essay_writer.assignment_gate import (
    _RANGE,
    _RANK,
    _content_words,
    _covers,
    _numeric_band,
    assignment_gate,
    mark_scheme_text,
)
from essay_writer.confirm import refuse
from essay_writer.diary import cover_residue_line, title_page_cover_lines
from essay_writer.known_paths import research_file
from essay_writer.paths import module_code, work_dir
from essay_writer.state import load_state
from essay_writer.textutil import citation_tokens, count_words

ROLE = "eqr"

HIGHEST_BAND = (
    "When a mark scheme is present, the review judges the highest band, "
    "the top descriptors, not merely a pass."
)

# Stage prompt for the independent review. No house labels.
REVIEW_PROMPT = """
You are reviewing one essay sequence. You are not the writer.
Do not draft the essay and do not gather sources.
When the finished Word file is already saved, open that file. The review note is not a new essay.
Record the independent review for stage 6 before the first Word file, or stage 9 before the final Word file.
Stage 6 answers only whether the draft answers the assignment question. It does not check format. It does not rewrite the draft. It does not fill gaps. Format and citations stay with the writer.
Stage 9 checks the same requirement list on the opened Word file. The constitution is not that check. A pass with no opened Word file does not clear. Stage 9 refuses a pass when the opened Word file misses a word band, a cover field, or a diary heading already recorded from the brief, the marking table, or the diary. Nothing is added when those documents did not record that ask.
Stage 9 marks a mark scheme table row for argument, evidence, or critical thinking at the top band. It does not mark format. It does not mark citations. Stage 6 does not mark those rows.
Say whether the essay is ready for that Word file, or whether it needs more work. Draft, show it, and approve the save only after an explicit yes.
Do not invent sources in the review note.
When a mark scheme is present, judge the highest band, the top descriptors, not merely a pass.
The findings note opens with the north-star judgement: what the mark scheme scores, the format, and whether the draft does the assignment. Claim checks against the opened PDFs follow that section. A pass with no findings note does not clear the Word file.
When that Word file also has a copy under Documents, open the copy as well. The findings note names each place that was opened. The findings Word file is Documents/<project>/<project>-findings.docx, in the project folder, not the stage 10 folder. When Documents is not mounted, do not invent that file.
Review facts, structure, citation completeness, and what the brief asks. Do not polish voice, elegance, or parallelism. Keep the student's voice and the writer's voice.
The findings note may name a leftover stock line, or a line that could sit on any assignment. Naming that line is not a score. Do not rewrite those lines. Naming them does not fail the review.
The findings note may name a weighted section the draft does not carry, and a heading the brief does not weight. Naming either is not a score and does not fail the review by itself. Do not invent a band descriptor, a criterion, or a weight the brief or the mark scheme does not print.
""".strip()

REVIEW_LIMIT = (
    "Review facts, structure, citation completeness, and what the brief asks. "
    "Do not polish voice, elegance, or parallelism. "
    "Keep the student's voice and the writer's voice."
)


def review_limit_sha() -> str:
    """Stable id for the reviewer voice limit. PR bodies cite this."""
    import hashlib

    return hashlib.sha256(REVIEW_LIMIT.encode("utf-8")).hexdigest()


def finished_word_path(essay_id: str):
    """The Word file the student receives for this sequence, when it has been saved."""

    return work_dir(essay_id) / "output" / f"{essay_id}.docx"


def _read_docx(path) -> dict[str, Any]:
    """Open one Word file. A missing file is not invented."""
    import hashlib
    from zipfile import BadZipFile

    from docx.opc.exceptions import PackageNotFoundError

    if not path.is_file():
        return {
            "exists": False,
            "opened": False,
            "text": "",
            "paragraphs": 0,
            "sha256": None,
            "excerpt": "",
            "unreadable": False,
        }
    try:
        from docx import Document

        document = Document(str(path))
        texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError):
        return {
            "exists": True,
            "opened": False,
            "text": "",
            "paragraphs": 0,
            "sha256": None,
            "excerpt": "",
            "unreadable": True,
        }
    return {
        "exists": True,
        "opened": True,
        "text": "\n".join(texts),
        "paragraphs": len(texts),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "excerpt": texts[0][:240] if texts else "",
        "unreadable": False,
    }


def open_finished_word(essay_id: str, *, documents=None) -> dict[str, Any]:
    """Open the Word file the student receives, and the Documents copy when it exists."""
    from essay_writer.user_place import finished_word_copies

    absent = {
        "exists": False,
        "opened": False,
        "text": "",
        "paragraphs": 0,
        "sha256": None,
        "excerpt": "",
        "places": [],
        "copies": [],
    }
    copies = finished_word_copies(essay_id, documents=documents)
    if not copies:
        return absent
    opened_copies: list[dict[str, Any]] = []
    for copy in copies:
        read = _read_docx(copy["path"])
        if read.get("unreadable"):
            return {
                "exists": True,
                "opened": False,
                "text": "",
                "paragraphs": 0,
                "sha256": None,
                "excerpt": "",
                "unreadable": True,
                "places": [],
                "copies": [],
            }
        if not read.get("opened"):
            continue
        opened_copies.append(
            {
                "role": copy["role"],
                "place": copy["place"],
                "text": read["text"],
                "paragraphs": read["paragraphs"],
                "sha256": read["sha256"],
                "excerpt": read["excerpt"],
            }
        )
    if not opened_copies:
        return absent
    primary = next((item for item in opened_copies if item["role"] == "student"), opened_copies[0])
    return {
        "exists": True,
        "opened": True,
        "text": "\n".join(item["text"] for item in opened_copies),
        "paragraphs": primary["paragraphs"],
        "sha256": primary["sha256"],
        "excerpt": primary["excerpt"],
        "unreadable": False,
        "places": [item["place"] for item in opened_copies],
        "copies": opened_copies,
    }


def claim_check_opened(essay_id: str, text: str) -> dict[str, Any]:
    """Compare the opened Word text with gathered files already stored. No new look-up."""
    if not text.strip():
        return {"ran": False, "matches": []}
    from essay_writer.gates import unquoted_exact
    from essay_writer.invent import read_gathered_text

    gathered: list[str] = []
    pdfs: list[str] = []
    state = load_state(essay_id)
    research = state.get("research") if isinstance(state.get("research"), dict) else {}
    for depth in ("summary", "full"):
        record = research.get(depth) if isinstance(research, dict) else None
        sources = record.get("sources") if isinstance(record, dict) else None
        if not isinstance(sources, list):
            continue
        for source in sources:
            if not isinstance(source, dict):
                continue
            locator = source.get("locator") or source.get("path")
            if not isinstance(locator, str) or not locator.strip():
                continue
            path = research_file(essay_id, locator.strip())
            if path is None or not path.is_file():
                continue
            try:
                gathered.append(read_gathered_text(path, essay_id))
            except (OSError, ValueError):
                continue
            if path.suffix.lower() == ".pdf":
                pdfs.append(path.name)
    if not gathered:
        return {"ran": False, "matches": [], "pdfs": pdfs}
    return {"ran": True, "matches": unquoted_exact(text, gathered), "pdfs": pdfs}


NORTH_STAR = "North-star judgement"
CLAIM_CHECKS = "Claim checks"
BOTH_PLACES = "Opened from both places."


def _draft_does_assignment(record: dict[str, Any], opened_text: str) -> str:

    question = record.get("question") if isinstance(record.get("question"), str) else ""
    if not opened_text.strip():
        if question:
            return (
                "The draft is not saved, so it does not do the assignment yet. "
                f"Question: {question}"
            )
        return "The draft is not saved, so it does not do the assignment yet."
    does = True
    named = record.get("format") if isinstance(record.get("format"), str) else ""
    if named and named not in opened_text.casefold():
        does = False
    if question and not _covers(opened_text, _content_words(question)):
        does = False
    if question:
        if does:
            return f"The draft does the assignment. Question: {question}"
        return f"The draft does not do the assignment. Question: {question}"
    if named and does:
        return "The draft does the assignment."
    if named:
        return "The draft does not do the assignment."
    return "No assignment question is on the brief, so the draft is not judged against one."


def build_findings(essay_id: str, opened: dict[str, Any], claim: dict[str, Any]) -> str:
    """North-star judgement first, then claim checks against opened PDFs."""
    from essay_writer.specialists.grok_bot import basics_pass_lines

    record = essay_assignment_gate.read_assignment(essay_id)
    scores, format_line, _question_line = basics_pass_lines(essay_id)
    lines: list[str] = [
        NORTH_STAR,
        scores,
        format_line,
        _draft_does_assignment(record, opened.get("text") or ""),
        "",
        CLAIM_CHECKS,
        "Independent review.",
        "Claim checks against the opened PDFs.",
    ]
    copies_raw = opened.get("copies")
    copies: list[object] = copies_raw if isinstance(copies_raw, list) else []
    student_opened = any(isinstance(item, dict) and item.get("role") == "student" for item in copies)
    if opened.get("opened") and (student_opened or not copies):
        lines.append(f"Opened the saved Word file ({opened['paragraphs']} paragraphs).")
        if opened.get("excerpt"):
            lines.append(str(opened["excerpt"]))
    elif opened.get("unreadable"):
        lines.append("The saved Word file could not be opened.")
    elif not opened.get("opened"):
        lines.append("The finished Word file is not saved yet.")
    for name in claim.get("pdfs") or []:
        lines.append(f"Opened PDF {name}.")
    if claim.get("ran"):
        if claim.get("matches"):
            lines.append("A gathered sentence appears in the Word file without quotation marks.")
        else:
            lines.append("No unquoted gathered sentence was found in the opened Word file.")
    return with_opened_places("\n".join(lines), opened)


def _shared_opening_shape(line: str) -> bool:
    """Shapes standard_opening can print when the assignment id is not at hand."""
    folded = line.casefold()
    if folded.startswith("how to score high:"):
        return True
    return "mark scheme" in folded


def findings_note_clears(findings: object, essay_id: str | None = None) -> bool:
    """A PASS clears Word only when the note opens on the shared standard line.

    With an assignment id, that line is standard_opening for this record:
    How to score high and the printed column for a diary table, what the
    mark scheme scores for an essay or a report, and the no-scheme line
    when there is none. A later type uses that same function. The format
    line, the assignment line, and the claim checks stay required.
    A one-line verdict, or a pass with no findings note, does not clear.
    """
    if not isinstance(findings, str):
        return False
    lines = [line.strip() for line in findings.splitlines() if line.strip()]
    if not lines or lines[0] != NORTH_STAR:
        return False
    if CLAIM_CHECKS not in lines:
        return False
    body: list[str] = []
    for line in lines[1:]:
        if line == CLAIM_CHECKS:
            break
        body.append(line)
    if len(body) < 3:
        return False
    opening = body[0]
    if isinstance(essay_id, str) and essay_id.strip():
        from essay_writer.specialists.grok_bot import standard_opening

        if opening != standard_opening(essay_id):
            return False
    elif not _shared_opening_shape(opening):
        return False
    if not body[1].casefold().startswith("format:"):
        return False
    if "assignment" not in body[2].casefold():
        return False
    return True


def with_opened_places(findings: str, opened: dict[str, Any]) -> str:
    """Name each opened place. When the draft was copied, say both were opened."""
    copies_raw = opened.get("copies")
    copies: list[object] = copies_raw if isinstance(copies_raw, list) else []
    documents = [
        item
        for item in copies
        if isinstance(item, dict) and item.get("role") == "documents" and item.get("place")
    ]
    roles = {item.get("role") for item in copies if isinstance(item, dict)}
    if not documents:
        return findings
    extras: list[str] = []
    for item in documents:
        place = str(item["place"])
        line = f"Opened {place}."
        if line not in findings:
            extras.append(line)
        excerpt = item.get("excerpt") if isinstance(item.get("excerpt"), str) else ""
        if excerpt and excerpt not in findings:
            extras.append(excerpt)
    if "student" in roles and "documents" in roles and BOTH_PLACES not in findings:
        extras.append(BOTH_PLACES)
    if not extras:
        return findings
    lines = findings.splitlines()
    insert_at = len(lines)
    for index, line in enumerate(lines):
        if line.startswith("Opened the saved Word file"):
            insert_at = index + 1
            break
    excerpt = opened.get("excerpt") if isinstance(opened.get("excerpt"), str) else ""
    if excerpt and insert_at < len(lines) and lines[insert_at] == excerpt:
        insert_at += 1
    for offset, extra in enumerate(extras):
        lines.insert(insert_at + offset, extra)
    return "\n".join(lines)


def _findings_refused(*, chat_docx: dict[str, Any] | None = None, mac_copy: dict[str, Any] | None = None) -> dict[str, Any]:
    body = {
        "ok": False,
        "wrote": False,
        "place": "",
        "reason": "mac-not-mounted",
        "refused": "documents",
        "delivery": "chat",
    }
    if chat_docx is not None:
        body["chat_docx"] = chat_docx
    if mac_copy is not None:
        body["mac_copy"] = mac_copy
    return body


def _save_findings_docx(path, findings: str) -> None:
    from docx import Document

    from essay_writer.word import _add_heading, _font, _style_document

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _style_document(document)
    _add_heading(document, "Findings", level=1)
    for line in findings.splitlines():
        if not line.strip():
            continue
        if line.strip() in {NORTH_STAR, CLAIM_CHECKS}:
            _add_heading(document, line.strip(), level=2)
            continue
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line.strip())
        _font(run, 12)
    document.save(str(path))


def write_findings_note(essay_id: str, findings: str, *, documents=None) -> dict[str, Any]:
    """Write the findings Word file for chat, Mac, and Documents when mounted.

    The sequence copy under output/ is always written. Documents/<project>/ is
    written when that folder is visible. The stage 10 folder is not this file.
    """
    from essay_writer.first_run import copy_deliverable_to_mac
    from essay_writer.late_chain import documents_root, findings_place
    from essay_writer.research_packs import chat_docx

    project = module_code(essay_id)
    place = findings_place(essay_id)
    filename = f"{project}-findings.docx"
    sequence_path = work_dir(essay_id) / "output" / filename
    _save_findings_docx(sequence_path, findings)
    chat = chat_docx(sequence_path)
    mac_copy = copy_deliverable_to_mac(sequence_path)
    root = documents_root(documents)
    if root is None or (root.exists() and not root.is_dir()):
        return _findings_refused(chat_docx=chat, mac_copy=mac_copy)
    folder = root / project
    folder.mkdir(parents=True, exist_ok=True)
    documents_path = folder / filename
    _save_findings_docx(documents_path, findings)
    if mac_copy.get("copied") is not True:
        mac_copy = copy_deliverable_to_mac(sequence_path)
    return {
        "ok": True,
        "wrote": True,
        "place": place,
        "reason": "",
        "chat_docx": chat,
        "mac_copy": mac_copy,
    }


def stage6_scope() -> dict[str, Any]:
    """Stage 6 judges the assignment question. Format, rewrite, and gap-fill stay off."""
    return {
        "judges": "answers-the-question",
        "format": False,
        "freeflow": False,
        "gap_fill": False,
    }


def answers_assignment_question(essay_id: str, text: str) -> bool:
    """True when the opened text answers the assignment question.

    No question on the brief leaves this check quiet. The named shape is not read.
    """

    record = essay_assignment_gate.read_assignment(essay_id)
    question_raw = record.get("question")
    question = question_raw if isinstance(question_raw, str) else ""
    if not question.strip():
        return True
    body_text = text if isinstance(text, str) else ""
    return _covers(body_text, _content_words(question))


def submit_word_path(essay_id: str):
    """The sequence Word file a later pass has to match."""

    return work_dir(essay_id) / "output" / f"{essay_id}.docx"


def read_submit_word(essay_id: str) -> dict[str, Any]:
    """Open the submit Word file and say whether it already has citations."""

    read = _read_docx(submit_word_path(essay_id))
    text_raw = read.get("text")
    text = text_raw if isinstance(text_raw, str) else ""
    read["citations"] = bool(citation_tokens(text))
    return read


def _locked(reason: str, eqr: str = "missing") -> dict[str, Any]:
    return {
        "unlocked": False,
        "reason": reason,
        "eqr": eqr,
        "remarked": False,
        "freeflow": False,
        "gap_fill": False,
    }


def submit_pass_unlocked(essay_id: str) -> dict[str, Any]:
    """Whether a stage 9 pass clears the current submit Word file.

    A pass with no file, or a pass whose file hash no longer matches, does not
    clear. Citations have to be on that file. This does not mark again.
    """
    from essay_writer.review import eqr_record

    record = eqr_record(essay_id, "9")
    if not isinstance(record, dict):
        return _locked("missing")
    verdict = record.get("verdict")
    if verdict == "FAIL":
        return _locked("fail", "FAIL")
    if verdict != "PASS":
        return _locked("missing")
    if not findings_note_clears(record.get("findings"), essay_id):
        return _locked("missing")
    stored = record.get("submit_sha256")
    if record.get("word_opened") is not True or not isinstance(stored, str) or not stored:
        return _locked("false")
    current = read_submit_word(essay_id)
    if current.get("opened") is not True or current.get("sha256") != stored:
        return _locked("stale")
    if current.get("citations") is not True:
        return _locked("citations")
    return {
        "unlocked": True,
        "reason": "pass",
        "eqr": "PASS",
        "remarked": False,
        "freeflow": False,
        "gap_fill": False,
    }


_ROW_ORDER = ("argument", "evidence", "critical thinking")
_SKIP_ORDER = ("format", "citations")
_ROW_MATCH = (
    ("critical thinking", re.compile(r"(?i)\bcritical\s+thinking\b")),
    ("evidence", re.compile(r"(?i)\bevidence\b")),
    ("argument", re.compile(r"(?i)\barguments?\b")),
)
_WRITER_OWNED = (
    ("format", re.compile(r"(?i)\b(?:formats?|formatting|presentation|layout)\b")),
    ("citations", re.compile(r"(?i)\b(?:citations?|referenc\w*|harvard)\b")),
)


def _clean_cell(cell: str) -> str:
    text = cell.strip().strip("*").strip()
    return re.sub(r"\s+", " ", text)


def _table_rows(text: str) -> list[list[list[str]]]:
    tables: list[list[list[str]]] = []
    current: list[list[str]] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and stripped.count("|") >= 2:
            cells = [_clean_cell(cell) for cell in stripped.strip("|").split("|")]
            current.append(cells)
            continue
        if current:
            tables.append(current)
            current = []
    if current:
        tables.append(current)
    return tables


def _rule_row(cells: list[str]) -> bool:
    if not cells:
        return False
    return all(not cell or set(cell) <= set("-: ") for cell in cells)


def _header_score(cell: str) -> float:
    """Highest band written in one header cell. A name that is not a band scores nothing."""

    label = cell.strip()
    if not label:
        return -1.0
    parsed = _numeric_band(label)
    if parsed:
        return parsed[0]
    best = -1.0
    numbers = _RANGE.search(label)
    if numbers:
        best = float(max(int(numbers.group(1)), int(numbers.group(2))))
    folded = label.casefold()
    for name, score in sorted(_RANK.items(), key=lambda item: len(item[0]), reverse=True):
        if name in folded:
            best = max(best, float(score))
            break
    return best


def _row_name(criterion: str) -> str | None:
    for name, pattern in _ROW_MATCH:
        if pattern.search(criterion):
            return name
    return None


def _writer_owned(criterion: str) -> list[str]:
    return [name for name, pattern in _WRITER_OWNED if pattern.search(criterion)]


def _top_column(header: list[str]) -> tuple[int, str] | None:
    best_score = -1.0
    best_index: int | None = None
    for index in range(1, len(header)):
        score = _header_score(header[index])
        if score < 0:
            continue
        if best_index is None or score >= best_score:
            best_score = score
            best_index = index
    if best_index is None:
        return None
    return best_index, header[best_index]


def read_top_band_rows(scheme: str) -> dict[str, Any]:
    """Top-band cells for argument, evidence, and critical thinking.

    Format and citation rows are skipped. A row the table does not print
    is not invented. A table with no band column is not given a band.
    """
    found: dict[str, dict[str, str]] = {}
    skipped: set[str] = set()
    if not isinstance(scheme, str) or not scheme.strip():
        return {"rows": [], "skipped": []}
    for table in _table_rows(scheme):
        body = [cells for cells in table if not _rule_row(cells)]
        if len(body) < 2:
            continue
        picked = _top_column(body[0])
        if picked is None:
            continue
        index, band = picked
        for cells in body[1:]:
            if not cells or not cells[0]:
                continue
            owned = _writer_owned(cells[0])
            if owned:
                skipped.update(owned)
                continue
            name = _row_name(cells[0])
            if name is None or name in found:
                continue
            descriptor = cells[index] if index < len(cells) else ""
            if _rule_row([descriptor]):
                descriptor = ""
            found[name] = {"row": name, "band": band, "descriptor": descriptor}
    return {
        "rows": [found[name] for name in _ROW_ORDER if name in found],
        "skipped": [name for name in _SKIP_ORDER if name in skipped],
    }


def _descriptor_met(text: str, descriptor: str) -> bool:
    """True when the opened text uses the descriptor's own words."""

    words = _content_words(descriptor)
    if not words:
        return False
    return _covers(text if isinstance(text, str) else "", words)


def mark_top_band_rows(essay_id: str, text: str) -> dict[str, Any]:
    """Mark the top-band table rows. Format and citations stay unmarked.

    A scheme with none of those rows does not invent them, and it does not
    block. This does not rewrite the draft.
    """

    found = read_top_band_rows(mark_scheme_text(essay_id))
    rows: list[dict[str, Any]] = []
    missed: list[str] = []
    for item in found["rows"]:
        met = _descriptor_met(text, item["descriptor"])
        rows.append({**item, "met": met})
        if not met:
            missed.append(item["row"])
    return {
        "marked": bool(rows),
        "top_band": not missed,
        "rows": rows,
        "missed": missed,
        "skipped": list(found["skipped"]),
        "format": False,
        "citations": False,
    }


def with_top_band_marks(findings: str, mark: dict[str, Any] | None) -> str:
    """Append the top-band rows under the findings note. A quiet scheme adds nothing."""
    if not isinstance(mark, dict) or mark.get("marked") is not True:
        return findings
    rows = mark.get("rows")
    if not isinstance(rows, list) or not rows:
        return findings
    lines = [findings.rstrip(), "", "Top band"]
    for row in rows:
        if not isinstance(row, dict):
            continue
        name = str(row.get("row") or "").strip()
        if not name:
            continue
        label = name[:1].upper() + name[1:]
        band = str(row.get("band") or "").strip()
        descriptor = str(row.get("descriptor") or "").strip()
        state = "Met." if row.get("met") is True else "Not met."
        lines.append(f"{label} ({band}): {descriptor}. {state}")
    return "\n".join(lines)


_COVER_ALSO = {
    "counted words": ("counted words", "word count"),
    "registration number": (
        "registration number",
        "registration no",
        "student number",
        "student id",
    ),
    "module title": ("module title",),
    "module code": ("module code",),
    "assignment title": ("assignment title",),
    "name": ("name",),
}


def _norm_phrase(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _phrase_in(folded: str, phrase: str) -> bool:
    needle = phrase.casefold().strip()
    if not needle:
        return False
    if " " in needle:
        return needle in folded
    return re.search(rf"(?<![a-z0-9]){re.escape(needle)}(?![a-z0-9])", folded) is not None


def _cover_line_shown(folded: str, line: str) -> bool:
    """True when this parsed cover line, or the same ask in other words, is in the file."""
    wanted = line.strip()
    if not wanted:
        return True
    if wanted.casefold() in folded:
        return True
    label = wanted[:-1].strip() if wanted.endswith(":") else wanted
    phrases = _COVER_ALSO.get(label.casefold(), (label,))
    return any(_phrase_in(folded, phrase) for phrase in phrases)


def _body_for_word_band(text: str) -> str:
    """Opened body for a word-count ask. Cover lines and the reference list stay out."""

    kept: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.casefold() == "references":
            break
        if cover_residue_line(stripped):
            continue
        kept.append(stripped)
    return "\n".join(kept)


def _missed_word_band(essay_id: str, text: str) -> dict[str, Any] | None:
    """The parsed word band when the opened body sits outside it. No band is added."""
    from essay_writer.specialists.writer import document_word_band

    band = document_word_band(essay_id)
    low = band.get("min") if isinstance(band, dict) else None
    high = band.get("max") if isinstance(band, dict) else None
    if not isinstance(low, int) or not isinstance(high, int):
        return None
    counted = count_words(
        _body_for_word_band(text),
        exclude_citations=bool(band.get("exclude_in_text_citations")),
    )
    if low <= counted <= high:
        return None
    if counted > high:
        from essay_writer.rewrite_baseline import rewrite_vs_opened_word, stricter_invented_max

        if stricter_invented_max(essay_id, band, high=high) is not None:
            decision = rewrite_vs_opened_word(
                essay_id,
                text,
                band,
                counted=counted,
                high=high,
                counter="words",
                exclude_citations=bool(band.get("exclude_in_text_citations")),
            )
            if decision["waive_word"] and not decision["document_still_fails"]:
                return None
    return {"min": low, "max": high, "counted": counted}


def _missed_cover_fields(essay_id: str, text: str) -> list[str]:
    """Cover lines the title-page parse already asks for and the file does not show."""

    record = essay_assignment_gate.read_assignment(essay_id)
    if record.get("active") is not True or record.get("mapped") is False:
        return []
    spec = record.get("title_page")
    if not isinstance(spec, dict):
        return []
    fields = spec.get("fields") if isinstance(spec.get("fields"), list) else []
    labels = spec.get("labels") if isinstance(spec.get("labels"), list) else []
    if (
        not fields
        and not labels
        and spec.get("registration") is not True
        and spec.get("name") != "required"
        and spec.get("programme") is not True
    ):
        return []
    lines = title_page_cover_lines(
        essay_id,
        {},
        {"title": "", "word_count": {}},
        always_programme=False,
    )
    folded = text.casefold() if isinstance(text, str) else ""
    missed: list[str] = []
    seen: set[str] = set()
    for line in lines:
        if not isinstance(line, str) or not line.strip():
            continue
        if _cover_line_shown(folded, line):
            continue
        label = line.strip()
        if label.endswith(":"):
            label = label[:-1].strip()
        key = label.casefold()
        if not label or key in seen:
            continue
        seen.add(key)
        missed.append(label)
    return missed


def _missed_diary_headings(essay_id: str, text: str) -> list[str]:
    """Diary headings the brief named and the opened file does not show.

    Essay and report parts are not turned into headings here. A count of
    entries with no names does not invent Entry 1.
    """

    record = essay_assignment_gate.read_assignment(essay_id)
    if record.get("writing_type") != "diary" and record.get("lock") != "entries":
        return []
    stored = record.get("named_entries")
    if not isinstance(stored, list):
        return []
    folded = _norm_phrase(text if isinstance(text, str) else "")
    missed: list[str] = []
    for name in stored:
        if not isinstance(name, str) or not name.strip():
            continue
        if _norm_phrase(name) not in folded:
            missed.append(name.strip())
    return missed


def engagement_misses(essay_id: str, text: str) -> dict[str, Any]:
    """Parsed engagement asks the opened Word text does not meet.

    Word band, cover fields, and diary headings come from the parsers the
    brief, the marking table, and the diary already use. An empty result
    means those documents did not record that ask.
    """
    submitted = text if isinstance(text, str) else ""
    missed: dict[str, Any] = {}
    band = _missed_word_band(essay_id, submitted)
    if band:
        missed["word_band"] = band
    covers = _missed_cover_fields(essay_id, submitted)
    if covers:
        missed["cover_fields"] = covers
    headings = _missed_diary_headings(essay_id, submitted)
    if headings:
        missed["diary_headings"] = headings
    return missed


def _engagement_reason(missed: dict[str, Any]) -> str:
    parts: list[str] = []
    band = missed.get("word_band")
    if isinstance(band, dict):
        parts.append(
            "The opened Word file is outside the word band "
            f"({band.get('counted')} counted, {band.get('min')}–{band.get('max')})."
        )
    covers = missed.get("cover_fields")
    if isinstance(covers, list) and covers:
        parts.append(
            "The opened Word file is missing cover fields the brief asked for: "
            + ", ".join(str(item) for item in covers)
            + "."
        )
    headings = missed.get("diary_headings")
    if isinstance(headings, list) and headings:
        parts.append(
            "The opened Word file is missing headings the brief asked for: "
            + ", ".join(str(item) for item in headings)
            + "."
        )
    if parts:
        return " ".join(parts)
    return "The opened Word file misses an engagement ask."


def opened_word_gate(essay_id: str, text: str) -> dict[str, Any] | None:
    """Refuse a PASS when the opened Word text misses the requirement list.

    This is the same assignment_gate check the writer save uses, plus the
    engagement asks already parsed for a word band, cover fields, and diary
    headings. Daily use returns None. The constitution is not a substitute
    for the opened file. Cursor and Grok Build call this same function.
    """

    report = assignment_gate(essay_id, text, seat="eqr")
    if report.get("ok") is not True:
        return refuse("assignment-gate", str(report.get("reason") or ""), assignment=report)
    missed = engagement_misses(essay_id, text if isinstance(text, str) else "")
    if not missed:
        return None
    return refuse("engagement", _engagement_reason(missed), engagement=missed)


def review_gate(essay_id: str, text: str) -> dict[str, Any] | None:
    """Same check as opened_word_gate. Prefer opened_word_gate in new call sites."""
    return opened_word_gate(essay_id, text)


_REVIEW = {
    "essay_review",
    "eqr_verdict",
    "PERSON_BLOCK",
    "writer_clears_left",
    "REVIEW_INSTRUCTIONS",
    "STAGES",
    "WRITER_CLEARS_LEFT",
    "eqr_record",
}


def __getattr__(name: str):
    if name not in _REVIEW:
        raise AttributeError(name)
    from essay_writer import review

    return getattr(review, name)
