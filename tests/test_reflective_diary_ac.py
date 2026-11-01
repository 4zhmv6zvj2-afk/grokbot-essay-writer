"""Reflective diary is a third writing type. Essay and report stay as they are."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from helpers import ESSAY, pack, picked

from essay_writer.assignment_gate import read_assignment
from essay_writer.diary import (
    blocked_identity_keys,
    diary_signals,
    layout_for,
    module_id_from_brief,
    section_word_plan,
)
from essay_writer.gates import essay_gates
from essay_writer.harvard_4a import in_text_uses_et_al
from essay_writer.paths import essay_dir, essay_id_error
from essay_writer.projects import essay_sequence_new
from essay_writer.specialists.grok_bot import basics_pass_lines
from essay_writer.specialists.writer import body_is_justified, document_word_band
from essay_writer.stage_llm import brief_genre, stage_settings, stage_title
from essay_writer.stage_track import essay_stage
from essay_writer.student_voice import student_surface_hits
from essay_writer.textutil import genre_from_brief_text, word_band_from_brief_text
from essay_writer.word import cover_lines, write_docx

REPO = Path(__file__).resolve().parents[1]
DIARY = "RPT5001AA"
SPELL = (
    "In the text, write out every author when there are more than two, "
    "and spell them out again in the reference list."
)

DIARY_BRIEF = """
Assignment title: Individual reflective diary

Module title: Practice module
Module code: RPT5001AA

Aim: Understand the practice by reflecting on your own decisions.

Assignment task: Write four entries. Each entry names one decision and critically analyses it with theory and evidence.

The main body of the report is the diary. First person is allowed.

One entry for each occasion: North occasion, South occasion, East occasion, and West occasion.

Each entry is about 500 words, plus or minus 10%.
The whole diary is 2,000 words, plus or minus 10%.

Title page: module title, module code, assignment title, registration number, and the word count.
The name does not appear on the title page or anywhere else in the file.

Pages: Arial, size 12. Headings bold and left-aligned. Body text fully justified. 1.5 line spacing. 2.5 cm margins on all sides. Page numbers in the lower right.

# Marking criteria/feedback form

| Criterion | 0-39 | 40-49 | 50-59 | 60-69 | 70+ |
| --- | --- | --- | --- | --- | --- |
| Clarity of the decisions | thin | partial | sound | clear | exact and evidenced |

Participant number
Comments
Strengths
Areas for improvement

""" + SPELL + "\n"

ESSAY_BRIEF = (
    "Genre: essay\n"
    "Assignment title: Essay on the corridor\n"
    "Question: How do chokepoints constrain the corridor?\n"
    "Include a reflective section on what you learned.\n"
)
REPORT_BRIEF = (
    "Genre: report\n"
    "Assignment title: Individual business report\n"
    "Question: How do chokepoints constrain the corridor?\n"
    "Reflection\n"
    "One short reflective section on the process.\n"
)


def _open(module: str) -> str:
    draft = essay_sequence_new(module, "draft")
    assert draft["ok"] is True, draft
    opened = essay_sequence_new(module, "confirm")
    assert opened["ok"] is True, opened
    return opened["sequence_id"]


def _write_brief(module: str, text: str) -> None:
    folder = essay_dir(module)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(text, encoding="utf-8")


def test_ac1_diary_activates_and_essay_report_stay(workspace):
    assert diary_signals(DIARY_BRIEF) is True
    assert genre_from_brief_text(DIARY_BRIEF) == "diary"
    assert genre_from_brief_text("This assignment is a reflective diary.\n") is None
    assert diary_signals("This assignment is a logbook of weekly entries.\n") is False
    assert genre_from_brief_text(ESSAY_BRIEF) == "essay"
    assert genre_from_brief_text(REPORT_BRIEF) == "report"
    assert genre_from_brief_text("**Genre:** report\n") == "report"
    assert genre_from_brief_text("*Genre:* essay\n") == "essay"

    _write_brief(DIARY, DIARY_BRIEF)
    _write_brief(ESSAY, "")
    seq = _open(DIARY)
    assert brief_genre(seq) == "diary"
    assert brief_genre(ESSAY) == "report"
    assert stage_title("2", "diary") == "Establish the entries"
    assert stage_title("3", "diary") == "Write a diary plan"
    assert stage_title("8", "diary") == "Edit and redraft your diary"
    assert stage_title("2", "essay") == "Establish a thesis"
    assert stage_title("3", "essay") == "Write an essay plan"
    assert stage_title("2", "report") == "Establish purpose and scope"
    assert stage_title("3", "report") == "Write a report plan"
    assert stage_title("8", "essay") == "Edit and redraft"
    assert stage_title("8", "report") == "Edit and redraft"

    diary_plan = stage_settings("3", "diary")["prompt_overlay"]
    assert "not a rule the brief wrote" in diary_plan
    assert "trade-off" in diary_plan
    assert "A diary has no thesis." in stage_settings("2", "diary")["prompt_overlay"]
    assert "locked thesis" not in stage_settings("2", "diary")["prompt_overlay"]
    report_plan = stage_settings("3", "report")["prompt_overlay"]
    assert "central theme" in report_plan
    assert "trade-off" not in report_plan
    essay_two = stage_settings("2", "essay")["prompt_overlay"]
    assert "thesis" in essay_two
    assert "Establish the entries" not in essay_two
    for stage in ("2", "3", "5", "6", "7", "8"):
        assert student_surface_hits(stage_settings(stage, "diary")["prompt_overlay"]) == []

    router = (REPO / "docs" / "academic-writing-helpers.md").read_text(encoding="utf-8")
    helper = (REPO / "docs" / "reflective-diary-helpers.md").read_text(encoding="utf-8")
    assert "docs/reflective-diary-helpers.md" in router
    assert "docs/essay-writing-helpers.md" in router
    assert "docs/report-writing-helpers.md" in router
    assert "reflective section" in router.casefold()
    assert student_surface_hits(router) == []
    assert student_surface_hits(helper) == []
    for banned in ("RPT5001AA", "Environmental sustainability", "28 Nov", "Stakeholder mapping"):
        assert banned not in helper
    assert "University of Sydney" in helper or "Sydney Learning Centre" in helper
    assert "Reflective writing" in helper
    assert "Harvard Referencing" in helper
    assert "et al." in helper


def test_ac2_brief_table_is_extracted_and_the_top_band_is_printed(workspace):
    _write_brief(DIARY, DIARY_BRIEF)
    seq = _open(DIARY)
    seen = read_assignment(seq)
    table = essay_dir(DIARY) / "marking-table.md"
    assert table.is_file()
    copied = table.read_text(encoding="utf-8")
    assert "70+" in copied
    assert "exact and evidenced" in copied
    for field in ("Participant number", "Comments", "Strengths", "Areas for improvement"):
        assert field not in copied
    assert seen["mark_scheme"] is True
    assert seen["highest_band"] == "70+"
    assert seen["format"] == "reflective diary"
    for invented in ("Distinction", "Merit", "First"):
        assert invented not in (seen["highest_band"] or "")
        assert invented not in copied

    lines = basics_pass_lines(seq)
    assert lines[0].startswith("How to score high: 70+")
    assert "exact and evidenced" in lines[0]
    assert lines[1] == "Format: reflective diary."
    for invented in ("Distinction", "Merit", "First"):
        assert invented not in lines[0]
    note = "\n".join(lines)
    drafted = essay_stage(seq, "1", "draft", note)
    assert drafted["ok"] is True, drafted
    assert essay_stage(seq, "1", "confirm", note, **picked(drafted))["stored"] is True

    bare = (
        "Marking criteria\n\n"
        "Clarity of the decisions is one criterion.\n"
        "There is no band column on this table.\n"
    )
    from essay_writer.assignment_gate import _highest_band

    assert _highest_band(bare) is None
    assert _highest_band("High Distinction (80-100): evidenced\nMerit (60-69): sound\n") == "High Distinction"

    folder = essay_dir(ESSAY)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(
        "Genre: essay\nQuestion: Should the council publish the figure?\n",
        encoding="utf-8",
    )
    (folder / "mark-scheme.md").write_text(
        "High Distinction (80-100): evidenced\nPass (50-59): adequate\n",
        encoding="utf-8",
    )
    essay = _open(ESSAY)
    essay_seen = read_assignment(essay)
    assert essay_seen["highest_band"] == "High Distinction"
    assert essay_seen["format"] == "essay"
    assert not (essay_dir(ESSAY) / "marking-table.md").exists()
    assert basics_pass_lines(essay)[0].startswith("The mark scheme scores")


def test_ac3_aim_and_assignment_task_are_the_ask(workspace):
    _write_brief(DIARY, DIARY_BRIEF)
    seq = _open(DIARY)
    seen = read_assignment(seq)
    question = seen["question"]
    assert "Understand the practice by reflecting on your own decisions." in question
    assert "Each entry names one decision and critically analyses it with theory and evidence." in question
    assert "Question:" not in DIARY_BRIEF.split("Assignment task", 1)[0]
    line = basics_pass_lines(seq)[2]
    assert line.startswith("Question: ")
    assert "Understand the practice" in line
    assert "critically analyses it with theory" in line

    _write_brief("RPT6001AA", ESSAY_BRIEF)
    essay = _open("RPT6001AA")
    essay_seen = read_assignment(essay)
    assert essay_seen["question"] == "How do chokepoints constrain the corridor?"
    assert essay_seen["format"] == "essay"


def test_ac4_total_and_section_words_and_a_silent_plan(workspace):
    plan = section_word_plan(DIARY_BRIEF)
    assert plan["total"]["limit"] == 2000
    assert plan["total"]["min"] == 1800
    assert plan["total"]["max"] == 2200
    assert plan["clash"] is None
    assert plan["plan_not_brief_rule"] is False
    assert [row["name"] for row in plan["sections"]] == [
        "North occasion",
        "South occasion",
        "East occasion",
        "West occasion",
    ]
    assert all(row["from_brief"] is True for row in plan["sections"])
    assert all(row["limit"] == 500 for row in plan["sections"])
    assert word_band_from_brief_text(DIARY_BRIEF)["limit"] == 2000
    assert word_band_from_brief_text(DIARY_BRIEF)["min"] == 1800

    _write_brief(DIARY, DIARY_BRIEF)
    assert document_word_band(DIARY)["limit"] == 2000

    clash = section_word_plan(
        "Assignment title: Individual reflective diary\n"
        "Write four entries. Each entry names one decision and analyses it with theory.\n"
        "One entry for each occasion: North occasion, South occasion, East occasion, and West occasion.\n"
        "Each entry is about 800 words.\n"
        "The whole diary is 2,000 words.\n"
    )
    assert clash["clash"]
    assert clash["total"]["limit"] == 2000
    assert clash["sections"][0]["limit"] == 800
    assert "do not add up" in clash["clash"].casefold() or "do not add up" in clash["clash"]

    silent = (
        "Assignment title: Individual reflective diary\n"
        "Write two entries. Each entry names one decision and analyses it with theory.\n"
        "One entry for each occasion: North occasion and South occasion.\n"
        "The whole diary is 2,000 words, plus or minus 10%.\n"
    )
    planned = section_word_plan(silent)
    assert planned["plan_not_brief_rule"] is True
    assert planned["clash"] is None
    assert len(planned["sections"]) == 2
    assert all(row["from_brief"] is False for row in planned["sections"])
    assert sum(row["limit"] for row in planned["sections"]) == 2000
    assert "rule the brief wrote" not in str(planned["sections"])

    essay_band = word_band_from_brief_text("Word count: 800 words (720 to 880).\n")
    assert essay_band["min"] == 720
    assert essay_band["max"] == 880
    assert word_band_from_brief_text("Word count: 500 words.\n")["limit"] == 500


def test_ac5_module_id_accepts_a_brief_code_or_mints(workspace):
    assert essay_id_error(DIARY) is None
    assert essay_id_error("PBS5001C") is not None
    assert module_id_from_brief(DIARY_BRIEF) == DIARY
    assert module_id_from_brief(DIARY_BRIEF) != "RPT5001BB"
    assert module_id_from_brief(DIARY_BRIEF) != "PBS5001C"
    seq = _open(DIARY)
    assert seq == f"{DIARY}-1"

    untitled = (
        "Module title: Coastal Practice Diary\n"
        "Assignment title: Individual reflective diary\n"
        "Write two entries. Each entry names one decision and analyses it with theory.\n"
    )
    minted = module_id_from_brief(untitled)
    assert minted
    assert not minted.endswith("HK")
    assert essay_id_error(minted) is None
    assert essay_id_error("MOD1004AA") is None
    opened = _open(minted)
    assert opened == f"{minted}-1"


def test_ac6_title_page_follows_the_brief_and_blocks_a_forbidden_name(workspace):
    _write_brief(DIARY, DIARY_BRIEF)
    seq = _open(DIARY)
    body = pack(DIARY)
    body["assessment"] = "Individual reflective diary"
    body["module_title"] = "Practice module"
    body["registration_number"] = "445566"
    body["sections"] = [
        {"heading": "North occasion", "body": "One decision, with theory."},
        {"heading": "South occasion", "body": "A later decision, with theory."},
    ]
    report = {
        "title": "Individual reflective diary",
        "word_count": {"counted": 40, "stated": 2000},
        "sections": body["sections"],
        "figures": [],
        "tables": [],
        "sources": [],
    }
    lines = cover_lines(seq, body, report)
    joined = "\n".join(lines)
    assert "Practice module" in joined
    assert DIARY in joined
    assert "Individual reflective diary" in joined
    assert "Registration number: 445566" in joined
    assert "Teaching partner" not in joined
    assert "PROG01" not in joined
    assert "Conclusion:" not in joined
    assert blocked_identity_keys(seq, body) == []

    named = dict(body)
    named["student_name"] = "Ada Example"
    assert "student_name" in blocked_identity_keys(seq, named)
    named_lines = "\n".join(cover_lines(seq, named, report))
    assert "Ada Example" not in named_lines
    assert "Registration number: 445566" in named_lines

    gates = essay_gates(seq, body)
    identity = next(item for item in gates["gates"] if item["name"] == "identity")
    assert identity["passed"] is True
    named_gates = essay_gates(seq, named)
    named_identity = next(item for item in named_gates["gates"] if item["name"] == "identity")
    assert named_identity["passed"] is False

    essay_body = pack(ESSAY)
    essay_body["student_number"] = "445566"
    assert "student_number" in blocked_identity_keys(ESSAY, essay_body)

    assert layout_for(seq)["margin_cm"] == 2.5
    assert layout_for(seq)["page_number"] == "right"
    assert layout_for(seq)["heading_pt"] == 12
    assert body_is_justified(seq) is True
    assert layout_for(ESSAY) == {}

    path = workspace / "Essay-writer" / DIARY / "sequences" / "1" / "output" / "diary.docx"
    write_docx(seq, body, report, path)
    document = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Registration number: 445566" in text
    assert "Ada Example" not in text
    assert "Teaching partner" not in text
    assert "Participant number" not in text
    section = document.sections[0]
    assert abs(section.left_margin.cm - 2.5) < 0.05
    assert section.footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    assert headings
    assert headings[0].runs[0].font.size.pt == 12
    prose = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("One decision")
    ]
    assert prose
    assert prose[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_ac7_harvard_et_al_ignores_a_spell_all_authors_line():
    assert in_text_uses_et_al(2, SPELL) is False
    assert in_text_uses_et_al(3, SPELL) is False
    assert in_text_uses_et_al(4, SPELL) is True
    guide = (REPO / "docs" / "referencing-harvard.md").read_text(encoding="utf-8")
    helper = (REPO / "docs" / "reflective-diary-helpers.md").read_text(encoding="utf-8")
    assert "does not replace this rule" in guide
    assert "Four or more" in guide
    assert "spell every author" in helper.casefold()
    assert "et al." in helper
    assert student_surface_hits(guide) == []
