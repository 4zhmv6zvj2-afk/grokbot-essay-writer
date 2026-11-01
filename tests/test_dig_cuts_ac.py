"""Dig cuts prevention. Stable ids: AC1–AC5.

AC1: Similarly-shaped leading transitions — gathered pass, ungathered invent-forbid, lone Similarly stays a name.
AC2: Pack validate and the Word cite path read the same citation_tokens (no drift).
AC3: A file outside the word band names the counted number beside the band.
AC4: Thin 6→10 Word journey — cite parse, validate, stage 10 refs_from log, copy-to-mac refuse reason.
AC5: Build covered. Cursor and Grok Build share this module. Cut-surface hashes pinned.
No separate harness.
"""

from __future__ import annotations

import hashlib
import inspect
import json
import re
from pathlib import Path

from docx import Document
from helpers import (
    ESSAY,
    TOPIC,
    open_sequence,
    pack,
    report_sections,
    write_upload,
)

from essay_writer import build_office, cursor_office
from essay_writer import validate as validate_mod
from essay_writer.build_office import opened_word_gate as build_opened
from essay_writer.build_office import stage_launch
from essay_writer.cursor_office import cursor_stage_launch
from essay_writer.cursor_office import opened_word_gate as cursor_opened
from essay_writer.debug_log import debug_log_path, log_word_step, ml_events_path
from essay_writer.first_run import main as first_run_main
from essay_writer.gates import essay_gates
from essay_writer.late_chain import write_stage10
from essay_writer.ml_events import handoff_outcome, handoff_refuse_reason
from essay_writer.paths import essay_dir, work_dir
from essay_writer.research import essay_research_summary
from essay_writer.server import HANDLERS
from essay_writer.specialists.eqr import opened_word_gate, read_submit_word
from essay_writer.specialists.writer import document_word_band
from essay_writer.state import load_state, save_state
from essay_writer.textutil import citation_tokens
from essay_writer.validate import essay_pack_validate
from essay_writer.word_references import cite_reference_gaps

REPO = Path(__file__).resolve().parents[1]
SIMILARLY_FIXTURE = REPO / "tests" / "fixtures" / "dig_cuts_similarly_cite_v1.json"

# Pinned dig-cut surfaces. A silent rewrite fails here before behaviour drifts.
CITATION_TOKENS_SHA = (
    "7034ca75a129de1b4e55dc7875b608f350d9bbbf41a5bb17eeb5c169f9d090db"
)
CITE_REFERENCE_GAPS_SHA = (
    "30c9a4f08715e0d726b8cdde1028adeeb97e66a786aae243ffe3811c871ace38"
)
LOG_WORD_STEP_REFS_SHA = (
    "7542c3c50ebcbaeebdb85cd7a47369b9afa12e215bee62344c07c1ff182e56ae"
)

QUESTION = "How do chokepoints constrain the corridor?"
YULIUS_QUOTE = "The Strait handles a large share of flows."
SHIU_QUOTE = "One named supply chain on the corridor is in scope."
SAVED_LINE = "Yulius, A. (2024) 'Saved reference text', Practice Journal, 1(1), pp. 1-4."
BAND_BRIEF = (
    "Genre: report\n"
    f"Question: {QUESTION}\n"
    "Word count: 2000 words.\n"
)
NARROW_BRIEF = (
    "Genre: report\n"
    f"Question: {QUESTION}\n"
    "Word count: 800 words (720 to 880).\n"
)


def _codes(items: list[dict]) -> list[str]:
    return [item["code"] for item in items]


def _gate(report: dict, name: str) -> dict:
    for item in report["gates"]:
        if item["name"] == name:
            return item
    raise AssertionError(name)


def _brief(module: str, text: str) -> None:
    folder = essay_dir(module)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(text, encoding="utf-8")


def _source(module: str, stem: str, author: str, year: str, quote: str) -> dict:
    return {
        "id": f"src-{stem}",
        "origin": "upload",
        "locator": f"uploads/{module}/{stem}.md",
        "author": author,
        "year": year,
        "title": quote,
        "quote": quote,
    }


def _cited_pack(seq: str, sections: list[dict[str, str]], sources: list[dict]) -> dict:
    body = pack(seq)
    body["conclusion"] = "open"
    body["sections"] = sections
    body["sources"] = sources
    return body


def _rows(essay_id: str | None = None) -> list[dict]:
    path = ml_events_path(essay_id)
    if not path.is_file():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _reference_paragraphs(path: Path) -> list[str]:
    texts = [paragraph.text.strip() for paragraph in Document(str(path)).paragraphs if paragraph.text.strip()]
    if "References" not in texts:
        return []
    return texts[texts.index("References") + 1 :]


def _load_similarly_fixture() -> list[dict]:
    payload = json.loads(SIMILARLY_FIXTURE.read_text(encoding="utf-8"))
    return list(payload["samples"])


def _pin_cut_surface_hashes() -> None:
    assert (
        hashlib.sha256(inspect.getsource(citation_tokens).encode()).hexdigest()
        == CITATION_TOKENS_SHA
    )
    assert (
        hashlib.sha256(inspect.getsource(cite_reference_gaps).encode()).hexdigest()
        == CITE_REFERENCE_GAPS_SHA
    )
    step_src = inspect.getsource(log_word_step)
    start = step_src.index("if refs_from:")
    end = step_src.index("if refused:")
    assert hashlib.sha256(step_src[start:end].encode()).hexdigest() == LOG_WORD_STEP_REFS_SHA


def test_ac1_leading_transition_is_not_the_author(workspace):
    """AC1. Similarly, Yulius (2024) and Similarly, Shiu (2023) pass; an unknown author fails."""
    for sample in _load_similarly_fixture():
        assert citation_tokens(sample["text"]) == sample["tokens"]

    yulius_line = "Yulius, A. (2024) 'Choice', Practice Journal, 1(1), pp. 1-2."
    yulius_sentence = f"Similarly, Yulius (2024) shows {YULIUS_QUOTE}"
    shiu_sentence = f"Similarly, Shiu (2023) shows {SHIU_QUOTE}"
    assert cite_reference_gaps(
        [yulius_line],
        [{"heading": "Analysis", "body": yulius_sentence}],
    ) == []
    missing = cite_reference_gaps(
        [yulius_line],
        [{"heading": "Analysis", "body": "Similarly, Patel (2019) shows a gap in the corridor."}],
    )
    assert missing == ["(patel, 2019)"]
    assert all("similarly" not in item for item in missing)

    _brief(ESSAY, f"Genre: report\nQuestion: {QUESTION}\n")
    write_upload(workspace, ESSAY, "yulius.md", f"Yulius, A. (2024)\n{YULIUS_QUOTE}\n")
    write_upload(workspace, ESSAY, "shiu.md", f"Shiu, M. (2023)\n{SHIU_QUOTE}\n")
    seq = open_sequence()
    yulius = _source(ESSAY, "yulius", "Yulius, A.", "2024", YULIUS_QUOTE)
    shiu = _source(ESSAY, "shiu", "Shiu, M.", "2023", SHIU_QUOTE)
    gathered = _cited_pack(
        seq,
        report_sections(
            f"{yulius_sentence} {QUESTION}",
            f"{shiu_sentence} {QUESTION}",
            f"{yulius_sentence} One named supply chain stays in this report.",
            f"{shiu_sentence} {YULIUS_QUOTE}",
        ),
        [yulius, shiu],
    )
    essay_research_summary(seq, TOPIC, "draft", [yulius, shiu])
    essay_research_summary(seq, TOPIC, "confirm", [yulius, shiu])
    passed = essay_pack_validate(seq, gathered, "draft")
    assert passed["ok"] is True, passed
    assert passed.get("refused") != "invent-forbid"
    assert "invent-forbid" not in _codes(passed["blockers"])

    unknown = _cited_pack(
        seq,
        report_sections(
            f"Similarly, Patel (2019) shows a gap. {QUESTION}",
            f"{yulius_sentence} {QUESTION}",
            f"{yulius_sentence} One named supply chain stays in this report.",
            f"{shiu_sentence} {YULIUS_QUOTE}",
        ),
        [yulius, shiu],
    )
    blocked = essay_pack_validate(seq, unknown, "draft")
    assert blocked["ok"] is False, blocked
    assert blocked["refused"] == "invent-forbid"
    assert "patel, 2019" in blocked["reason"]
    assert "similarly" not in blocked["reason"].casefold()


def test_ac2_sibling_cite_sites_share_citation_tokens(workspace):
    """AC2. Pack validate and Word cite gaps call the same citation_tokens."""
    assert validate_mod.citation_tokens is citation_tokens
    gaps_src = inspect.getsource(cite_reference_gaps)
    assert "citation_tokens(" in gaps_src
    assert "def citation_tokens" not in gaps_src

    probe = "Similarly, Yulius (2024) cites one point. Similarly, Patel (2019) cites another."
    tokens = citation_tokens(probe)
    assert tokens == ["yulius, 2024", "patel, 2019"]

    yulius_line = "Yulius, A. (2024) 'Choice', Practice Journal, 1(1), pp. 1-2."
    gaps = cite_reference_gaps(
        [yulius_line],
        [{"heading": "Body", "body": probe}],
    )
    assert gaps == ["(patel, 2019)"]

    _brief(ESSAY, f"Genre: report\nQuestion: {QUESTION}\n")
    write_upload(workspace, ESSAY, "yulius.md", f"Yulius, A. (2024)\n{YULIUS_QUOTE}\n")
    seq = open_sequence()
    yulius = _source(ESSAY, "yulius", "Yulius, A.", "2024", YULIUS_QUOTE)
    body = _cited_pack(
        seq,
        [{"heading": "Analysis", "body": probe}],
        [yulius],
    )
    essay_research_summary(seq, TOPIC, "draft", [yulius])
    essay_research_summary(seq, TOPIC, "confirm", [yulius])
    blocked = essay_pack_validate(seq, body, "draft")
    assert blocked["ok"] is False, blocked
    assert blocked["refused"] == "invent-forbid"
    assert "patel, 2019" in blocked["reason"]
    assert gaps[0].strip("()").casefold() in blocked["reason"].casefold()


def test_ac3_word_band_refuse_names_the_counted_number(workspace):
    """AC3. Outside 1800–2200, or another brief band, the refuse names the count beside the band."""
    module = "MOD1011AA"
    _brief(module, BAND_BRIEF)
    seq = open_sequence(module)
    band = document_word_band(seq)
    assert band["min"] == 1800
    assert band["max"] == 2200
    short = f"This is a report. {QUESTION}"
    refused = opened_word_gate(seq, short)
    assert refused is not None
    assert refused["refused"] == "engagement"
    counted = refused["engagement"]["word_band"]["counted"]
    assert isinstance(counted, int)
    assert counted < 1800
    reason = refused["reason"]
    assert re.search(rf"{counted} counted,\s*1800.2200", reason), reason
    assert cursor_opened(seq, short) == build_opened(seq, short) == refused

    outside = essay_gates(
        seq,
        {"essay_id": seq, "title": "Corridor", "genre": "report", "conclusion": "open"},
        body=short,
    )
    gate_reason = _gate(outside, "word-band")["reason"]
    assert _gate(outside, "word-band")["passed"] is False
    assert re.search(rf"{len(short.split())} counted,\s*1800.2200", gate_reason), gate_reason

    narrow_module = "RPT1008AA"
    _brief(narrow_module, NARROW_BRIEF)
    narrow = open_sequence(narrow_module)
    assert document_word_band(narrow)["min"] == 720
    assert document_word_band(narrow)["max"] == 880
    narrow_text = f"This is a report. {QUESTION}"
    narrow_gate = opened_word_gate(narrow, narrow_text)
    assert narrow_gate is not None
    narrow_count = narrow_gate["engagement"]["word_band"]["counted"]
    assert re.search(rf"{narrow_count} counted,\s*720.880", narrow_gate["reason"]), narrow_gate["reason"]


def test_ac4_stage6_to_stage10_word_journey_smoke(workspace, monkeypatch):
    """AC4. Cite parse, validate, stage 10 refs_from log, copy-to-mac refuse reason."""
    monkeypatch.setenv("ESSAY_WRITER_DEBUG", "1")
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", "")

    for stage in ("6", "10"):
        cursor_settings = cursor_stage_launch(stage)["settings"]
        build_settings = stage_launch(stage)["settings"]
        assert cursor_settings == build_settings
        assert cursor_settings["fast"] is False
        assert cursor_settings["model"] == "grok-4.7"

    _brief(ESSAY, f"Genre: report\nQuestion: {QUESTION}\n")
    write_upload(workspace, ESSAY, "yulius.md", f"Yulius, A. (2024)\n{YULIUS_QUOTE}\n")
    seq = open_sequence()
    yulius = _source(ESSAY, "yulius", "Yulius, A.", "2024", YULIUS_QUOTE)
    yulius_sentence = f"Similarly, Yulius (2024) shows {YULIUS_QUOTE}"
    assert citation_tokens(yulius_sentence) == ["yulius, 2024"]

    gathered = _cited_pack(
        seq,
        report_sections(
            f"{yulius_sentence} {QUESTION}",
            f"{yulius_sentence} {QUESTION}",
            f"{yulius_sentence} One named supply chain stays in this report.",
            f"{yulius_sentence} {YULIUS_QUOTE}",
        ),
        [yulius],
    )
    essay_research_summary(seq, TOPIC, "draft", [yulius])
    essay_research_summary(seq, TOPIC, "confirm", [yulius])
    validated = essay_pack_validate(seq, gathered, "draft")
    assert validated["ok"] is True, validated

    saved = {
        "essay_id": seq,
        "title": "Corridor flows",
        "assessment": "Corridor flows",
        "referencing": "harvard",
        "conclusion": "open",
        "references_markdown": SAVED_LINE,
        "sections": gathered["sections"],
        "sources": [yulius],
    }
    state = load_state(seq)
    state["saved_pack"] = saved
    save_state(seq, state)

    out_docx = work_dir(seq) / "output" / f"{seq}.docx"
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(yulius_sentence)
    document.save(str(out_docx))
    read = read_submit_word(seq)
    assert read["citations"] is True

    final = write_stage10(seq)
    assert final["ok"] is True, final
    assert "refs_from" not in final
    written = _reference_paragraphs(out_docx)
    assert written == [SAVED_LINE]
    log_text = debug_log_path(seq).read_text(encoding="utf-8")
    assert f"event=essay.product.write_stage10 essay_id={seq} refs_from=saved-path" in log_text
    assert f"event=essay.product.write_stage10 essay_id={seq} sha=" in log_text

    handoff_dir = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output"
    handoff_dir.mkdir(parents=True, exist_ok=True)
    (handoff_dir / "annotation.docx").write_bytes(b"ann-bytes")
    code = first_run_main(["--copy-to-mac", str(handoff_dir)])
    assert code == 0
    refused = [
        row
        for row in _rows(seq)
        if row.get("flag") == "--copy-to-mac" and row.get("outcome") == "refused"
    ]
    assert len(refused) == 1
    assert refused[0]["reason"] == "unmounted"
    assert refused[0]["reason"].strip()
    assert handoff_outcome({"ok": True, "copied": False, "files": [{"ok": True, "copied": False, "reason": "mac-not-mounted"}]}) == "refused"
    assert handoff_refuse_reason({"ok": True, "copied": False, "files": [{"ok": True, "copied": False, "reason": "mac-not-mounted"}]}) == "unmounted"


def test_ac5_build_covered_cursor_and_grok_build(workspace):
    """AC5. Build covered. One slice for Cursor and Grok Build. Fast stays off. Hashes pinned."""
    _pin_cut_surface_hashes()
    assert "opened_word_gate as _gate" in inspect.getsource(cursor_opened)
    assert "opened_word_gate as _gate" in inspect.getsource(build_opened)
    cursor_src = Path(cursor_office.__file__).read_text(encoding="utf-8")
    build_src = Path(build_office.__file__).read_text(encoding="utf-8")
    for name in ("def citation_tokens", "def write_stage10", "def note_cli_handoff"):
        assert name not in cursor_src
        assert name not in build_src
    assert "build_deliverable_handoff" in inspect.getsource(first_run_main)
    assert HANDLERS["essay.pack.validate"].__name__ == "essay_pack_validate"
    for stage in ("6", "10"):
        cursor_settings = cursor_stage_launch(stage)["settings"]
        build_settings = stage_launch(stage)["settings"]
        assert cursor_settings == build_settings
        assert cursor_settings["fast"] is False
        assert cursor_settings["model"] == "grok-4.7"
