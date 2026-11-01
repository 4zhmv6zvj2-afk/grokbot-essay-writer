"""Layer 4. Writer owns Word body shape and document word bands."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from helpers import ESSAY, open_sequence, pack, source, write_upload

from essay_writer.paths import essay_dir
from essay_writer.research import essay_research_summary
from essay_writer.specialists.writer import (
    HIGHEST_BAND,
    body_is_justified,
    document_word_band,
)
from essay_writer.textutil import word_rules_for
from essay_writer.validate import essay_pack_validate
from essay_writer.word import essay_pack_render_word


def _write(workspace, body=None, essay_id=ESSAY):
    write_upload(workspace, essay_id=essay_id)
    seq = open_sequence(essay_id)
    essay_research_summary(seq, "chokepoints", "draft", [source(essay_id)])
    essay_research_summary(seq, "chokepoints", "confirm", [source(essay_id)])
    payload = body or pack(essay_id)
    essay_pack_validate(seq, payload, "draft")
    essay_pack_validate(seq, payload, "confirm")
    essay_pack_render_word(seq, payload, "draft")
    written = essay_pack_render_word(seq, payload, "confirm")
    path = workspace / "Essay-writer" / essay_id / "sequences" / "1" / "output" / f"{seq}.docx"
    return written, path


def test_body_is_left_aligned_and_hashes_become_headings(workspace):
    payload = pack()
    payload["sections"][1]["heading"] = "## Supply chain context"
    payload["sections"][1]["body"] = (
        "One named supply chain on the corridor is in scope for this report (Smith, 2020).\n"
        "### Evidence\n"
        "The Strait handles a large share of flows through that chain (Smith, 2020)."
    )
    written, path = _write(workspace, payload)
    assert written["wrote"] is True
    document = Document(str(path))
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
    assert "Supply chain context" in headings
    assert "Evidence" in headings
    assert all(not text.startswith("#") for text in headings)
    plain = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Normal"
        and paragraph.text.strip()
        and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER
    ]
    assert plain
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT for paragraph in plain)
    assert "##" not in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "###" not in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_stated_justified_alignment_wins(workspace):
    folder = essay_dir(ESSAY)
    folder.mkdir(parents=True)
    (folder / "brief.md").write_text(
        "Genre: report\n"
        "Question: How do chokepoints constrain the corridor?\n"
        "Alignment: justified\n",
        encoding="utf-8",
    )
    assert body_is_justified(ESSAY) is True
    _written, path = _write(workspace)
    document = Document(str(path))
    prose = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Normal"
        and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER
        and "strait" in paragraph.text.casefold()
    ]
    assert prose
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY for paragraph in prose)


def test_mark_scheme_word_band_wins_over_the_brief_and_the_pack(workspace):
    folder = essay_dir(ESSAY)
    folder.mkdir(parents=True)
    (folder / "brief.md").write_text("Word count: 800 words (720 to 880).\n", encoding="utf-8")
    (folder / "mark-scheme.md").write_text("Word count: 500 words.\n", encoding="utf-8")
    band = document_word_band(ESSAY)
    assert band["min"] == 450
    assert band["max"] == 550
    ruled = word_rules_for(ESSAY, {"word_rules": {"min": 10, "max": 20}})
    assert ruled["min"] == 450
    assert ruled["max"] == 550
    assert "highest band" in HIGHEST_BAND
    assert "not merely a pass" in HIGHEST_BAND


def test_brief_band_wins_over_a_pack_budget(workspace):
    folder = essay_dir("RPT6001AA")
    folder.mkdir(parents=True)
    (folder / "brief.md").write_text("Word count: 800 words (720 to 880).\n", encoding="utf-8")
    ruled = word_rules_for("RPT6001AA", {"word_rules": {"min": 10, "max": 20}})
    assert ruled["min"] == 720
    assert ruled["max"] == 880


def test_no_document_band_does_not_invent_one_for_an_unknown_module(workspace):
    assert document_word_band("RPT6001AA") == {}
    assert word_rules_for("RPT6001AA", {}) == {}
