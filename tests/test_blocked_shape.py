"""Blocked shape locks."""

import json
from pathlib import Path

from docx import Document
from helpers import (
    ESSAY,
    TOPIC,
    mac_handoff_sequence_dir,
    open_sequence,
    pack,
    picked,
    plant_gathered_source,
    source,
    write_upload,
)

from essay_writer.databases import database_view, names_from_brief, remember_databases
from essay_writer.invent import gather_source
from essay_writer.originals import essay_originals
from essay_writer.research import essay_research_summary
from essay_writer.review import essay_review
from essay_writer.stage7 import claim_delta, note_edits, note_questions
from essay_writer.stage_llm import stage_settings
from essay_writer.stage_track import essay_stage
from essay_writer.student_voice import say, student_surface_hits
from essay_writer.user_place import CHAT_PLACE, NAMED_PLACE, chat_place
from essay_writer.validate import essay_pack_validate
from essay_writer.word import essay_pack_render_word
from essay_writer.writing_basics import CHECKLIST_WALK, CHECKLISTS

REPO = Path(__file__).resolve().parents[1]
STAGES = ("1", "2", "3", "4a", "4b", "5", "6", "7", "8", "9", "10")
BEFORE = (
    "GHG means greenhouse gases. "
    "Available at: doi:10.1000/abc. "
    "The firm cut emissions in 2024. "
    "The limit is the sample year."
)
AFTER = (
    "GHG means greenhouse gas. "
    "Available at: doi:10.1000/def. "
    "The firm cut emissions in 2023. "
    "The limit is one site."
)


def _advance(seq: str, stop: str, look: str = "Kept.") -> dict:
    last: dict = {}
    for stage in STAGES:
        if stage == stop:
            return last
        draft = essay_stage(seq, stage, "draft", look)
        assert draft["ok"] is True, draft
        if stage == "4a":
            plant_gathered_source(seq)
        if stage == "5":
            assert draft["stored"] is True, draft
            last = draft
            continue
        last = essay_stage(seq, stage, "confirm", look, **picked(draft))
        assert last["ok"] is True, last
    return last


def _write_docx(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for sentence in text.split(". "):
        line = sentence.strip()
        if not line:
            continue
        if not line.endswith("."):
            line += "."
        document.add_paragraph(line)
    document.save(str(path))


def test_user_facing_place_and_stage6_master(workspace):
    assert chat_place(None) == CHAT_PLACE
    assert chat_place("~/Documents/Essays/MOD1004AA") == "Documents/Essays/MOD1004AA"
    assert chat_place("/Volumes/ResearchDrive/out") == NAMED_PLACE
    assert chat_place("/Volumes/ArchiveDrive/notes/foo") == NAMED_PLACE
    assert chat_place("/Users/student/BackupDrive/notes") == NAMED_PLACE
    spoken = say(f"The place is {CHAT_PLACE}.")
    assert CHAT_PLACE in spoken
    assert "ResearchDrive" not in spoken
    assert CHECKLIST_WALK.index("never shows") < CHECKLIST_WALK.index(CHAT_PLACE)
    overlay = stage_settings("6")["prompt_overlay"]
    assert "first read uses chat delivery, or a folder the student already named" in overlay
    assert CHAT_PLACE in overlay
    assert "missing or renamed" in overlay
    assert "Do not guess" in overlay
    assert student_surface_hits(overlay) == []

    seq = open_sequence()
    saved6 = _advance(seq, "6")
    assert saved6["stage"] == "5"
    output = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output"
    _write_docx(output / "renamed-draft.docx", "The firm cut emissions in 2024.")
    asked = essay_stage(seq, "6", "draft", "Kept.")
    assert asked["ok"] is True
    assert asked["master_ask"] is True
    assert asked["place"] == CHAT_PLACE
    assert asked["choices"] == ["renamed-draft.docx"]
    assert asked["chain_started"] is False
    assert "next_run" not in asked
    blocked = essay_stage(seq, "6", "confirm", "Kept.", **picked(asked))
    assert blocked["ok"] is False
    assert blocked["refused"] == "master"
    assert blocked["choices"] == ["renamed-draft.docx"]
    assert CHAT_PLACE in blocked["reason"]
    assert "ResearchDrive" not in blocked["reason"]
    assert "SSD" not in blocked["reason"]
    named = essay_stage(seq, "6", "draft", "Kept.", master="renamed-draft.docx")
    assert named["ok"] is True
    assert named["master_ask"] is False
    kept = essay_stage(seq, "6", "confirm", "Kept.", master="renamed-draft.docx", **picked(named))
    assert kept["ok"] is True
    assert kept["current"] == "7"
    assert kept["chain_started"] is False
    assert "next_run" not in kept


def test_brief_named_database_and_student_file_slot(workspace):
    folder = workspace / "Essay-writer" / ESSAY
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(
        "Databases: Statista, Lloyd's List, and Clarksons.\n"
        "The library list also names scopus.\n",
        encoding="utf-8",
    )
    assert names_from_brief((folder / "brief.md").read_text(encoding="utf-8")) == [
        "Statista",
        "Lloyd's List",
        "Clarksons",
        "scopus",
    ]
    seq = open_sequence()
    listed = essay_originals(seq, "list")
    assert listed["databases"]["names"] == ["Statista", "Lloyd's List", "Clarksons", "scopus"]
    assert listed["databases"]["slot"] == "student-held file"
    assert listed["databases"]["path_stored"] is False
    remembered = remember_databases(seq, "/Volumes/ResearchDrive/secret/statista.csv")
    assert remembered["held_file"] == "statista.csv"
    assert remembered["path_stored"] is False
    stored = json.loads(
        (workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "tracking" / "databases.json").read_text(
            encoding="utf-8"
        )
    )
    dumped = json.dumps(stored)
    assert "ResearchDrive" not in dumped
    assert "secret" not in dumped
    assert stored["path_stored"] is False
    assert database_view(seq)["held_file"] == "statista.csv"
    _item, issues = gather_source(seq, {"database": "scopus", "origin": "upload"}, 0)
    assert issues[0]["code"] == "invent-forbid"
    stage1 = stage_settings("1")["prompt_overlay"]
    stage4 = stage_settings("4a")["prompt_overlay"]
    assert "slot for a file the student holds" in stage1
    assert "not a look-up" in stage4
    assert "Inditex" not in stage1


def test_stage7_asks_then_edits_and_lists_every_claim(workspace):
    from essay_writer.stage7 import load_stage7, save_stage7

    seq = open_sequence()
    opened = _advance(seq, "7")
    assert opened["card"]["stage"] == "7"
    assert opened["questions"]
    assert load_stage7(seq).get("asked") is True
    record = load_stage7(seq)
    record["asked"] = False
    save_stage7(seq, record)
    early = essay_stage(
        seq,
        "7",
        "draft",
        "Kept.",
        answers="Use the later year.",
        edited=AFTER,
        webCheck={"currency": True, "tone": True},
    )
    assert early["ok"] is False
    assert "Ask" in early["reason"]
    asked = essay_stage(seq, "7", "draft", "Kept.")
    assert asked["ok"] is True
    assert asked["file_work"] == "chat"
    assert asked["chain_started"] is False
    assert "next_run" not in asked
    assert len(asked["questions"]) >= 4
    output = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output" / f"{seq}.docx"
    _write_docx(output, BEFORE)
    thin = essay_stage(
        seq,
        "7",
        "draft",
        "Kept.",
        answers="Use the later year.",
        edited=AFTER,
        webCheck={"currency": True},
    )
    assert thin["ok"] is False
    edit = essay_stage(
        seq,
        "7",
        "draft",
        "Kept.",
        answers="Use 2023 and one site.",
        edited=AFTER,
        webCheck={"currency": True, "tone": True},
    )
    assert edit["ok"] is True, edit
    assert edit["file_work"] == "chat"
    delta = edit["delta"]
    assert "The firm cut emissions in 2024." in delta
    assert "The limit is the sample year." in delta
    assert "The firm cut emissions in 2023." in delta
    assert "The limit is one site." in delta
    assert all("means" not in line.casefold() for line in delta)
    assert all("doi:" not in line.casefold() for line in delta)
    saved = essay_stage(
        seq,
        "7",
        "confirm",
        "Kept.",
        answers="Use 2023 and one site.",
        edited=AFTER,
        webCheck={"currency": True, "tone": True},
        **picked(edit),
    )
    assert saved["ok"] is True, saved
    assert saved["current"] == "8"
    assert saved["executed_runs"] == ["7"]
    assert saved["next_run"] == {
        "stage": "8",
        "run": "edit",
        "separate": True,
        "runs_in_this_step": 1,
    }
    assert saved["chain_armed"] is True
    overlay = stage_settings("7")["prompt_overlay"]
    rules = stage_settings("7")["rules"]
    assert overlay.index("ask questions") < overlay.index("delta lists every sentence")
    assert "Grok Bot reads the stage 6 draft" in rules
    assert "Cursor does heavy file work only" in rules
    assert student_surface_hits("\n".join(CHECKLISTS["7"])) == []
    only_gloss = claim_delta("GHG means greenhouse gases.", "GHG means greenhouse gas.")
    assert only_gloss == []


def test_split_chain_shows_blocked_review_then_final_word(workspace, monkeypatch):
    documents = workspace / "Documents"
    documents.mkdir()
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    seq = open_sequence()
    stage6 = _advance(seq, "7")
    assert stage6["chain_started"] is False
    assert "next_run" not in stage6
    output = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output" / f"{seq}.docx"
    _write_docx(output, BEFORE + " Smith (2020) supports the corridor.")
    essay_stage(seq, "7", "draft", "Kept.")
    web = {"currency": True, "clarity": True}
    edited = essay_stage(seq, "7", "draft", "Kept.", answers="Use 2023 and one site.", edited=AFTER, webCheck=web)
    saved7 = essay_stage(
        seq,
        "7",
        "confirm",
        "Kept.",
        answers="Use 2023 and one site.",
        edited=AFTER,
        webCheck=web,
        **picked(edited),
    )
    assert saved7["current"] == "8"
    quiet = essay_stage(seq, "8", "draft", "Kept.")
    assert quiet["ok"] is True
    assert quiet["stored"] is False
    assert "Open the Confirm card to save Stage 8" in quiet["next"]
    assert "Not yet" in quiet["next"]
    assert "Say yes" not in quiet["next"]
    assert quiet["card"]["actions"][0]["label"] == "Confirm"
    saved8 = essay_stage(seq, "8", "confirm", "Kept.", **picked(quiet))
    assert saved8["current"] == "9"
    assert saved8["executed_runs"] == ["8"]
    assert saved8["next_run"]["stage"] == "9"
    assert saved8["next_run"]["runs_in_this_step"] == 1
    assert "10" not in saved8["executed_runs"]
    saved9 = essay_stage(seq, "9", "draft", "Kept.")
    assert saved9["stored"] is False
    kept9 = essay_stage(seq, "9", "confirm", "Kept.", **picked(saved9))
    assert kept9["current"] == "10"
    assert kept9["blocked"] is True
    assert kept9["eqr"] == "missing"
    assert kept9["next_run"]["run"] == "review"
    assert "independent review" in kept9["next"].casefold()
    drafted10 = essay_stage(seq, "10", "draft", "Kept.")
    blocked = essay_stage(seq, "10", "confirm", "Kept.", **picked(drafted10))
    assert blocked["ok"] is False
    assert blocked["refused"] == "eqr"
    assert blocked["blocked"] is True
    assert not (mac_handoff_sequence_dir(documents, seq) / f"{ESSAY}-final.docx").exists()
    essay_review(seq, "9", "PASS", "draft")
    essay_review(seq, "9", "PASS", "confirm")
    from helpers import ensure_word_refs

    ensure_word_refs(seq)
    final = essay_stage(seq, "10", "confirm", "Kept.", **picked(drafted10))
    assert final["ok"] is True, final
    assert final["stored"] is True
    assert final["submitted"] is False
    assert final["place"] == f"Documents/Essay Bot/{ESSAY}/{seq}/{ESSAY}-final.docx"
    assert final["place"] in final["next"]
    assert "first-draft.md" not in final["next"]
    assert str(documents) not in final["next"]
    folder = mac_handoff_sequence_dir(documents, seq)
    assert (folder / f"{ESSAY}-final.docx").is_file()
    assert (folder / f"{ESSAY}-comments.docx").is_file()
    assert f"{seq}.docx" != f"{ESSAY}-final.docx"
    comments = "\n".join(paragraph.text for paragraph in Document(str(folder / f"{ESSAY}-comments.docx")).paragraphs)
    assert "independent review" in comments.casefold()
    assert "2024" in comments or "sample year" in comments


def test_word_save_rewrites_sequence_draft_note(workspace):
    write_upload(workspace)
    seq = open_sequence()
    essay_research_summary(seq, TOPIC, "draft", [source()])
    essay_research_summary(seq, TOPIC, "confirm", [source()])
    payload = pack()
    essay_pack_validate(seq, payload, "draft")
    essay_pack_validate(seq, payload, "confirm")
    note = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "first-draft.md"
    missing = essay_pack_render_word(seq, payload, "draft", stage="7")
    assert missing["ok"] is False
    assert not note.exists()
    note_questions(seq)
    note_edits(
        seq,
        answers="Keep the corridor claim.",
        edited="The corridor claim stays.",
        web_check={"currency": True, "tone": True},
    )
    essay_pack_render_word(seq, payload, "draft", stage="7")
    written7 = essay_pack_render_word(seq, payload, "confirm", stage="7")
    assert written7["wrote"] is True
    assert written7["file_work"] == "cursor"
    assert written7["draft_note_rewritten"] is True
    assert note.is_file()
    assert "first-draft.md" not in written7["next"]
    assert "Introduction" in note.read_text(encoding="utf-8")
    brief = workspace / "Essay-writer" / ESSAY / "brief.md"
    brief.write_text(
        "Genre: report\nQuestion: How do chokepoints constrain the corridor?\n",
        encoding="utf-8",
    )
    note_text = "This is a report. How do chokepoints constrain the corridor?"
    essay_review(seq, "9", "PASS", "draft", constitution=note_text)
    essay_review(seq, "9", "PASS", "confirm", constitution=note_text)
    note.write_text("stale\n", encoding="utf-8")
    essay_pack_render_word(seq, payload, "draft", stage="9")
    written9 = essay_pack_render_word(seq, payload, "confirm", stage="9")
    assert written9["wrote"] is True, {key: written9.get(key) for key in ("refused", "reason", "eqr", "ok")}
    assert written9["draft_note_rewritten"] is True
    assert "stale" not in note.read_text(encoding="utf-8")
    assert "first-draft.md" not in written9["next"]
    assert "sequences/" not in written9["next"]


def test_helpers_keep_the_stage6_question_line():
    helpers = (REPO / "docs" / "academic-writing-helpers.md").read_text(encoding="utf-8")
    assert student_surface_hits(helpers) == []
    assert "Read the stage 6 draft, ask questions, then edit from those answers" in helpers
    assert "Cursor" not in helpers
    assert "Grok Bot" not in helpers
