"""Shared Word door for stages 6, 7, 9, and 10.

AC1: Stage 6 render_word uses word_manifest, the same cover and references
entry Stage 10 uses.
AC2: Stage 7 and Stage 9 Word writes call that same entry. They do not keep
a second pack builder.
AC3: A body-only pack still prints the house cover (sequence id, brief Title,
counted words) and the Reference List when cites exist.
"""

import inspect
from pathlib import Path

from docx import Document
from helpers import ESSAY, open_sequence, pack, report_sections, write_upload

from essay_writer.late_chain import write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.research import essay_research_summary
from essay_writer.validate import essay_pack_validate
from essay_writer.word import (
    cover_lines,
    essay_pack_render_word,
    word_manifest,
    write_docx,
)

MODULE = ESSAY
TOPIC = "How should the firm manage corridor risk?"
BRIEF_TITLE = "How can a mid-size retailer manage supply risk?"
REPORT_BRIEF = (
    "Genre: report\n"
    f"Question: {TOPIC}\n"
    f"Title: {BRIEF_TITLE}\n"
)
SCHEME = "High Distinction (80-100): the question is answered with evidence.\n"


def _brief() -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(REPORT_BRIEF, encoding="utf-8")
    (folder / "mark-scheme.md").write_text(SCHEME, encoding="utf-8")


def _cover(path: Path) -> list[str]:
    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.startswith("Heading"):
            break
        lines.append(text)
    return lines


def _references_body(path: Path) -> list[str]:
    texts = [
        paragraph.text.strip()
        for paragraph in Document(str(path)).paragraphs
        if paragraph.text.strip()
    ]
    if "References" not in texts:
        return []
    return texts[texts.index("References") + 1 :]


def _report_pack(seq: str) -> dict:
    cite_hasan = "(Hasan and Sarker, 2023)"
    cite_imma = "(Immadisetty, 2022)"
    body = pack(seq)
    body["assessment"] = "Individual supply chain report"
    body["sections"] = report_sections(
        f"High Distinction. This report asks how the firm should manage corridor risk {cite_hasan}.",
        f"One named chain is in scope {cite_imma}.",
        f"Evidence supports the finding {cite_hasan}.",
        f"The conclusion follows the sources {cite_imma}.",
    )
    body["sources"] = [
        {
            "id": "src-hasan",
            "origin": "upload",
            "locator": f"uploads/{MODULE}/hasan.md",
            "author": "Hasan and Sarker",
            "year": "2023",
            "title": "Corridor risk is material for the chain",
            "quote": "Corridor risk is material for the chain.",
        },
        {
            "id": "src-imma",
            "origin": "upload",
            "locator": f"uploads/{MODULE}/immadisetty.md",
            "author": "Immadisetty",
            "year": "2022",
            "title": "One named chain is in scope for the report",
            "quote": "One named chain is in scope for the report.",
        },
    ]
    return body


def _confirm_research(seq: str, body: dict) -> None:
    sources = body["sources"]
    draft = essay_research_summary(seq, TOPIC, "draft", sources)
    assert draft["ok"] is True, draft
    essay_research_summary(seq, TOPIC, "confirm", sources)


def _spy(monkeypatch):
    seen: list[str] = []
    real = word_manifest

    def spy(essay_id, pack, **kwargs):
        seen.append(str(essay_id))
        return real(essay_id, pack, **kwargs)

    monkeypatch.setattr("essay_writer.word.word_manifest", spy)
    return seen


def test_ac1_stage6_render_word_uses_shared_cover_refs_door(workspace, monkeypatch):
    """AC1. Stage 6 render_word and Stage 10 call the same word_manifest door."""
    render = inspect.getsource(essay_pack_render_word)
    stage10 = inspect.getsource(write_stage10)
    manifest = inspect.getsource(word_manifest)
    assert render.count("word_manifest(") == 1
    assert stage10.count("word_manifest(") == 1
    assert "_word_pack(" in manifest
    assert "cover_lines(" in manifest
    assert "cover_field_rows(" in inspect.getsource(cover_lines)
    assert "cover_field_rows(" in inspect.getsource(cover_lines)
    assert "cover_lines_from_field_rows(" in inspect.getsource(cover_lines)

    seen = _spy(monkeypatch)
    seq = open_sequence(MODULE)
    body = pack(seq)
    preview = essay_pack_render_word(seq, body, "draft", stage="6")
    assert seen == [seq]
    prepared, report = word_manifest(seq, body)
    assert preview["cover"] == report["cover"]
    assert preview["references"] == report["references"]
    assert prepared["essay_id"] == seq


def test_ac2_stage7_and_stage9_use_the_same_door(workspace, monkeypatch):
    """AC2. Stage 7 and Stage 9 do not keep a pack builder that skips cover or refs."""
    render = inspect.getsource(essay_pack_render_word)
    stage10 = inspect.getsource(write_stage10)
    assert "_word_pack(" not in render
    assert "_word_pack(" not in stage10
    assert "word_manifest(" in render
    assert "word_manifest(" in stage10

    seen = _spy(monkeypatch)
    seq = open_sequence(MODULE)
    body = pack(seq)
    essay_pack_render_word(seq, body, "draft", stage="7")
    essay_pack_render_word(seq, body, "draft", stage="9")
    assert seen == [seq, seq]


def test_ac3_body_only_pack_keeps_house_cover_and_refs(workspace, monkeypatch, tmp_path):
    """AC3. Body-only recovery still prints sequence id, Title, counted words, and refs."""
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    _brief()
    write_upload(
        workspace,
        MODULE,
        "hasan.md",
        (
            "Hasan and Sarker (2023) How should the firm manage corridor risk? "
            "Corridor risk is material for the chain.\n"
        ),
    )
    write_upload(
        workspace,
        MODULE,
        "immadisetty.md",
        (
            "Immadisetty (2022) How should the firm manage corridor risk? "
            "One named chain is in scope for the report.\n"
        ),
    )
    seq = open_sequence(MODULE)
    full = _report_pack(seq)
    _confirm_research(seq, full)
    draft_validate = essay_pack_validate(seq, full, "draft")
    assert draft_validate["ok"] is True, draft_validate
    essay_pack_validate(seq, full, "confirm")

    body_only = {
        "essay_id": seq,
        "title": "Draft",
        "referencing": "harvard",
        "conclusion": "noted",
        "sections": full["sections"],
    }
    from essay_writer.state import load_state, save_state

    state = load_state(seq)
    state["saved_pack"] = body_only
    save_state(seq, state)

    seen = _spy(monkeypatch)
    final = write_stage10(seq)
    assert seen == [seq]
    assert final.get("refused") != "title-page", final
    assert final["sequence_written"] is True, final

    stage_path = work_dir(seq) / "output" / f"{seq}.docx"
    lines = _cover(stage_path)
    assert seq in lines
    assert BRIEF_TITLE in lines
    assert any(item.startswith("Counted words:") for item in lines)
    refs = _references_body(stage_path)
    joined = "\n".join(refs)
    assert "Hasan" in joined
    assert "Immadisetty" in joined

    prepared, report = word_manifest(seq, body_only)
    assert seq in report["cover"]
    assert BRIEF_TITLE in report["cover"]
    assert any(item.startswith("Counted words:") for item in report["cover"])
    assert any("Hasan" in item for item in report["references"])
    assert prepared.get("sources")


def test_build_covered_cursor_and_grok_build_share_word_manifest():
    """Cursor and Grok Build mount the same Word door. Neither launch writes its own."""
    from essay_writer.build_office import stage_launch
    from essay_writer.cursor_office import cursor_stage_launch
    from essay_writer.server import HANDLERS
    from essay_writer.shells import cursor_shell, front_office
    from essay_writer.stage_track import _render_stage10_before_close

    shared = cursor_shell()["handlers"]
    assert front_office()["handlers"] is shared
    assert shared["essay.pack.render_word"] is essay_pack_render_word
    assert HANDLERS["essay.pack.render_word"] is essay_pack_render_word
    assert "word_manifest(" in inspect.getsource(essay_pack_render_word)
    assert "write_stage10(" in inspect.getsource(_render_stage10_before_close)
    assert "word_manifest(" in inspect.getsource(write_stage10)
    assert "def word_manifest" not in inspect.getsource(stage_launch)
    assert "def word_manifest" not in inspect.getsource(cursor_stage_launch)
    assert "write_docx(" not in inspect.getsource(stage_launch)
    assert "write_docx(" not in inspect.getsource(cursor_stage_launch)
    assert "write_docx(" in inspect.getsource(write_docx)
