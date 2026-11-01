import json
import shutil
from pathlib import Path

from docx import Document
from helpers import picked

from essay_writer.health import essay_health
from essay_writer.intake import essay_intake, essay_unassigned_collect
from essay_writer.paths import (
    kb1_module_dir,
    kb1_sequence_dir,
    recommended_paths,
    sequence_dir,
    uploads_dir,
)
from essay_writer.projects import (
    _replace_token,
    essay_project_archive,
    essay_project_delete,
    essay_project_unarchive,
    essay_projects,
    essay_sequence_new,
    save_sequence_state,
)
from essay_writer.server import INSTRUCTIONS
from essay_writer.setup import SETUP_PROMPT, essay_setup

MODULE = "MOD1011AA"


def _docx(path: Path, paragraphs: list[str], headings: list[str] | None = None) -> None:
    document = Document()
    for heading in headings or []:
        document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _confirm_new(parent: str | None = None) -> dict:
    draft = essay_sequence_new(MODULE, "draft", parent)
    assert draft["ok"] is True
    assert draft["wrote"] is False
    return essay_sequence_new(MODULE, "confirm", parent)


def test_instructions_ask_for_a_sequence_before_writing():
    text = INSTRUCTIONS.casefold()
    assert "essay.projects" in text
    assert "mod1011aa-5" in text
    assert "essay.sequence.new" in text


def test_sequence_rewrite_leaves_a_longer_number_alone():
    text = "MOD1011AA-1 and MOD1011AA-10 sequences/1/output sequences/10"
    updated = _replace_token(text, "MOD1011AA-1", "MOD1011AA-2")
    updated = _replace_token(updated, "sequences/1", "sequences/2")
    assert updated == "MOD1011AA-2 and MOD1011AA-10 sequences/2/output sequences/10"


def test_setup_prompt_does_not_clone_and_names_the_cloud_folders():
    text = SETUP_PROMPT.casefold()
    assert "nothing is cloned" in text
    assert "git clone" not in text
    assert "grok build" not in text
    assert "cursor" not in text
    assert "/workspace/essay-writer/app" in text
    assert "/workspace/essay-writer/kb1" in text
    assert "grok version" in text
    assert "install_grokbuild true" not in text


def test_setup_creates_folders_without_installing(tmp_path, monkeypatch):
    monkeypatch.setenv("ESSAY_WRITER_ROOT", str(tmp_path / "ledger"))
    app = tmp_path / "app"
    kb1 = tmp_path / "KB1"
    refused = essay_setup(install_grokbuild=True, app_path=str(app), kb1_path=str(kb1))
    assert refused["ok"] is False
    assert not app.exists()
    draft = essay_setup(
        app_path=str(app),
        kb1_path=str(kb1),
        confirmPhase="draft",
        acknowledge_skipped_searches=True,
    )
    assert draft["ok"] is True
    assert draft["wrote"] is False
    assert not app.exists()
    assert not kb1.exists()
    assert not (tmp_path / "setup.json").exists()
    result = essay_setup(
        app_path=str(app),
        kb1_path=str(kb1),
        confirmPhase="confirm",
        acknowledge_skipped_searches=True,
    )
    assert result["ok"] is True
    assert "clone" not in result
    assert "repo" not in result
    assert result["app_path"] == str(app.resolve())
    assert app.is_dir()
    assert (app / "Essay-writer").is_dir()
    assert (app / "uploads").is_dir()
    assert (app / "unassigned").is_dir()
    assert kb1.is_dir()
    assert result["grokbuild"]["ran_install"] is False
    assert result["grokbuild"]["where"] == "user-local"
    assert result["where"] == "user-local"
    record = json.loads((tmp_path / "setup.json").read_text(encoding="utf-8"))
    assert record["app_path"] == str(app.resolve())
    assert record["app"] == str(app.resolve())
    assert record["app_path"] != "open"
    assert record["kb1"] == str(kb1.resolve())
    assert not any(kb1.iterdir())


def test_recommended_paths_use_the_cloud_workspace(monkeypatch, tmp_path):
    cloud = tmp_path / "workspace"
    cloud.mkdir()
    monkeypatch.setattr("essay_writer.paths.CLOUD_WORKSPACE", cloud)
    app, kb1 = recommended_paths()
    assert app == cloud / "Essay-writer" / "app"
    assert kb1 == cloud / "Essay-writer" / "KB1"


def test_recommended_paths_without_a_cloud_workspace(monkeypatch, tmp_path):
    monkeypatch.setattr("essay_writer.paths.CLOUD_WORKSPACE", tmp_path / "missing")
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: tmp_path))
    assert recommended_paths() is None
    assert not (tmp_path / "Documents").exists()


def test_roll_forward_parent_allocates_last_plus_one_not_parent_plus_one(workspace):
    for parent_number in (1, 4):
        if parent_number == 1:
            assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
        else:
            save_sequence_state(
                MODULE,
                parent_number,
                {
                    "essay_id": f"{MODULE}-{parent_number}",
                    "sequence": parent_number,
                    "sequence_id": f"{MODULE}-{parent_number}",
                    "research": {"summary": None, "full": None},
                    "validations": {},
                    "word": {},
                    "revisions": [],
                    "shelf": {},
                },
            )
            (sequence_dir(MODULE, parent_number)).mkdir(parents=True, exist_ok=True)
    for number in range(5, 16):
        (sequence_dir(MODULE, number)).mkdir(parents=True, exist_ok=True)
        (sequence_dir(MODULE, number) / "state.json").write_text("{}\n", encoding="utf-8")
    pick = essay_projects("roll-forward", MODULE)
    assert pick["sequence_pick"] is True
    assert pick["next_sequence_id"] == f"{MODULE}-16"
    assert f"{MODULE}-4" in pick["sequences"]
    draft = essay_sequence_new(MODULE, "draft", f"{MODULE}-4")
    assert draft["parent"] == f"{MODULE}-4"
    assert draft["sequence_id"] == f"{MODULE}-16"
    opened = essay_sequence_new(MODULE, "confirm", f"{MODULE}-4")
    assert opened["sequence_id"] == f"{MODULE}-16"
    child = json.loads((sequence_dir(MODULE, 16) / "state.json").read_text(encoding="utf-8"))
    assert child["branched_from"] == f"{MODULE}-4"


def test_continue_lists_latest_sequence_only(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    assert _confirm_new()["sequence_id"] == f"{MODULE}-2"
    report = essay_projects("continue")
    row = report["rows"][0]
    assert row["essay_id"] == f"{MODULE}-2"
    assert "module" not in row
    assert "sequence_id" not in row
    assert row["sequence"] == 2
    assert row["stage"] == "open"
    assert row["roll_forward_ask"] is True
    listed = essay_projects("archive")
    assert listed["modules"] == [MODULE]
    assert all("-" not in name for name in listed["modules"])


def test_clone_keeps_uploads_shared_and_branches_sequence_files(workspace):
    _confirm_new()
    upload = uploads_dir(MODULE) / "paper.md"
    upload.parent.mkdir(parents=True)
    upload.write_text("shared paper\n", encoding="utf-8")
    note = kb1_module_dir(MODULE) / "paper.md"
    note.parent.mkdir(parents=True)
    note.write_text("vault paper\n", encoding="utf-8")
    essay = kb1_sequence_dir(MODULE, 1) / "essay.md"
    essay.parent.mkdir(parents=True)
    essay.write_text("sequence one\n", encoding="utf-8")
    state_path = sequence_dir(MODULE, 1) / "state.json"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["research"]["summary"] = {"topic": "ports"}
    state_path.write_text(json.dumps(state), encoding="utf-8")
    cloned = _confirm_new()
    assert cloned["sequence_id"] == f"{MODULE}-2"
    assert upload.read_text(encoding="utf-8") == "shared paper\n"
    assert list(uploads_dir(MODULE).rglob("*")) == [upload]
    child = json.loads((sequence_dir(MODULE, 2) / "state.json").read_text(encoding="utf-8"))
    parent = json.loads(state_path.read_text(encoding="utf-8"))
    assert child["research"]["summary"]["topic"] == "ports"
    assert child["branched_from"] == f"{MODULE}-1"
    assert parent["sequence_id"] == f"{MODULE}-1"
    assert (kb1_sequence_dir(MODULE, 2) / "essay.md").read_text(encoding="utf-8") == "sequence one\n"
    assert (kb1_module_dir(MODULE) / "paper.md").read_text(encoding="utf-8") == "vault paper\n"


def test_clone_rewrites_ids_inside_copied_files(workspace):
    _confirm_new()
    output = sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx"
    _docx(
        output,
        [f"{MODULE}-1", f"{MODULE}-10", "sequences/1/output", "sequences/10"],
    )
    note = kb1_sequence_dir(MODULE, 1) / "essay.md"
    note.parent.mkdir(parents=True)
    note.write_text(f"see {MODULE}-1 and {MODULE}-10\n", encoding="utf-8")
    assert _confirm_new()["sequence_id"] == f"{MODULE}-2"
    child = sequence_dir(MODULE, 2) / "output" / f"{MODULE}-2.docx"
    lines = [paragraph.text for paragraph in Document(str(child)).paragraphs if paragraph.text.strip()]
    assert lines == [f"{MODULE}-2", f"{MODULE}-10", "sequences/2/output", "sequences/10"]
    assert (kb1_sequence_dir(MODULE, 2) / "essay.md").read_text(encoding="utf-8") == (
        f"see {MODULE}-2 and {MODULE}-10\n"
    )
    parent = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
    assert f"{MODULE}-1" in parent
    assert "sequences/1/output" in parent


def test_archive_hides_module_until_unarchived(workspace):
    _confirm_new()
    draft = essay_project_archive(MODULE, "draft")
    assert draft["wrote"] is False
    assert archived_marker_absent()
    essay_project_archive(MODULE, "confirm")
    assert essay_projects("continue")["rows"] == []
    assert essay_projects("unarchive")["modules"] == [MODULE]
    essay_project_unarchive(MODULE, "draft")
    essay_project_unarchive(MODULE, "confirm")
    assert essay_projects("continue")["rows"][0]["essay_id"] == f"{MODULE}-1"


def archived_marker_absent() -> bool:
    return not (sequence_dir(MODULE, 1).parent.parent / "archived.json").is_file()


def test_delete_waits_for_confirm_and_refuses_a_changed_list(workspace):
    _confirm_new()
    upload = uploads_dir(MODULE) / "paper.md"
    upload.parent.mkdir(parents=True)
    upload.write_text("keep until confirm\n", encoding="utf-8")
    (kb1_module_dir(MODULE) / "paper.md").parent.mkdir(parents=True)
    (kb1_module_dir(MODULE) / "paper.md").write_text("vault\n", encoding="utf-8")
    draft = essay_project_delete(MODULE, "draft")
    assert draft["wrote"] is False
    assert draft["sequences"] == [f"{MODULE}-1"]
    assert upload.is_file()
    _confirm_new()
    refused = essay_project_delete(MODULE, "confirm")
    assert refused["ok"] is False
    assert upload.is_file()
    fresh = essay_project_delete(MODULE, "draft")
    deleted = essay_project_delete(MODULE, "confirm", **picked(fresh))
    assert deleted["deleted"] is True
    assert not sequence_dir(MODULE, 1).exists()
    assert not uploads_dir(MODULE).exists()
    assert not kb1_module_dir(MODULE).exists()


def test_loose_file_is_held_then_a_paper_is_shared(workspace):
    readme = workspace / "uploads" / "README.md"
    readme.write_text("stay\n", encoding="utf-8")
    loose = workspace / "uploads" / "note.md"
    loose.write_text("a paper\n", encoding="utf-8")
    draft = essay_unassigned_collect("draft")
    assert draft["wrote"] is False
    assert loose.is_file()
    assert readme.is_file()
    held = essay_unassigned_collect("confirm")
    assert held["files"] == ["note.md"]
    assert not loose.exists()
    assert readme.read_text(encoding="utf-8") == "stay\n"
    draft = essay_intake("note.md", "draft", MODULE)
    assert draft["kind"] == "supporting"
    assert draft["wrote"] is False
    filed = essay_intake("note.md", "confirm", MODULE)
    assert filed["wrote"] is True
    assert (uploads_dir(MODULE) / "note.md").read_text(encoding="utf-8") == "a paper\n"
    rows = essay_projects("continue")["rows"]
    assert rows == [
        {
            "essay_id": MODULE,
            "next": f"I can't find a sequence for {MODULE}. Start a new one?",
        }
    ]


def test_changed_word_draft_branches_and_updates_references(workspace):
    _confirm_new()
    original = [
        "The port handles containers.",
        "Reference List",
        "Smith, A. (2020) Port pricing. Journal of Ports.",
    ]
    output = sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx"
    _docx(output, original, headings=["Findings"])
    save_sequence_state(
        MODULE,
        1,
        {
            "essay_id": MODULE,
            "sequence": 1,
            "sequence_id": f"{MODULE}-1",
            "research": {
                "summary": {
                    "sources": [{"id": "src-1", "title": "Port pricing", "authors": ["Smith"]}]
                },
                "full": None,
            },
            "validations": {},
            "word": {
                "pack-1": {
                    "path": f"Essay-writer/{MODULE}/sequences/1/output/{MODULE}-1.docx",
                }
            },
            "revisions": [],
            "shelf": {},
            "reference_index": [],
            "unresolved_references": [],
        },
    )
    parent_hash = output.read_bytes()
    incoming = workspace / "unassigned" / "draft.docx"
    _docx(
        incoming,
        [
            "The port handles containers and barges.",
            "Reference List",
            "Smith, A. (2020) Port pricing. Journal of Ports.",
            "Unknown, B. (1999) Missing paper.",
        ],
        headings=["Findings", "Barges"],
    )
    draft = essay_intake("draft.docx", "draft", MODULE)
    assert draft["kind"] == "output"
    assert draft["wrote"] is False
    assert draft["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in draft
    assert draft["summary"]["headings_added"] == ["Barges"]
    filed = essay_intake("draft.docx", "confirm", MODULE)
    assert filed["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in filed
    assert filed["wrote"] is True
    assert output.read_bytes() == parent_hash
    child_state = json.loads((sequence_dir(MODULE, 2) / "state.json").read_text(encoding="utf-8"))
    assert child_state["branched_from"] == f"{MODULE}-1"
    assert child_state["reference_index"][0]["source_id"] == "src-1"
    assert child_state["unresolved_references"] == ["Unknown, B. (1999) Missing paper."]
    assert child_state["word"]["upload"]["path"].endswith(f"{MODULE}-2.docx")
    assert child_state["word"]["pack-1"]["path"].endswith(f"sequences/2/output/{MODULE}-2.docx")
    assert (uploads_dir(MODULE) / "draft.docx").exists() is False


def test_reference_list_paper_stays_on_the_module(workspace):
    _confirm_new()
    output = sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx"
    _docx(
        output,
        [
            "The port handles containers.",
            "Counted words: 1200",
            "Reference List",
            "Smith, A. (2020) Port pricing.",
        ],
    )
    incoming = workspace / "unassigned" / "paper.docx"
    _docx(
        incoming,
        [
            "A journal article about hull design.",
            "Reference List",
            "Lee, C. (2019) Hull form.",
        ],
    )
    draft = essay_intake("paper.docx", "draft", MODULE)
    assert draft["kind"] == "supporting"
    assert draft["sequence_id"] is None
    assert draft["wrote"] is False


def test_health_follows_the_latest_sequence_and_hides_archived(workspace):
    _confirm_new()
    state_path = sequence_dir(MODULE, 1) / "state.json"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["research"]["summary"] = {"topic": "ports"}
    state_path.write_text(json.dumps(state), encoding="utf-8")
    report = essay_health()
    assert report["opening"] == ["new", "continue", "roll-forward", "archive", "unarchive", "delete"]
    assert report["packs"][0]["essay_id"] == f"{MODULE}-1"
    assert "sequence_id" not in report["packs"][0]
    assert report["packs"][0]["stage"] == "research"
    essay_project_archive(MODULE, "draft")
    essay_project_archive(MODULE, "confirm")
    assert essay_health()["essay_ids"] == []


def test_unchanged_word_file_does_not_start_a_sequence(workspace):
    _confirm_new()
    output = sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx"
    _docx(output, ["The port handles containers.", "Reference List"])
    incoming = workspace / "unassigned" / "same.docx"
    incoming.parent.mkdir()
    shutil.copy2(output, incoming)
    report = essay_intake("same.docx", "confirm", MODULE)
    assert report["summary"]["same_file"] is True
    assert report["wrote"] is False
    assert not sequence_dir(MODULE, 2).exists()
