"""Post-smoke checks for citations, seats, and setup."""

import json
import re
import zipfile
from pathlib import Path

from helpers import ESSAY, TOPIC, open_sequence, pack, source, write_upload

from essay_writer.doors import DOORS
from essay_writer.health import essay_health
from essay_writer.intake import HONESTY, essay_intake
from essay_writer.paths import workspace_root
from essay_writer.research import essay_research_summary
from essay_writer.review import (
    PERSON_BLOCK,
    REVIEW_INSTRUCTIONS,
    essay_review,
    writer_clears_left,
)
from essay_writer.seats import seat_reading
from essay_writer.server import INSTRUCTIONS, house_instructions
from essay_writer.setup import (
    CURSOR_SETUP_PROMPT,
    SETUP_PROMPT,
    essay_setup,
    mid_office,
    setup_prompt,
)
from essay_writer.textutil import citation_tokens
from essay_writer.validate import essay_pack_validate
from essay_writer.word import essay_pack_render_word

REPO = Path(__file__).resolve().parents[1]
HELPERS_MARK = "When the brief does not name a module-specific writing guide"
HARVARD_MARK = "EndNote style to select"


def test_narrative_three_author_cite_is_one_citation():
    tokens = citation_tokens("A, B and C (2024) studied the corridor.")
    assert tokens == ["a, b and c, 2024"]
    assert "c, 2024" not in tokens

    both = citation_tokens("Both Smith (2024) agrees.")
    assert both == ["smith, 2024"]
    assert all(not token.startswith("both") for token in both)

    leading = citation_tokens("In Jones (2024) the point is narrow.")
    assert leading == ["jones, 2024"]
    assert all(not token.startswith("in ") for token in leading)

    wrapped = citation_tokens("Both A, B and C (2024) agree.")
    assert wrapped == ["a, b and c, 2024"]
    assert "c, 2024" not in wrapped
    pair = citation_tokens("Smith and Jones (2020) argue.")
    assert pair == ["smith and jones, 2020"]
    assert "jones, 2020" not in pair


def test_later_tools_load_app_path_not_recommended_base(tmp_path, monkeypatch):
    cloud = tmp_path / "cloud"
    cloud.mkdir()
    monkeypatch.setattr("essay_writer.paths.CLOUD_WORKSPACE", cloud)
    monkeypatch.delenv("ESSAY_WRITER_ROOT", raising=False)
    app = tmp_path / "real-app"
    (app / "Essay-writer").mkdir(parents=True)
    decoy = tmp_path / "decoy"
    decoy.mkdir()
    record_dir = cloud / "Essay-writer"
    record_dir.mkdir()
    (record_dir / "setup.json").write_text(
        json.dumps(
            {
                "app_path": str(app),
                "app": str(decoy),
                "recommended_base": str(decoy),
                "kb1": str(tmp_path / "kb1"),
            }
        ),
        encoding="utf-8",
    )

    def boom():
        raise AssertionError("recommended_base was read")

    monkeypatch.setattr("essay_writer.paths.recommended_base", boom)
    assert workspace_root() == app.resolve()
    report = essay_health()
    assert report["app_path"] == str(app.resolve())
    assert report["app_path"] != "open"


def test_setup_does_not_clone_or_store_open(tmp_path, monkeypatch):
    monkeypatch.setenv("ESSAY_WRITER_ROOT", str(tmp_path / "ledger"))
    assert "git clone" not in SETUP_PROMPT.casefold()
    assert "nothing is cloned" in SETUP_PROMPT.casefold()
    assert "Grok Build" not in SETUP_PROMPT
    assert "Cursor" not in SETUP_PROMPT
    opened = tmp_path / "open"
    vault = tmp_path / "KB1"
    refused = essay_setup(app_path=str(opened), kb1_path=str(vault), confirmPhase="draft")
    assert refused["ok"] is False
    assert not opened.exists()
    app = tmp_path / "app"
    draft = essay_setup(
        app_path=str(app), kb1_path=str(vault), confirmPhase="draft", acknowledge_skipped_searches=True
    )
    assert draft["wrote"] is False
    assert "clone" not in draft
    result = essay_setup(
        app_path=str(app), kb1_path=str(vault), confirmPhase="confirm", acknowledge_skipped_searches=True
    )
    assert result["ok"] is True
    assert result["wrote"] is True
    assert "clone" not in result
    record = json.loads((tmp_path / "setup.json").read_text(encoding="utf-8"))
    assert record["app_path"] == str(app.resolve())
    assert Path(record["app_path"]).name != "open"
    assert essay_health()["app_path"] == str((tmp_path / "ledger").resolve())
    assert essay_health()["app_path"] != "open"


def test_student_chat_and_word_omit_house_jargon(workspace):
    assert "invent-forbid" in INSTRUCTIONS
    assert "look-up" in HONESTY or "source card" in HONESTY
    for banned in ("dig", "door", "house", "Confirm", "EQR", "Grok Build", "Grok Bot", "SuperGrok", "Cursor", "PASS"):
        assert banned not in HONESTY

    write_upload(workspace)
    seq = open_sequence()
    essay_research_summary(seq, TOPIC, "draft", [source()])
    essay_research_summary(seq, TOPIC, "confirm", [source()])
    payload = pack(conclusion="open")
    payload["sections"][0]["body"] += (
        " Keep essay.dig and dig-card.json."
        " The dig door house Confirm EQR pack shelf mid office front office seat smoke tip GO PASS Grok Build Grok Bot SuperGrok Cursor."
        " invent-forbid: do not invent a source."
    )
    essay_pack_validate(seq, payload, "draft")
    essay_pack_validate(seq, payload, "confirm")
    essay_pack_render_word(seq, payload, "draft")
    written = essay_pack_render_word(seq, payload, "confirm")
    assert written["wrote"] is True
    path = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output" / f"{seq}.docx"
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    plain = re.sub(r"<[^>]+>", "", xml)
    assert re.search(r"\bdig\b", plain, re.IGNORECASE) is None
    assert "essay.dig" not in plain.casefold()
    assert "dig-card" not in plain.casefold()
    assert "look-up" in plain
    assert "invent-forbid" not in plain.casefold()
    assert "do not invent" not in plain.casefold()
    for banned in (
        "door",
        "house",
        "Confirm",
        "EQR",
        "shelf",
        "mid office",
        "front office",
        "seat",
        "smoke",
        "Grok Build",
        "Grok Bot",
        "SuperGrok",
        "Cursor",
        " PASS",
        " GO",
    ):
        assert banned not in plain


def test_mid_office_shell_is_persisted_and_not_an_env_flag(tmp_path, monkeypatch):
    monkeypatch.delenv("ESSAY_WRITER_CURSOR", raising=False)
    monkeypatch.delenv("ESSAY_WRITER_ROOT", raising=False)
    cloud = tmp_path / "workspace"
    cloud.mkdir()
    monkeypatch.setattr("essay_writer.paths.CLOUD_WORKSPACE", cloud)
    assert mid_office() == "Cursor"
    assert "Cursor is the mid office by default" in INSTRUCTIONS
    assert "subscription" in INSTRUCTIONS.casefold()
    assert "Cursor" not in SETUP_PROMPT
    assert "Grok Build" not in SETUP_PROMPT
    assert "Cursor" not in HONESTY
    monkeypatch.setenv("ESSAY_WRITER_CURSOR", "1")
    assert mid_office() == "Cursor"
    assert "agent --version" in setup_prompt(None).casefold()

    record_dir = cloud / "Essay-writer"
    record_dir.mkdir()
    record = record_dir / "setup.json"
    record.write_text(json.dumps({"app_path": str(tmp_path / "app")}), encoding="utf-8")
    assert mid_office() == "Cursor"
    assert "Cursor does not own this chat" in house_instructions(REVIEW_INSTRUCTIONS)
    assert "agent --version" in house_instructions().casefold()
    assert "grok version" not in house_instructions().casefold()

    record.write_text(
        json.dumps({"app_path": str(tmp_path / "app"), "mid_office_shell": "cursor"}),
        encoding="utf-8",
    )
    monkeypatch.setenv("ESSAY_WRITER_CURSOR", "0")
    assert mid_office() == "Cursor"
    writer = house_instructions()
    reviewer = house_instructions(REVIEW_INSTRUCTIONS)
    assert "Cursor is the mid office" in writer
    assert "Grok Build" not in writer
    assert "agent --version" in writer
    assert "do not install grok" in writer.casefold()
    assert "Cursor does not own this chat" in reviewer
    assert "Grok Build" not in reviewer
    assert "Grok Bot" in reviewer
    assert "agent --version" not in reviewer

    record.write_text(
        json.dumps({"app_path": str(tmp_path / "app"), "mid_office_shell": "grokbuild"}),
        encoding="utf-8",
    )
    assert mid_office() == "Grok Build"


def test_setup_pick_persists_and_branches_the_prompt(tmp_path, monkeypatch):
    from essay_writer.student_voice import for_word, say

    monkeypatch.setenv("ESSAY_WRITER_ROOT", str(tmp_path / "ledger"))
    assert setup_prompt(None) == CURSOR_SETUP_PROMPT
    assert setup_prompt("Cursor") == CURSOR_SETUP_PROMPT
    assert setup_prompt("SuperGrok") == SETUP_PROMPT
    assert setup_prompt("grok_build") == SETUP_PROMPT
    assert "grok version" in SETUP_PROMPT.casefold()
    assert "Cursor" not in SETUP_PROMPT
    cursor_prompt = setup_prompt(None)
    assert cursor_prompt == CURSOR_SETUP_PROMPT
    assert "agent --version" in cursor_prompt
    assert "do not install grok" in cursor_prompt.casefold()
    assert "grok version" not in cursor_prompt.casefold()
    assert "worker host is not a save folder" in cursor_prompt.casefold()

    spoken = say("Save on the Grok Bot cloud computer with SuperGrok, Grok Build, or Cursor.")
    assert "Grok Bot" not in spoken
    assert "SuperGrok" not in spoken
    assert "Grok Build" not in spoken
    assert "Cursor" not in spoken
    assert "cloud computer" in spoken
    worded = for_word("Grok Bot runs a dig for SuperGrok.")
    assert "Grok Bot" not in worded
    assert "SuperGrok" not in worded
    assert "dig" not in worded.casefold()
    assert "look-up" in worded
    named = for_word("Keep essay.dig and dig-card.json.")
    assert "dig" not in named.casefold()
    assert "look-up" in named

    app = tmp_path / "app"
    kb1 = tmp_path / "KB1"
    volume = tmp_path / "ResearchDrive" / "app"
    refused = essay_setup(app_path=str(volume), kb1_path=str(kb1), confirmPhase="draft")
    assert refused["ok"] is False
    assert "named save path on this computer" in refused["reason"]
    for banned in ("Grok Bot", "Grok Build", "Cursor", "SuperGrok"):
        assert banned not in refused["reason"]

    bad = essay_setup(
        app_path=str(app),
        kb1_path=str(kb1),
        confirmPhase="draft",
        mid_office_shell="grokbuild",
    )
    assert bad["ok"] is False
    assert not app.exists()
    assert not (tmp_path / "setup.json").exists()

    draft = essay_setup(
        app_path=str(app), kb1_path=str(kb1), confirmPhase="draft", acknowledge_skipped_searches=True
    )
    assert draft["ok"] is True
    assert draft["mid_office_shell"] == "cursor"
    assert draft["wrote"] is False
    confirmed = essay_setup(
        app_path=str(app), kb1_path=str(kb1), confirmPhase="confirm", acknowledge_skipped_searches=True
    )
    assert confirmed["ok"] is True
    assert confirmed["grokbuild"]["ran_install"] is False
    record_path = tmp_path / "setup.json"
    record = json.loads(record_path.read_text(encoding="utf-8"))
    assert record["mid_office_shell"] == "cursor"
    assert record["grokbuild"] == "grok"
    assert record["where"] == "user-local"
    assert confirmed["cursor_cli"]["command"] == "agent --version"

    essay_setup(
        app_path=str(app),
        kb1_path=str(kb1),
        confirmPhase="draft",
        mid_office_shell="SuperGrok",
        acknowledge_skipped_searches=True,
    )
    again = essay_setup(
        app_path=str(app),
        kb1_path=str(kb1),
        confirmPhase="confirm",
        mid_office_shell="SuperGrok",
        acknowledge_skipped_searches=True,
    )
    assert again["ok"] is True
    assert json.loads(record_path.read_text(encoding="utf-8"))["mid_office_shell"] == "grok_build"

    essay_setup(
        app_path=str(app),
        kb1_path=str(kb1),
        confirmPhase="draft",
        mid_office_shell="Cursor",
        acknowledge_skipped_searches=True,
    )
    picked = essay_setup(
        app_path=str(app),
        kb1_path=str(kb1),
        confirmPhase="confirm",
        mid_office_shell="Cursor",
        acknowledge_skipped_searches=True,
    )
    assert picked["ok"] is True
    assert picked["mid_office_shell"] == "cursor"
    assert picked["cursor_cli"]["ran_install"] is False
    assert picked["cursor_cli"]["worker_save_root"] is False
    assert picked["cursor_cli"]["command"] == "agent --version"
    assert picked["grokbuild"]["ran_install"] is False
    stored = json.loads(record_path.read_text(encoding="utf-8"))
    assert stored["mid_office_shell"] == "cursor"
    assert stored["grokbuild"] == "grok"
    assert stored["where"] == "user-local"

    for door in DOORS:
        for banned in ("Grok Bot", "Grok Build", "SuperGrok", "Cursor"):
            assert banned not in door["description"]
    helpers = (REPO / "docs" / "academic-writing-helpers.md").read_text(encoding="utf-8")
    assert "Cursor" not in helpers
    assert "the writer uses the" in helpers.casefold()


def test_fail_demo_writer_cannot_clear_a_review(workspace):
    assert writer_clears_left() == 0
    write_upload(workspace)
    brief = workspace / "Essay-writer" / ESSAY / "brief.md"
    brief.parent.mkdir(parents=True, exist_ok=True)
    brief.write_text(
        "Genre: report\nQuestion: How do chokepoints constrain the corridor?\n",
        encoding="utf-8",
    )
    seq = open_sequence()
    essay_research_summary(seq, TOPIC, "draft", [source()])
    essay_research_summary(seq, TOPIC, "confirm", [source()])
    essay_pack_validate(seq, pack(), "draft")
    essay_pack_validate(seq, pack(), "confirm")

    failed = essay_review(seq, "6", "FAIL", "draft")
    assert failed["wrote"] is False
    assert PERSON_BLOCK not in failed["next"]
    assert "writer_clears_left" not in failed
    stored = essay_review(seq, "6", "FAIL", "confirm")
    assert stored["stored"] is True
    assert PERSON_BLOCK not in json.dumps(stored)

    preview = essay_pack_render_word(seq, pack(), "draft", stage="6")
    assert preview["wrote"] is False
    assert preview["writer_clears_left"] == 0
    assert PERSON_BLOCK in preview["next"]
    blocked = essay_pack_render_word(seq, pack(), "confirm", stage="6")
    assert blocked["ok"] is False
    assert blocked["refused"] == "eqr"
    assert blocked["wrote"] is False
    assert blocked["writer_clears_left"] == 0
    assert PERSON_BLOCK in blocked["reason"]
    assert not (workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output" / f"{seq}.docx").exists()

    writer = essay_review(seq, "6", "PASS", "confirm", seat="writer")
    assert writer["ok"] is False
    assert writer["writer_clears_left"] == 0
    assert PERSON_BLOCK in writer["reason"]
    assert writer_clears_left() == 0

    constitution = "This is a report. How do chokepoints constrain the corridor?"
    passed = essay_review(seq, "6", "PASS", "draft", constitution=constitution)
    assert passed["seat"] == "reviewer"
    assert PERSON_BLOCK not in passed["next"]
    assert "writer_clears_left" not in passed
    confirmed = essay_review(seq, "6", "PASS", "confirm", constitution=constitution)
    assert confirmed["stored"] is True
    assert PERSON_BLOCK not in json.dumps(confirmed)
    written = essay_pack_render_word(seq, pack(), "confirm", stage="6")
    assert written["wrote"] is True


def test_research_and_write_seats_read_the_notes_and_reviewer_does_not():
    helpers = (REPO / "docs" / "academic-writing-helpers.md").read_text(encoding="utf-8")
    harvard = (REPO / "docs" / "referencing-harvard.md").read_text(encoding="utf-8")
    assert HELPERS_MARK in helpers
    assert HARVARD_MARK in harvard
    assert HARVARD_MARK not in helpers
    assert HELPERS_MARK in INSTRUCTIONS
    assert HARVARD_MARK in INSTRUCTIONS
    research = next(door["description"] for door in DOORS if door["tool"] == "essay.research.summary")
    full = next(door["description"] for door in DOORS if door["tool"] == "essay.research.full")
    write = next(door["description"] for door in DOORS if door["tool"] == "essay.pack.render_word")
    validate = next(door["description"] for door in DOORS if door["tool"] == "essay.pack.validate")
    for description in (research, full, write, validate):
        assert HELPERS_MARK not in description
        assert HARVARD_MARK not in description
        assert "server instructions" in description
    assert HELPERS_MARK not in REVIEW_INSTRUCTIONS
    assert HARVARD_MARK not in REVIEW_INSTRUCTIONS
    assert HELPERS_MARK not in house_instructions(REVIEW_INSTRUCTIONS)
    assert seat_reading("reviewer") == ""
    assert HELPERS_MARK in seat_reading("research")
    assert HARVARD_MARK not in seat_reading("research")
    assert HARVARD_MARK in seat_reading("write")


def test_missing_excerpt_list_is_measurement_not_invent_forbid(workspace):
    write_upload(workspace)
    seq = open_sequence()
    bare = source()
    bare.pop("quote")
    essay_research_summary(seq, TOPIC, "draft", [bare])
    essay_research_summary(seq, TOPIC, "confirm", [bare])
    payload = pack()
    payload["sources"] = [bare]
    result = essay_pack_validate(seq, payload, "draft")
    assert result["ok"] is False
    assert result["refused"] == "measurement"
    assert result["invent_forbid"] == "pass"
    assert result["stored"] is False
    assert result["wrote"] is False
    assert all(item["code"] != "invent-forbid" for item in result["blockers"])


def test_excerpt_not_in_the_file_is_invent_forbid(workspace):
    write_upload(workspace)
    seq = open_sequence()
    essay_research_summary(seq, TOPIC, "draft", [source()])
    essay_research_summary(seq, TOPIC, "confirm", [source()])
    payload = pack()
    payload["excerpts"] = ["This excerpt was never in the gathered file."]
    result = essay_pack_validate(seq, payload, "draft")
    assert result["refused"] == "invent-forbid"
    assert result["invent_forbid"] == "refuse"
    assert result["stored"] is False
    locked = pack()
    locked["excerpts"] = ["The Strait handles a large share of flows."]
    ok = essay_pack_validate(seq, locked, "draft")
    assert ok["passed"] is True
    assert ok.get("refused") is None


def test_intake_still_names_the_new_essay_without_a_dig_label(workspace):
    incoming = workspace / "unassigned" / "fresh.docx"
    from docx import Document

    document = Document()
    document.add_paragraph("The corridor depends on a narrow strait and a named supply chain.")
    document.add_paragraph("Reference List")
    document.add_paragraph("Patel, A. (2019) Chokepoints and crude.")
    incoming.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(incoming))
    draft = essay_intake("fresh.docx", "draft", "MOD1011AA", create=True)
    assert "No saved research trail" in draft["next"]
    assert "look-up" in draft["next"] or "source card" in draft["next"]
    assert "dig" not in draft["next"].casefold()
