"""Reviewer seat. Stable ids: AC1, AC2, AC3, AC4, AC5.

Independent-review findings land in Documents/<project>/, not the stage 10 folder.
A verdict-only pass does not clear. An unmounted Documents folder invents no file.
"""

from pathlib import Path

from docx import Document
from helpers import (
    ESSAY,
    ensure_word_refs,
    mac_handoff_sequence_dir,
    open_sequence,
    pack,
    source,
    write_upload,
)

from essay_writer.build_office import stage_launch
from essay_writer.cursor_office import cursor_stage_launch
from essay_writer.late_chain import findings_place, write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.review import essay_review, writer_clears_left
from essay_writer.server import HANDLERS, REVIEW_HANDLERS, build_review_server
from essay_writer.specialists.eqr import (
    BOTH_PLACES,
    CLAIM_CHECKS,
    NORTH_STAR,
    REVIEW_PROMPT,
    findings_note_clears,
)
from essay_writer.stage_llm import stage_settings
from essay_writer.state import load_state, save_state
from essay_writer.validate import essay_pack_validate
from essay_writer.word import essay_pack_render_word

REPO = Path(__file__).resolve().parents[1]
QUESTION = "How do chokepoints constrain the corridor?"
CONSTITUTION = f"This is a report. {QUESTION}"
STOLEN = "The opened paper states a measured corridor limit today."


def _brief() -> None:
    folder = essay_dir(ESSAY)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(
        f"Genre: report\nQuestion: {QUESTION}\n",
        encoding="utf-8",
    )


def _docx(path: Path, paragraphs: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for line in paragraphs:
        document.add_paragraph(line)
    document.save(str(path))


def _pdf(sentence: str) -> bytes:
    stream = f"({sentence})".encode("latin-1")
    return (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Page /Contents 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream\nendobj\n"
        b"trailer\n<<>>\n%%EOF\n"
    )


def _plant_pdf(workspace, seq: str) -> None:
    pdf = workspace / "uploads" / ESSAY / "paper.pdf"
    pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf.write_bytes(_pdf(STOLEN))
    state = load_state(seq)
    state["research"] = {
        "summary": {
            "sources": [
                {
                    "id": "src-pdf",
                    "locator": f"uploads/{ESSAY}/paper.pdf",
                    "title": "Opened paper",
                }
            ]
        }
    }
    save_state(seq, state)


def _student_word(seq: str, paragraphs: list[str]) -> Path:
    path = work_dir(seq) / "output" / f"{seq}.docx"
    _docx(path, paragraphs)
    return path


def test_ac1_findings_docx_under_project_folder_not_stage10(workspace, monkeypatch, tmp_path):
    """E2E #15. Findings Word is Documents/<project>/, not Documents/<project>-10/."""
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    _brief()
    seq = open_sequence()
    _student_word(seq, [CONSTITUTION + " (Smith, 2020).", "Student copy of the finished draft."])
    draft = essay_review(seq, "9", "PASS", "draft", constitution=CONSTITUTION)
    assert draft["ok"] is True, draft
    stored = essay_review(seq, "9", "PASS", "confirm", constitution=CONSTITUTION)
    assert stored["stored"] is True, stored
    note = documents / ESSAY / f"{ESSAY}-findings.docx"
    assert note.is_file()
    assert findings_place(seq) == f"Documents/{ESSAY}/{ESSAY}-findings.docx"
    assert stored["findings_file"]["place"] == findings_place(seq)
    seq_mac = mac_handoff_sequence_dir(documents, seq)
    assert (seq_mac / f"{ESSAY}-findings.docx").is_file()

    ensure_word_refs(seq)
    final = write_stage10(seq)
    assert final["copied"] is True, final
    assert (seq_mac / f"{ESSAY}-final.docx").is_file()
    assert (seq_mac / f"{ESSAY}-comments.docx").is_file()
    assert (seq_mac / f"{ESSAY}-findings.docx").is_file()
    assert note.is_file()
    assert final["findings_place"] == findings_place(seq)
    text = "\n".join(paragraph.text for paragraph in Document(str(note)).paragraphs)
    assert text.splitlines()[0] == "Findings"
    assert NORTH_STAR in text
    assert f"Documents/Essay Bot/{ESSAY}/{seq}/{ESSAY}-final.docx" in text
    assert BOTH_PLACES in text
    assert "the saved Word file" in text


def test_ac2_review_opens_student_word_documents_copy_and_pdfs(workspace, monkeypatch, tmp_path):
    """E2E #16. Review opens both Word files, claim-checks opened PDFs, and writes the note."""
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    _brief()
    seq = open_sequence()
    _plant_pdf(workspace, seq)
    student = [
        "Student copy of the finished draft.",
        CONSTITUTION,
        STOLEN,
    ]
    _student_word(seq, student)
    copied = essay_review(seq, "9", "PASS", "draft", constitution=CONSTITUTION)
    assert copied["ok"] is False, copied
    assert copied["stored"] is not True
    assert STOLEN in " ".join(copied.get("matches") or [])

    _student_word(seq, ["Student copy of the finished draft.", CONSTITUTION + " (Smith, 2020)."])
    _docx(
        mac_handoff_sequence_dir(documents, seq) / f"{ESSAY}-final.docx",
        ["Harbour copy of the finished draft.", CONSTITUTION],
    )
    draft = essay_review(seq, "9", "PASS", "draft", constitution=CONSTITUTION)
    assert draft["ok"] is True, draft
    assert draft["word_opened"] is True
    findings = draft["findings"]
    rows = [line.strip() for line in findings.splitlines() if line.strip()]
    assert rows[0] == NORTH_STAR
    assert "mark scheme" in rows[1].casefold()
    assert rows[2].startswith("Format:")
    assert "assignment" in rows[3].casefold()
    assert rows.index(CLAIM_CHECKS) > 3
    assert "Opened the saved Word file" in findings
    assert "Student copy of the finished draft." in findings
    assert f"Opened Documents/Essay Bot/{ESSAY}/{seq}/{ESSAY}-final.docx." in findings
    assert "Harbour copy of the finished draft." in findings
    assert BOTH_PLACES in findings
    assert "Opened PDF paper.pdf." in findings
    assert "Claim checks against the opened PDFs." in findings
    stored = essay_review(seq, "9", "PASS", "confirm", constitution=CONSTITUTION)
    assert stored["stored"] is True, stored
    note = documents / ESSAY / f"{ESSAY}-findings.docx"
    assert note.is_file()
    assert (mac_handoff_sequence_dir(documents, seq) / f"{ESSAY}-findings.docx").is_file()
    file_text = "\n".join(paragraph.text for paragraph in Document(str(note)).paragraphs)
    assert BOTH_PLACES in file_text
    assert "Opened PDF paper.pdf." in file_text
    assert "Harbour copy of the finished draft." in file_text
    assert "Student copy of the finished draft." in file_text


def test_ac3_verdict_only_pass_does_not_clear(workspace):
    """A one-line or verdict-only stored PASS does not clear the Word file."""
    assert findings_note_clears("PASS") is False
    assert findings_note_clears("The independent review recorded a pass.") is False
    bare_north = (
        f"{NORTH_STAR}\n"
        "No mark scheme is on this module.\n"
        "Format: the brief does not name one.\n"
        "The draft does the assignment.\n"
    )
    assert findings_note_clears(bare_north) is False
    write_upload(workspace)
    _brief()
    seq = open_sequence()
    from essay_writer.research import essay_research_summary

    essay_research_summary(seq, "chokepoints", "draft", [source()])
    essay_research_summary(seq, "chokepoints", "confirm", [source()])
    body = pack()
    assert essay_pack_validate(seq, body, "draft")["ok"] is True
    assert essay_pack_validate(seq, body, "confirm")["stored"] is True
    state = load_state(seq)
    state["eqr"] = {
        "9": {"verdict": "PASS", "seat": "reviewer", "stage": "9", "findings": "PASS"},
        "6": {
            "verdict": "PASS",
            "seat": "reviewer",
            "stage": "6",
            "findings": "The independent review recorded a pass.",
        },
    }
    save_state(seq, state)
    from essay_writer.review import eqr_verdict

    assert eqr_verdict(seq, "9") is None
    assert eqr_verdict(seq, "6") is None
    waiting = essay_pack_render_word(seq, body, "draft", stage="9")
    assert waiting["wrote"] is False
    blocked = essay_pack_render_word(seq, body, "confirm", stage="9")
    assert blocked["wrote"] is False
    assert blocked["refused"] == "eqr"
    assert not (work_dir(seq) / "output" / f"{seq}.docx").exists()


def test_ac4_documents_unmounted_refuses_without_inventing_a_file(workspace, monkeypatch):
    """Documents not mounted: refuse the findings file and do not invent one."""
    monkeypatch.delenv("ESSAY_WRITER_MAC_DIR", raising=False)
    _brief()
    seq = open_sequence()
    _student_word(seq, [CONSTITUTION + " (Smith, 2020)."])
    drafted = essay_review(seq, "9", "PASS", "draft", constitution=CONSTITUTION)
    assert drafted["ok"] is True, drafted
    assert drafted["word_opened"] is True
    stored = essay_review(seq, "9", "PASS", "confirm", constitution=CONSTITUTION)
    assert stored["stored"] is True, stored
    findings_file = stored["findings_file"]
    assert findings_file["wrote"] is False
    assert findings_file["ok"] is False
    assert findings_file["refused"] == "documents"
    assert findings_file["reason"] == "mac-not-mounted"
    assert findings_file["place"] == ""
    ensure_word_refs(seq)
    final = write_stage10(seq)
    assert final["reason"] == "mac-not-mounted"
    assert final["findings_place"] == ""
    assert final["findings_file"]["refused"] == "documents"
    assert final["copied"] is False
    sequence_note = work_dir(seq) / "output" / f"{ESSAY}-findings.docx"
    assert sequence_note.is_file()
    assert findings_file.get("chat_docx", {}).get("name") == f"{ESSAY}-findings.docx"
    assert list(workspace.rglob("*findings.docx")) == [sequence_note]
    assert not (workspace / "Documents").exists()


def test_ac5_build_covered_reviewer_seat_writer_cannot_self_pass(workspace):
    """Build covered. The writer server cannot self-pass. essay.review stays the reviewer seat."""
    assert "essay.review" not in HANDLERS
    assert set(REVIEW_HANDLERS).isdisjoint(HANDLERS)
    assert REVIEW_HANDLERS["essay.review"] is essay_review
    assert writer_clears_left() == 0
    settings = stage_settings("eqr")
    assert REVIEW_PROMPT in settings["rules"]
    assert "Documents/<project>/<project>-findings.docx" in settings["rules"]
    assert settings["fast"] is False
    grok = stage_launch("eqr")
    cursor = cursor_stage_launch("eqr")
    assert grok["source"] == cursor["source"]
    assert "Documents/<project>/<project>-findings.docx" in " ".join(grok["command"])
    assert cursor["settings"]["rules"] == settings["rules"]
    reviewer = build_review_server()
    tools = list(reviewer._tool_manager.list_tools())
    assert [tool.name for tool in tools] == ["essay.review"]
    assert "Documents/<project>/<project>-findings.docx" in tools[0].description
    seq = open_sequence()
    refused = essay_review(seq, "9", "PASS", "confirm", seat="writer")
    assert refused["ok"] is False
    assert refused["refused"] == "seat"
    assert refused.get("stored") is not True
