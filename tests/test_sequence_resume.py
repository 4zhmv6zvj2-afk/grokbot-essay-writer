"""Last-saved resume, bare-module refusal, atomic clone, and document id rewrite."""

import json
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from helpers import ESSAY, TOPIC, open_sequence, pack, source, write_upload

from essay_writer.health import essay_health
from essay_writer.intake import essay_intake
from essay_writer.paths import essay_dir, read_last_saved, sequence_dir
from essay_writer.projects import (
    essay_project_archive,
    essay_project_unarchive,
    essay_projects,
    essay_sequence_new,
    save_sequence_state,
)
from essay_writer.research import essay_research_summary
from essay_writer.shelf import essay_shelf_put
from essay_writer.validate import essay_pack_validate
from essay_writer.word import essay_pack_render_word

MODULE = "MOD1011AA"


def _confirm_new(parent: str | None = None) -> dict:
    draft = essay_sequence_new(MODULE, "draft", parent)
    assert draft["ok"] is True
    assert draft["wrote"] is False
    return essay_sequence_new(MODULE, "confirm", parent)


def _add_textbox(paragraph, text: str) -> None:
    paragraph._p.append(
        parse_xml(
            f"""
            <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                 xmlns:v="urn:schemas-microsoft-com:vml">
              <w:pict>
                <v:shape style="width:120pt;height:24pt">
                  <v:textbox>
                    <w:txbxContent>
                      <w:p>
                        <w:r><w:t>{text}</w:t></w:r>
                      </w:p>
                    </w:txbxContent>
                  </v:textbox>
                </v:shape>
              </w:pict>
            </w:r>
            """
        )
    )


def _rich_docx(path: Path, token: str, longer: str) -> None:
    document = Document()
    document.core_properties.title = token
    document.core_properties.subject = longer
    document.core_properties.keywords = f"see {token}"
    document.add_paragraph(token)
    split = document.add_paragraph()
    split.add_run(token[:-1])
    split.add_run(token[-1])
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = f"{token} and {longer} sequences/1/output sequences/10"
    section = document.sections[0]
    section.header.paragraphs[0].text = token
    section.footer.paragraphs[0].text = token
    box = document.add_paragraph()
    _add_textbox(box, token)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _visible_text(path: Path) -> str:
    document = Document(str(path))
    chunks = [
        document.core_properties.title or "",
        document.core_properties.subject or "",
        document.core_properties.keywords or "",
    ]
    parts = [document.element]
    for rel in document.part.rels.values():
        if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"):
            parts.append(rel.target_part.element)
    tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    for part in parts:
        for node in part.iter(tag):
            chunks.append(node.text or "")
    return "\n".join(chunks)


def test_continue_lists_every_module_and_shows_a_shared_id_once(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    assert _confirm_new()["sequence_id"] == f"{MODULE}-2"
    (workspace / "uploads" / "MOD1008AA").mkdir()
    rows = essay_projects("continue")["rows"]
    assert rows[0] == {
        "essay_id": "MOD1008AA",
        "next": "I can't find a sequence for MOD1008AA. Start a new one?",
    }
    assert rows[1]["essay_id"] == f"{MODULE}-2"
    assert "module" not in rows[1]
    assert "sequence_id" not in rows[1]
    assert rows[1]["sequence"] == 2
    assert rows[1]["stage"] == "open"
    assert rows[1]["roll_forward_ask"] is True
    assert "last_sequential_id" not in rows[1]


def test_continue_and_health_follow_last_saved_not_the_highest_folder(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    assert _confirm_new()["sequence_id"] == f"{MODULE}-2"
    ten = sequence_dir(MODULE, 10)
    ten.mkdir(parents=True)
    (ten / "state.json").write_text("{}\n", encoding="utf-8")
    save_sequence_state(
        MODULE,
        2,
        {
            "essay_id": f"{MODULE}-2",
            "sequence": 2,
            "sequence_id": f"{MODULE}-2",
            "research": {"summary": {"topic": "the one in hand"}, "full": None},
            "validations": {},
            "word": {},
            "revisions": [],
            "shelf": {},
        },
    )
    assert read_last_saved(MODULE) == 2
    row = essay_projects("continue")["rows"][0]
    assert row["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in row
    assert row["sequence"] == 2
    assert row["stage"] == "research"
    assert row["last_sequential_id"] == f"{MODULE}-10"
    assert row["last_sequential"] == 10
    pack_row = essay_health()["packs"][0]
    assert pack_row["stage"] == "research"
    assert pack_row["essay_id"] == f"{MODULE}-2"
    assert "module" not in pack_row
    assert "sequence_id" not in pack_row

    draft = essay_sequence_new(MODULE, "draft")
    assert draft["parent"] == f"{MODULE}-2"
    assert draft["sequence_id"] == f"{MODULE}-11"
    opened = essay_sequence_new(MODULE, "confirm")
    assert opened["sequence_id"] == f"{MODULE}-11"
    assert opened["parent"] == f"{MODULE}-2"
    child = json.loads((sequence_dir(MODULE, 11) / "state.json").read_text(encoding="utf-8"))
    assert child["branched_from"] == f"{MODULE}-2"
    assert read_last_saved(MODULE) == 11


def test_corrupt_state_is_unreadable_on_continue_and_health(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    (sequence_dir(MODULE, 1) / "state.json").write_text("{", encoding="utf-8")
    row = essay_projects("continue")["rows"][0]
    assert row["essay_id"] == f"{MODULE}-1"
    assert "sequence_id" not in row
    assert row["state"] == "unreadable"
    assert "stage" not in row
    assert "won't treat it as ready" in row["next"]
    pack_row = essay_health()["packs"][0]
    assert pack_row["state"] == "unreadable"
    assert pack_row["essay_id"] == f"{MODULE}-1"
    assert "sequence_id" not in pack_row
    assert pack_row.get("stage") != "open"


def test_unknown_and_bare_module_codes_do_not_write_off_the_sequence_tree(workspace):
    unknown = essay_research_summary("NOT-A-MODULE", TOPIC, "draft", [source(MODULE)])
    assert unknown["ok"] is False
    assert unknown["wrote"] is False
    assert "Start a new sequence" in unknown["reason"]
    assert not essay_dir("NOT-A-MODULE").exists()

    bare = essay_pack_render_word(MODULE, pack(MODULE), "draft")
    assert bare["ok"] is False
    assert bare["wrote"] is False
    assert f"I can't find a sequence for {MODULE}" in bare["reason"]
    assert not (essay_dir(MODULE) / "output").exists()
    assert not (essay_dir(MODULE) / "state.json").exists()
    assert not (essay_dir(MODULE) / ".drafts").exists()

    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    again = essay_research_summary(MODULE, TOPIC, "draft", [source(MODULE)])
    assert again["ok"] is False
    assert again["wrote"] is False
    assert f"{MODULE}-1" in again["reason"]
    assert "Tell me that id" in again["reason"]
    state = json.loads((sequence_dir(MODULE, 1) / "state.json").read_text(encoding="utf-8"))
    assert state["research"]["summary"] is None
    assert not (essay_dir(MODULE) / "state.json").exists()
    assert not (essay_dir(MODULE) / "output" / f"{MODULE}.docx").exists()
    research_drafts = list((essay_dir(MODULE) / ".drafts").glob("essay.research.summary/**"))
    assert research_drafts == []


def test_failed_clone_does_not_leave_a_folder_that_could_become_latest(workspace, monkeypatch):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"

    def boom(*_args, **_kwargs):
        raise RuntimeError("disk stopped")

    monkeypatch.setattr("essay_writer.projects._rewrite_tree", boom)
    draft = essay_sequence_new(MODULE, "draft")
    assert draft["ok"] is True
    refused = essay_sequence_new(MODULE, "confirm")
    assert refused["ok"] is False
    assert refused["wrote"] is False
    assert "didn't keep a partial copy" in refused["reason"]
    sequences = essay_dir(MODULE) / "sequences"
    names = sorted(path.name for path in sequences.iterdir())
    assert names == ["1"]
    assert read_last_saved(MODULE) == 1
    row = essay_projects("continue")["rows"][0]
    assert row["essay_id"] == f"{MODULE}-1"
    assert "sequence_id" not in row
    assert "last_sequential_id" not in row


def test_clone_rewrites_docx_id_in_table_header_footer_textbox_and_core(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    parent = sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx"
    _rich_docx(parent, f"{MODULE}-1", f"{MODULE}-10")
    assert _confirm_new()["sequence_id"] == f"{MODULE}-2"
    child = sequence_dir(MODULE, 2) / "output" / f"{MODULE}-2.docx"
    text = _visible_text(child)
    assert f"{MODULE}-10" in text
    assert "sequences/2/output" in text
    assert "sequences/10" in text
    assert "sequences/1/output" not in text
    opened = Document(str(child))
    body = [paragraph.text for paragraph in opened.paragraphs if paragraph.text.strip()]
    assert body[0] == f"{MODULE}-2"
    assert body[1] == f"{MODULE}-2"
    assert opened.tables[0].cell(0, 0).text == (
        f"{MODULE}-2 and {MODULE}-10 sequences/2/output sequences/10"
    )
    assert opened.sections[0].header.paragraphs[0].text == f"{MODULE}-2"
    assert opened.sections[0].footer.paragraphs[0].text == f"{MODULE}-2"
    assert opened.core_properties.title == f"{MODULE}-2"
    assert opened.core_properties.subject == f"{MODULE}-10"
    assert opened.core_properties.keywords == f"see {MODULE}-2"
    box = opened.element.xpath(".//w:txbxContent//w:t")
    assert [node.text for node in box] == [f"{MODULE}-2"]
    parent_text = _visible_text(parent)
    assert f"{MODULE}-1" in parent_text
    assert f"{MODULE}-2" not in parent_text


def test_intake_roll_forward_does_not_leave_the_parent_id_on_the_child_cover(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    output = sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx"
    document = Document()
    document.add_paragraph("The port handles containers.")
    document.add_paragraph("Reference List")
    document.add_paragraph("Smith, A. (2020) Port pricing. Journal of Ports.")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    incoming = workspace / "unassigned" / "draft.docx"
    _rich_docx(incoming, f"{MODULE}-1", f"{MODULE}-10")
    # Body paragraphs still have to look like an edited draft of the latest output.
    edited = Document(str(incoming))
    edited.add_paragraph("The port handles containers and barges.")
    edited.add_paragraph("Reference List")
    edited.add_paragraph("Smith, A. (2020) Port pricing. Journal of Ports.")
    edited.add_paragraph("Unknown, B. (1999) Missing paper.")
    edited.save(str(incoming))

    draft = essay_intake("draft.docx", "draft", MODULE)
    assert draft["kind"] == "output"
    assert draft["wrote"] is False
    filed = essay_intake("draft.docx", "confirm", MODULE)
    assert filed["ok"] is True
    assert filed["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in filed
    child = sequence_dir(MODULE, 2) / "output" / f"{MODULE}-2.docx"
    opened = Document(str(child))
    assert opened.core_properties.title == f"{MODULE}-2"
    assert opened.core_properties.subject == f"{MODULE}-10"
    assert [paragraph.text for paragraph in opened.paragraphs if paragraph.text.strip()][1] == f"{MODULE}-2"
    assert opened.tables[0].cell(0, 0).text.startswith(f"{MODULE}-2 ")
    assert f"{MODULE}-10" in opened.tables[0].cell(0, 0).text
    assert f"{MODULE}-1 " not in opened.tables[0].cell(0, 0).text
    assert opened.sections[0].header.paragraphs[0].text == f"{MODULE}-2"
    assert opened.sections[0].footer.paragraphs[0].text == f"{MODULE}-2"
    box = opened.element.xpath(".//w:txbxContent//w:t")
    assert [node.text for node in box] == [f"{MODULE}-2"]


def test_archived_module_kb1_notes_are_not_live_resume_targets(workspace):
    note = essay_dir(MODULE) / "kb1" / "ports.md"
    note.parent.mkdir(parents=True)
    note.write_text("A module note.\n", encoding="utf-8")
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    assert essay_health()["local_packs"] == [
        {"essay_id": MODULE, "name": f"Essay-writer/{MODULE}/kb1/ports.md"}
    ]
    essay_project_archive(MODULE, "draft")
    essay_project_archive(MODULE, "confirm")
    assert essay_health()["local_packs"] == []
    assert essay_health()["essay_ids"] == []
    assert essay_projects("continue")["rows"] == []
    essay_project_unarchive(MODULE, "draft")
    essay_project_unarchive(MODULE, "confirm")
    assert essay_health()["local_packs"] == [
        {"essay_id": MODULE, "name": f"Essay-writer/{MODULE}/kb1/ports.md"}
    ]


def test_shelf_draft_names_the_sequence_shelf(workspace):
    write_upload(workspace, ESSAY)
    seq = open_sequence(ESSAY)
    essay_research_summary(seq, TOPIC, "draft", [source()])
    essay_research_summary(seq, TOPIC, "confirm", [source()])
    essay_pack_validate(seq, pack(), "draft")
    essay_pack_validate(seq, pack(), "confirm")
    essay_pack_render_word(seq, pack(), "draft")
    written = essay_pack_render_word(seq, pack(), "confirm")
    assert written["path"] == f"Essay-writer/{ESSAY}/sequences/1/output/{seq}.docx"
    assert not (essay_dir(ESSAY) / "output" / f"{ESSAY}.docx").exists()
    draft = essay_shelf_put(seq, pack(), "draft")
    assert draft["wrote"] is False
    assert draft["path"] == f"Essay-writer/{ESSAY}/sequences/1/shelf/{draft['pack_id']}.json"
    assert not (essay_dir(ESSAY) / "shelf").exists()
