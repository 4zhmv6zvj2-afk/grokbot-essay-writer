"""Basics lines before confirm. Stable ids: AC1, AC2, AC3, AC4."""

from pathlib import Path

from docx import Document
from helpers import (
    ESSAY,
    open_sequence,
    pack,
    picked,
    plant_gathered_source,
    source,
    write_upload,
)

from essay_writer.build_office import stage_launch
from essay_writer.cursor_office import cursor_stage_launch
from essay_writer.late_chain import chat_leaks, completed_draft_chat
from essay_writer.paths import essay_dir, work_dir
from essay_writer.review import eqr_verdict, essay_review
from essay_writer.server import HANDLERS
from essay_writer.specialists.eqr import (
    CLAIM_CHECKS,
    NORTH_STAR,
    findings_note_clears,
    write_findings_note,
)
from essay_writer.specialists.grok_bot import BASICS_FIRST_LINES, basics_pass_lines
from essay_writer.specialists.writer import FIRST_DRAFT_NOTE
from essay_writer.stage_llm import stage_settings
from essay_writer.stage_track import essay_stage
from essay_writer.state import load_state, save_state
from essay_writer.validate import essay_pack_validate
from essay_writer.word import essay_pack_render_word
from essay_writer.writing_basics import CHECKLISTS

REPO = Path(__file__).resolve().parents[1]
QUESTION = "How should the firm manage the corridor?"
OTHER = "RPT6001AA"
OTHER_QUESTION = "Should the council publish the recycling figure?"
QUOTE = "The Strait handles a large share of flows."
COVERS = (
    "Smith, J. (2020) Chokepoints and crude.\n"
    f"{QUOTE}\n"
    "The review judges aim, method, and a single-corridor limit.\n"
)


def _brief(module: str, text: str) -> None:
    folder = essay_dir(module)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(text, encoding="utf-8")


def test_ac1_basics_first_lines_before_confirm(workspace):
    overlay = stage_settings("4a")["prompt_overlay"]
    rules = stage_settings("4a")["rules"]
    assert overlay.startswith("Before the save, the chat opens with a short Basics pass.")
    assert overlay.startswith(BASICS_FIRST_LINES)
    walk_at = overlay.index("Walk this Basics checklist internally before the save.")
    head = overlay[:walk_at]
    assert "mark scheme scores" in head
    assert "format the brief names" in head
    assert "assignment question" in head
    assert rules.index("mark scheme scores") < rules.index(
        "Walk this Basics checklist internally before the save."
    )
    bot = (REPO / "docs" / "templates" / "bot.md").read_text(encoding="utf-8")
    assert "short Basics pass" in bot
    assert "Chat for this step stays quiet" in " ".join(CHECKLISTS["5"])
    assert "Do not paste this checklist into the chat" in " ".join(CHECKLISTS["5"])

    _brief(ESSAY, f"Genre: report\nQuestion: {QUESTION}\n")
    seq = open_sequence()
    lines = basics_pass_lines(seq)
    assert [line.split(":", 1)[0] for line in lines] == [
        "No mark scheme is on this module.",
        "Format",
        "Question",
    ]
    assert lines[1] == "Format: report."
    assert lines[2] == f"Question: {QUESTION}"
    opened = "\n".join(lines)
    assert chat_leaks(opened) is False
    assert completed_draft_chat(opened) == opened

    buried = f"Delivered annotation.docx. This is a report. {QUESTION}"
    seeded_buried = essay_stage(seq, "1", "draft", buried)
    assert seeded_buried["ok"] is True, seeded_buried
    assert seeded_buried.get("basics_seeded") is True
    assert seeded_buried["basics_pass"] == list(lines)

    alone = essay_stage(seq, "1", "draft", "annotation.docx")
    assert alone["ok"] is True, alone
    assert alone.get("basics_seeded") is True
    assert not (work_dir(seq) / "output" / "annotation.docx").exists()

    draft = essay_stage(seq, "1", "draft", opened)
    assert draft["ok"] is True, draft
    assert draft["basics_pass"] == list(lines)
    saved = essay_stage(seq, "1", "confirm", opened, **picked(draft))
    assert saved["stored"] is True, saved

    for stage in ("2", "3", "4a", "4b"):
        drafted = essay_stage(seq, stage, "draft", opened)
        assert drafted["ok"] is True, drafted
        assert drafted["basics_pass"] == list(lines)
        if stage == "4a":
            plant_gathered_source(seq)
        kept = essay_stage(seq, stage, "confirm", opened, **picked(drafted))
        assert kept["stored"] is True, kept
    quiet = essay_stage(seq, "5", "draft", opened)
    assert quiet["stored"] is True, quiet
    assert quiet["quiet"] is True
    assert quiet["checklist_in_chat"] is False
    assert "basics_pass" not in quiet
    assert "Essay introduction moves from general to specific" not in quiet["next"]
    assert "Walk this Basics checklist" not in quiet["next"]
    assert quiet["chat"] == quiet["next"]


def test_ac2_findings_first_section_is_north_star(workspace, monkeypatch, tmp_path):
    write_upload(workspace)
    _brief(ESSAY, "Genre: report\nQuestion: How do chokepoints constrain the corridor?\n")
    seq = open_sequence()
    from essay_writer.research import essay_research_summary

    essay_research_summary(seq, "chokepoints", "draft", [source()])
    essay_research_summary(seq, "chokepoints", "confirm", [source()])
    body = pack()
    assert essay_pack_validate(seq, body, "draft")["ok"] is True
    assert essay_pack_validate(seq, body, "confirm")["stored"] is True

    writer = essay_review(seq, "6", "PASS", "confirm", seat="writer")
    assert writer["refused"] == "seat"
    assert writer["stored"] is not True

    state = load_state(seq)
    state["eqr"] = {
        "6": {"verdict": "PASS", "seat": "reviewer", "stage": "6", "findings": ""}
    }
    save_state(seq, state)
    assert findings_note_clears("") is False
    assert findings_note_clears("North-star judgement\n") is False
    assert eqr_verdict(seq, "6") is None
    waiting = essay_pack_render_word(seq, body, "draft", stage="6")
    assert waiting["wrote"] is False
    blocked = essay_pack_render_word(seq, body, "confirm", stage="6")
    assert blocked["wrote"] is False
    assert blocked["refused"] == "eqr"
    assert not (work_dir(seq) / "output" / f"{seq}.docx").exists()

    path = work_dir(seq) / "output" / f"{seq}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(
        "This is a report. How do chokepoints constrain the corridor? (Smith, 2020)."
    )
    document.save(str(path))
    constitution = "This is a report. How do chokepoints constrain the corridor?"
    draft = essay_review(seq, "9", "PASS", "draft", constitution=constitution)
    assert draft["ok"] is True, draft
    findings = draft["findings"]
    rows = [line.strip() for line in findings.splitlines() if line.strip()]
    assert rows[0] == NORTH_STAR
    assert "mark scheme" in rows[1].casefold()
    assert rows[2].startswith("Format:")
    assert "assignment" in rows[3].casefold()
    assert rows.index(CLAIM_CHECKS) > 3
    assert findings_note_clears(findings) is True
    assert "Claim checks against the opened PDFs." in findings
    assert "Opened the saved Word file" in findings

    unmounted = write_findings_note(seq, findings)
    assert unmounted["wrote"] is False
    assert unmounted["reason"] == "mac-not-mounted"

    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    stored = essay_review(seq, "9", "PASS", "confirm", constitution=constitution)
    assert stored["stored"] is True, stored
    assert eqr_verdict(seq, "9") == "PASS"
    note = documents / ESSAY / f"{ESSAY}-findings.docx"
    assert note.is_file()
    texts = [paragraph.text.strip() for paragraph in Document(str(note)).paragraphs if paragraph.text.strip()]
    assert texts[0] == "Findings"
    assert texts[1] == NORTH_STAR
    assert texts[2].strip()
    assert texts[2] != NORTH_STAR
    assert "mark scheme" in texts[2].casefold()
    assert texts[3].startswith("Format:")
    assert "assignment" in texts[4].casefold()
    assert texts.index(CLAIM_CHECKS) > 4


def test_ac3_stage5_argument_not_paper_digest(workspace):
    assert "argument that answers the question" in CHECKLISTS["5"][0]
    assert CHECKLISTS["5"][0] in stage_settings("5")["prompt_overlay"]
    assert "essay argument" in FIRST_DRAFT_NOTE
    assert "Do not paste them in as the body." in FIRST_DRAFT_NOTE
    assert "source material" in FIRST_DRAFT_NOTE
    assert "MOD1004AA" not in FIRST_DRAFT_NOTE

    write_upload(
        workspace,
        OTHER,
        text=(
            "Smith, J. (2020) Chokepoints and crude.\n"
            f"{QUOTE}\n"
            "Publish the recycling figure.\n"
        ),
    )
    _brief(OTHER, f"Genre: report\nQuestion: {OTHER_QUESTION}\n")
    scheme = essay_dir(OTHER) / "mark-scheme.md"
    scheme.write_text(
        "\n".join(
            [
                "Sections:",
                "- Purpose",
                "- Argument",
                "- Evidence",
                "High Distinction (80-100): the question is answered with evidence",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    seq = open_sequence(OTHER)
    from essay_writer.research import essay_research_summary

    essay_research_summary(seq, "recycling", "draft", [source(OTHER)])
    essay_research_summary(seq, "recycling", "confirm", [source(OTHER)])

    def reviewed(**extra):
        body = pack(OTHER)
        body["genre"] = "report"
        body["assessment"] = "Individual report"
        body["critical_reviews"] = [{"paper_id": "src-1", "markdown": COVERS}]
        body.update(extra)
        return body

    digest = reviewed(
        conclusion="open",
        title="High Distinction. Should the council publish the recycling figure?",
        assessment="Individual report",
        sections=[
            {"heading": "Purpose", "body": COVERS},
            {"heading": "Argument", "body": COVERS},
            {"heading": "Evidence", "body": COVERS},
        ],
    )
    refused = essay_pack_validate(seq, digest, "draft")
    assert refused["ok"] is False, refused
    assert refused["refused"] == "argument"
    assert refused["stored"] is False
    assert "does not answer the assignment question" in refused["reason"]
    assert "MOD1004AA" not in refused["reason"]

    cite = "(Smith, 2020)"
    argument = reviewed(
        conclusion="open",
        sections=[
            {
                "heading": "Purpose",
                "body": (
                    "The purpose of this report is that the council should publish "
                    "the recycling figure. High Distinction."
                ),
            },
            {
                "heading": "Argument",
                "body": (
                    "The argument is that the council should publish the recycling figure "
                    "because the strait concentrates flows."
                ),
            },
            {
                "heading": "Evidence",
                "body": (
                    f"{QUOTE} {cite} That evidence sits inside the argument that "
                    "the council should publish the recycling figure."
                ),
            },
        ],
    )
    wrong = reviewed(
        sections=[
            {"heading": "Paper one", "body": argument["sections"][0]["body"]},
            {"heading": "Paper two", "body": argument["sections"][1]["body"]},
            {"heading": "Paper three", "body": argument["sections"][2]["body"]},
        ],
    )
    headings = essay_pack_validate(seq, wrong, "draft")
    assert headings["refused"] == "argument"
    assert "mark-scheme section titles" in headings["reason"]
    assert "MOD1004AA" not in headings["reason"]

    invented = reviewed(sections=argument["sections"])
    invented["sources"] = [source(OTHER)]
    invented["sources"][0]["quote"] = "A sentence nobody gathered."
    fake = essay_pack_validate(seq, invented, "draft")
    assert fake["ok"] is False
    assert fake["refused"] in {"invent-forbid", "critical-review"}
    assert fake["stored"] is False

    passed = essay_pack_validate(seq, argument, "draft")
    assert passed["ok"] is True, passed
    assert passed["passed"] is True
    assert passed["stored"] is False


def test_ac4_build_covered_shares_essay_stage():
    assert HANDLERS["essay.stage"] is essay_stage
    item = stage_settings("5")
    build = stage_launch("5")
    cursor = cursor_stage_launch("5")
    assert build["source"] == "essay_writer.stage_llm.stage_settings"
    assert cursor["source"] == "essay_writer.stage_llm.stage_settings"
    assert build["settings"]["rules"] == item["rules"]
    assert cursor["settings"]["rules"] == item["rules"]
    assert item["prompt_overlay"].startswith(BASICS_FIRST_LINES)
    assert BASICS_FIRST_LINES in "\n".join(build["command"])
    assert cursor["settings"]["prompt_overlay"] == item["prompt_overlay"]
    assert "argument that answers the question" in item["rules"]
