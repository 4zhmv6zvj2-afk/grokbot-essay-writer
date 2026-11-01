"""RPT5001AA-2 cuts 1+2+3: refs gate, annotation stubs, assignment handoff.

AC1–AC5: references and Stage 4a annotation shape.
AC6–AC8: chat vs Mac vs papers.zip handoff (global).
"""

import base64
import json
from pathlib import Path

from docx import Document
from helpers import (
    SAMPLE_REFERENCE_LINE,
    TOPIC,
    open_sequence,
    picked,
    plant_gathered_source,
    write_upload,
)

from essay_writer.brief_contract import remember_registration
from essay_writer.harvard_4a import prose_issues
from essay_writer.late_chain import write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.research import essay_research_summary
from essay_writer.research_packs import ship_4a
from essay_writer.stage_track import essay_stage, load_pointer
from essay_writer.state import load_state, save_state

MODULE = "RPT5001AA"
NUMBER = "445566"
BRIEF = """
Assignment title: Individual reflective diary

Module title: Practice module
Module code: RPT5001AA

Aim: Understand the practice by reflecting on your own decisions.

Assignment task: Write four entries. Each entry names one decision and critically analyses it with theory and evidence.

Title page: module title, module code, assignment title, and the word count.
Include your student ID.
The name stays off the title page and the file.

Pages: Arial, size 12. 2.5 cm margins on all sides. Page numbers in the lower right.
"""
LOOK = "The diary reflects on one decision named in the brief."
JOURNAL_TEXT = (
    "Qin, Y. (2021) 'Stakeholder mapping', Journal of Practice, 8(1), pp. 2-9. "
    "doi:10.1000/stake\n"
    "The paper maps stakeholders for one corridor decision.\n"
)
GOOD = (
    "Qin (2021) mapped stakeholders for one corridor decision. "
    "The paper found that engagement shapes the outcome. "
    "That finding helps this assignment on reflective practice."
)
CITATION = (
    "Qin, Y. (2021) 'Stakeholder mapping', Journal of Practice, 8(1), pp. 2-9. "
    "Available at: doi:10.1000/stake."
)


def _references(path: Path) -> list[str]:
    texts = [
        paragraph.text.strip()
        for paragraph in Document(str(path)).paragraphs
        if paragraph.text.strip()
    ]
    if "References" not in texts:
        return []
    return texts[texts.index("References") + 1 :]


def _annotation_path(workspace, seq: str) -> Path:
    number = seq.rsplit("-", 1)[-1]
    return workspace / "Essay-writer" / MODULE / "sequences" / number / "output" / "annotation.docx"


def _diary_pack(seq: str, *, references_markdown: str | None = None) -> dict:
    pack = {
        "essay_id": seq,
        "title": "Individual reflective diary",
        "referencing": "harvard",
        "assessment": "Individual reflective diary",
        "conclusion": "noted",
        "sections": [
            {
                "heading": "One decision",
                "body": "A decision, with theory, on the practice the brief names.",
            }
        ],
    }
    if references_markdown is not None:
        pack["references_markdown"] = references_markdown
    return pack


def _brief() -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(BRIEF, encoding="utf-8")


def _park_stage10(seq: str) -> None:
    pointer = load_pointer(seq)
    pointer["current"] = "10"
    pointer["saved"] = ["1", "2", "3", "4a", "4b", "5", "6", "7", "8", "9"]
    pointer["closed"] = False
    pointer["chain_armed"] = False
    pointer["waiting_stage"] = None
    pointer["waiting_digest"] = None
    pointer["room_grant"] = None
    pointer["room_lock"] = False
    path = work_dir(seq) / "tracking" / "stage.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(pointer, indent=2) + "\n", encoding="utf-8")


def _plant_stage10_pack(seq: str, pack: dict) -> None:
    state = load_state(seq)
    state["saved_pack"] = pack
    save_state(seq, state)


def _advance_to_stage10(seq: str) -> None:
    plant_gathered_source(seq)
    for stage in ("1", "2", "3", "4a", "4b", "5", "6", "7", "8", "9"):
        # Stage 7 "all fine" opens Confirm. The confirm call must not send it again.
        draft_kwargs = {"answers": "all fine"} if stage == "7" else {}
        draft = essay_stage(seq, stage, "draft", LOOK, **draft_kwargs)
        assert draft["ok"] is True, (stage, draft)
        saved = essay_stage(seq, stage, "confirm", LOOK, **picked(draft))
        assert saved.get("stored") is True or stage == "5", (stage, saved)


def test_ac1_stage10_confirm_refuses_empty_references(workspace, monkeypatch, tmp_path):
    """AC1. No sources and no references_markdown: stage 10 does not close or write heading-only refs."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    _brief()
    seq = open_sequence(MODULE)
    remember_registration(seq, NUMBER)
    _park_stage10(seq)
    _plant_stage10_pack(seq, _diary_pack(seq))
    pointer = load_pointer(seq)
    assert pointer["current"] == "10"

    draft = essay_stage(seq, "10", "draft", LOOK)
    assert draft["ok"] is True, draft
    refused = essay_stage(seq, "10", "confirm", LOOK, **picked(draft))
    assert refused["ok"] is False, refused
    assert refused["refused"] == "empty-references"
    assert refused.get("closed") is False
    assert load_pointer(seq)["current"] == "10"

    stage_doc = work_dir(seq) / "output" / f"{seq}.docx"
    if stage_doc.is_file():
        assert _references(stage_doc) == []

    direct = write_stage10(seq)
    assert direct["ok"] is False, direct
    assert direct["refused"] == "empty-references"
    assert direct.get("sequence_written") is not True


def test_ac2_stage10_word_uses_references_markdown(workspace, monkeypatch, tmp_path):
    """AC2. Pack references_markdown prints under References when sources are empty."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    seq = open_sequence(MODULE)
    _advance_to_stage10(seq)
    lines = "\n".join([SAMPLE_REFERENCE_LINE, "Jones, A. (2019) 'Second line', Practice, 1(1), pp. 1-2."])
    _plant_stage10_pack(seq, _diary_pack(seq, references_markdown=lines))

    final = write_stage10(seq)
    assert final["ok"] is True, final
    assert final.get("sequence_written") is True

    stage_doc = work_dir(seq) / "output" / f"{seq}.docx"
    refs = _references(stage_doc)
    assert SAMPLE_REFERENCE_LINE in refs[0] or any(SAMPLE_REFERENCE_LINE in item for item in refs)
    assert any("Jones, A. (2019)" in item for item in refs)


def test_ac3_stage10_word_keeps_source_ctr_lines(workspace, monkeypatch, tmp_path):
    """AC3. With no references_markdown, Word still prints gathered Harvard lines."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)
    seq = open_sequence(MODULE)
    source = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Stakeholder mapping",
        "container": "Journal of Practice",
        "volume": "8",
        "issue": "1",
        "pages": "2-9",
        "doi": "10.1000/stake",
        "annotation": GOOD,
    }
    essay_research_summary(seq, TOPIC, "draft", [source])
    essay_research_summary(seq, TOPIC, "confirm", [source])
    _advance_to_stage10(seq)
    pack = _diary_pack(seq)
    pack["sources"] = [source]
    _plant_stage10_pack(seq, pack)

    final = write_stage10(seq)
    assert final["ok"] is True, final
    refs = _references(work_dir(seq) / "output" / f"{seq}.docx")
    joined = "\n".join(refs)
    assert "Qin, Y. (2021)" in joined
    assert "Stakeholder mapping" in joined


def _qin_source(**extra):
    row = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Stakeholder mapping",
        "container": "Journal of Practice",
        "volume": "8",
        "issue": "1",
        "pages": "2-9",
        "doi": "10.1000/stake",
        "quote": "The paper maps stakeholders for one corridor decision.",
    }
    row.update(extra)
    return row


def test_ac4_stage4a_confirm_refuses_surname_stub(workspace, monkeypatch, tmp_path):
    """AC4. Surname-only annotation refuses before annotation.docx is written."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)
    seq = open_sequence(MODULE)
    row = _qin_source(annotation="Qin.")
    state = load_state(seq)
    state["research"] = {"summary": {"sources": [row]}, "full": None}
    save_state(seq, state)

    for stage in ("1", "2", "3"):
        draft = essay_stage(seq, stage, "draft", LOOK)
        saved = essay_stage(seq, stage, "confirm", LOOK, **picked(draft))
        assert saved["stored"] is True, saved

    draft4a = essay_stage(seq, "4a", "draft", LOOK)
    assert draft4a["ok"] is True, draft4a
    refused = essay_stage(seq, "4a", "confirm", LOOK, **picked(draft4a))
    assert refused["ok"] is False, refused
    assert refused["refused"] == "annotation-prose"
    assert load_pointer(seq)["current"] == "4a"
    assert not _annotation_path(workspace, seq).exists()


def test_annotation_page_range_with_an_em_dash_is_checked():
    """An em-dash page range is read as the whole range, not only its first page."""
    source = {"author": "Qin, Y.", "year": "2021"}
    gathered = "The paper maps stakeholders for one corridor decision on page 4."
    prose = "Qin (2021, pp. 4—9) maps the stakeholders for one corridor decision."
    issues = prose_issues("src-1 annotation", source, gathered, prose)
    assert any("4—9" in item["reason"] for item in issues)
    hyphen = prose.replace("4—9", "4-9")
    assert any("4-9" in item["reason"] for item in prose_issues("src-1 annotation", source, gathered, hyphen))


def test_annotation_reads_a_page_cite():
    """A page-bearing cite for another author is not gathered for this source."""
    source = {"author": "Qin, Y.", "year": "2021"}
    gathered = "The paper maps stakeholders for one corridor decision on page 4."
    prose = (
        "Jones (2021, p. 4) maps stakeholders for one corridor decision. "
        "The paper explains the strait and the corridor decision."
    )
    issues = prose_issues("src-1 annotation", source, gathered, prose, annotation=True)
    assert any("jones, 2021" in item["reason"].casefold() for item in issues)


def test_ac5_stage4a_confirm_ships_annotation_docx(workspace, monkeypatch, tmp_path):
    """AC5. Full-sentence annotations still ship citation line plus prose paragraph."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)
    seq = open_sequence(MODULE)
    row = _qin_source(annotation=GOOD)
    essay_research_summary(seq, TOPIC, "draft", [row])
    essay_research_summary(seq, TOPIC, "confirm", [row])

    for stage in ("1", "2", "3"):
        draft = essay_stage(seq, stage, "draft", LOOK)
        saved = essay_stage(seq, stage, "confirm", LOOK, **picked(draft))
        assert saved["stored"] is True, saved

    draft4a = essay_stage(seq, "4a", "draft", LOOK)
    saved4a = essay_stage(seq, "4a", "confirm", LOOK, **picked(draft4a))
    assert saved4a["ok"] is True, saved4a
    assert saved4a["stored"] is True, saved4a
    path = _annotation_path(workspace, seq)
    assert path.is_file()
    text = "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    assert CITATION.split("'")[0].strip() in text or "Stakeholder mapping" in text
    assert "helps this assignment" in text
    assert text.count("Stakeholder mapping") >= 1


def test_ac6_stage10_and_4a_surface_word_in_chat(workspace, monkeypatch, tmp_path):
    """AC6. Final sequence Word and annotation.docx include chat docx payloads."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)
    seq = open_sequence(MODULE)
    remember_registration(seq, NUMBER)
    row = _qin_source(annotation=GOOD)
    essay_research_summary(seq, TOPIC, "draft", [row])
    essay_research_summary(seq, TOPIC, "confirm", [row])
    delivery = ship_4a(seq, [row], {row["id"]: GOOD})
    assert delivery["annotation"]["deliver"] == "chat"
    assert "docx_base64" in delivery["annotation"]

    _park_stage10(seq)
    lines = SAMPLE_REFERENCE_LINE
    _plant_stage10_pack(seq, _diary_pack(seq, references_markdown=lines))
    final = write_stage10(seq)
    assert final["ok"] is True, final
    chat = final["delivery"]["word"]
    assert chat["deliver"] == "chat"
    assert "docx_base64" in chat
    stage_doc = work_dir(seq) / "output" / f"{seq}.docx"
    assert base64.b64decode(chat["docx_base64"]) == stage_doc.read_bytes()

def test_ac7_unmounted_mac_still_delivers_word_in_chat(workspace, monkeypatch):
    """AC7. No Mac folder: no silent Mac copy, sequence Word still in chat."""
    monkeypatch.delenv("ESSAY_WRITER_MAC_DIR", raising=False)
    _brief()
    seq = open_sequence(MODULE)
    remember_registration(seq, NUMBER)
    _park_stage10(seq)
    _plant_stage10_pack(seq, _diary_pack(seq, references_markdown=SAMPLE_REFERENCE_LINE))
    final = write_stage10(seq)
    assert final["ok"] is True, final
    assert final.get("copied") is False
    assert final.get("reason") == "mac-not-mounted"
    chat = final["delivery"]["word"]
    assert "docx_base64" in chat
    assert not (workspace / "Documents").exists()


def test_ac8_papers_zip_is_path_only_not_chat_attached(workspace):
    """AC8. papers.zip is named on disk only; chat does not attach zip bytes."""
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)
    seq = open_sequence(MODULE)
    row = _qin_source(annotation=GOOD)
    delivery = ship_4a(seq, [row], {row["id"]: GOOD})
    papers = delivery["papers_zip"]
    assert papers.get("deliver") == "path"
    assert "zip_base64" not in papers
    assert papers.get("path")
    assert papers.get("name") == "papers.zip"
