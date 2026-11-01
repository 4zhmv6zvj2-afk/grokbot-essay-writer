"""Blocked shape. Stable ids: AC1, AC2, AC3, AC4.

Stage 4a annotations are a few full sentences. The Harvard citation line
stays exact. A fragment or a citation line with no sentences is refused.
Mark-scheme words stay; they are not scored away.
"""

import inspect

from docx import Document
from helpers import ESSAY, TOPIC, open_sequence, write_upload

from essay_writer.build_office import stage_launch, write_stage_deliverables
from essay_writer.cursor_office import cursor_stage_launch
from essay_writer.harvard_4a import (
    ANNOTATION_SENTENCES,
    FOUR_A_FORMAT,
    harvard_4a_line,
    harvard_4a_parts,
)
from essay_writer.research import essay_research_summary
from essay_writer.research_packs import ship_4a, write_annotation_docx
from essay_writer.server import HANDLERS
from essay_writer.shells import cursor_shell, front_office
from essay_writer.specialists import researcher
from essay_writer.stage_llm import stage_settings
from essay_writer.student_voice import (
    annotation_prose,
    prose_for_surface,
    student_surface_hits,
)
from essay_writer.writing_basics import basics_lines

JOURNAL_TEXT = (
    "Smith, J. (2020) Chokepoints and crude. Maritime Policy 12(3) pp. 4-9. doi:10.1000/strait\n"
    "The Strait handles a large share of flows.\n"
    "The authors confirm the limit in the gathered file.\n"
)
GOOD = (
    "Smith (2020) measured chokepoint risk on one corridor. "
    "The paper found that the strait carries a large share of flows. "
    "That finding helps this assignment on how a firm should manage the corridor. "
    "Scope 3 emissions stay in the note (indirect emissions the brief names)."
)
CITATION = (
    "Smith, J. (2020) 'Chokepoints and crude', Maritime Policy, 12(3), pp. 4-9. "
    "Available at: doi:10.1000/strait."
)


def _journal(**extra):
    row = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{ESSAY}/smith.md",
        "author": "Smith, J.",
        "year": "2020",
        "title": "Chokepoints and crude",
        "container": "Maritime Policy",
        "volume": "12",
        "issue": "3",
        "pages": "4-9",
        "doi": "10.1000/strait",
        "quote": "The Strait handles a large share of flows.",
    }
    row.update(extra)
    return row


def _paragraphs(path) -> list[str]:
    return [paragraph.text for paragraph in Document(str(path)).paragraphs]


def _annotation(workspace, seq: str):
    number = seq.rsplit("-", 1)[-1]
    return workspace / "Essay-writer" / ESSAY / "sequences" / number / "output" / "annotation.docx"


def _confirm(workspace, annotation: str):
    write_upload(workspace, text=JOURNAL_TEXT)
    seq = open_sequence()
    row = _journal(annotation=annotation)
    draft = essay_research_summary(seq, TOPIC, "draft", [row])
    assert draft["ok"] is True, draft
    saved = essay_research_summary(seq, TOPIC, "confirm", [row])
    assert saved["ok"] is True, saved
    return seq, saved


def test_ac1_full_sentence_annotation_prose(workspace):
    """What the paper did, what it found, and how it helps this assignment."""
    lines = basics_lines("4a")
    assert lines[1] == ANNOTATION_SENTENCES
    assert "what the paper did" in ANNOTATION_SENTENCES
    assert "what it found" in ANNOTATION_SENTENCES
    assert "how it helps this assignment" in ANNOTATION_SENTENCES
    assert "gloss" in ANNOTATION_SENTENCES
    assert "Do not delete those words." in ANNOTATION_SENTENCES
    assert "what the paper did" in FOUR_A_FORMAT
    assert "Do not delete those words." in FOUR_A_FORMAT
    overlay = stage_settings("4a")["prompt_overlay"]
    assert ANNOTATION_SENTENCES in overlay
    assert FOUR_A_FORMAT in overlay
    assert student_surface_hits(overlay) == []
    assert annotation_prose(GOOD) == GOOD
    assert prose_for_surface(GOOD, "annotation") == GOOD
    assert "Scope 3" in annotation_prose(GOOD)
    assert "chokepoint" in annotation_prose(GOOD)

    seq, saved = _confirm(workspace, GOOD)
    assert saved["sources"][0]["annotation"] == GOOD
    text = "\n".join(_paragraphs(_annotation(workspace, seq)))
    assert "measured chokepoint risk on one corridor." in text
    assert "The paper found that the strait carries a large share of flows." in text
    assert "helps this assignment" in text
    assert "Scope 3" in text
    assert "indirect emissions the brief names" in text
    assert "simple_words" not in inspect.getsource(annotation_prose)
    assert "simple_words" not in inspect.getsource(essay_research_summary)


def test_ac2_citation_line_unchanged(workspace):
    """Harvard Referencing via harvard_4a_parts. The sentences do not edit that line."""
    source = _journal()
    assert harvard_4a_line(source) == CITATION
    parts = harvard_4a_parts(source)
    assert "".join(piece for piece, _italic in parts).strip() == CITATION
    assert any(italic and piece == "Maritime Policy" for piece, italic in parts)

    seq, _saved = _confirm(workspace, GOOD)
    path = _annotation(workspace, seq)
    document = Document(str(path))
    citation = next(paragraph for paragraph in document.paragraphs if "Maritime Policy" in paragraph.text)
    assert citation.text == CITATION
    italic = [
        run.text
        for run in citation.runs
        if run.italic and "Maritime Policy" in (run.text or "")
    ]
    assert italic == ["Maritime Policy"]
    body = next(paragraph for paragraph in document.paragraphs if "helps this assignment" in paragraph.text)
    assert CITATION not in body.text
    assert harvard_4a_line(source) == CITATION


def test_ac3_no_stub_bibliography_refuses_fragments_and_cite_only(workspace):
    """A fragment or a citation line is refused. A later save keeps the sentences."""
    write_upload(workspace, text=JOURNAL_TEXT)
    seq = open_sequence()
    fragment = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_journal(annotation="narrow corridor and crude flows")],
    )
    assert fragment["ok"] is False
    assert fragment["refused"] == "annotation-prose"
    assert "fragment" in fragment["reason"]
    assert fragment["wrote"] is False
    assert not _annotation(workspace, seq).exists()

    cite_only = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_journal(annotation=CITATION)],
    )
    assert cite_only["ok"] is False
    assert cite_only["refused"] == "annotation-prose"
    assert "citation line" in cite_only["reason"]
    assert "full sentences" in cite_only["reason"]
    assert not _annotation(workspace, seq).exists()

    stub = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_journal(annotation="Smith, J. (2020) 'Chokepoints and crude'.")],
    )
    assert stub["ok"] is False
    assert stub["refused"] == "annotation-prose"
    assert "citation line" in stub["reason"]
    assert not _annotation(workspace, seq).exists()

    kept = _confirm(workspace, GOOD)
    saved_seq = kept[0]
    path = _annotation(workspace, saved_seq)
    from essay_writer.specialists.researcher import deliver_bibliography

    again = deliver_bibliography(saved_seq)
    assert again is not None
    text = "\n".join(_paragraphs(path))
    assert "measured chokepoint risk on one corridor." in text
    assert "helps this assignment" in text
    assert text.count(CITATION) == 1
    names = write_stage_deliverables(saved_seq)
    assert "annotation.docx" in names
    rebuilt = "\n".join(_paragraphs(path))
    assert "The paper found that the strait carries a large share of flows." in rebuilt
    assert rebuilt.count(CITATION) == 1


def test_ac4_build_covered_shared_annotation_path():
    """Cursor and Grok Build share the 4a rules and the annotation writer."""
    rules = stage_settings("4a")["rules"]
    assert stage_launch("4a")["settings"]["rules"] == rules
    assert cursor_stage_launch("4a")["settings"]["rules"] == rules
    assert ANNOTATION_SENTENCES in rules
    assert FOUR_A_FORMAT in rules
    assert HANDLERS["essay.research.summary"] is essay_research_summary
    assert cursor_shell()["handlers"]["essay.research.summary"] is HANDLERS["essay.research.summary"]
    assert front_office()["handlers"]["essay.research.summary"] is HANDLERS["essay.research.summary"]
    assert researcher.ship_4a is ship_4a
    assert "write_annotation_docx" in inspect.getsource(ship_4a)
    assert "deliver_saved_sources" in inspect.getsource(write_stage_deliverables)
    assert write_annotation_docx.__module__ == "essay_writer.research_packs"
    assert ship_4a.__module__ == "essay_writer.research_packs"
