"""Word deliverables. Stable ids: AC1, AC2, AC3."""

import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from helpers import ESSAY, open_sequence, pack, source, write_upload

from essay_writer.paths import essay_dir, sequence_number
from essay_writer.research import essay_research_summary
from essay_writer.research_packs import write_annotation_docx
from essay_writer.server import HANDLERS
from essay_writer.specialists.eqr import NORTH_STAR, write_findings_note
from essay_writer.specialists.writer import body_is_justified
from essay_writer.validate import essay_pack_validate
from essay_writer.word import TOOL, essay_pack_render_word


def _zip_xml(path) -> str:
    return zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")


def _essay_output_path(workspace, essay_id: str, seq: str):
    number = sequence_number(seq) or 1
    return (
        workspace
        / "Essay-writer"
        / essay_id
        / "sequences"
        / str(number)
        / "output"
        / f"{seq}.docx"
    )


def _write_essay_docx(workspace, payload=None, essay_id=ESSAY, *, seq=None, setup=True):
    if setup:
        write_upload(workspace, essay_id=essay_id)
        seq = open_sequence(essay_id)
        essay_research_summary(seq, "chokepoints", "draft", [source(essay_id)])
        essay_research_summary(seq, "chokepoints", "confirm", [source(essay_id)])
    if seq is None:
        raise ValueError("seq is required when setup is false")
    body = payload or pack(essay_id)
    essay_pack_validate(seq, body, "draft")
    essay_pack_validate(seq, body, "confirm")
    essay_pack_render_word(seq, body, "draft")
    written = essay_pack_render_word(seq, body, "confirm")
    path = _essay_output_path(workspace, essay_id, seq)
    return written, path, seq


def test_ac1_word_deliverables_use_heading_styles_not_markdown_hashes(workspace, tmp_path):
    payload = pack()
    payload["sections"][1]["heading"] = "## Supply chain context"
    payload["sections"][1]["body"] = (
        "One named supply chain on the corridor is in scope for this report (Smith, 2020).\n"
        "### Evidence\n"
        "The Strait handles a large share of flows through that chain (Smith, 2020)."
    )
    written, essay_path, _seq = _write_essay_docx(workspace, payload)
    assert written["wrote"] is True
    document = Document(str(essay_path))
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
    assert "Supply chain context" in headings
    assert "Evidence" in headings
    joined = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "##" not in joined
    assert "###" not in joined
    essay_xml = _zip_xml(essay_path)
    assert 'w:val="Heading1"' in essay_xml
    assert 'w:val="Heading2"' in essay_xml or "Evidence" in headings

    annotation_path = tmp_path / "annotation.docx"
    write_annotation_docx(
        annotation_path,
        [{"parts": [("Smith (2020) Title.", False)], "paragraph": "A short note."}],
        empty_note="No PDF was saved.",
    )
    annotation_xml = _zip_xml(annotation_path)
    assert "Annotated bibliography" in annotation_xml
    assert 'w:val="Heading1"' in annotation_xml
    assert "##" not in re.sub(r"<[^>]+>", "", annotation_xml)

    findings = write_findings_note(ESSAY, f"{NORTH_STAR}\nA line of findings.", documents=tmp_path)
    assert findings["wrote"] is True
    findings_xml = _zip_xml(tmp_path / "MOD1004AA" / "MOD1004AA-findings.docx")
    assert 'w:val="Heading1"' in findings_xml
    assert "Findings" in findings_xml


def test_ac2_house_left_align_and_brief_mark_scheme_win(workspace):
    folder = essay_dir(ESSAY)
    folder.mkdir(parents=True, exist_ok=True)
    assert body_is_justified(ESSAY) is False
    _written, path, seq = _write_essay_docx(workspace)
    document = Document(str(path))
    prose = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Normal"
        and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER
        and "strait" in paragraph.text.casefold()
    ]
    assert prose
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT for paragraph in prose)

    (folder / "brief.md").write_text(
        "Genre: report\n"
        "Question: How do chokepoints constrain the corridor?\n"
        "Alignment: justified\n",
        encoding="utf-8",
    )
    assert body_is_justified(ESSAY) is True
    _written, justified_path, _seq = _write_essay_docx(workspace, seq=seq, setup=False)
    assert _written["wrote"] is True
    justified = Document(str(justified_path))
    body = [
        paragraph
        for paragraph in justified.paragraphs
        if paragraph.style.name == "Normal"
        and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER
        and "strait" in paragraph.text.casefold()
    ]
    assert body
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY for paragraph in body)

    (folder / "mark-scheme.md").write_text("Alignment: left\n", encoding="utf-8")
    assert body_is_justified(ESSAY) is False
    _written, scheme_path, _seq = _write_essay_docx(workspace, seq=seq, setup=False)
    assert _written["wrote"] is True
    again = Document(str(scheme_path))
    reset = [
        paragraph
        for paragraph in again.paragraphs
        if paragraph.style.name == "Normal"
        and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER
        and "strait" in paragraph.text.casefold()
    ]
    assert reset
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT for paragraph in reset)


def test_ac3_build_covered_shared_word_render_path():
    from essay_writer.specialists.writer import essay_pack_render_word as writer_export
    from essay_writer.word import essay_pack_render_word as word_impl

    assert HANDLERS["essay.pack.render_word"] is writer_export
    assert writer_export is word_impl
    assert TOOL == "essay.pack.render_word"
