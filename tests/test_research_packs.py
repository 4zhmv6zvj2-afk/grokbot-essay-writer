"""4a Word-in-chat, PDF-only zip, 4b Word zip, and the Harvard Referencing line."""

import base64
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from helpers import CLARIFY_OPEN, ESSAY, TOPIC, open_sequence, source, write_upload

from essay_writer.doors import DOORS
from essay_writer.harvard_4a import (
    BOOK_SHAPE,
    FOUR_A_FORMAT,
    JOURNAL_SHAPE,
    WEB_SHAPE,
    harvard_4a_line,
)
from essay_writer.research import essay_research_full, essay_research_summary
from essay_writer.research_packs import write_papers_zip
from essay_writer.server import house_instructions
from essay_writer.stage_llm import stage_settings
from essay_writer.student_voice import student_surface_hits
from essay_writer.writing_basics import basics_lines

PDF = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n"
JOURNAL_LINE = (
    "Smith, J. (2020) 'Chokepoints and crude', Maritime Policy, 12(3), pp. 4-9. "
    "Available at: doi:10.1000/strait."
)
JOURNAL_TEXT = (
    "Smith, J. (2020) Chokepoints and crude. Maritime Policy 12(3) pp. 4-9. doi:10.1000/strait\n"
    "The Strait handles a large share of flows.\n"
    "The authors confirm the limit in the gathered file.\n"
)
ANN = (
    'Smith (2020) argues the corridor is narrow. The authors confirm the limit. '
    'The paper says "The Strait handles a large share of flows."'
)
REPO = Path(__file__).resolve().parents[1]


def _output(workspace: Path) -> Path:
    return workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output"


def _docx(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def _docx_bytes(payload: bytes) -> str:
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(payload)).paragraphs)


def _journal(**extra):
    row = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{ESSAY}/smith.md",
        "author": "Smith, J.",
        "year": "2020",
        "title": "Chokepoints and crude",
        "container": "Maritime Policy",
        "volume": "12",
        "issue": "3",
        "pages": "4-9",
        "doi": "10.1000/strait",
        "quote": "The Strait handles a large share of flows.",
    }
    row.update(extra)
    return row


def test_stage_prompt_bakes_the_harvard_line():
    lines = basics_lines("4a")
    assert FOUR_A_FORMAT in lines
    overlay = stage_settings("4a")["prompt_overlay"]
    assert JOURNAL_SHAPE in overlay
    assert BOOK_SHAPE in overlay
    assert WEB_SHAPE in overlay
    bullets = "\n".join(f"- {line}" for line in lines)
    assert bullets in overlay
    assert overlay.startswith("Before the save, the chat opens with a short Basics pass.")
    assert "Walk this Basics checklist internally before the save." in overlay
    summary = next(door["description"] for door in DOORS if door["tool"] == "essay.research.summary")
    full = next(door["description"] for door in DOORS if door["tool"] == "essay.research.full")
    assert JOURNAL_SHAPE in summary
    assert "annotation.docx" in summary
    assert "real PDFs only" in summary
    assert "either mid-office pick" in summary
    assert "Grok Build" not in summary and "Cursor" not in summary
    assert "critical-reviews.zip" in full.casefold()
    assert "word file per paper" in full.casefold()
    assert "either mid-office pick" in full
    assert "Grok Build" not in full and "Cursor" not in full
    workflow = (REPO / "docs" / "workflow.md").read_text(encoding="utf-8")
    assert "Cursor" in workflow and "Grok Build" in workflow
    assert "real PDFs only" in workflow or "real PDFs only" in workflow.casefold()
    helpers = (REPO / "docs" / "academic-writing-helpers.md").read_text(encoding="utf-8")
    assert FOUR_A_FORMAT in helpers


def test_harvard_journal_line_omits_nothing_it_was_given():
    assert harvard_4a_line(_journal()) == JOURNAL_LINE


def test_delivery_briefing_covers_both_workplaces(monkeypatch):
    cursor = house_instructions()
    assert "writes these files for either file-work host" in cursor
    assert "either mid-office pick" not in cursor
    assert "annotation.docx" in cursor
    assert "real PDFs only" in cursor
    assert "critical-reviews.zip" in cursor.casefold()
    assert "word file per paper" in cursor.casefold()
    monkeypatch.setattr("essay_writer.server.mid_office", lambda: "Grok Build")
    grok = house_instructions()
    assert "Grok Build writes these files for either file-work host" in grok
    assert "either mid-office pick" not in grok
    assert "real PDFs only" in grok
    assert "critical-reviews.zip" in grok.casefold()
    assert "word file per paper" in grok.casefold()
    for name in ("harvard_4a.py", "research_packs.py"):
        text = (REPO / "src" / "essay_writer" / name).read_text(encoding="utf-8")
        assert "mid_office" not in text
        assert "SHELL_CURSOR" not in text


def test_papers_zip_keeps_real_pdfs_and_drops_markdown(tmp_path):
    pdf = tmp_path / "paper.pdf"
    pdf.write_bytes(PDF)
    notes = tmp_path / "notes.md"
    notes.write_text("# fluff\n", encoding="utf-8")
    fake = tmp_path / "page.pdf"
    fake.write_bytes(b"<html>not a paper</html>")
    dest = tmp_path / "papers.zip"
    assert write_papers_zip(dest, [pdf, notes, fake]) == ["paper.pdf"]
    archive = zipfile.ZipFile(dest)
    assert archive.namelist() == ["paper.pdf"]
    assert archive.read("paper.pdf").startswith(b"%PDF-")


def test_summary_confirm_delivers_word_in_chat_and_a_pdf_zip(workspace):
    write_upload(workspace, text=JOURNAL_TEXT)
    seq = open_sequence()
    draft = essay_research_summary(seq, TOPIC, "draft", [_journal(annotation=ANN)])
    assert draft["ok"] is True
    assert draft["wrote"] is False
    assert "delivery" not in draft
    assert not (_output(workspace) / "annotation.docx").exists()
    confirmed = essay_research_summary(seq, TOPIC, "confirm", [_journal(annotation=ANN)])
    assert confirmed["ok"] is True
    assert confirmed["stored"] is True
    assert confirmed["wrote"] is False
    folder = _output(workspace)
    annotation = folder / "annotation.docx"
    text = _docx(annotation)
    assert "Annotated bibliography" in text
    assert JOURNAL_LINE in text
    assert "The authors confirm the limit." in text
    assert "The Strait handles a large share of flows." in text
    assert "Reference List" not in text
    assert re.search(r"\bdig\b", text, re.IGNORECASE) is None
    assert student_surface_hits(text) == []
    italic = False
    for paragraph in Document(str(annotation)).paragraphs:
        for run in paragraph.runs:
            if run.italic and "Maritime Policy" in (run.text or ""):
                italic = True
    assert italic
    chat = confirmed["delivery"]["annotation"]
    assert chat["deliver"] == "chat"
    assert chat["name"] == "annotation.docx"
    assert chat["media_type"].endswith("wordprocessingml.document")
    assert base64.b64decode(chat["docx_base64"]) == annotation.read_bytes()
    listed = folder / "download-list.docx"
    assert listed.is_file()
    assert listed.suffix == ".docx"
    assert listed.read_bytes()[:2] == b"PK"
    assert confirmed["delivery"]["download_list"]["format"] == "docx"
    assert not (folder / "download-list.md").exists()
    archive = zipfile.ZipFile(folder / "papers.zip")
    assert archive.namelist() == []
    assert all(not name.endswith(".md") for name in archive.namelist())


def test_annotation_quote_and_jargon_are_refused_and_draft_writes_nothing(workspace):
    write_upload(workspace, text=JOURNAL_TEXT)
    seq = open_sequence()
    invented = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_journal(annotation='The paper says "This sentence was never gathered."')],
    )
    assert invented["ok"] is False
    assert invented["refused"] == "invent-forbid"
    assert invented["invent_forbid"] == "refuse"
    assert invented["wrote"] is False
    assert not (_output(workspace) / "annotation.docx").exists()
    jargon = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_journal(annotation=ANN + " Then dig into the method.")],
    )
    assert jargon["ok"] is False
    assert jargon["refused"] == "house-jargon"
    assert jargon["wrote"] is False
    assert not (_output(workspace) / "annotation.docx").exists()


def test_confirm_without_the_same_draft_writes_no_pack(workspace):
    write_upload(workspace, text=JOURNAL_TEXT)
    seq = open_sequence()
    essay_research_summary(seq, TOPIC, "draft", [_journal(annotation=ANN)])
    changed = essay_research_summary(
        seq,
        TOPIC,
        "confirm",
        [_journal(annotation="Smith (2020) argues a different paragraph.")],
    )
    assert changed["ok"] is False
    assert changed["refused"] == "confirm"
    assert changed["wrote"] is False
    assert not (_output(workspace) / "annotation.docx").exists()
    silent = essay_research_summary(seq, TOPIC, "confirm", [_journal()])
    assert silent["ok"] is False
    assert not (_output(workspace) / "papers.zip").exists()


def test_full_confirm_zips_one_word_file_per_paper(workspace):
    write_upload(workspace, text=JOURNAL_TEXT)
    write_upload(
        workspace,
        name="lee.md",
        text="Lee, A. 2019 Corridor depth\nDepth varies by season.\n",
    )
    second = {
        "id": "src-2",
        "origin": "upload",
        "locator": f"uploads/{ESSAY}/lee.md",
        "author": "Lee, A.",
        "year": "2019",
        "title": "Corridor depth",
        "quote": "Depth varies by season.",
    }
    reviews = [
        {"paper_id": "src-1", "markdown": ANN},
        {
            "paper_id": "src-2",
            "markdown": 'Lee (2019) measures depth. "Depth varies by season."',
        },
    ]
    seq = open_sequence()
    draft = essay_research_full(
        seq,
        TOPIC,
        "draft",
        [_journal(), second],
        clarify_context=CLARIFY_OPEN,
        reviews=reviews,
    )
    assert draft["ok"] is True
    assert draft["wrote"] is False
    assert not (_output(workspace) / "critical-reviews.zip").exists()
    confirmed = essay_research_full(
        seq,
        TOPIC,
        "confirm",
        [_journal(), second],
        clarify_context=CLARIFY_OPEN,
        reviews=reviews,
    )
    assert confirmed["stored"] is True
    delivery = confirmed["delivery"]
    assert delivery["stage"] == "4b"
    assert [item["paper_id"] for item in delivery["reviews"]] == ["src-1", "src-2"]
    review_folder = _output(workspace) / "reviews"
    assert (review_folder / "src-1.md").is_file()
    assert (review_folder / "src-2.md").is_file()
    archive = zipfile.ZipFile(_output(workspace) / "critical-reviews.zip")
    assert archive.namelist() == ["src-1.docx", "src-2.docx"]
    assert all(name.endswith(".docx") for name in archive.namelist())
    assert not any(name.endswith(".md") for name in archive.namelist())
    first = _docx_bytes(archive.read("src-1.docx"))
    assert JOURNAL_LINE in first
    assert "The authors confirm the limit." in first
    assert re.search(r"\bdig\b", first, re.IGNORECASE) is None
    assert student_surface_hits(first) == []
    assert "Depth varies by season." in _docx_bytes(archive.read("src-2.docx"))
    duplicate = essay_research_full(
        seq,
        TOPIC,
        "draft",
        [_journal()],
        clarify_context=CLARIFY_OPEN,
        reviews=[reviews[0], {"paper_id": "src-1", "markdown": ANN + " Again."}],
    )
    assert duplicate["ok"] is False
    assert duplicate["refused"] == "critical-review"
    assert "one markdown per paper" in duplicate["reason"].casefold()
    missing = essay_research_full(
        seq,
        TOPIC,
        "draft",
        [_journal(), second],
        clarify_context=CLARIFY_OPEN,
        reviews=[reviews[0]],
    )
    assert missing["ok"] is False
    assert missing["refused"] == "critical-review"
    invented = essay_research_full(
        seq,
        TOPIC,
        "draft",
        [_journal()],
        clarify_context=CLARIFY_OPEN,
        reviews=[{"paper_id": "src-1", "markdown": '"This sentence was never gathered."'}],
    )
    assert invented["refused"] == "invent-forbid"
    assert invented["wrote"] is False


def test_full_without_reviews_still_writes_one_note_per_paper(workspace):
    write_upload(workspace)
    seq = open_sequence()
    essay_research_full(seq, TOPIC, "draft", [source()], clarify_context=CLARIFY_OPEN)
    confirmed = essay_research_full(seq, TOPIC, "confirm", [source()], clarify_context=CLARIFY_OPEN)
    assert (_output(workspace) / "reviews" / "src-1.md").is_file()
    archive = zipfile.ZipFile(_output(workspace) / "critical-reviews.zip")
    assert archive.namelist() == ["src-1.docx"]
    text = _docx_bytes(archive.read("src-1.docx"))
    assert "The Strait handles a large share of flows." in text
    assert "Chokepoints and crude" in text
    assert re.search(r"\bdig\b", text, re.IGNORECASE) is None
    assert student_surface_hits(text) == []
    assert confirmed["delivery"]["reviews_zip"]["entries"] == ["src-1.docx"]
