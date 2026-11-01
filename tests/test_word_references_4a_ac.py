"""RPT5001AA-2 A+C+B: Word References and stage 4a annotation notes.

AC1: references_markdown is the References section when it has entries.
AC2: stage 10 does not close on a heading-only list, or when cites are missing.
AC3: stage 4a Confirm rejects a surname stub and an empty annotation body.
AC4: Basics, then 4a Confirm, then stage 10 Word, on one sequence.
"""

import inspect
import json
from pathlib import Path

from docx import Document
from helpers import ESSAY, open_sequence, picked, write_upload

from essay_writer.brief_contract import remember_registration
from essay_writer.build_office import stage_launch
from essay_writer.cursor_office import cursor_stage_launch
from essay_writer.late_chain import write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.research import essay_research_summary
from essay_writer.stage_track import (
    _render_stage10_before_close,
    _save_confirm,
    essay_stage,
    load_pointer,
)
from essay_writer.state import load_state, save_state

MODULE = "RPT5001AA"
NUMBER = "445566"
TOPIC = "How does one decision sit against the gathered paper?"
LOOK = "The diary reflects on one decision named in the brief."
BRIEF = """
Assignment title: Individual reflective diary

Module title: Practice module
Module code: RPT5001AA

Aim: Understand the practice by reflecting on your own decisions.

Assignment task: Write four entries. Each entry names one decision and critically analyses it with theory and evidence.

Title page: module title, module code, assignment title, and the word count.
Include your student ID.
The name stays off the title page and the file.

Pages: Arial, size 12. 2.5 cm margins on all sides. Page numbers in the lower right.
"""
JOURNAL_TEXT = (
    "Qin, Y. (2021) 'Stakeholder mapping', Journal of Practice, 8(1), pp. 2-9. "
    "doi:10.1000/stake\n"
    "The paper maps stakeholders for one corridor decision.\n"
)
GOOD = (
    "Qin (2021) mapped stakeholders for one corridor decision. "
    "The paper found that engagement shapes the outcome. "
    "That finding helps this assignment on reflective practice."
)
QIN_LINE = (
    "Qin, Y. (2021) 'Stakeholder mapping', Journal of Practice, 8(1), pp. 2-9. "
    "Available at: doi:10.1000/stake."
)
JONES_LINE = "Jones, A. (2019) 'Second line', Practice, 1(1), pp. 1-2."
MARKDOWN = "\n".join([QIN_LINE, JONES_LINE])


def _references(path: Path) -> list[str]:
    texts = [
        paragraph.text.strip()
        for paragraph in Document(str(path)).paragraphs
        if paragraph.text.strip()
    ]
    if "References" not in texts:
        return []
    return texts[texts.index("References") + 1 :]


def _brief() -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(BRIEF, encoding="utf-8")


def _qin_source(**extra):
    row = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Stakeholder mapping",
        "container": "Journal of Practice",
        "volume": "8",
        "issue": "1",
        "pages": "2-9",
        "doi": "10.1000/stake",
        "quote": "The paper maps stakeholders for one corridor decision.",
    }
    row.update(extra)
    return row


def _diary_pack(seq: str, *, references_markdown: str | None = None, body: str | None = None) -> dict:
    pack = {
        "essay_id": seq,
        "title": "Individual reflective diary",
        "referencing": "harvard",
        "assessment": "Individual reflective diary",
        "conclusion": "noted",
        "sections": [
            {
                "heading": "One decision",
                "body": body
                or "A decision, with theory, on the practice the brief names (Qin, 2021; Jones, 2019).",
            }
        ],
    }
    if references_markdown is not None:
        pack["references_markdown"] = references_markdown
    return pack


def _park_stage10(seq: str) -> None:
    pointer = load_pointer(seq)
    pointer["current"] = "10"
    pointer["saved"] = ["1", "2", "3", "4a", "4b", "5", "6", "7", "8", "9"]
    pointer["closed"] = False
    pointer["chain_armed"] = False
    pointer["waiting_stage"] = None
    pointer["waiting_digest"] = None
    pointer["room_grant"] = None
    pointer["room_lock"] = False
    path = work_dir(seq) / "tracking" / "stage.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(pointer, indent=2) + "\n", encoding="utf-8")


def _plant_stage10_pack(seq: str, pack: dict) -> None:
    state = load_state(seq)
    state["saved_pack"] = pack
    save_state(seq, state)


def _markdown_lines() -> list[str]:
    return [line.strip() for line in MARKDOWN.splitlines() if line.strip()]


def _save_basics(seq: str) -> None:
    for stage in ("1", "2", "3"):
        draft = essay_stage(seq, stage, "draft", LOOK)
        assert draft["ok"] is True, (stage, draft)
        saved = essay_stage(seq, stage, "confirm", LOOK, **picked(draft))
        assert saved["stored"] is True, (stage, saved)


def _plant_research(seq: str, row: dict) -> None:
    state = load_state(seq)
    state["research"] = {"summary": {"sources": [row]}, "full": None}
    save_state(seq, state)


def test_ac1_stage10_word_references_match_markdown(workspace, monkeypatch, tmp_path):
    """AC1. references_markdown is printed even when gathered sources would print a shorter list."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    _brief()
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)
    seq = open_sequence(MODULE)
    remember_registration(seq, NUMBER)
    source = _qin_source(annotation=GOOD)
    essay_research_summary(seq, TOPIC, "draft", [source])
    saved = essay_research_summary(seq, TOPIC, "confirm", [source])
    assert saved["ok"] is True, saved
    pack = _diary_pack(seq, references_markdown=MARKDOWN)
    pack["sources"] = [source]
    _plant_stage10_pack(seq, pack)

    final = write_stage10(seq)
    assert final["ok"] is True, final
    assert final.get("sequence_written") is True
    refs = _references(work_dir(seq) / "output" / f"{seq}.docx")
    assert refs == _markdown_lines()
    assert JONES_LINE in refs


def test_ac2_stage10_close_refuses_empty_or_unmatched_references(workspace, monkeypatch, tmp_path):
    """AC2. Heading-only References cannot close, and a cite missing from the list cannot close."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    _brief()
    seq = open_sequence(MODULE)
    remember_registration(seq, NUMBER)
    _park_stage10(seq)
    _plant_stage10_pack(seq, _diary_pack(seq, body="A decision, with theory, on the practice the brief names."))

    draft = essay_stage(seq, "10", "draft", LOOK)
    assert draft["ok"] is True, draft
    refused = essay_stage(seq, "10", "confirm", LOOK, **picked(draft))
    assert refused["ok"] is False, refused
    assert refused["refused"] == "empty-references"
    assert refused.get("closed") is False
    assert "no entries" in str(refused.get("reason") or "").casefold()
    assert load_pointer(seq)["current"] == "10"
    stage_doc = work_dir(seq) / "output" / f"{seq}.docx"
    if stage_doc.is_file():
        assert _references(stage_doc) == []
    direct = write_stage10(seq)
    assert direct["ok"] is False, direct
    assert direct["refused"] == "empty-references"
    assert direct.get("sequence_written") is not True

    other = open_sequence(MODULE)
    remember_registration(other, NUMBER)
    _park_stage10(other)
    _plant_stage10_pack(
        other,
        _diary_pack(
            other,
            references_markdown=JONES_LINE,
            body="One decision used (Qin, 2021) as evidence for this diary.",
        ),
    )
    other_draft = essay_stage(other, "10", "draft", LOOK)
    assert other_draft["ok"] is True, other_draft
    mismatched = essay_stage(other, "10", "confirm", LOOK, **picked(other_draft))
    assert mismatched["ok"] is False, mismatched
    assert mismatched["refused"] == "references-cites"
    assert mismatched.get("closed") is False
    assert "qin" in str(mismatched.get("reason") or "").casefold()
    assert load_pointer(other)["current"] == "10"
    assert load_pointer(other)["closed"] is False
    other_doc = work_dir(other) / "output" / f"{other}.docx"
    if other_doc.is_file():
        assert _references(other_doc) == []


def test_ac3_stage4a_confirm_rejects_surname_stub_and_empty_body(workspace, monkeypatch, tmp_path):
    """AC3. A surname stub or an empty annotation body blocks stage 4a Confirm."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    _brief()
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)

    seq = open_sequence(MODULE)
    _save_basics(seq)
    _plant_research(seq, _qin_source(annotation="Qin."))
    draft = essay_stage(seq, "4a", "draft", LOOK)
    assert draft["ok"] is True, draft
    refused = essay_stage(seq, "4a", "confirm", LOOK, **picked(draft))
    assert refused["ok"] is False, refused
    assert refused["refused"] == "annotation-prose"
    assert "surname" in str(refused.get("reason") or "").casefold()
    assert load_pointer(seq)["current"] == "4a"
    number = seq.rsplit("-", 1)[-1]
    annotation = (
        workspace / "Essay-writer" / MODULE / "sequences" / number / "output" / "annotation.docx"
    )
    assert not annotation.exists()

    _plant_research(seq, _qin_source(annotation=""))
    empty_draft = essay_stage(seq, "4a", "draft", LOOK)
    assert empty_draft["ok"] is True, empty_draft
    empty = essay_stage(seq, "4a", "confirm", LOOK, **picked(empty_draft))
    assert empty["ok"] is False, empty
    assert empty["refused"] == "annotation-prose"
    assert "no sentences" in str(empty.get("reason") or "").casefold()
    assert load_pointer(seq)["current"] == "4a"
    assert not annotation.exists()


def test_ac4_basics_4a_confirm_stage10_word_journey(workspace, monkeypatch, tmp_path):
    """AC4. Basics, a stubbed 4a Confirm, then real notes, then stage 10 Word references."""
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    for stage in ("4a", "10"):
        cursor = cursor_stage_launch(stage)
        build = stage_launch(stage)
        assert cursor["settings"]["fast"] is False
        assert build["settings"]["fast"] is False
        assert cursor["settings"]["model"] == build["settings"]["model"]
    assert "stage_bibliography" in inspect.getsource(_save_confirm)
    assert "write_stage10" in inspect.getsource(_render_stage10_before_close)

    _brief()
    write_upload(workspace, MODULE, "qin.md", JOURNAL_TEXT)
    seq = open_sequence(MODULE)
    remember_registration(seq, NUMBER)
    _save_basics(seq)
    _plant_research(seq, _qin_source(annotation="Qin."))
    stub_draft = essay_stage(seq, "4a", "draft", LOOK)
    stub = essay_stage(seq, "4a", "confirm", LOOK, **picked(stub_draft))
    assert stub["ok"] is False, stub
    assert stub["refused"] == "annotation-prose"
    assert load_pointer(seq)["current"] == "4a"

    _plant_research(seq, _qin_source(annotation=GOOD))
    good_draft = essay_stage(seq, "4a", "draft", LOOK)
    good = essay_stage(seq, "4a", "confirm", LOOK, **picked(good_draft))
    assert good["ok"] is True, good
    assert good["stored"] is True, good
    number = seq.rsplit("-", 1)[-1]
    annotation = (
        workspace / "Essay-writer" / MODULE / "sequences" / number / "output" / "annotation.docx"
    )
    assert annotation.is_file()
    note = "\n".join(paragraph.text for paragraph in Document(str(annotation)).paragraphs)
    assert "helps this assignment" in note

    _park_stage10(seq)
    _plant_stage10_pack(seq, _diary_pack(seq, references_markdown=MARKDOWN))
    draft = essay_stage(seq, "10", "draft", LOOK)
    assert draft["ok"] is True, draft
    saved = essay_stage(seq, "10", "confirm", LOOK, **picked(draft))
    assert saved["ok"] is True, saved
    assert saved.get("stored") is True, saved
    assert saved.get("closed") is True, saved
    refs = _references(work_dir(seq) / "output" / f"{seq}.docx")
    assert refs == _markdown_lines()


SAVED_QUESTION_TITLE = (
    "Smith, J. (2020) Is the strait safe? Bristol: Harbour Press."
)
SAVED_ABBREVIATION_TITLE = (
    "Smith, J. (2020) U.S. maritime policy. Bristol: Harbour Press."
)
SAVED_JOURNAL_LINE = (
    "Smith, J. (2020) 'Chokepoints and crude', Maritime Policy, 12(3), pp. 4-9."
)


def _saved_runs(line: str) -> list[tuple[str, bool]]:
    from essay_writer.word_references import styled_reference_runs

    return styled_reference_runs(line, [])


def test_ac5_saved_title_italic_span_stops_at_the_title_end(workspace, tmp_path):
    """AC5. A title ending in ? or holding an initial keeps its whole span.

    The publisher stays outside the italic run, and the article title in a
    saved journal line stays outside it too.
    """
    question = _saved_runs(SAVED_QUESTION_TITLE)
    italic = "".join(text for text, marked in question if marked)
    plain = "".join(text for text, marked in question if not marked)
    assert italic == "Is the strait safe?"
    assert "Harbour Press" in plain
    assert "Harbour Press" not in italic

    initial = _saved_runs(SAVED_ABBREVIATION_TITLE)
    italic = "".join(text for text, marked in initial if marked)
    plain = "".join(text for text, marked in initial if not marked)
    assert italic == "U.S. maritime policy"
    assert "Harbour Press" in plain
    assert "Harbour Press" not in italic

    journal = _saved_runs(SAVED_JOURNAL_LINE)
    italic = "".join(text for text, marked in journal if marked)
    assert italic == "Maritime Policy"
    assert "Chokepoints and crude" not in italic

    from essay_writer.word import write_docx

    lines = [SAVED_QUESTION_TITLE, SAVED_JOURNAL_LINE]
    target = tmp_path / "saved-references.docx"
    write_docx(
        ESSAY,
        {"assessment": "Individual report", "references_markdown": "\n".join(lines)},
        {
            "title": "Saved reference lines",
            "word_count": {"counted": 12},
            "sections": [{"heading": "Introduction", "body": "The corridor is narrow."}],
            "figures": [],
            "tables": [],
            "sources": [],
        },
        target,
    )
    paragraphs = [
        paragraph
        for paragraph in Document(str(target)).paragraphs
        if paragraph.text.strip()
    ]
    printed = {paragraph.text.strip(): paragraph for paragraph in paragraphs}
    assert SAVED_QUESTION_TITLE in printed
    italic_runs = [run.text for run in printed[SAVED_QUESTION_TITLE].runs if run.italic]
    assert "Is the strait safe?" in italic_runs
    assert "Harbour Press" not in "".join(italic_runs)
    journal_italic = [run.text for run in printed[SAVED_JOURNAL_LINE].runs if run.italic]
    assert "Maritime Policy" in journal_italic
    assert "Chokepoints and crude" not in "".join(journal_italic)


SAVED_CHAPTER_LINE = (
    "Smith, J. (2020) 'Chapter title', in Jones, K. (ed.) Book Title. "
    "5th edn. Place: Publisher, pp. 1-2."
)


def test_ac6_saved_chapter_italic_span_is_the_book_title_only(workspace):
    """AC6. A chapter line italicises the book title, not the editor or publisher."""
    runs = _saved_runs(SAVED_CHAPTER_LINE)
    italic = "".join(text for text, marked in runs if marked)
    plain = "".join(text for text, marked in runs if not marked)
    assert italic == "Book Title"
    assert "in Jones, K. (ed.)" in plain
    assert "5th edn." in plain
    assert "Publisher" in plain
    assert "Chapter title" not in italic


MISMATCH_LINE = "Smith, J. (2020) The 2021 reform. Bristol: Harbour Press."


def test_ac7_cite_year_has_to_sit_with_its_name_in_the_list(workspace):
    """AC7. A year printed elsewhere in an entry does not cover that cite."""
    from essay_writer.word_references import cite_reference_gaps

    def gaps(body: str) -> list[str]:
        return cite_reference_gaps([MISMATCH_LINE], [{"heading": "Body", "body": body}])

    assert gaps("The point is old (Smith, 2021).") == ["(smith, 2021)"]
    assert gaps("Smith (2021) writes.") == ["(smith, 2021)"]
    assert gaps("The point is old (Smith, 2020).") == []
    assert gaps("Smith (2020) writes.") == []

    two = "Smith, J. and Jones, K. (2020) 'Joint work', Practice, 1, pp. 1-2."
    assert cite_reference_gaps(
        [two], [{"heading": "Body", "body": "Both agree (Smith and Jones, 2020)."}]
    ) == []
    lettered = "Smith, J. (2020a) 'First', Practice, 1, pp. 1-2."
    assert cite_reference_gaps(
        [lettered], [{"heading": "Body", "body": "One point (Smith, 2020a)."}]
    ) == []


def test_ac8_withheld_line_is_matched_on_the_family_name(workspace):
    """AC8. A keyed initial does not keep a line that only covers another author."""
    import json

    from essay_writer.cite_fit import _record_path
    from essay_writer.word_references import without_withheld_reference_lines

    essay_id = f"{MODULE}-1"
    path = _record_path(essay_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps({"ok": [], "rejected": [], "withheld": ["jones, 2020"]}),
        encoding="utf-8",
    )
    line = "Jones, J. (2020) 'Withheld work', Practice, 1, pp. 1-2."
    kept = [{"author": "Smith, J.", "year": "2020"}]
    assert without_withheld_reference_lines(essay_id, [line], kept) == []

    own_line = "Smith, J. (2020) 'Kept work', Practice, 2, pp. 3-4."
    assert without_withheld_reference_lines(essay_id, [own_line], kept) == [own_line]


PAGE_CITE_LINE = "Smith, J. (2020) 'Chokepoints and crude', Maritime Policy, 12(3), pp. 4-9."


def test_ac9_a_cite_with_a_page_still_needs_its_reference_line(workspace):
    """AC9. A page locator does not hide the author-date pair from the close gate."""
    from essay_writer.word_references import cite_reference_gaps

    def gaps(body: str) -> list[str]:
        return cite_reference_gaps([PAGE_CITE_LINE], [{"heading": "Body", "body": body}])

    assert gaps("The point (Patel, 2019, p. 4) stands.") == ["(patel, 2019)"]
    assert gaps("Patel (2019, pp. 4-5) writes.") == ["(patel, 2019)"]
    assert gaps("The point (Smith, 2020, p. 4) stands.") == []
    assert gaps("Smith (2020, pp. 4-5) writes.") == []
    assert gaps("Two points (Patel, 2019, p. 4; Smith, 2020, p. 5).") == ["(patel, 2019)"]



def test_ac10_a_narrative_et_al_cite_needs_its_reference_line(workspace):
    """AC10. `Smith et al. (2020)` is read as the same author-date pair."""
    from essay_writer.word_references import cite_reference_gaps

    line = "Patel, A. (2019) 'Other work', Practice, 1, pp. 1-2."
    assert cite_reference_gaps(
        [line], [{"heading": "Body", "body": "Smith et al. (2020) argue this."}]
    ) == ["(smith, 2020)"]
    smith = "Smith, J., Lee, C., Patel, A., and Wu, B. (2020) 'Team', Practice, 2, pp. 3-4."
    assert cite_reference_gaps(
        [smith], [{"heading": "Body", "body": "Smith et al. (2020) argue this."}]
    ) == []


JOINT_CITE_LINE = "Smith, J., Jones, A. and Brown, B. (2020) 'Team work', Practice, 2, pp. 3-4."
JONES_ONLY_LINE = "Jones, A. (2020) 'Second author alone', Practice, 1, pp. 1-2."


def test_ac11_a_joint_cite_is_read_on_its_first_named_author(workspace):
    """AC11. A joint cite is matched on the family name that comes first."""
    from essay_writer.word_references import cite_reference_gaps

    body = "The team agree (Smith, Jones and Brown, 2020)."

    def gaps(lines: list[str], text: str) -> list[str]:
        return cite_reference_gaps(lines, [{"heading": "Body", "body": text}])

    assert gaps([JONES_ONLY_LINE], body) == ["(smith, jones and brown, 2020)"]
    assert gaps([JOINT_CITE_LINE], body) == []
    narrative = "Smith, Jones and Brown (2020) agree on the point."
    assert gaps([JONES_ONLY_LINE], narrative) == ["(smith, jones and brown, 2020)"]
    assert gaps([JOINT_CITE_LINE], narrative) == []


SAVED_ABBREVIATION_JOURNAL = (
    "Smith, J. (2021) 'Article title', Journal of A.I. Research. "
    "Available at: doi:10.1/example"
)


def test_ac12_a_saved_journal_title_keeps_its_whole_span(workspace):
    """AC12. A dot inside an abbreviation does not end the journal title."""
    runs = _saved_runs(SAVED_ABBREVIATION_JOURNAL)
    italic = "".join(text for text, marked in runs if marked)
    plain = "".join(text for text, marked in runs if not marked)
    assert italic == "Journal of A.I. Research"
    assert plain.startswith("Smith, J. (2021) 'Article title', ")
    assert plain.endswith(". Available at: doi:10.1/example")
    short = _saved_runs("Smith, J. (2021) 'Article title', Practice. Available at: doi:10.1/x")
    assert "".join(text for text, marked in short if marked) == "Practice"


def test_ac13_a_stub_fills_only_from_the_same_author(workspace):
    """AC13. Another author with the same family name does not fill a stub."""
    from essay_writer.textutil import complete_citation_sources

    research = {
        "other.md": {
            "id": "r1",
            "locator": f"uploads/{MODULE}/other.md",
            "author": "Smith, A.",
            "year": "2019",
            "title": "Another study",
            "pages": "1-2",
        }
    }
    stub = {"author": "Smith, J.", "locator": f"uploads/{MODULE}/smith.md"}
    filled = complete_citation_sources([stub], research)[0]
    assert filled["author"] == "Smith, J."
    assert not filled.get("year")
    assert not filled.get("title")
    assert not filled.get("pages")

    surname_only = {"author": "Smith", "locator": f"uploads/{MODULE}/smith.md"}
    same = complete_citation_sources([surname_only], research)[0]
    assert same["author"] == "Smith, A."
    assert same["year"] == "2019"


CURLY_LINE = "O’Brien, S. (2020) 'Curly title', Practice, 1, pp. 1-2."
STRAIGHT_LINE = "O'Brien, S. (2020) 'Straight title', Practice, 1, pp. 1-2."


def test_ac14_a_curly_apostrophe_is_the_same_name(workspace):
    """AC14. A name with a curly apostrophe is the same name as a straight one."""
    from essay_writer.word_references import cite_reference_gaps

    def gaps(lines: list[str], body: str) -> list[str]:
        return cite_reference_gaps(lines, [{"heading": "Body", "body": body}])

    assert gaps([CURLY_LINE], "The point (O'Brien, 2020) stands.") == []
    assert gaps([STRAIGHT_LINE], "The point (O’Brien, 2020) stands.") == []
    assert gaps([CURLY_LINE], "The point (O'Brien, 2019) stands.") == ["(o'brien, 2019)"]
    assert gaps([CURLY_LINE], "The point (Brown, 2020) stands.") == ["(brown, 2020)"]

def test_ac15_a_name_in_a_title_is_not_the_entry_author(workspace):
    """AC15. A surname printed in a title does not cover that citation."""
    from essay_writer.word_references import cite_reference_gaps

    def gaps(lines: list[str], body: str) -> list[str]:
        return cite_reference_gaps(lines, [{"heading": "Body", "body": body}])

    title_name = "Jones, A. (2020) 'Smith (2020) and the port', Practice, 1, pp. 1-2."
    assert gaps([title_name], "The point (Smith, 2020) stands.") == ["(smith, 2020)"]
    assert gaps([title_name], "The point (Jones, 2020) stands.") == []
    own = "Smith, J. (2020) 'Own work', Practice, 1, pp. 1-2."
    assert gaps([own], "The point (Smith, 2020) stands.") == []

