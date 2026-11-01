"""Citation shape locks."""

import json
from pathlib import Path

from docx import Document
from helpers import ESSAY, TOPIC, open_sequence

from essay_writer.cite_open import apply_opened_citation, compound_author, in_text_page_issues
from essay_writer.harvard_4a import harvard_4a_line
from essay_writer.invent import read_gathered_text
from essay_writer.research import essay_research_summary
from essay_writer.stage_llm import stage_settings
from essay_writer.textutil import reference_line
from essay_writer.word import write_docx

REPO = Path(__file__).resolve().parents[1]
QUOTE = "The Strait handles a large share of flows."
FIRM = "The firm uses the data."


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _pdf(pages: list[str]) -> bytes:
    objects: dict[int, bytes] = {}
    kids: list[str] = []
    next_id = 3
    for text in pages:
        page_id = next_id
        content_id = next_id + 1
        next_id += 2
        stream = f"BT /F1 12 Tf 72 720 Td ({_escape(text)}) Tj ET".encode()
        objects[content_id] = (
            b"<< /Length "
            + str(len(stream)).encode()
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )
        objects[page_id] = f"<< /Type /Page /Parent 2 0 R /Contents {content_id} 0 R >>".encode()
        kids.append(f"{page_id} 0 R")
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    chunks = [b"%PDF-1.4\n"]
    for number in sorted(objects):
        chunks.append(f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n")
    chunks.append(b"%%EOF\n")
    return b"".join(chunks)


def _pad_over_2mb(data: bytes) -> bytes:
    extra = 2_000_001 - len(data)
    extra = max(extra, 1)
    return data.replace(b"%%EOF\n", b"%" + (b"0" * extra) + b"\n%%EOF\n", 1)


def _row(name: str, **extra: object) -> dict:
    item = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{ESSAY}/{name}",
        "author": "Smith, J.",
        "year": "2020",
        "title": "Chokepoints and crude",
        "quote": QUOTE,
    }
    item.update(extra)
    return item


def test_overlays_and_notes_record_citation_shape():
    overlay = stage_settings("9")["prompt_overlay"]
    for phrase in (
        "larger than 2MB",
        "pdf-read-gaps.json",
        "no extractable text",
        "Fosso Wamba",
        "Wamba, S.F.",
        "article-number slot",
        "PDF viewer index",
        "Available at: doi:",
        "heading is References",
    ):
        assert phrase in overlay
    note = (REPO / "docs" / "referencing-harvard.md").read_text(encoding="utf-8")
    assert "[online]" not in note
    assert "not under `Available at:`" not in note
    printer = note.split("## What the Word file prints today", 1)[1].split("## Local subject notes", 1)[0]
    assert "**References**" in printer
    assert "Available at: doi:" in printer
    assert "stays blank" in printer
    assert "[online]" not in (REPO / "src" / "essay_writer" / "textutil.py").read_text(encoding="utf-8")


def test_citation_check_opens_a_large_pdf(workspace):
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    path = folder / "smith.pdf"
    path.write_bytes(
        _pad_over_2mb(
            _pdf(
                [
                    "Smith, J. (2020) Chokepoints and crude\n"
                    + QUOTE
                    + "\n32\n"
                ]
            )
        )
    )
    assert path.stat().st_size > 2_000_000
    assert "Chokepoints and crude" in read_gathered_text(path)
    seq = open_sequence()
    draft = essay_research_summary(seq, TOPIC, "draft", [_row("smith.pdf", pages="32")])
    assert draft["ok"] is True, draft
    source = draft["sources"][0]
    assert source["author"] == "Smith, J."
    assert source["year"] == "2020"
    assert source["pages"] == "32"
    refused = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_row("smith.pdf", pages="1", quote="This sentence was not in the file.")],
    )
    assert refused["ok"] is False
    assert "too large" not in refused["reason"].casefold()
    assert "cannot verify citations against .pdf" not in refused["reason"]
    pages = in_text_page_issues(
        [{"heading": "Introduction", "body": "The point is narrow (Smith, 2020, p. 1)."}],
        [source],
    )
    assert pages and "viewer index" in pages[0]["reason"]
    assert (
        in_text_page_issues(
            [{"heading": "Introduction", "body": "The point is narrow (Smith, 2020, p. 32)."}],
            [source],
        )
        == []
    )


def test_blank_pdf_page_is_recorded_and_not_quoted(workspace):
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    (folder / "smith.pdf").write_bytes(
        _pdf(
            [
                "Smith, J. (2020) Chokepoints and crude\n" + QUOTE + "\n",
                "",
            ]
        )
    )
    seq = open_sequence()
    kept = essay_research_summary(seq, TOPIC, "draft", [_row("smith.pdf")])
    assert kept["ok"] is True, kept
    assert "pages" not in kept["sources"][0]
    gaps = json.loads((folder / "pdf-read-gaps.json").read_text(encoding="utf-8"))
    blank = [gap for gap in gaps["gaps"] if gap.get("reason") == "no extractable text"]
    assert blank
    assert blank[0]["quoted"] is False
    assert blank[0]["part"] != kept["sources"][0].get("pages")
    refused = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_row("smith.pdf", quote="Words from the page with no text.")],
    )
    assert refused["ok"] is False
    assert "no extractable text" in refused["reason"]
    assert "not quoted" in refused["reason"]
    assert "pdf-read-gaps.json" in refused["reason"]


def test_existing_read_gaps_are_not_quoted(workspace):
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    name = "gapped.pdf"
    data = _pdf(
        [
            "Smith, J. (2020) Chokepoints and crude\n" + QUOTE + "\n",
            "TAIL-QUOTE-NOT-FOR-CITATION",
        ]
    )
    marker = b"TAIL-QUOTE-NOT-FOR-CITATION"
    offset = data.find(marker)
    assert offset > 0
    (folder / "pdf-read-gaps.json").write_text(
        json.dumps(
            {
                "bytes_leave_machine": False,
                "gaps": [
                    {
                        "source": name,
                        "part": 2,
                        "reason": "read past 20MB is not a separate readable PDF",
                        "bytes_on_machine": True,
                        "offset": offset,
                        "length": len(marker),
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    (folder / name).write_bytes(data)
    seq = open_sequence()
    kept = essay_research_summary(seq, TOPIC, "draft", [_row(name)])
    assert kept["ok"] is True, kept
    assert "TAIL-QUOTE-NOT-FOR-CITATION" not in read_gathered_text(folder / name)
    refused = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [_row(name, quote="TAIL-QUOTE-NOT-FOR-CITATION")],
    )
    assert refused["ok"] is False
    assert "not quoted" in refused["reason"]
    saved = json.loads((folder / "pdf-read-gaps.json").read_text(encoding="utf-8"))
    assert any("20MB" in str(gap.get("reason")) for gap in saved["gaps"])


def test_compound_surname_stays_on_the_author_line(workspace):
    western = "John Adam Smith\nPlease cite this article as: Smith, J.A., Port pricing.\n"
    assert compound_author(western) is None
    article = "\n".join(
        [
            "Samuel Fosso Wamba",
            "Fosso Wamba, Samuel",
            "Please cite this article as: Wamba, S.F., Big data analytics.",
            "Journal of Production 12(3)",
            "Article number: 104567",
            "2019",
            FIRM,
        ]
    )
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    (folder / "wamba.pdf").write_bytes(_pdf([article]))
    seq = open_sequence()
    row = _row(
        "wamba.pdf",
        author="Wamba, S.F.",
        year="2019",
        title="Big data analytics",
        container="Journal of Production",
        volume="12",
        issue="3",
        pages="1",
        quote=FIRM,
    )
    draft = essay_research_summary(seq, TOPIC, "draft", [row])
    assert draft["ok"] is True, draft
    source = draft["sources"][0]
    assert source["author"] == "Fosso Wamba, S."
    assert "pages" not in source
    assert source["article_number"] == "104567"
    line = reference_line(source)
    assert line.startswith("Fosso Wamba, S.")
    assert "(2019)" not in line
    assert "Wamba, S.F." not in line
    assert "104567" in line
    assert "pp." not in line
    confirmed = essay_research_summary(seq, TOPIC, "confirm", [row])
    assert confirmed["ok"] is True, confirmed
    annotation = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "output" / "annotation.docx"
    text = "\n".join(paragraph.text for paragraph in Document(str(annotation)).paragraphs)
    assert "Fosso Wamba, S." in text
    assert "Wamba, S.F." not in text
    assert "pp." not in text


def test_stage9_printer_emits_tutorial_lines(tmp_path):
    journal = {
        "author": "Smith, J.",
        "year": "2020",
        "title": "Chokepoints and crude",
        "container": "Maritime Policy",
        "volume": "12",
        "issue": "3",
        "pages": "4-9",
        "doi": "10.1000/strait",
    }
    report = {
        "author": "Lee, C.",
        "year": "2019",
        "title": "Port emissions inventory",
        "place": "Bristol",
        "publisher": "Harbour Press",
    }
    missing = {
        "author": "Patel, A.",
        "year": "2018",
        "title": "A corridor note",
        "container": "Maritime Policy",
        "volume": "4",
        "issue": "1",
    }
    assert reference_line(journal) == (
        "Smith, J. (2020) 'Chokepoints and crude', Maritime Policy, 12(3), pp. 4-9. "
        "Available at: doi:10.1000/strait."
    )
    assert "[online]" not in reference_line({**journal, "url": "https://example.test/smith"})
    gap_line = reference_line(missing)
    assert "doi" not in gap_line.casefold()
    assert "pp." not in gap_line
    assert "Available at" not in gap_line
    numbered = harvard_4a_line({**missing, "article_number": "104567"})
    assert "104567" in numbered
    assert "pp." not in numbered
    target = tmp_path / "stage9.docx"
    write_docx(
        ESSAY,
        {"assessment": "Individual report"},
        {
            "title": "Chokepoints on the corridor",
            "word_count": {"counted": 12},
            "sections": [{"heading": "Introduction", "body": "The corridor is narrow."}],
            "figures": [],
            "tables": [],
            "sources": [journal, report, missing],
        },
        target,
    )
    document = Document(str(target))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "References" in text
    assert "Reference List" not in text
    assert "[online]" not in text
    assert "Available at: doi:10.1000/strait." in text
    assert "doi:10.1000/strait" in text
    italic = {run.text for paragraph in document.paragraphs for run in paragraph.runs if run.italic}
    plain = {run.text for paragraph in document.paragraphs for run in paragraph.runs if not run.italic}
    assert any("Maritime Policy" in piece for piece in italic)
    assert any("Port emissions inventory" in piece for piece in italic)
    assert any("Chokepoints and crude" in piece for piece in plain)


def test_in_text_page_check_sees_each_source_in_one_bracket(workspace):
    """Every source named in one bracket is checked against its own opened file."""
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    (folder / "smith.pdf").write_bytes(
        _pdf(["Smith, J. (2020) Chokepoints and crude\n" + QUOTE + "\n32\n"])
    )
    (folder / "jones.pdf").write_bytes(
        _pdf(["Jones, K. (2019) Port limits\n" + FIRM + "\n4\n"])
    )
    (folder / "team.pdf").write_bytes(
        _pdf(["Smith, J., Lee, C., Patel, A., and Wu, B. (2020) Team work\n" + QUOTE + "\n7\n"])
    )
    seq = open_sequence()
    draft = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [
            _row("smith.pdf", pages="32"),
            _row("jones.pdf", id="src-2", author="Jones, K.", year="2019", title="Port limits", pages="4", quote=FIRM),
            _row(
                "team.pdf",
                id="src-3",
                author="Smith, J., Lee, C., Patel, A., and Wu, B.",
                year="2020",
                title="Team work",
                pages="7",
            ),
        ],
    )
    assert draft["ok"] is True, draft
    sources = draft["sources"]
    bad = in_text_page_issues(
        [{"heading": "Body", "body": "Two points (Smith, 2020, p. 32; Jones, 2019, p. 99)."}],
        sources,
    )
    assert bad, "the second source's page is not printed in that file"
    good = in_text_page_issues(
        [{"heading": "Body", "body": "Two points (Smith, 2020, p. 32; Jones, 2019, p. 4)."}],
        sources,
    )
    assert good == []
    team = [source for source in sources if source.get("locator", "").endswith("team.pdf")]
    assert team
    bad_team = in_text_page_issues(
        [{"heading": "Body", "body": "One point (Smith et al., 2020, p. 99)."}],
        team,
    )
    assert bad_team
    good_team = in_text_page_issues(
        [{"heading": "Body", "body": "One point (Smith et al., 2020, p. 7)."}],
        team,
    )
    assert good_team == []


def test_a_looked_up_single_page_is_not_dropped(workspace):
    """A single page from a matching look-up stays when the PDF prints no page line."""
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    (folder / "single.pdf").write_bytes(
        _pdf(["Smith, J. (2020) Chokepoints and crude\n" + QUOTE])
    )
    source = _row(
        "single.pdf",
        container="Maritime Policy",
        pages="27",
        field_sources={"pages": "lookup"},
    )
    opened = apply_opened_citation(source)
    assert opened["pages"] == "27"
    assert ", p. 27." in reference_line(opened)


def test_an_in_text_page_from_a_matching_look_up_is_not_refused(workspace):
    """A looked-up page supports an in-text page when the file prints no page line."""
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    (folder / "single.pdf").write_bytes(
        _pdf(["Smith, J. (2020) Chokepoints and crude\n" + QUOTE])
    )
    looked_up = _row(
        "single.pdf",
        container="Maritime Policy",
        pages="27",
        field_sources={"pages": "lookup"},
    )
    body = [{"heading": "Body", "body": "One point (Smith, 2020, p. 27)."}]
    assert in_text_page_issues(body, [looked_up]) == []

    spanned = {**looked_up, "pages": "27-29"}
    between = [{"heading": "Body", "body": "One point (Smith, 2020, p. 28)."}]
    assert in_text_page_issues(between, [spanned]) == []

    bare = _row("single.pdf", container="Maritime Policy", pages="27")
    assert in_text_page_issues(body, [bare])


def test_a_page_cite_after_a_lead_in_is_still_checked(workspace):
    """A short lead-in in the bracket does not hide the page claim."""
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    (folder / "lead.pdf").write_bytes(
        _pdf(["Smith, J. (2020) Chokepoints and crude\n" + QUOTE + "\n9\n"])
    )
    source = _row("lead.pdf", container="Maritime Policy", pages="9")
    bad = [{"heading": "Body", "body": "One point (see Smith, 2020, p. 99)."}]
    assert in_text_page_issues(bad, [source])
    good = [{"heading": "Body", "body": "One point (see Smith, 2020, p. 9)."}]
    assert in_text_page_issues(good, [source]) == []


def test_printer_does_not_repeat_a_page_or_doi_label():
    """A gathered page or DOI that carries its own label prints it once."""
    journal = {
        "author": "Smith, J.",
        "year": "2020",
        "title": "Chokepoints and crude",
        "container": "Maritime Policy",
        "volume": "12",
        "issue": "3",
        "pages": "pp. 4-9",
        "doi": "https://doi.org/10.1000/strait",
    }
    line = reference_line(journal)
    assert line == (
        "Smith, J. (2020) 'Chokepoints and crude', Maritime Policy, 12(3), pp. 4-9. "
        "Available at: doi:10.1000/strait."
    )
    assert "pp. pp." not in line
    assert "doi:doi:" not in line
    assert "doi:https" not in line


def test_printer_uses_p_for_one_page_and_pp_for_a_range():
    """One page takes p.; a range or a page list takes pp."""
    journal = {
        "author": "Smith, J.",
        "year": "2020",
        "title": "Chokepoints and crude",
        "container": "Maritime Policy",
        "volume": "12",
        "issue": "3",
        "pages": "32",
    }
    line = reference_line(journal)
    assert ", p. 32." in line
    assert "pp. 32" not in line
    labelled = reference_line({**journal, "pages": "pp. 32"})
    assert ", p. 32." in labelled
    assert "pp. 32" not in labelled
    ranged = reference_line({**journal, "pages": "32-40"})
    assert ", pp. 32-40." in ranged
    listed = reference_line({**journal, "pages": "32, 35"})
    assert ", pp. 32, 35." in listed


def test_printer_keeps_an_article_id_out_of_the_page_slot():
    """An article id is never pp.; a known article number prints as article N."""
    journal = {
        "author": "Smith, J.",
        "year": "2020",
        "title": "Chokepoints and crude",
        "container": "Maritime Policy",
        "volume": "12",
        "issue": "3",
        "pages": "e12345",
    }
    unclear = reference_line(journal)
    assert "pp." not in unclear
    assert "e12345" not in unclear
    known = {key: value for key, value in journal.items() if key != "pages"}
    known["article_number"] = "104567"
    assert ", article 104567." in reference_line(known)
    labelled = {**known, "article_number": "Article 104567"}
    assert ", article 104567." in reference_line(labelled)
    elocator = {**known, "article_number": "e12345"}
    assert ", article e12345." in reference_line(elocator)


def test_printer_does_not_repeat_an_edition_full_stop():
    """An edition that already ends in a full stop prints one full stop."""
    book = {
        "author": "Smith, J.",
        "year": "2020",
        "title": "Port economics",
        "edition": "2nd edn.",
        "place": "Bristol",
        "publisher": "Harbour Press",
    }
    line = reference_line(book)
    assert line == (
        "Smith, J. (2020) Port economics. 2nd edn. Bristol: Harbour Press."
    )
    assert ".." not in line
    unauthorised = harvard_4a_line({**book, "edition": "2nd edn"})
    assert unauthorised == (
        "Smith, J. (2020) Port economics. 2nd edn. Bristol: Harbour Press."
    )
    with_url = harvard_4a_line({**book, "doi": "10.1000/port."})
    assert ".." not in with_url


def test_printer_keeps_the_stop_before_a_locator_after_publisher():
    """A book with a gathered web address keeps one stop before Available at."""
    book = {
        "author": "Smith, J.",
        "year": "2020",
        "title": "Port economics",
        "place": "Bristol",
        "publisher": "Harbour Press",
        "url": "https://example.test/port",
    }
    line = reference_line(book)
    assert line == (
        "Smith, J. (2020) Port economics. Bristol: Harbour Press. "
        "Available at: https://example.test/port."
    )
    assert "Harbour Press Available" not in line
    assert ".." not in line


def test_printer_writes_a_bare_ordinal_edition_as_edn():
    """A stored 2nd prints as 2nd edn., the handbook's form."""
    book = {
        "author": "Smith, J.",
        "year": "2020",
        "title": "Port economics",
        "edition": "2nd",
        "place": "Bristol",
        "publisher": "Harbour Press",
    }
    assert reference_line(book) == (
        "Smith, J. (2020) Port economics. 2nd edn. Bristol: Harbour Press."
    )
    assert "2nd." not in reference_line(book).replace("2nd edn.", "")
    assert "3rd edn." in harvard_4a_line({**book, "edition": "3RD"})


def test_in_text_page_check_reads_a_joint_narrative_label(workspace):
    """A joint or et al. narrative label is checked against its opened file."""
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    (folder / "smith.pdf").write_bytes(
        _pdf(["Smith, J. (2020) Chokepoints and crude\n" + QUOTE + "\n32\n"])
    )
    (folder / "jones.pdf").write_bytes(
        _pdf(["Jones, K. (2019) Port limits\n" + FIRM + "\n4\n"])
    )
    (folder / "team.pdf").write_bytes(
        _pdf(["Smith, J., Lee, C., Patel, A., and Wu, B. (2020) Team work\n" + QUOTE + "\n7\n"])
    )
    seq = open_sequence()
    draft = essay_research_summary(
        seq,
        TOPIC,
        "draft",
        [
            _row("smith.pdf", pages="32"),
            _row(
                "jones.pdf",
                id="src-2",
                author="Jones, K.",
                year="2019",
                title="Port limits",
                pages="4",
                quote=FIRM,
            ),
            _row(
                "team.pdf",
                id="src-3",
                author="Smith, J., Lee, C., Patel, A., and Wu, B.",
                year="2020",
                title="Team work",
                pages="7",
            ),
        ],
    )
    assert draft["ok"] is True, draft
    sources = draft["sources"]

    def issues(body: str) -> list[dict[str, str]]:
        return in_text_page_issues([{"heading": "Body", "body": body}], sources)

    assert issues("Smith and Jones (2020, p. 99) argue this.")
    assert issues("Smith and Jones (2020, p. 32) argue this.") == []
    assert issues("Smith et al. (2020, p. 99) argue this.")
    assert issues("Smith et al. (2020, p. 7) argue this.") == []
    assert issues("Smith and Jones (2020) summarise the whole study.") == []


def test_a_page_of_a_word_export_footer_is_a_printed_page(workspace):
    """`Page 32 of 40` in a footer counts as page 32, not as no printed page."""
    from essay_writer.cite_open import open_for_citation

    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True)
    path = folder / "footer.pdf"
    path.write_bytes(
        _pdf(["Smith, J. (2020) Chokepoints and crude\n" + QUOTE + "\nPage 32 of 40"])
    )
    opened = open_for_citation(path)
    assert "32" in opened.printed_pages
    assert "40" not in opened.printed_pages

    plain = folder / "plain.pdf"
    plain.write_bytes(_pdf(["Smith, J. (2020) Chokepoints and crude\n" + QUOTE + "\n9\n"]))
    assert "9" in open_for_citation(plain).printed_pages


def test_printer_does_not_repeat_a_volume_or_issue_label():
    """A volume or issue that carries its own label prints it once."""
    journal = {
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Stakeholder mapping",
        "container": "Journal of Practice",
        "volume": "vol. 12",
        "issue": "(3)",
        "pages": "2-9",
    }
    line = reference_line(journal)
    assert line == (
        "Qin, Y. (2021) 'Stakeholder mapping', Journal of Practice, 12(3), pp. 2-9."
    )
    assert "vol." not in line
    assert "((" not in line
    named = harvard_4a_line({**journal, "volume": "12", "issue": "no. 3"})
    assert ", 12(3)," in named
    supplement = harvard_4a_line({**journal, "volume": "12", "issue": "Suppl 1"})
    assert ", 12(Suppl 1)," in supplement
