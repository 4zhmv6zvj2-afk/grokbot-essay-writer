"""Title-page instruction prose is not a cover line, and the cover is one block.

AC1: stage 10 Word for a RPT5001AA-style brief has one title block.
AC2: brief instructions and a sample Draft page are not cover lines.
AC3: a real asked field still prints an empty label, and a banned name stays off.
AC4: registration, module, assignment title, and word count stay when the brief asks.
"""

import hashlib
from pathlib import Path

from docx import Document
from helpers import mac_handoff_sequence_dir, with_word_references

from essay_writer.brief_contract import pack_with_registration, remember_registration
from essay_writer.diary import diary_cover_lines, named_entries, title_page_spec
from essay_writer.late_chain import write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.specialists.grok_bot import front_office_call
from essay_writer.state import load_state, save_state
from essay_writer.word import write_docx

MODULE = "RPT5001AA"
NUMBER = "445566"
NAME = "Ada Example"
ADDRESS = "12 Quay Road"
ASSIGNMENT = "Individual reflective diary – Sustainability in business practice"
MODULE_TITLE = "The Challenge Module – Sustainable Business"
ENTRIES = ("North workshop", "South workshop", "East workshop", "West workshop")
COVER_SHA = "77f891455fdfc4ecd5108ba0e4c84ce8a40c43c671e87f3ffa96ef54164df047"

BRIEF = f"""
Assignment title: {ASSIGNMENT}

Module title: {MODULE_TITLE}
Module code: {MODULE}

Aim: Understand sustainability in business practice by reflecting on your own decisions.

Assignment task: Write four entries. Each entry names one decision and critically analyses it with theory and evidence.

One entry for each occasion: North workshop, South workshop, East workshop, and West workshop.

Each entry is about 500 words, plus or minus 10%.
The whole diary is 2,000 words, plus or minus 10%.
Title page, references, and appendices are outside the count.

Title page
It should specify the module title, the module code, the assignment title, student registration numbers, the word count (2,000 words), and the student address.
DO NOT include your name.
The cover of the file is marked Draft until you submit.
Draft
{MODULE_TITLE}
{MODULE}
Individual reflective diary
Registration number
Word count: 2,000
Use Arial, size 12, and 1.5 line spacing.
2.5 cm margins on all sides.
Page numbers in the lower right.
Body text fully justified.
Headings in bold and title case.
Harvard is the referencing style.

Pages: Arial, size 12. Headings bold and left-aligned. Body text fully justified. 1.5 line spacing. 2.5 cm margins on all sides. Page numbers in the lower right.

# Marking criteria/feedback form

| Criterion | 0-39 | 40-49 | 50-59 | 60-69 | 70+ |
| --- | --- | --- | --- | --- | --- |
| Structure and presentation | title page missing, name shown | partial | sound | clear | title page complete, Harvard followed, margins correct |
| Title page | thin | student address |
"""

_FRAGMENTS = (
    "It should specify the module title",
    "Student registration numbers",
    "000 words",
    "Draft",
    "Arial",
    "margins",
    "Harvard",
    "Structure",
    "Presentation",
    "title case",
    "Individual reflective diary:",
)


def _write(text: str) -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(text, encoding="utf-8")


def _kick() -> str:
    args = {"module": MODULE, "confirmPhase": "draft"}
    draft = front_office_call("essay.sequence.new", args)
    assert draft["ok"] is True, draft
    args["confirmPhase"] = "confirm"
    opened = front_office_call("essay.sequence.new", args)
    assert opened["ok"] is True, opened
    return opened["sequence_id"]


def _pack(seq: str) -> dict:
    return {
        "essay_id": seq,
        "title": "Draft",
        "assessment": "Draft",
        "referencing": "harvard",
        "module_title": MODULE_TITLE,
        "conclusion": "noted",
        "student_name": NAME,
        "sections": [
            {
                "heading": name,
                "body": "A decision, with theory, on the practice the brief names.",
            }
            for name in ENTRIES
        ],
    }


def _report(pack: dict) -> dict:
    return {
        "title": "Draft",
        "word_count": {"counted": 40, "stated": 2000},
        "sections": pack["sections"],
        "figures": [],
        "tables": [],
        "sources": [],
    }


def _cover(path: Path) -> list[str]:
    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading"):
            break
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    return lines


def _stable(lines: list[str]) -> str:
    shown = [
        "Counted words:" if line.startswith("Counted words:") else line
        for line in lines
    ]
    return "\n".join(shown)


def test_ac2_instruction_prose_is_not_a_cover_line():
    """AC2. Instruction prose, a sample Draft page, and marking comments are not fields."""
    spec = title_page_spec(BRIEF)
    assert spec["labels"] == ["Student address"]
    assert spec["name"] == "forbidden"
    assert "name" not in spec["fields"]
    blob = "\n".join(spec["labels"])
    for fragment in _FRAGMENTS:
        assert fragment.casefold() not in blob.casefold()
    assert named_entries(BRIEF) == list(ENTRIES)


def test_ac3_asked_label_stays_empty_and_name_stays_off(workspace):
    """AC3. Student address is labeled and empty. The banned name is not printed."""
    _write(BRIEF)
    seq = _kick()
    pack = _pack(seq)
    lines = diary_cover_lines(seq, pack, _report(pack))
    assert lines.count("Student address:") == 1
    assert all(line == "Student address:" for line in lines if "address" in line.casefold())
    assert "Registration number:" in lines
    assert NUMBER not in "\n".join(lines)
    assert NAME not in "\n".join(lines)
    assert not any(line.casefold().startswith("name:") for line in lines)
    assert "2,000" not in "\n".join(lines)
    assert "Draft" not in lines


def test_ac4_registration_module_assignment_and_word_count(workspace):
    """AC4. Module, code, assignment title, registration, and word count stay on the cover."""
    spec = title_page_spec(BRIEF)
    for key in ("module_title", "module_code", "assignment_title", "registration", "word_count"):
        assert key in spec["fields"]
    assert spec["registration"] is True
    _write(BRIEF)
    seq = _kick()
    assert remember_registration(seq, NUMBER) == NUMBER
    pack = pack_with_registration(seq, _pack(seq))
    lines = diary_cover_lines(seq, pack, _report(pack))
    assert lines[:6] == [
        MODULE_TITLE,
        MODULE,
        ASSIGNMENT,
        f"Registration number: {NUMBER}",
        "Counted words: 40",
        "Student address:",
    ]
    assert lines.count(ASSIGNMENT) == 1
    assert "Draft" not in lines


def test_ac1_stage10_word_has_one_cover_block(workspace, monkeypatch, tmp_path):
    """AC1. Stage 10 Word has one cover. The assignment title is on that cover, once."""
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    _write(BRIEF)
    seq = _kick()
    assert remember_registration(seq, NUMBER) == NUMBER
    pack = _pack(seq)
    report = _report(pack)
    stage = work_dir(seq) / "output" / f"{seq}.docx"
    write_docx(seq, pack, report, stage)
    first_cover = _cover(stage)
    assert "Student address:" in first_cover
    state = load_state(seq)
    bare = dict(pack)
    assert "registration_number" not in bare
    complete = dict(bare)
    complete["title_page_fields"] = {"Student address": ADDRESS}
    # Four short entries sit outside the 2,000-word band. The person has accepted that.
    complete["word_count_acknowledged"] = True
    state["saved_pack"] = with_word_references(complete)
    save_state(seq, state)
    final = write_stage10(seq)
    assert final["copied"] is True, final
    final_path = mac_handoff_sequence_dir(documents, seq) / f"{MODULE}-final.docx"
    assert final_path.is_file()
    assert seq == f"{MODULE}-1"
    for path in (stage, final_path):
        cover = _cover(path)
        assert _stable(cover) == "\n".join(
            [
                MODULE_TITLE,
                MODULE,
                ASSIGNMENT,
                f"Registration number: {NUMBER}",
                "Counted words:",
                ADDRESS,
                f"{MODULE}-1",
            ]
        )
        assert hashlib.sha256(_stable(cover).encode()).hexdigest() == COVER_SHA
        document = Document(str(path))
        texts = [paragraph.text for paragraph in document.paragraphs]
        assert texts.count(ASSIGNMENT) == 1
        assert not any(text.strip() == "Draft" for text in texts)
        assert NAME not in "\n".join(texts)
        headings = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.style is not None and paragraph.style.name.startswith("Heading")
        ]
        assert headings[: len(ENTRIES)] == list(ENTRIES)
        assert headings[-1] == "References"
        joined = "\n".join(texts)
        for fragment in _FRAGMENTS:
            assert fragment not in joined
