"""Module-seq E2E #2. Stage 10 Word cover and References on a report sequence.

AC1: report Word cover carries module or sequence id, assignment title, and counted words
when the brief stayed silent on the title page.
AC2: References list is non-empty when the body cites gathered sources, including after
a body-only pack recovery from the sequence Word file.
AC3: Build covered. Cursor and Grok Build share write_stage10, cover_lines, and write_docx.
"""

import inspect
from pathlib import Path

from docx import Document
from helpers import (
    ESSAY,
    mac_handoff_sequence_dir,
    open_sequence,
    pack,
    report_sections,
    write_upload,
)

from essay_writer.late_chain import write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.research import essay_research_summary
from essay_writer.validate import essay_pack_validate, inspect_pack
from essay_writer.word import cover_lines, write_docx

MODULE = ESSAY
TOPIC = "How should the firm manage corridor risk?"
ASSIGNMENT = "Individual supply chain report"
BAND = "High Distinction"
REPORT_BRIEF = (
    "Genre: report\n"
    f"Question: How should the firm manage corridor risk?\n"
    f"Assignment title: {ASSIGNMENT}\n"
)
SCHEME = "High Distinction (80-100): the question is answered with evidence.\n"


def _brief() -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(REPORT_BRIEF, encoding="utf-8")
    (folder / "mark-scheme.md").write_text(SCHEME, encoding="utf-8")


def _paragraphs(path: Path) -> list[str]:
    return [paragraph.text.strip() for paragraph in Document(str(path)).paragraphs if paragraph.text.strip()]


def _cover(path: Path) -> list[str]:
    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.startswith("Heading"):
            break
        lines.append(text)
    return lines


def _references_body(path: Path) -> list[str]:
    texts = _paragraphs(path)
    if "References" not in texts:
        return []
    index = texts.index("References")
    return texts[index + 1 :]


def _report_pack(seq: str) -> dict:
    cite_hasan = "(Hasan and Sarker, 2023)"
    cite_imma = "(Immadisetty, 2022)"
    body = pack(seq)
    body["assessment"] = ASSIGNMENT
    body["sections"] = report_sections(
        f"{BAND}. Corridor risk is material {cite_hasan}.",
        f"One named chain is in scope {cite_imma}.",
        f"Evidence supports the finding {cite_hasan}.",
        f"The conclusion follows the sources {cite_imma}.",
    )
    body["sources"] = [
        {
            "id": "src-hasan",
            "origin": "upload",
            "locator": f"uploads/{MODULE}/hasan.md",
            "author": "Hasan and Sarker",
            "year": "2023",
            "title": "Corridor risk is material for the chain",
            "quote": "Corridor risk is material for the chain.",
        },
        {
            "id": "src-imma",
            "origin": "upload",
            "locator": f"uploads/{MODULE}/immadisetty.md",
            "author": "Immadisetty",
            "year": "2022",
            "title": "One named chain is in scope for the report",
            "quote": "One named chain is in scope for the report.",
        },
    ]
    return body


def _confirm_research(seq: str, body: dict) -> None:
    sources = body["sources"]
    draft = essay_research_summary(seq, TOPIC, "draft", sources)
    assert draft["ok"] is True, draft
    essay_research_summary(seq, TOPIC, "confirm", sources)


def test_ac1_report_stage10_cover_has_house_fields(workspace, monkeypatch, tmp_path):
    """AC1. Silent title-page brief still prints id, assignment title, and counted words."""
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    _brief()
    write_upload(
        workspace,
        MODULE,
        "hasan.md",
        (
            "Hasan and Sarker (2023) How should the firm manage corridor risk? "
            "Corridor risk is material for the chain.\n"
        ),
    )
    write_upload(
        workspace,
        MODULE,
        "immadisetty.md",
        (
            "Immadisetty (2022) How should the firm manage corridor risk? "
            "One named chain is in scope for the report.\n"
        ),
    )
    parent = open_sequence(MODULE)
    from essay_writer.projects import essay_sequence_new

    draft = essay_sequence_new(MODULE, "draft", parent=parent)
    assert draft["ok"] is True, draft
    seq = essay_sequence_new(MODULE, "confirm", parent=parent)["sequence_id"]
    assert seq.endswith("-2")

    body = _report_pack(seq)
    _confirm_research(seq, body)
    draft_validate = essay_pack_validate(seq, body, "draft")
    assert draft_validate["ok"] is True, draft_validate
    confirmed = essay_pack_validate(seq, body, "confirm")
    assert confirmed["ok"] is True, confirmed

    report = inspect_pack(seq, body)
    cover = cover_lines(seq, body, report)
    assert seq in cover
    assert ASSIGNMENT in cover
    assert any(line.startswith("Counted words:") for line in cover)

    stage_path = work_dir(seq) / "output" / f"{seq}.docx"
    write_docx(seq, body, report, stage_path)
    from essay_writer.state import load_state, save_state

    state = load_state(seq)
    state.pop("saved_pack", None)
    save_state(seq, state)

    final = write_stage10(seq)
    assert final.get("refused") != "title-page", final
    assert final["sequence_written"] is True, final
    assert final["copied"] is True, final

    for path in (stage_path, mac_handoff_sequence_dir(documents, seq) / f"{MODULE}-final.docx"):
        lines = _cover(path)
        assert seq in lines
        assert ASSIGNMENT in lines
        assert any(item.startswith("Counted words:") for item in lines)


def test_ac2_stage10_references_follow_gathered_cites(workspace, monkeypatch, tmp_path):
    """AC2. Harvard entries print when the pack body was recovered without sources."""
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    _brief()
    write_upload(
        workspace,
        MODULE,
        "hasan.md",
        (
            "Hasan and Sarker (2023) How should the firm manage corridor risk? "
            "Corridor risk is material for the chain.\n"
        ),
    )
    write_upload(
        workspace,
        MODULE,
        "immadisetty.md",
        (
            "Immadisetty (2022) How should the firm manage corridor risk? "
            "One named chain is in scope for the report.\n"
        ),
    )
    seq = open_sequence(MODULE)
    body = _report_pack(seq)
    _confirm_research(seq, body)
    draft_validate = essay_pack_validate(seq, body, "draft")
    assert draft_validate["ok"] is True, draft_validate
    essay_pack_validate(seq, body, "confirm")

    report = inspect_pack(seq, body)
    stage_path = work_dir(seq) / "output" / f"{seq}.docx"
    write_docx(seq, body, report, stage_path)

    from essay_writer.state import load_state, save_state

    state = load_state(seq)
    state["saved_pack"] = {
        "essay_id": seq,
        "title": "Draft",
        "referencing": "harvard",
        "conclusion": "noted",
        "sections": body["sections"],
    }
    save_state(seq, state)

    final = write_stage10(seq)
    assert final["ok"] is True, final
    refs = _references_body(stage_path)
    joined = "\n".join(refs)
    assert "Hasan" in joined
    assert "Immadisetty" in joined
    assert refs, refs


def test_ac3_build_covered_on_shared_stage10_path():
    """AC3. Cursor and Grok Build share write_stage10 and write_docx."""
    from essay_writer.late_chain import write_stage10 as chain_write
    from essay_writer.word import word_manifest

    source = inspect.getsource(chain_write)
    assert "write_docx(" in source
    assert "word_manifest(" in source
    manifest = inspect.getsource(word_manifest)
    assert "_word_pack(" in manifest
    assert "cover_lines(" in manifest
    assert "cover_lines(" in inspect.getsource(write_docx)
