"""Stage 10 rewrite: one cover, no mark-scheme crumbs, fresh sequence -10 final.

AC1: a polluted sequence Word file is rewritten to exactly one cover block.
AC2: _pack_from_docx does not wrap the file under Draft, so covers do not stack.
AC3: brief instructions and mark-scheme lines are not cover labels, including
labels already stored on the assignment record.
AC4: registration, module, assignment, word count, NAME_OFF, and invent-forbid hold.
AC5: Stage 10 Confirm refreshes output/<module>-10/<module>-final.docx from the
live sequence Word file. A cloned parent final does not stay.

Build covered: Cursor and Grok Build share write_stage10, title_page_spec,
blank_field_labels, and write_docx.
"""

import hashlib
import inspect
import json
import os
from pathlib import Path

from docx import Document
from helpers import with_word_references

from essay_writer.assignment_gate import read_assignment
from essay_writer.brief_contract import pack_with_registration, remember_registration
from essay_writer.diary import blank_field_labels, diary_cover_lines, title_page_spec
from essay_writer.late_chain import _pack_from_docx, write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.projects import essay_sequence_new
from essay_writer.shells import front_office_call
from essay_writer.specialists.grok_bot import basics_pass_lines
from essay_writer.stage_track import essay_stage, load_pointer
from essay_writer.state import load_state, save_state
from essay_writer.word import cover_lines, write_docx

MODULE = "RPT5001AA"
NUMBER = "445566"
NAME = "Ada Example"
ADDRESS = "12 Quay Road"
MODULE_TITLE = "The Challenge Module"
ASSIGNMENT = "Individual reflective diary on sustainability in business practice"
ENTRIES = ("North workshop", "South workshop", "East workshop", "West workshop")
BODY = "A decision, with theory, on the practice the brief names."
CRUMBS = (
    "Length",
    "tutorials",
    "Fluent academic style",
    "It should specify",
    "Margins",
    "Marking criteria",
    "Structure",
)

BRIEF = f"""
Assignment title: {ASSIGNMENT}

Module title: {MODULE_TITLE}
Module code: {MODULE}

Aim: Understand sustainability in business practice by reflecting on your own decisions.

Assignment task: Write four entries. Each entry names one decision and critically analyses it with theory and evidence.

One entry for each occasion: North workshop, South workshop, East workshop, and West workshop.

Title page
It should specify the module title, the module code, the assignment title, student registration numbers, the word count, and the student address.
DO NOT include your name.
Length
The whole diary is 2,000 words, plus or minus 10%.
Each entry is about 500 words.
Structure
Fluent academic style.
Use tutorials to support the reflection.
Margins
2.5 cm margins on all sides.
Page numbers in the lower right.

# Marking criteria/feedback form

| Criterion | 0-39 | 70+ |
| --- | --- | --- |
| Length | short | full length |
| Structure | Fluent academic style | tutorials and title page missing |
| Marking criteria | weak | title page missing |
| Title page | thin | student address |
"""

INLINE = """
Title page: It should specify the module title, the module code, the assignment title, student registration numbers, the word count, and the student address. Length. Tutorials. Fluent academic style. Margins. Structure.
DO NOT include your name.
# Marking criteria

| Criterion | 0-39 | 70+ |
| --- | --- | --- |
| Length | short | full |
| Structure | Fluent academic style | tutorials and title page missing |
| Title page | thin | student address |
"""

CONFIRM_BRIEF = """
Assignment title: Individual reflective diary

Module title: Practice module
Module code: RPT5001AA

Aim: Understand the practice by reflecting on your own decisions.

Assignment task: Write four entries. Each entry names one decision and critically analyses it with theory and evidence.

Title page: module title, module code, assignment title, and the word count.
Include your student ID.
The name stays off the title page and the file.

Pages: Arial, size 12. 2.5 cm margins on all sides. Page numbers in the lower right.
"""


def _write(text: str) -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(text, encoding="utf-8")


def _kick(parent: str | None = None) -> str:
    draft = essay_sequence_new(MODULE, "draft", parent)
    assert draft["ok"] is True, draft
    opened = front_office_call(
        "essay.sequence.new",
        {"module": MODULE, "confirmPhase": "confirm", "parent": parent} if parent else {
            "module": MODULE,
            "confirmPhase": "confirm",
        },
    )
    assert opened["ok"] is True, opened
    return opened["sequence_id"]


def _pack(seq: str) -> dict:
    return {
        "essay_id": seq,
        "title": "Draft",
        "assessment": "Draft",
        "referencing": "harvard",
        "module_title": MODULE_TITLE,
        "conclusion": "noted",
        "title_page_fields": {"Student address": ADDRESS},
        "sections": [{"heading": name, "body": BODY} for name in ENTRIES],
    }


def _store_pack(seq: str, pack: dict) -> None:
    from essay_writer.state import load_state, save_state

    state = load_state(seq)
    state["saved_pack"] = with_word_references(dict(pack))
    save_state(seq, state)


def _report(pack: dict) -> dict:
    return {
        "title": "Draft",
        "word_count": {"counted": 40, "stated": 2000},
        "sections": pack["sections"],
        "figures": [],
        "tables": [],
        "sources": [],
    }


def _texts(path: Path) -> list[str]:
    document = Document(str(path))
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def _cover(path: Path) -> list[str]:
    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading"):
            break
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    return lines


def _headings(path: Path) -> list[str]:
    document = Document(str(path))
    return [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style is not None
        and paragraph.style.name.startswith("Heading")
        and paragraph.text.strip()
    ]


def _polluted(path: Path) -> None:
    """Two Draft headings and three cover blocks, the shape Stage 10 was rewriting."""
    document = Document()
    outer = [
        MODULE_TITLE,
        MODULE,
        ASSIGNMENT,
        "Registration number: 111",
        "Counted words: 10",
        "Student address:",
        "Length:",
        "tutorials:",
        "Fluent academic style:",
        "It should specify the module title, the module code, and the student address.",
        "Margins:",
        "Marking criteria:",
        "Structure:",
    ]
    inner = [
        MODULE_TITLE,
        "Registration number: 999",
        "Counted words: 8",
        "Length",
        "tutorials",
        "Fluent academic style",
        "It should specify the module title.",
        "Margins",
        "Structure",
        "Marking criteria",
    ]
    older = [
        "Registration number: 777",
        "Draft",
        MODULE_TITLE,
        "Margins on all sides",
    ]
    for line in outer:
        document.add_paragraph(line)
    document.add_page_break()
    document.add_heading("Draft", level=1)
    for line in inner:
        document.add_paragraph(line)
    document.add_heading("Draft", level=1)
    for line in older:
        document.add_paragraph(line)
    for name in ENTRIES:
        document.add_paragraph(name)
        document.add_paragraph(BODY)
    document.add_paragraph("References")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _assert_clean(path: Path) -> str:
    texts = _texts(path)
    cover = _cover(path)
    assert cover == [
        MODULE_TITLE,
        MODULE,
        ASSIGNMENT,
        f"Registration number: {NUMBER}",
        "Counted words: 40",
        ADDRESS,
        f"{MODULE}-1",
    ]
    assert texts.count(f"Registration number: {NUMBER}") == 1
    assert texts.count(ASSIGNMENT) == 1
    assert texts.count(MODULE_TITLE) == 1
    assert texts.count(ADDRESS) == 1
    assert texts.count(BODY) == len(ENTRIES)
    assert _headings(path) == [*ENTRIES, "References"]
    assert not any(text.casefold() == "draft" for text in texts)
    joined = "\n".join(texts)
    for crumb in CRUMBS:
        assert crumb.casefold() not in joined.casefold()
    assert NAME not in joined
    assert "111" not in joined
    assert "999" not in joined
    assert "777" not in joined
    digest = hashlib.sha256("\n".join(texts).encode()).hexdigest()
    return digest


def test_ac1_ac2_ac4_polluted_rewrite_has_one_cover(workspace):
    """AC1, AC2, AC4. Stage 10 rewrite of a stacked Word file yields one cover."""
    _write(BRIEF)
    seq = _kick()
    record = read_assignment(seq)
    assert record["lock"] == "entries"
    assert record["named_entries"] == list(ENTRIES)
    assert remember_registration(seq, NUMBER) == NUMBER
    stage = work_dir(seq) / "output" / f"{seq}.docx"
    _polluted(stage)
    recovered = _pack_from_docx(seq, stage)
    assert recovered is not None
    assert [item["heading"] for item in recovered["sections"]] == list(ENTRIES)
    assert [item["body"] for item in recovered["sections"]] == [BODY] * len(ENTRIES)
    assert "Draft" not in json.dumps(recovered["sections"])
    recovered["title_page_fields"] = {"Student address": ADDRESS}
    _store_pack(seq, recovered)
    final = write_stage10(seq)
    assert final["ok"] is True, final
    assert final["sequence_written"] is True
    digest = _assert_clean(stage)
    sequence_final = stage.parent / f"{MODULE}-10" / f"{MODULE}-final.docx"
    assert sequence_final.is_file()
    assert hashlib.sha256(sequence_final.read_bytes()).hexdigest() == hashlib.sha256(stage.read_bytes()).hexdigest()
    again = write_stage10(seq)
    assert again["ok"] is True, again
    assert _texts(stage) == _texts(sequence_final)
    assert _assert_clean(stage) == digest


def test_ac2_clean_docx_is_not_wrapped_under_draft(workspace):
    """AC2. A clean Word file keeps its headings. The next write does not add Draft."""
    _write(BRIEF)
    seq = _kick()
    assert remember_registration(seq, NUMBER) == NUMBER
    pack = _pack(seq)
    stage = work_dir(seq) / "output" / f"{seq}.docx"
    write_docx(seq, pack_with_registration(seq, pack), _report(pack), stage)
    recovered = _pack_from_docx(seq, stage)
    assert recovered is not None
    assert [item["heading"] for item in recovered["sections"]] == list(ENTRIES)
    assert all(item["heading"] != "Draft" for item in recovered["sections"])
    assert "Registration number" not in json.dumps(recovered["sections"])
    complete = dict(pack)
    complete["title_page_fields"] = {"Student address": ADDRESS}
    _store_pack(seq, complete)
    final = write_stage10(seq)
    assert final["ok"] is True, final
    _assert_clean(stage)
    assert _headings(stage).count("Draft") == 0


def test_ac3_instruction_and_mark_scheme_crumbs_are_not_cover_labels():
    """AC3. Length, tutorials, style, and marking lines are not cover fields."""
    spec = title_page_spec(BRIEF)
    assert spec["labels"] == ["Student address"]
    assert spec["name"] == "forbidden"
    assert spec["registration"] is True
    for key in ("module_title", "module_code", "assignment_title", "registration", "word_count"):
        assert key in spec["fields"]
    assert "name" not in spec["fields"]
    inline = title_page_spec(INLINE)
    assert inline["labels"] == ["Student address"]
    blob = "\n".join(spec["labels"] + inline["labels"])
    for crumb in CRUMBS:
        assert crumb.casefold() not in blob.casefold()


def test_ac3_cached_assignment_labels_stay_off_the_cover(workspace, monkeypatch):
    """AC3. Crumbs already stored on the assignment record are not printed."""
    _write(BRIEF)
    seq = _kick()
    assert remember_registration(seq, NUMBER) == NUMBER
    record = read_assignment(seq)
    page = dict(record["title_page"])
    page["labels"] = [
        "Length",
        "tutorials",
        "Fluent academic style",
        "It should specify the module title",
        "Margins",
        "Marking criteria",
        "Structure",
        "Student address",
        ASSIGNMENT,
    ]
    dirty = dict(record)
    dirty["title_page"] = page
    monkeypatch.setattr(
        "essay_writer.assignment_gate.read_assignment",
        lambda _essay_id: dirty,
    )
    pack_body = _pack(seq)
    pack_body.pop("title_page_fields", None)
    pack = pack_with_registration(seq, pack_body)
    report = _report(pack)
    assert blank_field_labels(page, BRIEF) == ["Student address:"]
    lines = diary_cover_lines(seq, pack, report)
    shared = cover_lines(seq, pack, report)
    for printed in (lines, shared):
        joined = "\n".join(printed)
        assert printed.count("Student address:") == 1
        assert f"Registration number: {NUMBER}" in printed
        assert MODULE_TITLE in printed
        assert MODULE in printed
        assert ASSIGNMENT in printed
        assert "Counted words:" in joined
        assert NAME not in joined
        for crumb in CRUMBS:
            assert crumb.casefold() not in joined.casefold()
        assert not any(line.casefold().startswith("name:") for line in printed)


def _confirm_pack(seq: str) -> dict:
    return {
        "essay_id": seq,
        "title": "Individual reflective diary",
        "referencing": "harvard",
        "assessment": "Individual reflective diary",
        "module_title": "Practice module",
        "conclusion": "noted",
        "student_name": NAME,
        "sections": [
            {
                "heading": "One decision",
                "body": "A decision, with theory, on the practice the brief names.",
            }
        ],
    }


def _park(seq: str) -> None:
    pointer = load_pointer(seq)
    pointer["current"] = "10"
    pointer["saved"] = ["1", "2", "3", "4a", "4b", "5", "6", "7", "8", "9"]
    pointer["closed"] = False
    pointer["chain_armed"] = False
    pointer["waiting_stage"] = None
    pointer["waiting_digest"] = None
    pointer["room_grant"] = None
    pointer["room_lock"] = False
    path = work_dir(seq) / "tracking" / "stage.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(pointer, indent=2) + "\n", encoding="utf-8")


def test_ac5_stage10_confirm_refreshes_the_sequence_final(workspace):
    """AC5. Confirm rewrites the sequence -10 final. The parent copy stays stale."""
    _write(CONFIRM_BRIEF)
    parent = _kick()
    pack = _confirm_pack(parent)
    parent_docx = work_dir(parent) / "output" / f"{parent}.docx"
    write_docx(parent, pack, {
        "title": "Individual reflective diary",
        "word_count": {"counted": 12, "stated": 2000},
        "sections": pack["sections"],
        "figures": [],
        "tables": [],
        "sources": [],
    }, parent_docx)
    state = load_state(parent)
    state["saved_pack"] = with_word_references(pack)
    save_state(parent, state)
    stale = work_dir(parent) / "output" / f"{MODULE}-10" / f"{MODULE}-final.docx"
    stale.parent.mkdir(parents=True, exist_ok=True)
    leftover = Document()
    leftover.add_paragraph("stale parent final from 17 Oct")
    leftover.save(str(stale))
    os.utime(stale, (1_700_000_000, 1_700_000_000))
    parent_bytes = stale.read_bytes()
    parent_mtime = stale.stat().st_mtime_ns

    child = _kick(parent)
    assert remember_registration(child, NUMBER) == NUMBER
    child_final = work_dir(child) / "output" / f"{MODULE}-10" / f"{MODULE}-final.docx"
    parked = work_dir(child) / "reference" / f"{MODULE}-final.docx"
    assert parked.read_bytes() == parent_bytes
    assert not child_final.exists()
    _park(child)
    look = "\n".join(basics_pass_lines(child))
    draft = essay_stage(child, "10", "draft", look)
    assert draft["ok"] is True, draft
    saved = front_office_call(
        "essay.stage",
        {
            "essay_id": child,
            "stage": "10",
            "confirmPhase": "confirm",
            "cardAction": "confirm",
            "cardDigest": draft["cardDigest"],
        },
    )
    assert saved["ok"] is True, saved
    assert saved["stored"] is True
    assert saved["closed"] is True
    live = work_dir(child) / "output" / f"{child}.docx"
    live_sha = hashlib.sha256(live.read_bytes()).hexdigest()
    final_sha = hashlib.sha256(child_final.read_bytes()).hexdigest()
    assert final_sha == live_sha
    assert child_final.read_bytes() != parent_bytes
    assert child_final.stat().st_mtime_ns > parent_mtime
    assert parked.read_bytes() == parent_bytes
    assert stale.read_bytes() == parent_bytes
    assert stale.stat().st_mtime_ns == parent_mtime
    assert f"Registration number: {NUMBER}" in "\n".join(_texts(child_final))
    assert NAME not in "\n".join(_texts(child_final))
    assert "copy2" not in inspect.getsource(write_stage10)
