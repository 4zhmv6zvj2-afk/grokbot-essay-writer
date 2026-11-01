"""Writer first-pass Harvard reference check before stage 10 is done.

essay-ref-ctl-a1: the writer path checks in-text citations and the reference
list against docs/referencing-harvard.md before the final Word file is done.
essay-ref-ctl-a2: a miss blocks that close and names the gap. A pass cannot
be marked without the check.
essay-ref-ctl-a3: the independent reviewer runs the same check and clears
only when it is green. The writer cannot clear a reviewer finding.
essay-ref-ctl-a4: the check reads structured fields and Word runs.
Build covered: Cursor and Grok Build share write_stage10 and write_docx.
"""

import inspect
import json
from pathlib import Path

from docx import Document
from helpers import open_sequence, picked, write_upload

from essay_writer import harvard_ref_check
from essay_writer.brief_contract import remember_registration
from essay_writer.harvard_ref_check import (
    HANDBOOK_NAME,
    claim_reference_pass,
    clear_reference_findings,
    compliance_misses,
    writer_reference_block,
)
from essay_writer.late_chain import write_stage10
from essay_writer.paths import essay_dir, work_dir
from essay_writer.review import PERSON_BLOCK, eqr_verdict, essay_review, writer_clears_left
from essay_writer.shells import cursor_shell, front_office
from essay_writer.stage_track import essay_stage, load_pointer
from essay_writer.state import load_state, save_state
from essay_writer.textutil import reference_line, reference_parts

MODULE = "RPT5001AA"
NUMBER = "445566"
LOOK = "The diary reflects on one decision named in the brief."
QUOTE = "The paper maps stakeholders for one corridor decision."
BRIEF = """
Assignment title: Individual reflective diary

Module title: Practice module
Module code: RPT5001AA

Aim: Understand the practice by reflecting on your own decisions.

Assignment task: Write four entries. Each entry names one decision and critically analyses it with theory and evidence.

Title page: module title, module code, assignment title, and the word count.
Include your student ID.
The name stays off the title page and the file.

Pages: Arial, size 12. 2.5 cm margins on all sides. Page numbers in the lower right.
"""
BAD_LINE = (
    "Qin, Y. (2021) 'Stakeholder mapping', Practice Notes, 8(1), pp. 99-100. "
    "Available at: doi:10.9999/made."
)


def _brief() -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(BRIEF, encoding="utf-8")


def _upload(workspace) -> None:
    write_upload(
        workspace,
        MODULE,
        "qin.md",
        "\n".join(
            [
                "Qin, Y.",
                "2021",
                "Stakeholder mapping",
                "Journal of Practice",
                "8",
                "1",
                "2-9",
                "10.1000/stake",
                "104567",
                QUOTE,
            ]
        ),
    )


def _source() -> dict:
    return {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Stakeholder mapping",
        "container": "Journal of Practice",
        "volume": "8",
        "issue": "1",
        "pages": "2-9",
        "doi": "10.1000/stake",
        "article_number": "104567",
        "quote": QUOTE,
    }


def _pack(seq: str, *, references_markdown: str | None) -> dict:
    pack = {
        "essay_id": seq,
        "title": "Individual reflective diary",
        "referencing": "harvard",
        "assessment": "Individual reflective diary",
        "conclusion": "noted",
        "sources": [_source()],
        "sections": [
            {
                "heading": "One decision",
                "body": "A decision, with theory, on the practice the brief names (Qin, 2021).",
            }
        ],
    }
    if references_markdown is not None:
        pack["references_markdown"] = references_markdown
    return pack


def _park(seq: str) -> None:
    pointer = load_pointer(seq)
    pointer["current"] = "10"
    pointer["saved"] = ["1", "2", "3", "4a", "4b", "5", "6", "7", "8", "9"]
    pointer["closed"] = False
    pointer["chain_armed"] = False
    pointer["waiting_stage"] = None
    pointer["waiting_digest"] = None
    pointer["room_grant"] = None
    pointer["room_lock"] = False
    path = work_dir(seq) / "tracking" / "stage.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(pointer, indent=2) + "\n", encoding="utf-8")


def _plant(seq: str, pack: dict) -> None:
    state = load_state(seq)
    state["research"] = {"summary": {"sources": [_source()]}, "full": None}
    state["saved_pack"] = pack
    save_state(seq, state)


def _open(workspace, monkeypatch, tmp_path) -> str:
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(tmp_path / "Documents"))
    _brief()
    _upload(workspace)
    seq = open_sequence(MODULE)
    remember_registration(seq, NUMBER)
    _park(seq)
    return seq


def _italic_journals(path: Path) -> list[str]:
    document = Document(str(path))
    return [
        run.text
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.italic and run.text
    ]


def test_essay_ref_ctl_a1_writer_checks_before_stage10_done(workspace, monkeypatch, tmp_path):
    """essay-ref-ctl-a1. Stage 10 stays open until the writer check is green."""
    seq = _open(workspace, monkeypatch, tmp_path)
    _plant(seq, _pack(seq, references_markdown=BAD_LINE))
    draft = essay_stage(seq, "10", "draft", LOOK)
    assert draft["ok"] is True, draft
    refused = essay_stage(seq, "10", "confirm", LOOK, **picked(draft))
    assert refused["ok"] is False, refused
    assert refused["refused"] == "harvard-references"
    assert refused.get("closed") is not True
    assert load_pointer(seq)["closed"] is False
    assert "10" not in load_pointer(seq)["saved"]
    direct = write_stage10(seq)
    assert direct["ok"] is False, direct
    assert direct.get("sequence_written") is not True
    assert direct["checked"] is True
    assert direct["handbook"] == HANDBOOK_NAME
    reason = str(direct.get("reason") or "")
    assert "Journal of Practice" in reason
    assert "not italic" in reason.casefold()
    assert "99-100" in reason
    assert "2-9" in reason
    assert "104567" in reason
    assert "10.9999/made" in reason
    assert "Qin (2021)" in reason
    unmatched = tmp_path / "unmatched.docx"
    _save(
        unmatched,
        "The corridor is narrow (Lee, 2018).",
        reference_line(_source()),
        parts=reference_parts(_source()),
    )
    unmatched_codes = {item["code"] for item in compliance_misses(unmatched, [_source()])}
    assert "author-date" in unmatched_codes

    _plant(seq, _pack(seq, references_markdown=None))
    final = write_stage10(seq)
    assert final["ok"] is True, final
    assert final.get("sequence_written") is True
    assert final["harvard_references"]["checked"] is True
    assert final["harvard_references"]["ok"] is True
    assert final["harvard_references"]["handbook"] == HANDBOOK_NAME
    written = work_dir(seq) / "output" / f"{seq}.docx"
    text = "\n".join(paragraph.text for paragraph in Document(str(written)).paragraphs)
    assert "pp. 2-9" in text
    assert "104567" in text
    assert "doi:10.1000/stake" in text
    assert "99-100" not in text
    assert any("Journal of Practice" in piece for piece in _italic_journals(written))
    again = essay_stage(seq, "10", "draft", LOOK)
    assert again["ok"] is True, again
    saved = essay_stage(seq, "10", "confirm", LOOK, **picked(again))
    assert saved["ok"] is True, saved
    assert saved.get("closed") is True
    assert load_pointer(seq)["closed"] is True
    stage_source = inspect.getsource(write_stage10)
    assert stage_source.index("writer_reference_block(") < stage_source.index(
        '["sequence_written"] = True'
    )
    assert cursor_shell()["handlers"]["essay.stage"] is front_office()["handlers"]["essay.stage"]
    assert "write_docx(" in stage_source
    assert "compliance_misses(" in inspect.getsource(writer_reference_block)


def test_essay_ref_ctl_a2_misses_block_and_writer_cannot_mark_pass(
    workspace, monkeypatch, tmp_path
):
    """essay-ref-ctl-a2. Misses name the gap. A marked pass does not skip the check."""
    seq = _open(workspace, monkeypatch, tmp_path)
    _plant(seq, _pack(seq, references_markdown=BAD_LINE))
    state = load_state(seq)
    state["harvard_reference_pass"] = {"ok": True, "checked": True}
    save_state(seq, state)
    marked = claim_reference_pass()
    assert marked["ok"] is False
    assert marked["checked"] is False
    assert marked.get("stored") is not True
    assert marked["writer_clears_left"] == 0
    assert "has not run" in marked["reason"]
    refused = write_stage10(seq)
    assert refused["ok"] is False, refused
    assert refused.get("sequence_written") is not True
    assert refused["refused"] == "harvard-references"
    codes = {item["code"] for item in refused["misses"]}
    assert "journal-italic" in codes
    assert "pages" in codes
    assert "article-number" in codes
    assert "doi" in codes
    assert load_state(seq).get("harvard_reference_pass") == {"ok": True, "checked": True}
    assert inspect.signature(writer_reference_block).parameters.keys() == {
        "essay_id",
        "path",
        "sources",
    }
    reviewer_mark = claim_reference_pass(seat="reviewer")
    assert reviewer_mark["ok"] is False
    assert reviewer_mark["checked"] is False


def test_essay_ref_ctl_a3_reviewer_reruns_and_writer_cannot_clear(
    workspace, monkeypatch, tmp_path
):
    """essay-ref-ctl-a3. The reviewer clears only when the same check is green."""
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    folder = essay_dir("MOD1004AA")
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(
        "Genre: report\nQuestion: How do chokepoints constrain the corridor?\n",
        encoding="utf-8",
    )
    seq = open_sequence("MOD1004AA")
    source = _source()
    source["locator"] = "uploads/MOD1004AA/qin.md"
    write_upload(workspace, "MOD1004AA", "qin.md", "Qin, Y.\n2021\nJournal of Practice\n")
    state = load_state(seq)
    state["research"] = {"summary": {"sources": [source]}, "full": None}
    state["saved_pack"] = {"essay_id": seq, "sources": [source]}
    save_state(seq, state)
    path = work_dir(seq) / "output" / f"{seq}.docx"
    _word(path, italic=False)
    blocked = essay_review(
        seq,
        "9",
        "PASS",
        "draft",
        constitution="This is a report. How do chokepoints constrain the corridor?",
    )
    assert blocked["ok"] is False, blocked
    assert blocked["refused"] == "harvard-references"
    assert blocked["checked"] is True
    assert blocked.get("cleared") is not True
    assert blocked.get("stored") is not True
    assert "Journal of Practice" in str(blocked.get("reason") or "")
    assert "eqr" not in load_state(seq)
    writer = essay_review(seq, "9", "PASS", "confirm", seat="writer")
    assert writer["ok"] is False
    assert writer["refused"] == "seat"
    assert writer["writer_clears_left"] == 0
    assert PERSON_BLOCK in writer["reason"]
    assert writer_clears_left() == 0
    own = clear_reference_findings(seq, path, seat="writer")
    assert own["ok"] is False
    assert own.get("cleared") is not True
    assert own["checked"] is False
    _word(path, italic=True)
    draft = essay_review(
        seq,
        "9",
        "PASS",
        "draft",
        constitution="This is a report. How do chokepoints constrain the corridor?",
    )
    assert draft["ok"] is True, draft
    stored = essay_review(
        seq,
        "9",
        "PASS",
        "confirm",
        constitution="This is a report. How do chokepoints constrain the corridor?",
    )
    assert stored["stored"] is True, stored
    assert eqr_verdict(seq, "9") == "PASS"
    cleared = clear_reference_findings(seq, path, seat="reviewer")
    assert cleared["cleared"] is True
    assert cleared["checked"] is True
    assert cleared["misses"] == []
    assert "compliance_misses(" in inspect.getsource(clear_reference_findings)
    assert "compliance_misses(" in inspect.getsource(writer_reference_block)
    after = json.dumps(load_state(seq).get("eqr"))
    again = essay_review(seq, "9", "PASS", "confirm", seat="writer")
    assert again["stored"] is not True
    assert json.dumps(load_state(seq).get("eqr")) == after


def test_essay_ref_ctl_a4_check_uses_word_runs_not_markdown(tmp_path):
    """essay-ref-ctl-a4. The same reference text fails when the journal run is not italic."""
    source = _source()
    line = reference_line(source)
    body = "The corridor is narrow (Qin, 2021)."
    plain = tmp_path / "plain.docx"
    styled = tmp_path / "styled.docx"
    _save(plain, body, line, parts=None)
    _save(styled, body, line, parts=reference_parts(source))
    plain_doc = Document(str(plain))
    styled_doc = Document(str(styled))
    assert _reference_text(plain_doc) == _reference_text(styled_doc) == line
    plain_misses = compliance_misses(plain, [source])
    assert [item["code"] for item in plain_misses] == ["journal-italic"]
    assert compliance_misses(styled, [source]) == []
    assert not any(run.italic for paragraph in plain_doc.paragraphs for run in paragraph.runs)
    assert any(
        run.italic and "Journal of Practice" in (run.text or "")
        for paragraph in styled_doc.paragraphs
        for run in paragraph.runs
    )
    checker = inspect.getsource(harvard_ref_check)
    assert "references_markdown" not in checker
    assert ".italic" in checker
    assert "compliance_misses(" in inspect.getsource(writer_reference_block)
    blocked = writer_reference_block("MOD1004AA-1", plain, [source])
    assert blocked is not None and blocked["ok"] is False
    assert writer_reference_block("MOD1004AA-1", styled, [source]) is None


def _save(path: Path, body: str, line: str, *, parts) -> None:
    document = Document()
    document.add_paragraph(body)
    document.add_paragraph("References")
    paragraph = document.add_paragraph()
    if parts is None:
        paragraph.add_run(line)
    else:
        for text, italic in parts:
            run = paragraph.add_run(text)
            run.italic = italic
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _reference_text(document) -> str:
    texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return texts[texts.index("References") + 1]


def _word(path: Path, *, italic: bool) -> None:
    source = _source()
    line = reference_line(source)
    document = Document()
    document.add_paragraph(
        "This is a report. How do chokepoints constrain the corridor? (Qin, 2021)."
    )
    document.add_paragraph("References")
    paragraph = document.add_paragraph()
    if italic:
        for text, flag in reference_parts(source):
            run = paragraph.add_run(text)
            run.italic = flag
    else:
        paragraph.add_run(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def test_harvard_check_reads_a_single_page_label(tmp_path):
    """A reference list that shows p. 32 for one page is not a page miss."""
    source = {**_source(), "pages": "32"}
    line = reference_line(source)
    path = tmp_path / "single-page.docx"
    _save(path, "The corridor is narrow (Qin, 2021).", line, parts=reference_parts(source))
    assert ", p. 32." in line
    assert compliance_misses(path, [source]) == []


def test_harvard_check_refuses_the_wrong_page_label(tmp_path):
    """A single page needs p.; a range needs pp. The first pass names the gap."""
    single = {**_source(), "pages": "32"}
    parts = [
        (text.replace("p. 32", "pp. 32"), italic)
        for text, italic in reference_parts(single)
    ]
    path = tmp_path / "wrong-single.docx"
    _save(path, "The corridor is narrow (Qin, 2021).", reference_line(single), parts=parts)
    misses = compliance_misses(path, [single])
    assert [item["code"] for item in misses] == ["pages-label"]
    assert "should show p. 32" in misses[0]["detail"]
    assert "pp." in misses[0]["detail"]

    ranged = {**_source(), "pages": "32-40"}
    ranged_parts = [
        (text.replace("pp. 32-40", "p. 32-40"), italic)
        for text, italic in reference_parts(ranged)
    ]
    ranged_path = tmp_path / "wrong-range.docx"
    _save(
        ranged_path,
        "The corridor is narrow (Qin, 2021).",
        reference_line(ranged),
        parts=ranged_parts,
    )
    range_misses = compliance_misses(ranged_path, [ranged])
    assert [item["code"] for item in range_misses] == ["pages-label"]
    assert "should show pp. 32-40" in range_misses[0]["detail"]


def test_harvard_check_refuses_a_bare_edition(tmp_path):
    """A book line that prints 2nd instead of 2nd edn. is a first-pass miss."""
    book = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Port economics",
        "edition": "2nd",
        "place": "Bristol",
        "publisher": "Harbour Press",
    }
    body = "The corridor is narrow (Qin, 2021)."
    correct = tmp_path / "edition-correct.docx"
    _save(correct, body, reference_line(book), parts=reference_parts(book))
    assert compliance_misses(correct, [book]) == []

    wrong_parts = [
        (text.replace("2nd edn.", "2nd"), italic)
        for text, italic in reference_parts(book)
    ]
    wrong = tmp_path / "edition-wrong.docx"
    _save(wrong, body, reference_line(book), parts=wrong_parts)
    misses = compliance_misses(wrong, [book])
    assert [item["code"] for item in misses] == ["edition"]
    assert "should show 2nd edn." in misses[0]["detail"]


def test_harvard_check_refuses_a_wrong_volume_or_issue(tmp_path):
    """A journal line with the wrong volume or issue is a first-pass miss."""
    source = _source()
    correct = tmp_path / "volume-correct.docx"
    _save(correct, "The corridor is narrow (Qin, 2021).", reference_line(source), parts=reference_parts(source))
    assert compliance_misses(correct, [source]) == []

    wrong_parts = [
        (text.replace("8(1)", "9(1)"), italic)
        for text, italic in reference_parts(source)
    ]
    wrong = tmp_path / "volume-wrong.docx"
    _save(wrong, "The corridor is narrow (Qin, 2021).", reference_line(source), parts=wrong_parts)
    misses = compliance_misses(wrong, [source])
    assert [item["code"] for item in misses] == ["volume-issue"]
    assert "should show 8(1)" in misses[0]["detail"]


def _web_source() -> dict:
    return {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Stakeholder mapping",
        "url": "https://example.org/stake",
        "accessed": "17 May 2022",
    }


def test_harvard_check_refuses_an_invented_web_locator(tmp_path):
    """A web address or access date the source did not gather is a first-pass miss."""
    web = _web_source()
    body = "The corridor is narrow (Qin, 2021)."
    correct = tmp_path / "web-correct.docx"
    _save(correct, body, reference_line(web), parts=reference_parts(web))
    assert compliance_misses(correct, [web]) == []

    wrong_url = [
        (
            text.replace("https://example.org/stake", "https://example.org/other"),
            italic,
        )
        for text, italic in reference_parts(web)
    ]
    url_path = tmp_path / "web-wrong-url.docx"
    _save(url_path, body, reference_line(web), parts=wrong_url)
    url_misses = compliance_misses(url_path, [web])
    assert [item["code"] for item in url_misses] == ["url"]
    assert "should show https://example.org/stake" in url_misses[0]["detail"]
    assert "example.org/other" in url_misses[0]["detail"]

    wrong_date = [
        (text.replace("17 May 2022", "18 May 2022"), italic)
        for text, italic in reference_parts(web)
    ]
    date_path = tmp_path / "web-wrong-date.docx"
    _save(date_path, body, reference_line(web), parts=wrong_date)
    date_misses = compliance_misses(date_path, [web])
    assert [item["code"] for item in date_misses] == ["accessed"]
    assert "should show (Accessed: 17 May 2022)" in date_misses[0]["detail"]
    assert "18 May 2022" in date_misses[0]["detail"]

    missing_path = tmp_path / "web-missing-locator.docx"
    missing_parts = [
        ("Qin, Y. (2021) ", False),
        ("Stakeholder mapping", True),
        (".", False),
    ]
    _save(missing_path, body, "Qin, Y. (2021) Stakeholder mapping.", parts=missing_parts)
    missing_misses = compliance_misses(missing_path, [web])
    assert [item["code"] for item in missing_misses] == ["url"]
    assert "should show https://example.org/stake" in missing_misses[0]["detail"]


def test_harvard_check_demands_an_access_date_when_a_web_page_has_a_publisher(tmp_path):
    """A place or publisher on a web entry does not cancel the access date."""
    web = {
        **_web_source(),
        "author": "World Bank",
        "title": "Corridor review",
        "place": "Washington, DC",
        "publisher": "World Bank",
    }
    body = "The corridor is narrow (World Bank, 2021)."
    line = reference_line(web)
    assert "Washington, DC: World Bank" in line
    correct = tmp_path / "web-publisher-correct.docx"
    _save(correct, body, line, parts=reference_parts(web))
    assert compliance_misses(correct, [web]) == []

    no_date = line.replace(" (Accessed: 17 May 2022)", "")
    no_date_parts = [
        (text.replace(" (Accessed: 17 May 2022)", ""), italic)
        for text, italic in reference_parts(web)
    ]
    assert "(Accessed:" not in no_date
    path = tmp_path / "web-publisher-no-date.docx"
    _save(path, body, no_date, parts=no_date_parts)
    misses = compliance_misses(path, [web])
    assert [item["code"] for item in misses] == ["accessed"]
    assert "should show (Accessed: 17 May 2022)" in misses[0]["detail"]


def test_harvard_check_reads_the_year_after_the_name(tmp_path):
    """A year printed in another entry's title is not the in-text year."""
    book = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Smith, J.",
        "year": "2019",
        "title": "The 2020 corridor review",
        "place": "London",
        "publisher": "Press",
    }
    line = reference_line(book)
    assert "Smith, J. (2019)" in line
    assert "2020" in line

    right = tmp_path / "year-right.docx"
    _save(right, "The corridor is narrow (Smith, 2019).", line, parts=reference_parts(book))
    assert compliance_misses(right, [book]) == []

    wrong = tmp_path / "year-in-title.docx"
    _save(wrong, "The corridor is narrow (Smith, 2020).", line, parts=reference_parts(book))
    misses = compliance_misses(wrong, [book])
    assert [item["code"] for item in misses] == ["author-date"]
    assert "2020" in misses[0]["detail"]


def test_harvard_check_reads_a_page_list(tmp_path):
    """A reference list that prints pp. 32, 34 matches a gathered page list."""
    source = {**_source(), "pages": "32, 34"}
    line = reference_line(source)
    assert ", pp. 32, 34." in line
    path = tmp_path / "page-list.docx"
    _save(path, "The corridor is narrow (Qin, 2021).", line, parts=reference_parts(source))
    assert compliance_misses(path, [source]) == []


def test_harvard_check_refuses_a_wrong_place_or_publisher(tmp_path):
    """A book line must print the gathered place and publisher."""
    book = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Port economics",
        "place": "Bristol",
        "publisher": "Harbour Press",
    }
    body = "The corridor is narrow (Qin, 2021)."
    line = reference_line(book)
    correct = tmp_path / "publisher-correct.docx"
    _save(correct, body, line, parts=reference_parts(book))
    assert compliance_misses(correct, [book]) == []

    wrong_publisher = tmp_path / "publisher-wrong.docx"
    wrong_publisher_parts = [
        (text.replace("Harbour Press", "Other Press"), italic)
        for text, italic in reference_parts(book)
    ]
    _save(
        wrong_publisher,
        body,
        line.replace("Harbour Press", "Other Press"),
        parts=wrong_publisher_parts,
    )
    publisher_misses = compliance_misses(wrong_publisher, [book])
    assert [item["code"] for item in publisher_misses] == ["publisher"]
    assert "should show Harbour Press" in publisher_misses[0]["detail"]

    wrong_place = tmp_path / "place-wrong.docx"
    wrong_place_parts = [
        (text.replace("Bristol", "Leeds"), italic)
        for text, italic in reference_parts(book)
    ]
    _save(wrong_place, body, line.replace("Bristol", "Leeds"), parts=wrong_place_parts)
    place_misses = compliance_misses(wrong_place, [book])
    assert [item["code"] for item in place_misses] == ["place"]
    assert "should show Bristol" in place_misses[0]["detail"]

    no_publisher = tmp_path / "publisher-missing.docx"
    no_publisher_parts = [
        ("Qin, Y. (2021) ", False),
        ("Port economics", True),
        (". Bristol.", False),
    ]
    _save(
        no_publisher,
        body,
        "Qin, Y. (2021) Port economics. Bristol.",
        parts=no_publisher_parts,
    )
    missing_misses = compliance_misses(no_publisher, [book])
    assert [item["code"] for item in missing_misses] == ["publisher"]
    assert "should show Harbour Press" in missing_misses[0]["detail"]


def test_harvard_check_refuses_a_surname_stub(tmp_path):
    """A reference entry prints the gathered initials, not a surname alone."""
    source = _source()
    body = "The corridor is narrow (Qin, 2021)."
    stub_parts = [
        (text.replace("Qin, Y.", "Qin"), italic)
        for text, italic in reference_parts(source)
    ]
    stub = tmp_path / "author-stub.docx"
    _save(stub, body, reference_line(source), parts=stub_parts)
    misses = compliance_misses(stub, [source])
    assert [item["code"] for item in misses] == ["author"]
    assert "should show Qin, Y." in misses[0]["detail"]


def test_harvard_check_refuses_an_article_number_as_a_page(tmp_path):
    """An article number in the page slot is not a page label."""
    source = {**_source(), "article_number": "e12345"}
    source.pop("pages", None)
    body = "The corridor is narrow (Qin, 2021)."
    correct = tmp_path / "article-correct.docx"
    _save(correct, body, reference_line(source), parts=reference_parts(source))
    assert "article e12345" in reference_line(source)
    assert compliance_misses(correct, [source]) == []

    wrong_parts = [
        (text.replace("article e12345", "pp. e12345"), italic)
        for text, italic in reference_parts(source)
    ]
    wrong = tmp_path / "article-as-page.docx"
    _save(wrong, body, reference_line(source), parts=wrong_parts)
    misses = compliance_misses(wrong, [source])
    assert [item["code"] for item in misses] == ["article-pages"]
    assert "pp. e12345" in misses[0]["detail"]
    assert "article number, not a page" in misses[0]["detail"]


def test_harvard_check_reads_a_stored_page_or_doi_label(tmp_path):
    """A label already on the gathered page or DOI is not a miss."""
    body = "The corridor is narrow (Qin, 2021)."
    page_source = {**_source(), "pages": "pp. 2-9"}
    page_line = reference_line(page_source)
    assert "pp. 2-9" in page_line
    page_path = tmp_path / "stored-page-label.docx"
    _save(page_path, body, page_line, parts=reference_parts(page_source))
    assert compliance_misses(page_path, [page_source]) == []

    for stored in ("https://doi.org/10.1000/stake", "doi:10.1000/stake"):
        doi_source = {**_source(), "doi": stored}
        line = reference_line(doi_source)
        assert "Available at: doi:10.1000/stake" in line
        path = tmp_path / f"stored-doi-{len(stored)}.docx"
        _save(path, body, line, parts=reference_parts(doi_source))
        assert compliance_misses(path, [doi_source]) == []


def test_harvard_check_does_not_ask_a_book_for_pages(tmp_path):
    """A book line the printer writes has no page range, so pages are not a miss."""
    book = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Port economics",
        "place": "Bristol",
        "publisher": "Harbour Press",
        "pages": "32-34",
    }
    body = "The corridor is narrow (Qin, 2021)."
    line = reference_line(book)
    assert "pp." not in line
    path = tmp_path / "book-pages.docx"
    _save(path, body, line, parts=reference_parts(book))
    assert compliance_misses(path, [book]) == []


def test_harvard_check_reads_a_page_cite_for_the_list(tmp_path):
    """An in-text cite with a page still needs a matching reference entry."""
    source = _source()
    other_line = "Jones, K. (2019) 'Port limits', Practice, 1(1), pp. 1-2."
    bodies = (
        "The corridor is narrow (Qin, 2021, p. 32).",
        "Qin (2021, p. 32) shows the narrow corridor.",
    )
    for index, body in enumerate(bodies):
        path = tmp_path / f"page-cite-{index}.docx"
        _save(path, body, other_line, parts=None)
        misses = compliance_misses(path, [source])
        assert [item["code"] for item in misses] == ["author-date"]

    covered = tmp_path / "page-cite-covered.docx"
    _save(
        covered,
        bodies[0],
        reference_line(source),
        parts=reference_parts(source),
    )
    assert compliance_misses(covered, [source]) == []


def test_harvard_check_does_not_ask_a_book_for_a_journal_slot(tmp_path):
    """A book line has no volume(issue) or article slot, so neither is a miss."""
    book = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Port economics",
        "place": "Bristol",
        "publisher": "Harbour Press",
        "volume": "2",
        "issue": "3",
        "article_number": "e12345",
    }
    body = "The corridor is narrow (Qin, 2021)."
    line = reference_line(book)
    assert "2(3)" not in line
    assert "e12345" not in line
    path = tmp_path / "book-journal-slot.docx"
    _save(path, body, line, parts=reference_parts(book))
    assert compliance_misses(path, [book]) == []


def test_harvard_check_refuses_a_reference_year_that_is_not_gathered(tmp_path):
    """The entry's author-date carries the gathered year, not another one."""
    source = _source()
    earlier = {**source, "year": "2019"}
    body = "The corridor is narrow (Qin, 2019)."
    path = tmp_path / "wrong-year.docx"
    _save(path, body, reference_line(earlier), parts=reference_parts(earlier))
    misses = compliance_misses(path, [source])
    assert [item["code"] for item in misses] == ["year"]
    assert "should show (2021) after the author" in misses[0]["detail"]

    right = tmp_path / "right-year.docx"
    _save(
        right,
        "The corridor is narrow (Qin, 2021).",
        reference_line(source),
        parts=reference_parts(source),
    )
    assert compliance_misses(right, [source]) == []


def test_harvard_check_refuses_an_unquoted_article_title(tmp_path):
    """An article title printed before the journal needs single quotation marks."""
    source = _source()
    body = "The corridor is narrow (Qin, 2021)."
    line = reference_line(source)
    correct = tmp_path / "article-quoted.docx"
    _save(correct, body, line, parts=reference_parts(source))
    assert compliance_misses(correct, [source]) == []

    plain = line.replace("'Stakeholder mapping'", "Stakeholder mapping")
    parts = [
        (text.replace("'Stakeholder mapping'", "Stakeholder mapping"), italic)
        for text, italic in reference_parts(source)
    ]
    path = tmp_path / "article-unquoted.docx"
    _save(path, body, plain, parts=parts)
    misses = compliance_misses(path, [source])
    assert [item["code"] for item in misses] == ["article-quotes"]
    assert "single quotation marks" in misses[0]["detail"]


def test_harvard_check_refuses_an_invented_doi_on_a_book(tmp_path):
    """A DOI the source did not gather is a miss on a book line too."""
    book = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Port economics",
        "place": "Bristol",
        "publisher": "Harbour Press",
    }
    body = "The corridor is narrow (Qin, 2021)."
    line = reference_line(book)
    assert "doi:" not in line
    clear = tmp_path / "book-no-doi.docx"
    _save(clear, body, line, parts=reference_parts(book))
    assert compliance_misses(clear, [book]) == []

    invented_line = line.rstrip(".") + ". Available at: doi:10.9999/made."
    path = tmp_path / "book-invented-doi.docx"
    invented_parts = [
        *reference_parts(book),
        (" Available at: doi:10.9999/made.", False),
    ]
    _save(path, body, invented_line, parts=invented_parts)
    misses = compliance_misses(path, [book])
    assert [item["code"] for item in misses] == ["invented-doi"]
    assert "10.9999/made" in misses[0]["detail"]


def test_harvard_check_refuses_a_plain_book_title(tmp_path):
    """A book title in the entry is italic. A roman title is a first-pass miss."""
    book = {
        "id": "src-1",
        "origin": "upload",
        "locator": f"uploads/{MODULE}/qin.md",
        "author": "Qin, Y.",
        "year": "2021",
        "title": "Port economics",
        "place": "Bristol",
        "publisher": "Harbour Press",
    }
    body = "The corridor is narrow (Qin, 2021)."
    correct = tmp_path / "book-title-italic.docx"
    _save(correct, body, reference_line(book), parts=reference_parts(book))
    assert compliance_misses(correct, [book]) == []

    roman_parts = [(text, False) for text, _italic in reference_parts(book)]
    path = tmp_path / "book-title-roman.docx"
    _save(path, body, reference_line(book), parts=roman_parts)
    misses = compliance_misses(path, [book])
    assert [item["code"] for item in misses] == ["title-italic"]
    assert "Port economics" in misses[0]["detail"]
