"""Word upload identify, roll-forward, new-from-Word honesty, and invent-forbid."""

import json
from pathlib import Path

from docx import Document
from helpers import ESSAY, TOPIC, open_sequence, source, write_upload

from essay_writer.doors import DOORS
from essay_writer.health import essay_health
from essay_writer.intake import essay_intake
from essay_writer.paths import essay_dir, kb1_module_dir, sequence_dir, uploads_dir
from essay_writer.projects import (
    essay_projects,
    essay_sequence_new,
    save_sequence_state,
    stage_name,
)
from essay_writer.research import essay_research_summary
from essay_writer.revise import essay_revise
from essay_writer.server import INSTRUCTIONS
from essay_writer.writing_basics import basics_lines

MODULE = "MOD1011AA"
OTHER = "MOD1008AA"


def _docx(path: Path, paragraphs: list[str], headings: list[str] | None = None) -> None:
    document = Document()
    for heading in headings or []:
        document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _confirm_new(module: str = MODULE, parent: str | None = None) -> dict:
    draft = essay_sequence_new(module, "draft", parent)
    assert draft["ok"] is True
    return essay_sequence_new(module, "confirm", parent)


def _port_output(module: str = MODULE, number: int = 1) -> Path:
    path = sequence_dir(module, number) / "output" / f"{module}-{number}.docx"
    _docx(
        path,
        [
            "The port handles containers at the quay.",
            "Ships wait offshore when the berth is full.",
        ],
    )
    return path


def _rewrite(path: Path) -> None:
    _docx(
        path,
        [
            "The port handles containers at the quay and alongside barges.",
            "Ships wait offshore when the berth is full of waiting traffic.",
            "Reference List",
            "Smith, A. (2020) Port pricing. Journal of Ports.",
            "Unknown, B. (1999) Missing paper.",
            "Lee, C. (2019) Hull form.",
        ],
    )


def _research(module: str = MODULE, number: int = 1) -> None:
    save_sequence_state(
        module,
        number,
        {
            "essay_id": module,
            "sequence": number,
            "sequence_id": f"{module}-{number}",
            "research": {
                "summary": {
                    "topic": "ports",
                    "sources": [
                        {"id": "src-1", "title": "Port pricing", "authors": ["Smith"]},
                        {
                            "id": "src-2",
                            "title": "Some other study",
                            "authors": ["Unknown"],
                            "gathered": False,
                        },
                        {"id": "src-3", "authors": ["Lee"]},
                    ],
                },
                "full": None,
            },
            "validations": {},
            "word": {},
            "revisions": [],
            "shelf": {},
            "reference_index": [],
            "unresolved_references": [],
        },
    )


def test_identify_then_confirm_rolls_forward(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    parent = _port_output()
    parent_bytes = parent.read_bytes()
    _research()
    incoming = workspace / "unassigned" / "rewrite.docx"
    _rewrite(incoming)

    draft = essay_intake("rewrite.docx", "draft")
    assert draft["ok"] is True
    assert draft["identified"] is True
    assert "module" not in draft
    assert draft["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in draft
    assert draft["sequence_pick"] is False
    assert draft["kind"] == "output"
    assert draft["wrote"] is False
    assert draft["research_started"] is False
    assert "roll forward" in draft["next"]
    assert not sequence_dir(MODULE, 2).exists()

    filed = essay_intake("rewrite.docx", "confirm", MODULE)
    assert filed["ok"] is True
    assert filed["wrote"] is True
    assert filed["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in filed
    assert filed["sequence_pick"] is False
    assert filed["research_started"] is False
    assert filed.get("stage") != "user-feedback"
    assert "rolled this forward" in filed["next"]
    assert parent.read_bytes() == parent_bytes
    assert (sequence_dir(MODULE, 2) / "output" / f"{MODULE}-2.docx").is_file()
    assert not (uploads_dir(MODULE) / "rewrite.docx").exists()
    child = json.loads((sequence_dir(MODULE, 2) / "state.json").read_text(encoding="utf-8"))
    assert child["branched_from"] == f"{MODULE}-1"
    assert child["research"]["summary"]["topic"] == "ports"
    assert child.get("user_feedback") is not True


def test_identify_confirm_without_repeating_the_module(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    _port_output()
    _research()
    _rewrite(workspace / "unassigned" / "rewrite.docx")
    draft = essay_intake("rewrite.docx", "draft")
    assert draft["identified"] is True
    assert "module" not in draft
    assert draft["essay_id"] == f"{MODULE}-2"
    filed = essay_intake("rewrite.docx", "confirm")
    assert filed["ok"] is True
    assert filed["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in filed
    assert filed["wrote"] is True


def test_unidentified_word_lists_modules_without_a_sequence_pick(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    assert _confirm_new()["sequence_id"] == f"{MODULE}-2"
    (workspace / "uploads" / OTHER).mkdir(parents=True)
    _docx(
        workspace / "unassigned" / "other.docx",
        [
            "Hull plating fails when the coating is thin.",
            "Paint systems differ across drydock yards.",
            "Class surveys record steel thickness each year.",
        ],
    )

    listed = essay_intake("other.docx", "draft")
    assert listed["ok"] is True
    assert listed["identified"] is False
    assert listed["sequence_pick"] is False
    assert listed["wrote"] is False
    assert listed["stored"] is False
    assert "sequence_id" not in listed
    assert listed["rows"] == [
        {"essay_id": OTHER},
        {"essay_id": f"{MODULE}-2"},
    ]
    for row in listed["rows"]:
        assert set(row) == {"essay_id"}
        assert "sequence_id" not in row
        assert "module" not in row
        assert "sequence" not in row
        assert "last_sequential_id" not in row
        assert "stage" not in row
    assert "roll this upload forward" in listed["next"]
    assert "pick a sequence" not in listed["next"].casefold()
    continue_row = next(
        row for row in essay_projects("continue")["rows"] if row["essay_id"] == f"{MODULE}-2"
    )
    assert "module" not in continue_row
    assert "sequence_id" not in continue_row
    assert not sequence_dir(MODULE, 3).exists()

    refused = essay_intake("other.docx", "confirm")
    assert refused["ok"] is False
    assert refused["wrote"] is False
    assert (workspace / "unassigned" / "other.docx").is_file()


def test_chosen_module_rolls_forward_without_a_sequence_pick(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    parent = _port_output()
    parent_bytes = parent.read_bytes()
    _docx(
        workspace / "unassigned" / "other.docx",
        [
            "Hull plating fails when the coating is thin.",
            "Paint systems differ across drydock yards.",
            "Class surveys record steel thickness each year.",
        ],
    )
    listed = essay_intake("other.docx", "draft")
    assert listed["identified"] is False
    assert f"{MODULE}-1" in [row["essay_id"] for row in listed["rows"]]

    draft = essay_intake("other.docx", "draft", MODULE, rollForward=True)
    assert draft["ok"] is True
    assert draft["kind"] == "output"
    assert draft["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in draft
    assert draft["sequence_pick"] is False
    assert draft["wrote"] is False
    assert "rows" not in draft
    assert not sequence_dir(MODULE, 2).exists()

    filed = essay_intake("other.docx", "confirm", MODULE, rollForward=True)
    assert filed["ok"] is True
    assert filed["wrote"] is True
    assert filed["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in filed
    assert filed["path"].endswith(f"sequences/2/output/{MODULE}-2.docx")
    assert not (uploads_dir(MODULE) / "other.docx").exists()
    assert parent.read_bytes() == parent_bytes
    assert (sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx").is_file()


def test_new_from_word_honesty_lands_on_user_feedback(workspace):
    incoming = workspace / "unassigned" / "fresh.docx"
    _docx(
        incoming,
        [
            "The corridor depends on a narrow strait and a named supply chain.",
            "Reference List",
            "Patel, A. (2019) Chokepoints and crude.",
        ],
    )
    draft = essay_intake("fresh.docx", "draft", OTHER, create=True)
    assert draft["ok"] is True
    assert draft["wrote"] is False
    assert draft["stored"] is False
    assert draft["new_from_word"] is True
    assert draft["essay_id"] == f"{OTHER}-1"
    assert "sequence_id" not in draft
    assert draft["research_started"] is False
    assert "No saved research trail" in draft["next"]
    assert "upload" in draft["next"]
    assert "look-up" in draft["next"] or "source card" in draft["next"]
    assert "dig" not in draft["next"].casefold()
    assert "unless you ask" in draft["next"]
    assert "first draft" in draft["next"]
    assert "feedback" in draft["next"]
    assert not sequence_dir(OTHER, 1).exists()

    filed = essay_intake("fresh.docx", "confirm", OTHER, create=True)
    assert filed["ok"] is True
    assert filed["wrote"] is True
    assert filed["stage"] == "user-feedback"
    assert filed["first_draft"] == "done"
    assert filed["research_started"] is False
    assert "No saved research trail" in filed["next"]
    assert not (uploads_dir(OTHER) / "fresh.docx").exists()
    state = json.loads((sequence_dir(OTHER, 1) / "state.json").read_text(encoding="utf-8"))
    assert stage_name(state) == "user-feedback"
    assert state["first_draft"] == "done"
    assert state["research"]["summary"] is None
    assert state["research"]["full"] is None
    assert state["research_on_request"] is True
    assert essay_projects("continue")["rows"][0]["stage"] == "user-feedback"
    assert essay_health()["packs"][0]["stage"] == "user-feedback"
    assert list(essay_dir(OTHER).rglob("essay.research.summary/**")) == []

    feedback = essay_revise(f"{OTHER}-1", "Please look at the structure.", "draft")
    assert feedback["ok"] is True
    assert feedback["wrote"] is False
    state = json.loads((sequence_dir(OTHER, 1) / "state.json").read_text(encoding="utf-8"))
    assert state["research"]["summary"] is None

    write_upload(workspace, OTHER)
    asked = essay_research_summary(f"{OTHER}-1", TOPIC, "draft", [source(OTHER)])
    assert asked["ok"] is True
    assert asked["stored"] is False


def test_rewritten_essay_is_not_a_shared_original_and_surnames_stay_unresolved(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    parent = _port_output()
    parent_bytes = parent.read_bytes()
    _research()
    incoming = workspace / "unassigned" / "rewrite.docx"
    _rewrite(incoming)
    # No exact paragraph is shared, and the file has no Counted words line.
    latest = [paragraph.text.strip() for paragraph in Document(str(parent)).paragraphs if paragraph.text.strip()]
    current = [paragraph.text.strip() for paragraph in Document(str(incoming)).paragraphs if paragraph.text.strip()]
    assert set(latest).isdisjoint(current)

    draft = essay_intake("rewrite.docx", "draft", MODULE)
    assert draft["kind"] == "output"
    assert draft["essay_id"] == f"{MODULE}-2"
    assert "sequence_id" not in draft
    filed = essay_intake("rewrite.docx", "confirm", MODULE)
    assert filed["ok"] is True
    assert filed["wrote"] is True
    assert filed["path"].startswith(f"Essay-writer/{MODULE}/sequences/2/")
    assert not (uploads_dir(MODULE) / "rewrite.docx").exists()
    assert parent.read_bytes() == parent_bytes
    assert filed["reference_index"] == [
        {
            "line": "Smith, A. (2020) Port pricing. Journal of Ports.",
            "source_id": "src-1",
        }
    ]
    assert filed["unresolved_references"] == [
        "Unknown, B. (1999) Missing paper.",
        "Lee, C. (2019) Hull form.",
    ]


def test_counted_words_and_thin_overlap_do_not_auto_fork(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    _docx(
        sequence_dir(MODULE, 1) / "output" / f"{MODULE}-1.docx",
        [
            "The port handles containers.",
            "Alpha discusses quay cranes and berth windows.",
            "Beta discusses pilotage and tug allocation.",
            "Gamma discusses yard density and dwell time.",
            "Delta discusses bunker cost and route choice.",
            "Counted words: 1200",
            "Reference List",
            "Smith, A. (2020) Port pricing. Journal of Ports.",
        ],
    )
    _docx(
        workspace / "unassigned" / "marker.docx",
        [
            "Warehouses in Europe handle different cargo.",
            "Tariffs change the route a ship takes today.",
            "Counted words: 1200",
            "Reference List",
            "Lee, C. (2019) Hull form.",
        ],
    )
    marker = essay_intake("marker.docx", "draft", MODULE)
    assert marker["kind"] == "supporting"
    assert marker["sequence_id"] is None
    assert marker["wrote"] is False
    filed = essay_intake("marker.docx", "confirm", MODULE)
    assert filed["wrote"] is True
    assert filed["path"] == f"uploads/{MODULE}/marker.docx"
    assert not sequence_dir(MODULE, 2).exists()

    _docx(
        workspace / "unassigned" / "thin.docx",
        [
            "Hull plating fails when the coating is thin.",
            "Paint systems differ across drydock yards.",
            "Class surveys record steel thickness.",
            "Cathodic protection slows corrosion on the hull.",
            "A surveyor signs the thickness report.",
            "Reference List",
            "Smith, A. (2020) Port pricing. Journal of Ports.",
        ],
    )
    thin = essay_intake("thin.docx", "draft", MODULE)
    assert thin["kind"] == "supporting"
    assert thin["sequence_id"] is None
    moved = essay_intake("thin.docx", "confirm", MODULE)
    assert moved["path"] == f"uploads/{MODULE}/thin.docx"
    assert not sequence_dir(MODULE, 2).exists()


def test_unchanged_word_stays_in_unassigned(workspace):
    assert _confirm_new()["sequence_id"] == f"{MODULE}-1"
    output = _port_output()
    incoming = workspace / "unassigned" / "same.docx"
    incoming.parent.mkdir(parents=True, exist_ok=True)
    incoming.write_bytes(output.read_bytes())

    report = essay_intake("same.docx", "confirm", MODULE)
    assert report["ok"] is True
    assert "refused" not in report
    assert report["wrote"] is False
    assert report["stored"] is False
    assert report["summary"]["same_file"] is True
    assert report["stays"] == "unassigned"
    assert "unassigned" in report["next"]
    assert "Nothing was written" in report["next"]
    assert incoming.is_file()
    assert not sequence_dir(MODULE, 2).exists()


def test_stage_six_copy_matches_roll_forward():
    bullet = next(
        line for line in basics_lines("6") if "rolls forward onto the next sequence" in line
    )
    assert "rolls forward onto the next sequence" in bullet
    assert "leaves the parent sequence" in bullet
    assert "replaces the draft" not in bullet
    revise = next(door["description"] for door in DOORS if door["tool"] == "essay.revise")
    intake = next(door["description"] for door in DOORS if door["tool"] == "essay.intake")
    assert "rolls forward onto the next sequence" in revise
    assert "Say the reply" in intake
    assert "does not own the chat" in intake
    for banned in ("Grok Bot", "Grok Build", "Cursor", "SuperGrok"):
        assert banned not in intake
    assert "no sequence id" in intake
    assert "unassigned" in intake
    text = INSTRUCTIONS
    assert "no saved research trail" in text
    assert "dig" in text
    assert "Grok Bot says these replies" in text
    assert "does not own the chat" in text


def test_kb1_vault_excerpt_is_not_a_gathered_source(workspace):
    note = kb1_module_dir(ESSAY) / "ports.md"
    note.parent.mkdir(parents=True)
    note.write_text("The quay was measured in 2019.\n", encoding="utf-8")
    module_note = essay_dir(ESSAY) / "kb1" / "ports.md"
    module_note.parent.mkdir(parents=True)
    module_note.write_text("The quay was measured in 2019.\n", encoding="utf-8")
    seq = open_sequence(ESSAY)

    lied = {
        "id": "src-lie",
        "origin": "kb1",
        "locator": f"kb1/{ESSAY}/ports.md",
        "read": True,
        "excerpt": "Nobody (1999) A paper that was never filed. The vault never said this.",
        "author": "Nobody",
        "title": "A paper that was never filed",
        "year": "1999",
        "quote": "The vault never said this.",
    }
    refused = essay_research_summary(seq, TOPIC, "draft", [lied])
    assert refused["ok"] is False
    assert refused["refused"] == "invent-forbid"
    assert refused["stored"] is False

    module_lie = dict(lied)
    module_lie["locator"] = f"Essay-writer/{ESSAY}/kb1/ports.md"
    module_refused = essay_research_summary(seq, TOPIC, "draft", [module_lie])
    assert module_refused["refused"] == "invent-forbid"

    real = {
        "id": "src-kb",
        "origin": "kb1",
        "locator": f"kb1/{ESSAY}/ports.md",
        "read": True,
        "excerpt": "This excerpt is not the file and must be ignored.",
        "title": "quay was measured",
        "year": "2019",
        "quote": "The quay was measured in 2019.",
    }
    opened = essay_research_summary(seq, TOPIC, "draft", [real])
    assert opened["ok"] is True
    assert opened["sources"][0]["quote"] == "The quay was measured in 2019."
    assert opened["sources"][0]["locator"] == f"kb1/{ESSAY}/ports.md"
