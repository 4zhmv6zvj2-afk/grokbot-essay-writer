"""E2E #6 — download list membership, Word shape, and the chat links card."""

import json
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from helpers import ESSAY, TOPIC, open_sequence, source, write_upload

from essay_writer.dig import essay_dig
from essay_writer.downloads import run_downloads
from essay_writer.research import essay_research_summary
from essay_writer.research_packs import ship_4a
from essay_writer.server import HANDLERS
from essay_writer.shells import cursor_shell, front_office, front_office_call
from essay_writer.student_voice import student_surface_hits

PDF = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n"
HTML = b"<html><h1>Abstract</h1><p>Publisher abstract page.</p></html>"
HTML_URL = "https://publisher.example/html"
PDF_URL = "https://publisher.example/good.pdf"
BAD_URL = "https://publisher.example/missing"
OPENED_URL = "https://publisher.example/opened-abstract"
ABSTRACT_URL = "https://publisher.example/abstract-only"
NEEDS_URL = "https://library.example/needs-a-pdf"
INTRO = (
    "Leftovers only: sources not already in the papers zip. "
    "Citation lines stay on annotation.docx. "
    "Fetch each file from the university library."
)
T0 = datetime(2026, 10, 13, 4, 0, tzinfo=timezone.utc)


def _output(workspace: Path, essay_id: str = f"{ESSAY}-1") -> Path:
    number = essay_id.rsplit("-", 1)[-1]
    return workspace / "Essay-writer" / ESSAY / "sequences" / number / "output"


def _docx(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def _plant(workspace: Path) -> str:
    seq = open_sequence()
    folder = workspace / "uploads" / ESSAY
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "ssrn-abstract-dump.pdf").write_bytes(PDF)
    state = workspace / "Essay-writer" / ESSAY / "sequences" / "1" / "download-run.json"
    state.parent.mkdir(parents=True, exist_ok=True)
    state.write_text(
        json.dumps(
            {
                "attempts": {},
                "seen_urls": [],
                "failures": [
                    {
                        "url": OPENED_URL,
                        "bucket": "html-abstract",
                        "title": "Opened corridor paper",
                        "doi": "10.1000/opened",
                    },
                    {
                        "url": ABSTRACT_URL,
                        "bucket": "html-abstract",
                        "title": "Abstract only runway",
                        "doi": "10.1000/abstract",
                    },
                ],
                "papers": [
                    {
                        "url": "https://cdn.example/ssrn-abstract-dump.pdf",
                        "bucket": "paper",
                        "title": "Opened corridor paper",
                        "doi": "10.1000/opened",
                        "path": f"uploads/{ESSAY}/ssrn-abstract-dump.pdf",
                    }
                ],
                "started_at": None,
                "last_useful_at": None,
                "stopped": False,
                "stop_reason": None,
            }
        ),
        encoding="utf-8",
    )
    return seq


def _sources() -> list[dict]:
    return [
        {
            "id": "opened",
            "author": "Ng, A.",
            "year": "2019",
            "title": "Opened corridor paper",
            "doi": "10.1000/opened",
            "locator": f"uploads/{ESSAY}/ssrn-abstract-dump.pdf",
        },
        {
            "id": "abstract-only",
            "author": "Karaosman, H.",
            "year": "2018",
            "title": "Abstract only runway",
            "doi": "10.1000/abstract",
            "url": ABSTRACT_URL,
        },
        {
            "id": "needs",
            "author": "Adey, P.",
            "year": "2018",
            "title": "Needs a library PDF",
            "url": NEEDS_URL,
        },
    ]


def test_opened_pdf_is_absent_and_leftovers_stay_on_the_download_list(workspace):
    seq = _plant(workspace)
    delivery = ship_4a(seq, _sources(), {"opened": "Ng argues the corridor is open."})
    listed = _docx(_output(workspace) / "download-list.docx")
    bibliography = _docx(_output(workspace) / "annotation.docx")
    archive = zipfile.ZipFile(_output(workspace) / "papers.zip")

    assert "Opened corridor paper" not in listed
    assert "10.1000/opened" not in listed
    assert OPENED_URL not in listed
    assert "ssrn-abstract-dump.pdf" in archive.namelist()
    assert "Opened corridor paper" in bibliography
    assert "Abstract only runway" not in bibliography

    assert "Abstract only runway" in listed or "10.1000/abstract" in listed
    assert "HTML abstract page" in listed
    assert ABSTRACT_URL in listed
    assert "not a paper" in listed.casefold()
    assert "Needs a library PDF" in listed
    assert "Still needs a PDF" in listed
    assert NEEDS_URL in listed
    assert "Fetch a PDF for this gathered source." in listed

    urls = [item["url"] for item in delivery["links"]["items"]]
    titles = [item["title"] for item in delivery["links"]["items"]]
    assert OPENED_URL not in urls
    assert "Opened corridor paper" not in titles
    assert ABSTRACT_URL in urls
    assert NEEDS_URL in urls
    assert "Abstract only runway" in titles
    assert "Needs a library PDF" in titles


def test_download_list_word_shape_is_left_aligned_with_links(workspace):
    seq = _plant(workspace)
    ship_4a(seq, _sources(), {})
    path = _output(workspace) / "download-list.docx"
    document = Document(str(path))
    body = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    assert body[0].text == "Sources to download"
    assert body[0].runs[0].bold
    assert body[1].text == INTRO
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT for paragraph in body)
    assert all(paragraph.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY for paragraph in body)
    needs = next(paragraph for paragraph in body if "Needs a library PDF" in paragraph.text)
    assert needs.paragraph_format.left_indent == Inches(0.5)
    assert needs.paragraph_format.first_line_indent == Inches(-0.5)
    assert "Fetch a PDF for this gathered source." in needs.text
    plain = "\n".join(paragraph.text for paragraph in body)
    assert student_surface_hits(plain) == []
    assert "Annotated bibliography" not in plain
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "w:hyperlink" in document_xml
    assert NEEDS_URL in rels
    assert ABSTRACT_URL in rels
    assert OPENED_URL not in rels
    assert "hyperlink" in rels


def test_ship_4a_and_dig_confirm_return_a_chat_links_card(workspace, monkeypatch):
    write_upload(workspace, text=(
        "Smith, J. (2020) Chokepoints and crude.\n"
        "The Strait handles a large share of flows.\n"
        f"{NEEDS_URL}\n"
    ))
    seq = open_sequence()
    row = source()
    row["url"] = NEEDS_URL
    row["annotation"] = (
        "Smith (2020) argues the corridor is narrow. "
        'The paper says "The Strait handles a large share of flows."'
    )
    essay_research_summary(seq, TOPIC, "draft", [row])
    confirmed = front_office_call(
        "essay.research.summary",
        {
            "essay_id": seq,
            "topic": TOPIC,
            "confirmPhase": "confirm",
            "sources": [row],
        },
    )
    links = confirmed["delivery"]["links"]
    assert links["deliver"] == "chat"
    assert links["kind"] == "links"
    assert links["name"] == "Sources to download"
    assert links["items"][0]["title"] == "Chokepoints and crude"
    assert links["items"][0]["url"] == NEEDS_URL
    assert "still need a pdf" in links["text"].casefold()
    assert confirmed["delivery"]["annotation"]["deliver"] == "chat"
    assert "docx_base64" in confirmed["delivery"]["annotation"]
    assert "docx_base64" not in confirmed["delivery"]["download_list"]
    assert front_office()["handlers"]["essay.research.summary"] is HANDLERS["essay.research.summary"]
    assert cursor_shell()["handlers"]["essay.research.summary"] is HANDLERS["essay.research.summary"]

    monkeypatch.setenv("ESSAY_WRITER_RESOLVERS", "live")
    monkeypatch.setenv("ESSAY_WRITER_UNPAYWALL_EMAIL", "reader@example.com")

    def fake_json(url, headers=None):
        if "api.openalex.org" in url:
            return {
                "results": [
                    {
                        "title": "HTML only",
                        "publication_year": 2020,
                        "doi": "https://doi.org/10.1/html",
                        "authorships": [],
                    },
                    {
                        "title": "Real PDF",
                        "publication_year": 2021,
                        "doi": "https://doi.org/10.1/good",
                        "authorships": [],
                    },
                ]
            }
        if "api.unpaywall.org" in url:
            if "html" in url:
                return {"best_oa_location": {"url_for_pdf": HTML_URL, "url": HTML_URL}}
            if "good" in url:
                return {"best_oa_location": {"url_for_pdf": PDF_URL}}
            return {}
        return {"results": []}

    def fake_fetch(url):
        if url == PDF_URL:
            return 200, "application/pdf", PDF
        return 200, "text/html", HTML

    monkeypatch.setattr("essay_writer.dig._get_json", fake_json)
    monkeypatch.setattr("essay_writer.downloads.fetch_url", fake_fetch)
    dig_seq = open_sequence()
    essay_dig(dig_seq, "zara supply chain", "draft")
    dug = front_office_call(
        "essay.dig",
        {"essay_id": dig_seq, "query": "zara supply chain", "confirmPhase": "confirm"},
    )
    card = dug["links"]
    assert card["deliver"] == "chat"
    assert card["kind"] == "links"
    assert [item["url"] for item in card["items"]] == [HTML_URL]
    assert all(item["title"] != "Real PDF" for item in card["items"])
    assert front_office()["handlers"]["essay.dig"] is cursor_shell()["handlers"]["essay.dig"]
    listed = _docx(_output(workspace, dig_seq) / "download-list.docx")
    assert "Real PDF" not in listed
    assert "10.1/good" not in listed
    assert HTML_URL in listed


def test_download_stop_returns_links_for_cannot_open_only(workspace):
    seq = open_sequence()

    def fetch(url: str):
        if url == PDF_URL:
            return 200, "application/pdf", PDF
        return 200, "text/html", HTML

    run_downloads(seq, [{"url": PDF_URL, "doi": "10.1/good", "title": "Real PDF"}], fetch=fetch, now=T0)
    for _ in range(2):
        run_downloads(seq, [{"url": BAD_URL, "doi": "10.1/bad", "title": "Missing paper"}], fetch=fetch, now=T0)
    stopped = run_downloads(
        seq,
        [{"url": BAD_URL, "doi": "10.1/bad", "title": "Missing paper"}],
        fetch=fetch,
        now=T0 + timedelta(minutes=1),
    )
    assert stopped["stopped"] is True
    assert stopped["stop_reason"] == "retries"
    assert stopped["links"]["deliver"] == "chat"
    assert stopped["links"]["kind"] == "links"
    assert [item["url"] for item in stopped["links"]["items"]] == [BAD_URL]
    assert all("Real PDF" not in item["title"] for item in stopped["links"]["items"])
    listed = _docx(_output(workspace) / "download-list.docx")
    assert "Missing paper" in listed or "10.1/bad" in listed
    assert "Real PDF" not in listed
    assert "10.1/good" not in listed
