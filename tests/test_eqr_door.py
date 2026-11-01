"""Layer 3. One EQR door opens the saved Word file when it exists."""

from docx import Document
from helpers import ESSAY, pack, source, write_upload

from essay_writer.paths import essay_dir, work_dir
from essay_writer.projects import essay_sequence_new
from essay_writer.review import essay_review
from essay_writer.server import REVIEW_HANDLERS
from essay_writer.specialists.eqr import HIGHEST_BAND, REVIEW_PROMPT
from essay_writer.specialists.eqr import essay_review as specialist_review
from essay_writer.student_voice import student_surface_hits
from essay_writer.word import essay_pack_render_word


def _open() -> str:
    essay_sequence_new(ESSAY, "draft")
    opened = essay_sequence_new(ESSAY, "confirm")
    assert opened["ok"] is True
    return opened["sequence_id"]


def test_review_without_a_word_file_does_not_invent_one(workspace):
    seq = _open()
    draft = essay_review(seq, "6", "PASS", "draft")
    assert draft["ok"] is True
    assert draft["word_opened"] is False
    assert "not saved yet" in draft["findings"]
    assert essay_review(seq, "6", "PASS", "confirm")["stored"] is True
    assert not (work_dir(seq) / "output" / f"{seq}.docx").exists()


def test_review_opens_the_saved_word_and_can_write_findings(workspace, monkeypatch, tmp_path):
    write_upload(workspace)
    folder = essay_dir(ESSAY)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(
        "Genre: report\nQuestion: How do chokepoints constrain the corridor?\n",
        encoding="utf-8",
    )
    seq = _open()
    from essay_writer.research import essay_research_summary

    essay_research_summary(seq, "chokepoints", "draft", [source()])
    essay_research_summary(seq, "chokepoints", "confirm", [source()])
    from essay_writer.validate import essay_pack_validate

    essay_pack_validate(seq, pack(), "draft")
    essay_pack_validate(seq, pack(), "confirm")
    constitution = "This is a report. How do chokepoints constrain the corridor?"
    essay_review(seq, "6", "PASS", "draft", constitution=constitution)
    essay_review(seq, "6", "PASS", "confirm", constitution=constitution)
    essay_pack_render_word(seq, pack(), "draft", stage="6")
    written = essay_pack_render_word(seq, pack(), "confirm", stage="6")
    assert written["wrote"] is True, written.get("reason")
    documents = tmp_path / "Documents"
    monkeypatch.setenv("ESSAY_WRITER_MAC_DIR", str(documents))
    draft = essay_review(seq, "9", "PASS", "draft", constitution=constitution)
    assert draft["word_opened"] is True
    assert "Opened the saved Word file" in draft["findings"]
    assert "Chokepoints" in draft["findings"] or "corridor" in draft["findings"].casefold()
    stored = essay_review(seq, "9", "PASS", "confirm", constitution=constitution)
    assert stored["stored"] is True
    assert stored["findings_file"]["wrote"] is True
    note = documents / ESSAY / f"{ESSAY}-findings.docx"
    assert note.is_file()
    assert not (documents / f"{ESSAY}-10").exists()
    text = "\n".join(paragraph.text for paragraph in Document(str(note)).paragraphs)
    assert "Findings" in text
    assert "Opened the saved Word file" in text
    assert REVIEW_HANDLERS["essay.review"] is specialist_review
    assert student_surface_hits(REVIEW_PROMPT) == []
    assert "highest band" in REVIEW_PROMPT
    assert "highest band" in HIGHEST_BAND


def test_unreadable_word_does_not_record_a_pass(workspace):
    seq = _open()
    path = work_dir(seq) / "output" / f"{seq}.docx"
    path.parent.mkdir(parents=True)
    path.write_bytes(b"not a word file")
    refused = essay_review(seq, "9", "PASS", "draft")
    assert refused["ok"] is False
    assert refused["stored"] is False


def test_writer_still_cannot_review(workspace):
    seq = _open()
    refused = essay_review(seq, "6", "PASS", "confirm", seat="writer")
    assert refused["ok"] is False
    assert refused["refused"] == "seat"


def test_mark_scheme_band_still_blocks_a_bare_pass(workspace):
    from docx import Document

    folder = essay_dir(ESSAY)
    folder.mkdir(parents=True)
    (folder / "mark-scheme.md").write_text(
        "High Distinction (80-100): evidence\nPass (50-59): adequate\n",
        encoding="utf-8",
    )
    seq = _open()
    stage6 = essay_review(seq, "6", "PASS", "draft", constitution="Pass.")
    assert stage6["ok"] is True
    path = work_dir(seq) / "output" / f"{seq}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("Pass.")
    document.save(str(path))
    from essay_writer.assignment_gate import assignment_gate

    bare = essay_review(seq, "9", "PASS", "draft", constitution="This is a report. evidence.")
    assert bare.get("refused") != "assignment-gate"
    assert assignment_gate(seq, "Pass.", seat="eqr")["ok"] is True
    writer_gate = assignment_gate(seq, "Pass.", seat="writer")
    assert writer_gate["ok"] is False
    assert "highest_band" in writer_gate["missing"]
    assert "High Distinction" in writer_gate["reason"]
