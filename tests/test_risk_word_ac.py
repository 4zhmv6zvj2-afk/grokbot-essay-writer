"""Regressions for Word final, recovery, withheld cites, and diary columns."""

from __future__ import annotations

import json

from docx import Document
from helpers import open_sequence

from essay_writer.diary import top_column_quote
from essay_writer.late_chain import write_stage10
from essay_writer.paths import essay_dir, module_code, work_dir
from essay_writer.state import load_state, save_state
from essay_writer.word import word_manifest, write_docx

MODULE = "MOD1004AA"
JONES_LINE = "Jones, A. (2021) 'Port calls', Harbour Journal, 2(1), pp. 1-3."
SMITH_LINE = "Smith, J. (2020) 'Chokepoints and crude', Maritime Policy, 12(3), pp. 4-9."
ALDEN_LINE = "Alden, P. (2019) Port calls remain seasonal."
BERRY_LINE = "Berry, K. (2018) Tides mark the channel."


def _brief(text: str) -> None:
    folder = essay_dir(MODULE)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(text, encoding="utf-8")


def _save(seq: str, pack: dict) -> None:
    state = load_state(seq)
    state["saved_pack"] = pack
    save_state(seq, state)


def _paragraphs(path) -> list[str]:
    document = Document(str(path))
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def _references(path) -> list[str]:
    texts = _paragraphs(path)
    if "References" not in texts:
        return []
    return texts[texts.index("References") + 1 :]


def test_word_count_outside_the_brief_band_is_not_written(workspace):
    """A count outside the brief band blocks Word and does not set the flag."""
    _brief(
        "Genre: report\n"
        "Assignment title: Individual report\n"
        "Question: How do chokepoints constrain the corridor?\n"
        "Word count: 2000 words, plus or minus 10%.\n"
    )
    seq = open_sequence(MODULE)
    pack = {
        "essay_id": seq,
        "title": "Individual report",
        "assessment": "Individual report",
        "referencing": "harvard",
        "conclusion": "noted",
        "references_markdown": ALDEN_LINE,
        "sections": [
            {
                "heading": "Notes",
                "body": "The corridor depends on a narrow strait.",
            }
        ],
    }
    _save(seq, pack)
    source = work_dir(seq) / "output" / f"{seq}.docx"
    final = work_dir(seq) / "output" / f"{MODULE}-10" / f"{MODULE}-final.docx"
    refused = write_stage10(seq)
    assert refused["ok"] is False
    assert refused["refused"] == "word-count"
    assert refused["sequence_written"] is False
    assert refused["copied"] is False
    assert not source.is_file()
    assert not final.exists()
    state = load_state(seq)
    assert state.get("word_count_acknowledged") is not True
    saved = state.get("saved_pack")
    assert isinstance(saved, dict)
    assert saved.get("word_count_acknowledged") is not True

    saved["word_count_acknowledged"] = True
    _save(seq, saved)
    allowed = write_stage10(seq)
    assert allowed["sequence_written"] is True, allowed
    assert source.is_file()
    assert final.is_file()
    assert final.read_bytes() == source.read_bytes()
    assert load_state(seq).get("word_count_acknowledged") is not True


def test_withheld_cite_leaves_the_body_and_the_reference_list(workspace):
    """A cite the gate already withheld is not printed in Word."""
    _brief(
        "Genre: report\n"
        "Assignment title: Individual report\n"
        "Question: How do chokepoints constrain the corridor?\n"
    )
    seq = open_sequence(MODULE)
    upload = workspace / "uploads" / MODULE
    upload.mkdir(parents=True, exist_ok=True)
    (upload / "smith.md").write_text(
        "Smith, J. (2020) Chokepoints and crude.\nThe Strait handles a large share of flows.\n",
        encoding="utf-8",
    )
    (upload / "jones.md").write_text(
        "Jones, A. (2021) Port calls.\nPort calls remain seasonal.\n",
        encoding="utf-8",
    )
    record = work_dir(seq) / "cite-fit.json"
    record.write_text(
        json.dumps({"ok": [], "rejected": [], "withheld": ["smith, 2020"]}) + "\n",
        encoding="utf-8",
    )
    pack = {
        "essay_id": seq,
        "title": "Individual report",
        "assessment": "Individual report",
        "referencing": "harvard",
        "conclusion": "noted",
        "references_markdown": f"{SMITH_LINE}\n{JONES_LINE}",
        "sections": [
            {
                "heading": "Notes",
                "body": "The strait matters (Smith, 2020). The port matters (Jones, 2021).",
            }
        ],
        "sources": [
            {
                "id": "src-smith",
                "origin": "upload",
                "locator": f"uploads/{MODULE}/smith.md",
                "author": "Smith, J.",
                "year": "2020",
                "title": "Chokepoints and crude",
                "quote": "The Strait handles a large share of flows.",
            },
            {
                "id": "src-jones",
                "origin": "upload",
                "locator": f"uploads/{MODULE}/jones.md",
                "author": "Jones, A.",
                "year": "2021",
                "title": "Port calls",
                "quote": "Port calls remain seasonal.",
            },
        ],
    }
    _prepared, report = word_manifest(seq, pack)
    assert "Smith, 2020" not in " ".join(section["body"] for section in report["sections"])
    assert all("Smith" not in line for line in report["references"])
    path = work_dir(seq) / "output" / f"{seq}.docx"
    write_docx(seq, _prepared, report, path)
    body = "\n".join(_paragraphs(path))
    assert "(Smith, 2020)" not in body
    assert "Smith, J. (2020)" not in body
    assert "(Jones, 2021)" in body
    refs = _references(path)
    assert any("Jones" in line for line in refs)
    assert all("Smith" not in line for line in refs)


def test_recovery_keeps_file_lines_and_does_not_invent(workspace):
    """A missing pack rebuilds from the docx and does not add research-only lines."""
    _brief(
        "Genre: report\n"
        "Assignment title: Individual report\n"
        "Question: How do chokepoints constrain the corridor?\n"
    )
    seq = open_sequence(MODULE)
    state = load_state(seq)
    state["saved_pack"] = None
    state["research"] = {
        "summary": {
            "sources": [
                {
                    "id": "src-smith",
                    "origin": "upload",
                    "locator": f"uploads/{MODULE}/smith.md",
                    "author": "Smith, J.",
                    "year": "2020",
                    "title": "Chokepoints and crude",
                    "quote": "The Strait handles a large share of flows.",
                }
            ]
        },
        "full": None,
    }
    save_state(seq, state)
    source = work_dir(seq) / "output" / f"{seq}.docx"
    source.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Notes", level=1)
    document.add_paragraph("The corridor depends on a narrow strait.")
    document.add_heading("References", level=1)
    document.add_paragraph(BERRY_LINE)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = ALDEN_LINE
    document.save(str(source))
    before = source.read_bytes()

    final = write_stage10(seq)
    assert final["sequence_written"] is True, final
    refs = _references(source)
    joined = "\n".join(refs)
    assert BERRY_LINE in joined
    assert ALDEN_LINE in joined
    assert "Smith" not in joined
    assert source.read_bytes() != before


def test_blank_registration_does_not_record_the_final(workspace):
    """An empty required registration number is refused and is not the final."""
    _brief(
        "Genre: report\n"
        "Assignment title: Individual report\n"
        "Question: How do chokepoints constrain the corridor?\n"
        "Title page: registration number.\n"
    )
    seq = open_sequence(MODULE)
    pack = {
        "essay_id": seq,
        "title": "Individual report",
        "assessment": "Individual report",
        "referencing": "harvard",
        "conclusion": "noted",
        "registration_number": "   ",
        "references_markdown": ALDEN_LINE,
        "sections": [
            {"heading": "Notes", "body": "The corridor depends on a narrow strait."}
        ],
    }
    _save(seq, pack)
    source = work_dir(seq) / "output" / f"{seq}.docx"
    final = work_dir(seq) / "output" / f"{module_code(seq)}-10" / f"{module_code(seq)}-final.docx"
    refused = write_stage10(seq)
    assert refused["ok"] is False, refused
    assert refused["refused"] == "title-page"
    assert refused["sequence_written"] is False
    assert refused["copied"] is False
    assert not source.is_file()
    assert not final.exists()


def test_unmounted_documents_is_not_created(workspace):
    """A missing Documents mount does not invent the folder or report a copy."""
    _brief(
        "Genre: report\n"
        "Assignment title: Individual report\n"
        "Question: How do chokepoints constrain the corridor?\n"
    )
    seq = open_sequence(MODULE)
    pack = {
        "essay_id": seq,
        "title": "Individual report",
        "assessment": "Individual report",
        "referencing": "harvard",
        "conclusion": "noted",
        "word_count_acknowledged": True,
        "references_markdown": ALDEN_LINE,
        "sections": [
            {"heading": "Notes", "body": "The corridor depends on a narrow strait."}
        ],
    }
    _save(seq, pack)
    missing = workspace / "Documents"
    result = write_stage10(seq, documents=missing)
    assert result["copied"] is False
    assert result.get("reason") == "mac-not-mounted"
    assert not missing.exists()
    assert not (workspace / "Documents").exists()


def test_diary_column_keeps_the_printed_name(workspace):
    """The quote uses the column header printed on the diary table."""
    del workspace
    table = (
        "| Clarity of the decisions | 70 + |\n"
        "| --- | --- |\n"
        "| Reflection | exact and evidenced |\n"
    )
    quoted = top_column_quote(table, "70+")
    assert quoted.startswith("70 +")
    assert "exact and evidenced" in quoted
    assert "mark scheme" not in quoted.casefold()


def test_rewritten_sequence_file_refreshes_the_final(workspace):
    """The sequence -10 final is the bytes just written, not a stale parent copy."""
    _brief(
        "Genre: report\n"
        "Assignment title: Individual report\n"
        "Question: How do chokepoints constrain the corridor?\n"
    )
    seq = open_sequence(MODULE)
    pack = {
        "essay_id": seq,
        "title": "Individual report",
        "assessment": "Individual report",
        "referencing": "harvard",
        "conclusion": "noted",
        "references_markdown": ALDEN_LINE,
        "sections": [
            {"heading": "Notes", "body": "The corridor depends on a narrow strait."}
        ],
    }
    _save(seq, pack)
    source = work_dir(seq) / "output" / f"{seq}.docx"
    final = work_dir(seq) / "output" / f"{MODULE}-10" / f"{MODULE}-final.docx"
    final.parent.mkdir(parents=True, exist_ok=True)
    stale = Document()
    stale.add_paragraph("stale parent final")
    stale.save(str(final))
    stale_bytes = final.read_bytes()
    result = write_stage10(seq)
    assert result["sequence_written"] is True, result
    assert source.is_file()
    assert final.read_bytes() == source.read_bytes()
    assert final.read_bytes() != stale_bytes


def test_empty_footer_does_not_record_the_sequence_final(workspace, monkeypatch):
    """An empty footer refuses the write and does not copy the sequence file to the final."""
    _brief(
        "Genre: report\n"
        "Assignment title: Individual report\n"
        "Question: How do chokepoints constrain the corridor?\n"
    )
    seq = open_sequence(MODULE)
    pack = {
        "essay_id": seq,
        "title": "Individual report",
        "assessment": "Individual report",
        "referencing": "harvard",
        "conclusion": "noted",
        "references_markdown": ALDEN_LINE,
        "sections": [
            {"heading": "Notes", "body": "The corridor depends on a narrow strait."}
        ],
    }
    _save(seq, pack)
    source = work_dir(seq) / "output" / f"{seq}.docx"
    final = work_dir(seq) / "output" / f"{MODULE}-10" / f"{MODULE}-final.docx"
    source.parent.mkdir(parents=True, exist_ok=True)
    planted = Document()
    planted.add_paragraph("sequence file with no page number")
    planted.save(str(source))
    before = source.read_bytes()
    final.parent.mkdir(parents=True, exist_ok=True)
    stale = Document()
    stale.add_paragraph("stale parent final")
    stale.save(str(final))
    stale_bytes = final.read_bytes()

    def _empty(*_args, **_kwargs):
        from essay_writer.word import EmptyFooter

        raise EmptyFooter()

    monkeypatch.setattr("essay_writer.word.write_docx", _empty)
    refused = write_stage10(seq)
    assert refused["ok"] is False
    assert refused["refused"] == "footer"
    assert refused["sequence_written"] is False
    assert refused["copied"] is False
    assert source.read_bytes() == before
    assert final.read_bytes() == stale_bytes
