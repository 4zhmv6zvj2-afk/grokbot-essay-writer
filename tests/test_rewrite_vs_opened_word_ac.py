"""Rewrite versus the opened Word file. Stable ids: AC1–AC5.

An opened Word baseline stops rewrite and EQR from failing only because the
draft is longer than an invented word band or per-section split that file
already exceeds. A stated brief band, a brief section budget, and Confirm
hard gates stay in place. No opened file leaves those gates as they were.
Cursor and Grok Build call the same function. Build covered.
"""

import inspect

from docx import Document
from helpers import open_sequence, source, write_upload

from essay_writer.build_office import rewrite_vs_opened_word as build_gate
from essay_writer.build_office import stage_launch
from essay_writer.cursor_office import cursor_stage_launch
from essay_writer.cursor_office import rewrite_vs_opened_word as cursor_gate
from essay_writer.diary import section_plan_for
from essay_writer.gates import essay_gates, evaluate_gates
from essay_writer.paths import essay_dir, work_dir
from essay_writer.review import essay_review
from essay_writer.rewrite_baseline import BUILD_COVERED_LINE, rewrite_vs_opened_word
from essay_writer.specialists.eqr import _missed_word_band
from essay_writer.specialists.grok_bot import confirm_gate
from essay_writer.specialists.writer import document_word_band
from essay_writer.validate import _word_gate, essay_pack_validate

STUDY = "MOD1008AA"
STATED = "MOD1004AA"
BUDGET = "RPT1008AA"
DIARY = "RPT5001AA"
INVENTED = {"limit": 2000, "tolerance": 0.10}
TIGHTER = {"min": 100, "max": 200}
QUESTION = "This is a report. How do chokepoints constrain the corridor?"
DIARY_BRIEF = (
    "Assignment title: Individual reflective diary\n"
    "Write two entries. Each entry names one decision and analyses it with theory.\n"
    "One entry for each occasion: North occasion and South occasion.\n"
    "The whole diary is 2,000 words, plus or minus 10%.\n"
)
SIMULATION = (
    "Genre: report\n"
    "Question: How do chokepoints constrain the corridor?\n"
    "The module names simulation events.\n"
    "The corridor drill stays on the brief.\n"
)


def _words(count: int) -> str:
    return " ".join(["flows"] * count)


def _brief(module: str, text: str) -> None:
    folder = essay_dir(module)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brief.md").write_text(text, encoding="utf-8")


def _docx(seq: str, paragraphs: list[str]) -> None:
    path = work_dir(seq) / "output" / f"{seq}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for line in paragraphs:
        document.add_paragraph(line)
    document.save(str(path))


def _docx_sections(seq: str, rows: list[tuple[str, str]]) -> None:
    path = work_dir(seq) / "output" / f"{seq}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for heading, body in rows:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    document.save(str(path))


def _gate(report: dict, name: str) -> dict:
    for item in report["gates"]:
        if item["name"] == name:
            return item
    raise AssertionError(name)


def _codes(items: list[dict]) -> list[str]:
    return [item["code"] for item in items]


def _study_pack(seq: str) -> dict:
    return {
        "essay_id": seq,
        "title": "Corridor flows",
        "genre": "report",
        "conclusion": "open",
        "word_rules": dict(INVENTED),
        "sections": [{"heading": "Analysis", "body": _words(2500)}],
    }


def test_ac1_opened_word_does_not_fail_a_looser_invented_band(workspace):
    """AC1. An opened Word file caps invented word-band and section pressure."""
    seq = open_sequence(STUDY)
    pack = _study_pack(seq)
    _docx(seq, [_words(2500), "(Smith, 2020)."])
    held = essay_gates(seq, pack, body=_words(2500))
    assert _gate(held, "word-band")["passed"] is True
    assert "already exceeds this invented band" in _gate(held, "word-band")["reason"]
    checked = essay_pack_validate(seq, pack, "draft")
    assert "word-count" not in _codes(checked["blockers"])
    assert any(
        item["code"] == "word-count" and "already exceeds this invented band" in item["reason"]
        for item in checked["notes"]
    )
    reviewed = essay_review(seq, "9", "PASS", "draft")
    assert reviewed["ok"] is True, reviewed
    assert reviewed.get("refused") != "engagement"
    offices = cursor_gate(seq, _words(2500), INVENTED, counter="words")
    assert offices["waive_word"] is True
    assert offices["document_still_fails"] is False
    assert build_gate(seq, _words(2500), INVENTED, counter="words") == offices

    _brief(
        STATED,
        "Genre: report\n"
        "Question: How do chokepoints constrain the corridor?\n"
        "Word count: 800 words (720 to 880).\n",
    )
    stated = open_sequence(STATED)
    _docx(stated, [_words(800)])
    inside = {
        "essay_id": stated,
        "title": "Corridor",
        "genre": "report",
        "conclusion": "open",
        "word_rules": dict(TIGHTER),
    }
    fit = essay_gates(stated, inside, body=_words(800))
    assert _gate(fit, "word-band")["passed"] is True
    assert document_word_band(stated)["max"] == 880
    _docx(stated, [_words(2000)])
    loose = essay_gates(stated, inside, body=_words(2000))
    assert _gate(loose, "word-band")["passed"] is False
    assert "720" in _gate(loose, "word-band")["reason"]
    stated_only = essay_gates(stated, {"essay_id": stated, "genre": "report", "conclusion": "open"}, body=_words(2000))
    assert _gate(stated_only, "word-band")["passed"] is False
    over = QUESTION + " " + _words(2000)
    _docx(stated, [over])
    refused = essay_review(stated, "9", "PASS", "draft")
    assert refused["ok"] is False
    assert refused["refused"] == "engagement"
    assert refused["engagement"]["word_band"]["counted"] > 880

    _brief(BUDGET, "Genre: report\nSection budgets: Introduction 40, Analysis 40\n")
    budget = open_sequence(BUDGET)
    sections = [
        {"heading": "Introduction", "body": _words(80)},
        {"heading": "Analysis", "body": _words(80)},
    ]
    budget_pack = {"essay_id": budget, "genre": "report", "conclusion": "open", "sections": sections}
    budget_gate = essay_gates(budget, budget_pack)
    assert _gate(budget_gate, "word-band")["passed"] is False
    assert "word budget" in _gate(budget_gate, "word-band")["reason"]

    _brief(DIARY, DIARY_BRIEF)
    diary = open_sequence(DIARY)
    plan = section_plan_for(diary)
    invented_rows = [row for row in plan["sections"] if row.get("from_brief") is False]
    assert plan["plan_not_brief_rule"] is True
    assert invented_rows
    high = invented_rows[0]["max"]
    long_entry = _words(high + 300)
    draft_sections = [{"heading": row["name"], "body": long_entry} for row in invented_rows]
    _docx_sections(diary, [(row["heading"], row["body"]) for row in draft_sections])
    section_decision = rewrite_vs_opened_word(diary, long_entry, {}, sections=draft_sections, counter="words")
    assert invented_rows[0]["name"] in section_decision["waived_sections"]
    diary_gate = essay_gates(
        diary,
        {"essay_id": diary, "genre": "diary", "conclusion": "open", "sections": draft_sections},
    )
    assert "word budget" not in _gate(diary_gate, "word-band")["reason"]
    diary_check = essay_pack_validate(
        diary,
        {
            "essay_id": diary,
            "title": "Individual reflective diary",
            "conclusion": "open",
            "sections": draft_sections,
        },
        "draft",
    )
    assert diary_check.get("refused") != "assignment-gate"
    assert "section-budget" not in _codes(diary_check["blockers"])


def test_ac2_no_baseline_keeps_word_band_and_brief_gates(workspace):
    """AC2. With no opened Word file, word-band and brief gates still fail."""
    seq = open_sequence(STUDY)
    pack = _study_pack(seq)
    assert not (work_dir(seq) / "output" / f"{seq}.docx").exists()
    open_gate = essay_gates(seq, pack, body=_words(2500))
    assert _gate(open_gate, "word-band")["passed"] is False
    assert "1800" in _gate(open_gate, "word-band")["reason"]
    checked = essay_pack_validate(seq, pack, "draft")
    assert "word-count" in _codes(checked["blockers"])
    absent = rewrite_vs_opened_word(seq, _words(2500), INVENTED, counter="words")
    assert absent["baseline"] is False
    assert absent["waive_word"] is False

    _brief(
        STATED,
        "Genre: report\nQuestion: How do chokepoints constrain the corridor?\nWord count: 800 words (720 to 880).\n",
    )
    stated = open_sequence(STATED)
    short = essay_gates(stated, {"essay_id": stated, "genre": "report", "conclusion": "open"}, body=_words(40))
    assert _gate(short, "word-band")["passed"] is False
    assert "720" in _gate(short, "word-band")["reason"]

    _brief(BUDGET, "Genre: report\nSection budgets: Introduction 40, Analysis 40\n")
    budget = open_sequence(BUDGET)
    sections = [
        {"heading": "Introduction", "body": _words(80)},
        {"heading": "Analysis", "body": _words(80)},
    ]
    budget_gate = essay_gates(
        budget,
        {"essay_id": budget, "genre": "report", "conclusion": "open", "sections": sections},
    )
    assert _gate(budget_gate, "word-band")["passed"] is False
    assert "word budget" in _gate(budget_gate, "word-band")["reason"]

    _brief(DIARY, DIARY_BRIEF)
    diary = open_sequence(DIARY)
    plan = section_plan_for(diary)
    invented_rows = [row for row in plan["sections"] if row.get("from_brief") is False]
    long_entry = _words(invented_rows[0]["max"] + 300)
    draft_sections = [{"heading": row["name"], "body": long_entry} for row in invented_rows]
    no_file = rewrite_vs_opened_word(diary, long_entry, {}, sections=draft_sections, counter="words")
    assert no_file["baseline"] is False
    assert no_file["waived_sections"] == []
    diary_gate = essay_gates(
        diary,
        {"essay_id": diary, "genre": "diary", "conclusion": "open", "sections": draft_sections},
    )
    assert "word budget" not in _gate(diary_gate, "word-band")["reason"]


def test_ac3_invent_forbid_and_confirm_hard_gates_stay(workspace):
    """AC3. Invent-forbid, citations, plagiarism, and Confirm stay hard with a baseline."""
    _brief(STATED, SIMULATION)
    seq = open_sequence(STATED)
    _docx(seq, [_words(2500)])
    invented = {
        "essay_id": seq,
        "title": "Corridor",
        "genre": "report",
        "conclusion": "open",
        "simulation_acknowledged": True,
        "word_rules": dict(TIGHTER),
        "sections": [
            {
                "heading": "Introduction",
                "body": (
                    "The simulation closed the strait for nine whole days (Smith, 2020). "
                    "This report asks how chokepoints constrain the corridor."
                ),
            }
        ],
    }
    blocked = essay_pack_validate(seq, invented, "draft")
    assert blocked["ok"] is False
    assert blocked["refused"] == "invent-forbid"

    study = open_sequence(STUDY)
    _docx(study, [_words(2500)])
    cited = essay_gates(
        study,
        {"essay_id": study, "genre": "report", "conclusion": "open", "word_rules": dict(INVENTED)},
        body="A later source says otherwise (Patel, 2019). " + _words(2500),
    )
    assert _gate(cited, "citations")["passed"] is False
    assert _gate(cited, "word-band")["passed"] is True

    write_upload(workspace, essay_id=STUDY)
    copied = "The Strait handles a large share of flows."
    plagiarised = essay_gates(
        study,
        {
            "essay_id": study,
            "genre": "report",
            "conclusion": "open",
            "word_rules": dict(INVENTED),
            "sources": [source(STUDY)],
        },
        body=copied + " " + _words(2500),
    )
    assert plagiarised["plagiarism"]["passed"] is False
    assert _gate(plagiarised, "word-band")["passed"] is True

    _brief(
        BUDGET,
        "Genre: report\nQuestion: How do chokepoints constrain the corridor?\n",
    )
    asked = open_sequence(BUDGET)
    _docx(asked, [QUESTION + " " + _words(400)])
    refused = confirm_gate(asked, "This note does not answer the assignment.")
    assert refused is not None
    assert refused["refused"] == "assignment-gate"
    assert refused["ok"] is False


def test_ac4_build_covered_cursor_and_grok_build(workspace):
    """AC4. Cursor and Grok Build call one rewrite gate. Build covered."""
    assert BUILD_COVERED_LINE == "Build covered"
    assert "rewrite_vs_opened_word as _gate" in inspect.getsource(cursor_gate)
    assert "rewrite_vs_opened_word as _gate" in inspect.getsource(build_gate)
    assert "Build covered" in inspect.getsource(cursor_gate)
    assert "Build covered" in inspect.getsource(build_gate)
    assert "rewrite_vs_opened_word" in inspect.getsource(evaluate_gates)
    assert "rewrite_vs_opened_word" in inspect.getsource(_word_gate)
    assert "rewrite_vs_opened_word" in inspect.getsource(_missed_word_band)
    seq = open_sequence(STUDY)
    _docx(seq, [_words(2500)])
    shared = rewrite_vs_opened_word(seq, _words(2500), INVENTED, counter="words")
    assert cursor_gate(seq, _words(2500), INVENTED, counter="words") == shared
    assert build_gate(seq, _words(2500), INVENTED, counter="words") == shared
    cursor_rules = cursor_stage_launch("8")["settings"]["rules"]
    build_rules = stage_launch("8")["settings"]["rules"]
    assert cursor_rules == build_rules
    assert "Do not tighten prose below what the opened Word file already says" in cursor_rules
    assert "Do not tighten prose below what the opened Word file already says" in stage_launch("6")["settings"]["rules"]


def test_ac5_ready_line_build_covered():
    """AC5. The Ready line for this slice is Build covered."""
    assert BUILD_COVERED_LINE == "Build covered"
    assert "Build covered" in inspect.getsource(cursor_gate)
    assert "Build covered" in inspect.getsource(build_gate)
